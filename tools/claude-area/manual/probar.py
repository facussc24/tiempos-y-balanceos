#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Prueba del generador del manual, en las dos direcciones (el caso bueno pasa, el malo se frena).

    python probar.py                     corre toda la prueba en una carpeta temporal:
                                         no toca capturas/ ni exports/
    python probar.py --de-prueba 03 04   deja en capturas/ capturas DE MENTIRA de esos numeros
                                         (un dibujo rotulado "CAPTURA DE PRUEBA"), para mirar una hoja armada.
                                         Despues hay que sacarlas: no son la app.

Las capturas de mentira tienen un boton VERDE justo donde contenido.json dice que va el recuadro rojo:
si el recuadro del manual no abraza al boton verde, el generador esta ubicando mal las marcas.
"""
import contextlib
import copy
import io
import json
import re
import sys
import tempfile
from pathlib import Path

from PIL import Image, ImageChops, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.lang import MSO_LANGUAGE_ID
from pptx.util import Pt

import generar_manual as g

AQUI = Path(__file__).resolve().parent
VERDE = (0, 170, 80)
TAMANOS = {"05": (900, 1100), "08": (1400, 800), "09": (1500, 1000), "15": (1200, 500)}


def _letra(px, negrita=True):
    return ImageFont.truetype(str(g.DIR_FUENTES / ("segoeuib.ttf" if negrita else "segoeui.ttf")), px)


def captura_de_prueba(ruta, cap):
    """Dibuja una ventana de mentira con un boton verde en cada recuadro y un punto azul en la punta de cada flecha."""
    num = Path(cap["archivo"]).name[:2]
    w, h = TAMANOS.get(num, (1920, 1080))
    im = Image.new("RGB", (w, h), (246, 247, 249))
    d = ImageDraw.Draw(im)
    d.rectangle([0, 0, w, int(h * 0.045)], fill=(52, 58, 70))                      # barra de la ventana
    for i, col in enumerate(((237, 106, 94), (245, 191, 79), (98, 197, 84))):
        d.ellipse([14 + i * 26, int(h * 0.012), 32 + i * 26, int(h * 0.012) + 18], fill=col)
    d.rectangle([0, int(h * 0.045), int(w * 0.15), h], fill=(232, 235, 239))       # barra del costado
    for i in range(9):                                                             # renglones de relleno
        y = int(h * (0.22 + i * 0.07))
        d.rectangle([int(w * 0.22), y, int(w * (0.55 + 0.03 * (i % 4))), y + max(6, h // 90)], fill=(214, 219, 225))
    d.text((w * 0.58, h * 0.45), "CAPTURA DE PRUEBA %s" % num, font=_letra(w // 16), fill=(150, 30, 30), anchor="mm")
    d.text((w * 0.58, h * 0.45 + w // 14), "no es la app: es un dibujo para probar el manual",
           font=_letra(w // 48, False), fill=(110, 110, 110), anchor="mm")
    for m in cap.get("marcas", []):
        if m["tipo"] == "recuadro":
            caja = [round(w * m["x"] / 100), round(h * m["y"] / 100),
                    round(w * (m["x"] + m["ancho"]) / 100) - 1, round(h * (m["y"] + m["alto"]) / 100) - 1]
            d.rectangle(caja, fill=VERDE)
            d.text(((caja[0] + caja[2]) / 2, (caja[1] + caja[3]) / 2), "BOTÓN", fill=(255, 255, 255), anchor="mm",
                   font=_letra(max(12, min((caja[3] - caja[1]) // 2, (caja[2] - caja[0]) // 5))))
        else:
            cx, cy, r = w * m["hasta"][0] / 100, h * m["hasta"][1] / 100, w / 110
            d.ellipse([cx - r, cy - r, cx + r, cy + r], fill=(0, 90, 200))
    ruta.parent.mkdir(parents=True, exist_ok=True)
    im.save(ruta)


def todas_las_capturas(datos):
    return [c for p in datos["paginas"] for c in p.get("capturas", [])]


def correr(args):
    salida = io.StringIO()
    with contextlib.redirect_stdout(salida):
        codigo = g.main(args)
    return codigo, salida.getvalue()


def caja_de_color(im, color):
    """Caja que encierra todos los pixeles de ese color exacto (o None)."""
    mascara = None
    for canal, valor in zip(im.split(), color):
        esta = canal.point(lambda p, valor=valor: 255 if p == valor else 0)
        mascara = esta if mascara is None else ImageChops.darker(mascara, esta)
    return mascara.getbbox()


def formas(shapes):
    """Todas las formas de una diapositiva, entrando a los grupos."""
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            yield from formas(sh.shapes)


def cuadros_de_texto(prs):
    for s in prs.slides:
        for sh in formas(s.shapes):
            if sh.has_text_frame:
                yield sh.text_frame
            elif getattr(sh, "has_table", False) and sh.has_table:
                for fila in sh.table.rows:
                    for celda in fila.cells:
                        yield celda.text_frame


def todo_el_texto(prs):
    return "\n".join(tf.text for tf in cuadros_de_texto(prs))


def revisar_pptx(prs, datos, dir_capturas, con_notas):
    """Lo que tiene que cumplir el PowerPoint para proyectarlo o recorrerlo en la PC. Devuelve los problemas."""
    problemas = []
    activas = [p for p in datos["paginas"] if p.get("activa", True)]
    diapos = list(prs.slides)
    if len(diapos) != len(activas) + 2:
        return ["tiene %d diapositivas y tendria que tener %d" % (len(diapos), len(activas) + 2)]
    esperados = [datos["portada"]["titulo"]] + [p["titulo"] for p in activas] + [datos["cierre"]["titulo"]]
    for i, (s, titulo) in enumerate(zip(diapos, esperados), 1):
        tit = s.shapes.title
        if tit is None or tit.text_frame.text != titulo:
            problemas.append("diapositiva %d: el titulo de verdad (el cuadro de titulo) no dice '%s'" % (i, titulo))
        if not s._element.xpath("./p:transition/p:fade"):
            problemas.append("diapositiva %d: sin fundido" % i)
        if s._element.xpath("./p:timing"):
            problemas.append("diapositiva %d: tiene animaciones" % i)
        for sh in formas(s.shapes):
            if sh.shape_type == 13 and not "".join(sh._element.xpath("./p:nvPicPr/p:cNvPr/@descr")).strip():
                problemas.append("diapositiva %d: la imagen '%s' no tiene texto alternativo" % (i, sh.name))
    velocidades = {v for s in diapos for v in s._element.xpath("./p:transition/@spd")}
    if len(velocidades) > 1:
        problemas.append("el fundido no es el mismo en todas las diapositivas: %s" % sorted(velocidades))
    if any(r.font.language_id != MSO_LANGUAGE_ID.SPANISH_ARGENTINA
           for tf in cuadros_de_texto(prs) for para in tf.paragraphs for r in para.runs):
        problemas.append("hay texto sin idioma castellano de Argentina (sin idioma puesto)")
    if prs.core_properties.language != "es-AR":
        problemas.append("el idioma del documento no es es-AR")
    for pag, s in zip(activas, diapos[1:-1]):
        grupos = {sh.name: sh for sh in s.shapes if sh.shape_type == 6}
        for c in pag.get("capturas", []):
            if not (Path(dir_capturas) / Path(c["archivo"]).name).exists():
                continue
            num = Path(c["archivo"]).name[:2]
            grupo = grupos.get("Captura %s con sus marcas" % num)
            if grupo is None:
                problemas.append("pagina %s: no esta el grupo de la captura %s" % (pag["numero"], num))
                continue
            foto = next(sh for sh in grupo.shapes if sh.shape_type == 13)
            alt = "".join(foto._element.xpath("./p:nvPicPr/p:cNvPr/@descr"))
            if alt != c["que_se_ve"]:
                problemas.append("pagina %s: el texto alternativo de la captura %s no es el 'que_se_ve'" % (pag["numero"], num))
            marcos = [sh for sh in grupo.shapes if sh.name == "Recuadro rojo"]
            recuadros = [m for m in c.get("marcas", []) if m["tipo"] == "recuadro"]
            if len(marcos) != len(recuadros):
                problemas.append("pagina %s: la captura %s tiene %d recuadros y tendria que tener %d"
                                 % (pag["numero"], num, len(marcos), len(recuadros)))
                continue
            medio = g.Mm(g.GROSOR_MARCA / 2)
            for marco, m in zip(marcos, recuadros):
                # el rojo va por fuera de lo marcado: la linea (centrada) queda media linea mas afuera
                esperado = (foto.left + foto.width * m["x"] / 100 - medio, foto.top + foto.height * m["y"] / 100 - medio,
                            foto.width * m["ancho"] / 100 + 2 * medio, foto.height * m["alto"] / 100 + 2 * medio)
                real = (marco.left, marco.top, marco.width, marco.height)
                if max(abs(a - b) for a, b in zip(esperado, real)) > 2000:          # 2000 EMU = 0,06 mm
                    problemas.append("pagina %s: un recuadro rojo de la captura %s no cae sobre lo que marca" % (pag["numero"], num))
        if pag.get("tabla"):
            tablas = [sh for sh in s.shapes if getattr(sh, "has_table", False) and sh.has_table]
            if len(tablas) != 1:
                problemas.append("pagina %s: la tabla no es una tabla de PowerPoint" % pag["numero"])
            else:
                t = tablas[0].table
                esperada = [pag["tabla"]["encabezado"]] + pag["tabla"]["filas"]
                real = [[c.text_frame.text for c in fila.cells] for fila in t.rows]
                if real != esperada or not t.first_row:
                    problemas.append("pagina %s: la tabla no dice lo mismo que contenido.json" % pag["numero"])
    if con_notas:
        for i, (s, pag) in enumerate(zip(diapos, [None] + activas + [None]), 1):
            debe = pag is None or pag.get("error_que_evita")
            if debe and not (s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip()):
                problemas.append("diapositiva %d: falta la nota del orador" % i)
    return problemas


def prueba_completa():
    fallas = []

    def control(nombre, bien, detalle=""):
        print("  [%s] %s%s" % ("OK " if bien else "MAL", nombre, (" — " + detalle) if detalle and not bien else ""))
        if not bien:
            fallas.append(nombre)

    original = json.loads((AQUI / "contenido.json").read_text(encoding="utf-8"))
    with tempfile.TemporaryDirectory(prefix="manual_prueba_") as tmp:
        tmp = Path(tmp)
        caps, sal = tmp / "capturas", tmp / "salida"
        confirmado = copy.deepcopy(original)
        for c in todas_las_capturas(confirmado):
            for m in c.get("marcas", []):
                m["posicion_confirmada"] = True
        ok_json = tmp / "confirmado.json"
        ok_json.write_text(json.dumps(confirmado, ensure_ascii=False), encoding="utf-8")
        for c in todas_las_capturas(original):
            captura_de_prueba(caps / Path(c["archivo"]).name, c)
        base = ["--capturas", str(caps), "--salida", str(sal)]
        n_pag = len([p for p in original["paginas"] if p.get("activa", True)])
        nombre = original["manual"]["archivo"]

        print("1. Con todas las capturas y las marcas confirmadas")
        cod, out = correr(base + ["--contenido", str(ok_json)])
        control("sale con 0", cod == 0, "salio con %s\n%s" % (cod, out))
        pdf, pptx = sal / (nombre + ".pdf"), sal / (nombre + ".pptx")
        control("deja el PDF y el PowerPoint", pdf.exists() and pptx.exists())
        hojas_pdf = len(re.findall(rb"/Type\s*/Page\b(?!s)", pdf.read_bytes())) if pdf.exists() else 0
        control("el PDF tiene %d paginas" % n_pag, hojas_pdf == n_pag, "tiene %d" % hojas_pdf)
        prs = Presentation(str(pptx))
        control("el PowerPoint tiene %d diapositivas (portada + %d paginas + cierre)" % (n_pag + 2, n_pag),
                len(prs.slides) == n_pag + 2, "tiene %d" % len(prs.slides))
        control("el PowerPoint es A4 apaisado", abs(prs.slide_width - g.Mm(297)) < 100 and abs(prs.slide_height - g.Mm(210)) < 100)
        paginas_ppt = list(prs.slides)[1:-1]            # las del manual: sin la portada ni el cierre
        tam = [[r.font.size for sh in formas(s.shapes) if sh.has_text_frame for p in sh.text_frame.paragraphs
                for r in p.runs if r.font.size] for s in prs.slides]
        control("cada diapositiva tiene el titulo en %d pt o mas" % 34, all(max(t) >= Pt(34) for t in tam))
        control("cada pagina tiene texto de %d pt" % g.TEXTO_PT, all(Pt(g.TEXTO_PT) in t for t in tam[1:-1]))
        control("la letra del texto es de 20 pt o mas y la del titulo de 34 o mas", g.TEXTO_PT >= 20 and g.TITULO_PT >= 34)
        fotos = [sum(1 for sh in formas(s.shapes) if sh.shape_type == 13) for s in paginas_ppt]
        esperadas = [len(p.get("capturas", [])) for p in original["paginas"] if p.get("activa", True)]
        control("cada diapositiva tiene sus capturas", fotos == esperadas, "%s contra %s" % (fotos, esperadas))

        print("1 bis. El PowerPoint para proyectar: titulo de verdad, texto alternativo, idioma, fundido, tabla, marcos")
        datos_ok = g.cargar(ok_json)
        problemas = revisar_pptx(prs, datos_ok, caps, con_notas=False)
        control("el PowerPoint pasa todos los controles", not problemas, "; ".join(problemas[:6]))
        # los casos MALOS: cada control tiene que poder dar mal
        malo = Presentation(str(pptx))
        malo.slides[3].shapes.title.text_frame.text = ""
        control("un titulo vacio se detecta", any("titulo" in p for p in revisar_pptx(malo, datos_ok, caps, False)))
        malo = Presentation(str(pptx))
        for sh in formas(malo.slides[2].shapes):
            if sh.shape_type == 13:
                sh._element.xpath("./p:nvPicPr/p:cNvPr")[0].set("descr", "")
        control("una captura sin texto alternativo se detecta",
                any("alternativo" in p for p in revisar_pptx(malo, datos_ok, caps, False)))
        malo = Presentation(str(pptx))
        malo.slides[4]._element.remove(malo.slides[4]._element.xpath("./p:transition")[0])
        control("una diapositiva sin fundido se detecta", any("fundido" in p for p in revisar_pptx(malo, datos_ok, caps, False)))
        malo = Presentation(str(pptx))
        marco = next(sh for sh in formas(malo.slides[5].shapes) if sh.name == "Recuadro rojo")
        marco.left = marco.left + g.Mm(0.5)
        control("un recuadro rojo corrido 0,5 mm se detecta",
                any("recuadro" in p for p in revisar_pptx(malo, datos_ok, caps, False)))
        malo = Presentation(str(pptx))
        for sh in formas(malo.slides[2].shapes):
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.language_id = None
        control("texto sin idioma se detecta", any("idioma" in p for p in revisar_pptx(malo, datos_ok, caps, False)))

        print("1 ter. La copia para mostrar: sin la franja 'a confirmar' y con las notas del orador")
        cod, out = correr(base + ["--contenido", str(ok_json), "--sin-avisos"])
        mostrar = sal / (nombre + " - para mostrar.pptx")
        control("deja la copia para mostrar", mostrar.exists(), "salio con %s\n%s" % (cod, out[:300]))
        if mostrar.exists():
            prs_m = Presentation(str(mostrar))
            control("la copia para mostrar no dice 'A confirmar'", "A confirmar" not in todo_el_texto(prs_m))
            control("el PowerPoint normal si dice 'A confirmar'", "A confirmar" in todo_el_texto(prs))
            problemas_m = revisar_pptx(prs_m, datos_ok, caps, con_notas=True)
            control("la copia para mostrar lleva las notas del orador en cada diapositiva", not problemas_m,
                    "; ".join(problemas_m[:6]))
            control("el PowerPoint normal no lleva notas", not any(s.has_notes_slide for s in prs.slides))

        print("2. El recuadro rojo cae sobre el boton verde y la captura no se deforma (pagina 1)")

        def medir(datos_medir):
            """(lo que le erra el recuadro al boton, en pixeles; proporcion ancho/alto de la captura en la hoja)"""
            hoja = g.armar_pdf(datos_medir, caps)[0]
            verde, rojo = caja_de_color(hoja, VERDE), caja_de_color(hoja, g.ROJO)
            # la ventana de mentira entera = su barra de arriba (todo el ancho) + su fondo (hasta abajo)
            barra, fondo = caja_de_color(hoja, (52, 58, 70)), caja_de_color(hoja, (246, 247, 249))
            if not (verde and rojo and barra and fondo):
                return None, None
            gpx = g.GROSOR_MARCA * g.PX_MM
            error = max(abs(v) for v in (verde[0] - rojo[0] - gpx, verde[1] - rojo[1] - gpx,
                                         rojo[2] - verde[2] - gpx, rojo[3] - verde[3] - gpx))
            return error, (barra[2] - barra[0]) / (fondo[3] - barra[1])

        datos = g.cargar(ok_json)
        error, proporcion = medir(datos)
        control("se ven el boton verde, el recuadro rojo y la captura entera", error is not None)
        if error is not None:
            # 6 pixeles = 0,6 mm: es lo que se desdibuja el borde del boton al agrandar la captura
            control("el rojo abraza al verde, pegado por los cuatro lados", error <= 6, "le erra por %.1f pixeles" % error)
            control("la captura conserva su proporcion (1920 x 1080)", abs(proporcion / (1920 / 1080) - 1) < 0.01,
                    "%.4f contra %.4f" % (proporcion, 1920 / 1080))
            # los dos casos MALOS: este control tiene que poder dar mal
            corrido = copy.deepcopy(datos)
            corrido["paginas"][0]["capturas"][0]["marcas"][0]["x"] += 2
            error_malo, _ = medir(corrido)
            control("un recuadro corrido 2 % se detecta", error_malo is not None and error_malo > 6,
                    "le erra por %s pixeles y no salto" % error_malo)
            encajar_bueno = g.encajar
            g.encajar = lambda iw, ih, x, y, w, h, arriba=False: (x, y, w, h)      # estira la captura a toda la caja
            try:
                _, proporcion_mala = medir(datos)
            finally:
                g.encajar = encajar_bueno
            control("una captura estirada se detecta", proporcion_mala is not None
                    and abs(proporcion_mala / (1920 / 1080) - 1) >= 0.01, "proporcion %s y no salto" % proporcion_mala)

        print("3. Si falta una captura")
        (caps / "03-boton-code.png").rename(tmp / "03.png")
        cod, out = correr(base + ["--contenido", str(ok_json)])
        control("sale con 1", cod == 1, "salio con %s" % cod)
        control("dice cual falta", "03-boton-code.png" in out and "FALTAN 1 CAPTURAS" in out)
        control("igual deja los archivos (con el cartel gris)", pdf.exists() and pptx.exists())
        (tmp / "03.png").rename(caps / "03-boton-code.png")

        print("4. Si la captura opcional no esta")
        (caps / "15-cartel-limite.png").rename(tmp / "15.png")
        cod, out = correr(base + ["--contenido", str(ok_json)])
        control("sale con 0 y avisa", cod == 0 and "OPCIONALES" in out, "salio con %s" % cod)
        (tmp / "15.png").rename(caps / "15-cartel-limite.png")

        print("5. Con las capturas puestas pero las marcas en posicion estimada")
        cod, out = correr(base)
        control("sale con 1 y lo dice", cod == 1 and "POSICION ESTIMADA" in out, "salio con %s" % cod)

        print("6. Contenido que no entra o esta mal escrito")
        for nombre_caso, cambio in (
            ("texto de 5 renglones", lambda d: d["paginas"][0]["texto"].extend(["Uno más.", "Y otro más."])),
            ("renglon larguisimo", lambda d: d["paginas"][0].__setitem__("texto", ["palabra " * 60, "Dos."])),
            ("titulo de 7 palabras", lambda d: d["paginas"][0].__setitem__("titulo", "Uno dos tres cuatro cinco seis siete")),
            ("marca fuera de la imagen", lambda d: d["paginas"][0]["capturas"][0]["marcas"][0].__setitem__("x", 95)),
            ("archivo mal nombrado", lambda d: d["paginas"][0]["capturas"][0].__setitem__("archivo", "capturas/foto.png")),
            ("portada sin titulo", lambda d: d["portada"].__setitem__("titulo", "")),
            ("cierre con un solo paso", lambda d: d["cierre"].__setitem__("pasos", ["Abrí Claude"])),
            ("paso del cierre que no entra", lambda d: d["cierre"]["pasos"].__setitem__(0, "palabra " * 40)),
            ("15 paginas", lambda d: d["paginas"].append(dict(copy.deepcopy(d["paginas"][0]), numero=15,
                                                              capturas=[dict(d["paginas"][0]["capturas"][0],
                                                                             archivo="capturas/21-otra.png")]))),
        ):
            malo = copy.deepcopy(confirmado)
            cambio(malo)
            (tmp / "malo.json").write_text(json.dumps(malo, ensure_ascii=False), encoding="utf-8")
            cod, out = correr(base + ["--contenido", str(tmp / "malo.json")])
            control("%s: sale con 2" % nombre_caso, cod == 2, "salio con %s\n%s" % (cod, out[:300]))

        print("7. Una pagina apagada (\"activa\": false) sale del manual y no pide su captura")
        sin13 = copy.deepcopy(confirmado)
        [p for p in sin13["paginas"] if p["numero"] == 13][0]["activa"] = False
        (tmp / "sin13.json").write_text(json.dumps(sin13, ensure_ascii=False), encoding="utf-8")
        (caps / "14-claude-pide-permiso.png").unlink()
        cod, out = correr(base + ["--contenido", str(tmp / "sin13.json")])
        control("sale con 0 y con una pagina menos", cod == 0 and "%d paginas" % (n_pag - 1) in out, "salio con %s" % cod)

        print("8. La cuadricula para ubicar las marcas")
        cod, out = correr(base + ["--grilla"])
        control("deja una imagen por captura", cod == 0 and len(list((caps / "_grilla").glob("*.png"))) == 15)

    print()
    if fallas:
        print("PRUEBA: %d control(es) MAL: %s" % (len(fallas), "; ".join(fallas)))
        return 1
    print("PRUEBA: todo bien.")
    return 0


def main():
    for flujo in (sys.stdout, sys.stderr):
        try:
            flujo.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass
    if len(sys.argv) > 1 and sys.argv[1] == "--de-prueba":
        datos = json.loads((AQUI / "contenido.json").read_text(encoding="utf-8"))
        pedidas = sys.argv[2:]
        for c in todas_las_capturas(datos):
            if Path(c["archivo"]).name[:2] in pedidas:
                destino = AQUI / c["archivo"]
                if destino.exists():
                    print("YA EXISTE, no la piso: %s" % destino)
                    continue
                captura_de_prueba(destino, c)
                print("captura DE MENTIRA: %s" % destino)
        return 0
    return prueba_completa()


if __name__ == "__main__":
    sys.exit(main())
