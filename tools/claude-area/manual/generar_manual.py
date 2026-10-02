#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Arma el manual de arranque de Claude en Barack: el PDF y el mismo en PowerPoint.

Lee  contenido.json  (las paginas, como datos)  +  capturas/  (las capturas de pantalla)
Deja exports/CLAUDES_POR_AREA_20261001/Manual de arranque - Claude Barack.pdf  y  .pptx

Uso:
    python generar_manual.py              arma los dos archivos
    python generar_manual.py --grilla     deja en capturas/_grilla/ cada captura con una cuadricula
                                          de 10 en 10 y las marcas actuales, para ubicar los recuadros
    python generar_manual.py --notas      ademas pone "para que esta la pagina" en las notas del PowerPoint
                                          (la copia --sin-avisos las lleva siempre)

Sale con:
    0  salio completo
    1  falta alguna captura, o hay una captura puesta cuya marca roja todavia esta en posicion estimada
    2  contenido.json tiene algo que no entra o esta mal escrito (lo dice en pantalla)

Solo usa python-pptx y Pillow. Las medidas de la hoja estan en milimetros (A4 apaisado) y se usan
igual para el PDF y para el PowerPoint: por eso las dos salidas son la misma pagina.

El PowerPoint lleva ademas (el PDF no): una portada con el logo y una diapositiva de cierre con los tres
pasos. Es un PowerPoint hecho para proyectar o recorrer en la PC: el titulo de cada diapositiva es el titulo
de verdad, cada captura tiene su texto alternativo (el "que_se_ve" de contenido.json), el idioma esta
puesto en castellano de Argentina, hay un fundido corto entre diapositivas y la tabla es una tabla de
PowerPoint. La letra es Segoe UI (viene con Windows): en una compu sin esa letra PowerPoint pone otra.
"""
import argparse
import contextlib
import datetime
import json
import math
import os
import re
import sys
from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.lang import MSO_LANGUAGE_ID
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Mm, Pt

AQUI = Path(__file__).resolve().parent
REPO = AQUI.parents[2]
SALIDA = REPO / "exports" / "CLAUDES_POR_AREA_20261001"
DIR_FUENTES = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"

# ------------------------------------------------------------------ medidas de la hoja (mm)
HOJA_W, HOJA_H = 297.0, 210.0          # A4 apaisado
DPI = 240                              # una captura de 1920 px entra sin achicarse
PX_MM = DPI / 25.4
MM_PT = 25.4 / 72.0                    # 1 punto en mm
MARGEN = 12.0
ANCHO = HOJA_W - 2 * MARGEN            # ancho util
BANDA_H = 26.0                         # franja azul del titulo
Y0 = BANDA_H + 6.0                     # donde empieza la imagen
TEXTO_FONDO = 196.0                    # hasta donde llega el texto de abajo
CHIP_FONDO = 197.5                     # hasta donde llega la franja "a confirmar"
PIE_Y = 200.5
SEP = 6.0                              # separacion entre la captura grande y la del costado
SANGRIA = 10.0                         # lugar de la viñeta
AIRE_IMAGEN_TEXTO = 5.0
ALTO_MINIMO_IMAGEN = 90.0              # si el texto deja menos que esto para la imagen, es un error

# ------------------------------------------------------------------ letra (puntos, a tamaño real)
TITULO_PT = 38
TEXTO_PT = 22
INTERLINEA = 1.3
ENTRE_PARRAFOS = 1.6                   # mm
TEXTO_LH = TEXTO_PT * INTERLINEA * MM_PT
CHIP_PT = 11
CHIP_LH = CHIP_PT * 1.25 * MM_PT
ROTULO_PT = 16
ROTULO_H = 8.5
TARJETA_PT = 20
TABLA_PT = 16
MAX_RENGLONES = 4                      # renglones IMPRESOS de texto por pagina
MAX_PAGINAS = 15
MAX_PALABRAS_TITULO = 6

assert TEXTO_PT >= 20 and TITULO_PT >= 34, "la letra no puede ser mas chica que 20 / 34 puntos"

# ------------------------------------------------------------------ colores
BLANCO = (255, 255, 255)
AZUL_OSC = (0x1F, 0x3A, 0x5F)
AZUL = (0x2F, 0x6F, 0xB0)
CELESTE = (0xEA, 0xF2, 0xFB)
TINTA = (0x1B, 0x27, 0x33)
GRIS = (0x6B, 0x77, 0x83)
GRIS_BORDE = (0xC9, 0xD2, 0xDC)
GRIS_FONDO = (0xE9, 0xED, 0xF1)
GRIS_FALTA = (0x4F, 0x5B, 0x67)
ROJO = (0xE1, 0x25, 0x1B)
AMBAR = (0x7A, 0x52, 0x00)
AMBAR_FONDO = (0xFF, 0xF0, 0xCC)
GROSOR_MARCA = 1.5                     # mm: se tiene que ver impreso
HALO = 0.6                             # borde blanco alrededor del rojo (por si la app es oscura)

# ------------------------------------------------------------------ PowerPoint
LOGO = REPO / "scripts" / "video" / "logo_hd.png"
FUENTE_PPT = "Segoe UI"                # viene con Windows; el PDF usa la misma
IDIOMA = "es-AR"
IDIOMA_ID = MSO_LANGUAGE_ID.SPANISH_ARGENTINA
# En un renglon de interlineado EXACTO, PowerPoint apoya la letra a esta fraccion del alto del renglon
# (medido con PowerPoint 16 sobre Segoe UI: 0,77 a 0,78). El PDF la apoya donde dicen las metricas de la letra.
# La diferencia (de 0,1 a 2 mm segun el tamaño) se corrige bajando el cuadro de texto: ver ajuste_ppt().
BASE_RENGLON_PPT = 0.767
PAD_TABLA = 2.2                        # mm de aire adentro de cada celda de la tabla
HOLGURA_PPT = 0.4                      # mm de mas de ancho en cada cuadro de texto de PowerPoint (ver LienzoPPT.texto)
TRANSICION = "fast"                    # fundido corto (0,5 s): lo mismo entre todas las diapositivas
VINETA_PUNTO_PCT = 54                  # tamaño del punto de las viñetas, en % de la letra (da 2,9 mm, como en el PDF)
VINETA_NUMERO_PCT = 120                # tamaño del circulo con numero, en % de la letra (6,8 mm)


class ErrorDeContenido(Exception):
    pass


# ================================================================== texto: medir y partir
_FUENTES = {}


def fuente(pt, negrita=False):
    clave = (pt, bool(negrita))
    if clave not in _FUENTES:
        archivo = DIR_FUENTES / ("segoeuib.ttf" if negrita else "segoeui.ttf")
        _FUENTES[clave] = ImageFont.truetype(str(archivo), max(1, round(pt * DPI / 72)))
    return _FUENTES[clave]


def ancho_txt(t, pt, negrita=False):
    return fuente(pt, negrita).getlength(t) / PX_MM


def trozos(texto):
    """'Elegí **Local** ya' -> [('Elegí ', False), ('Local', True), (' ya', False)]"""
    return [(t, i % 2 == 1) for i, t in enumerate(texto.split("**")) if t]


def separar_numero(parrafo):
    """'2. Apretá...' -> (2, 'Apretá...');  'Siempre esa' -> (None, 'Siempre esa')"""
    m = re.match(r"^(\d)\.\s+(.*)$", parrafo, re.S)
    return (int(m.group(1)), m.group(2)) if m else (None, parrafo)


def partir(texto, pt, ancho_mm, negrita=False):
    """Parte un parrafo en renglones que entren en ancho_mm.
    Devuelve renglones; cada renglon es una lista de palabras; cada palabra, una lista de (texto, negrita)."""
    palabras, actual = [], []
    for t, neg in trozos(texto):
        for i, pedazo in enumerate(t.split(" ")):
            if i > 0 and actual:
                palabras.append(actual)
                actual = []
            if pedazo:
                actual.append((pedazo, neg or negrita))
    if actual:
        palabras.append(actual)
    espacio = ancho_txt(" ", pt, negrita)
    renglones, linea, usado = [], [], 0.0
    for pal in palabras:
        w = sum(ancho_txt(t, pt, n) for t, n in pal)
        if linea and usado + espacio + w > ancho_mm:
            renglones.append(linea)
            linea, usado = [], 0.0
        if linea:
            usado += espacio
        linea.append(pal)
        usado += w
    if linea:
        renglones.append(linea)
    return renglones


def ancho_renglon(renglon, pt, negrita=False):
    espacio = ancho_txt(" ", pt, negrita)
    return sum(ancho_txt(t, pt, n) for pal in renglon for t, n in pal) + espacio * (len(renglon) - 1)


def alto_parrafos(parrafos, pt, ancho_mm, negrita=False, interlinea=INTERLINEA, entre=0.0, vinetas=False):
    w = ancho_mm - (SANGRIA if vinetas else 0.0)
    n = sum(len(partir(separar_numero(p)[1] if vinetas else p, pt, w, negrita)) for p in parrafos)
    return n * pt * interlinea * MM_PT + entre * (len(parrafos) - 1), n


def interlineado_ppt(pt, interlinea):
    """Interlineado exacto de PowerPoint, en puntos ENTEROS: PowerPoint redondea el interlineado y el espacio
    entre parrafos al punto entero (28,6 pt se dibuja como 29), asi que se pide ya redondeado."""
    return max(1, round(pt * interlinea))


def ajuste_ppt(pt, interlinea):
    """Cuantos mm hay que BAJAR un cuadro de texto de PowerPoint para que la letra quede donde la apoya el PDF.
    PDF: la letra se apoya a (alto del renglon - alto de la letra) / 2 + ascendente, desde arriba del renglon.
    PowerPoint, con interlineado exacto: a BASE_RENGLON_PPT del alto del renglon."""
    lh = pt * interlinea * MM_PT
    asc, desc = fuente(pt, False).getmetrics()
    base_pdf = ((lh * PX_MM - (asc + desc)) / 2 + asc) / PX_MM
    return base_pdf - BASE_RENGLON_PPT * interlineado_ppt(pt, interlinea) * MM_PT


# ================================================================== lienzo PDF (Pillow)
class LienzoPDF:
    def __init__(self):
        self.im = Image.new("RGB", (round(HOJA_W * PX_MM), round(HOJA_H * PX_MM)), BLANCO)
        self.d = ImageDraw.Draw(self.im)

    @staticmethod
    def px(mm):
        return int(round(mm * PX_MM))

    def rect(self, x, y, w, h, relleno=None, borde=None, grosor=0.0, punteado=False):
        caja = [self.px(x), self.px(y), self.px(x + w) - 1, self.px(y + h) - 1]
        g = max(1, self.px(grosor)) if borde and grosor > 0 else 0
        if punteado and g:
            if relleno:
                self.d.rectangle(caja, fill=relleno)
            self._punteado(caja, borde, g)
        else:
            self.d.rectangle(caja, fill=relleno, outline=borde if g else None, width=max(1, g))

    def _punteado(self, caja, color, g, raya=7.0, hueco=4.0):
        x0, y0, x1, y1 = caja
        r, h = self.px(raya), self.px(hueco)
        for (ax, ay, bx, by) in ((x0, y0, x1, y0), (x0, y1, x1, y1), (x0, y0, x0, y1), (x1, y0, x1, y1)):
            largo = max(bx - ax, by - ay)
            p = 0
            while p < largo:
                q = min(p + r, largo)
                if ay == by:
                    self.d.line([ax + p, ay, ax + q, ay], fill=color, width=g)
                else:
                    self.d.line([ax, ay + p, ax, ay + q], fill=color, width=g)
                p += r + h

    def elipse(self, x, y, w, h, relleno=None, borde=None, grosor=0.0):
        g = max(1, self.px(grosor)) if borde and grosor > 0 else 0
        self.d.ellipse([self.px(x), self.px(y), self.px(x + w), self.px(y + h)],
                       fill=relleno, outline=borde if g else None, width=g)

    def banda(self):
        """Franja azul del titulo."""
        self.rect(0, 0, HOJA_W, BANDA_H, relleno=AZUL_OSC)

    def titulo(self, x, y, w, h, texto, pt, color):
        self.texto(x, y, w, h, [texto], pt, color, negrita=True, anclar="medio", interlinea=1.0)

    @contextlib.contextmanager
    def grupo(self, nombre):
        """En el PowerPoint junta varias formas en un grupo (se mueven juntas). En el PDF no hace nada."""
        yield

    def tabla(self, x, y, anchos, altos, filas, descr=None):
        """filas: lista de filas; cada celda es un dict con texto, fondo, color, negrita."""
        yy = y
        for fila, alto in zip(filas, altos):
            xx = x
            for c, a in zip(fila, anchos):
                self.rect(xx, yy, a, alto, relleno=c["fondo"], borde=GRIS_BORDE, grosor=0.25)
                self.texto(xx + PAD_TABLA, yy, a - 2 * PAD_TABLA, alto, [c["texto"]], TABLA_PT, c["color"],
                           negrita=c["negrita"], anclar="medio", interlinea=1.22)
                xx += a
            yy += alto

    def texto(self, x, y, w, h, parrafos, pt, color=TINTA, negrita=False, alinear="izq",
              anclar="arriba", interlinea=INTERLINEA, entre=0.0, vinetas=False, deco=False, nombre=None):
        sangria = SANGRIA if vinetas else 0.0
        lh = pt * interlinea * MM_PT
        bloques = []
        for par in parrafos:
            num, cuerpo = separar_numero(par) if vinetas else (None, par)
            bloques.append((num, partir(cuerpo, pt, w - sangria, negrita)))
        alto = sum(len(r) for _, r in bloques) * lh + entre * (len(bloques) - 1)
        yy = y + {"arriba": 0.0, "medio": (h - alto) / 2, "abajo": h - alto}[anclar]
        asc, desc = fuente(pt, False).getmetrics()
        espacio = ancho_txt(" ", pt, negrita)
        for num, renglones in bloques:
            if vinetas:
                self._vineta(x, yy, lh, num)
            for reng in renglones:
                libre = w - sangria - ancho_renglon(reng, pt, negrita)
                xx = x + sangria + {"izq": 0.0, "centro": libre / 2, "der": libre}[alinear]
                base = yy * PX_MM + (lh * PX_MM - (asc + desc)) / 2 + asc
                for i, pal in enumerate(reng):
                    if i:
                        xx += espacio
                    for t, neg in pal:
                        f = fuente(pt, neg)
                        self.d.text((xx * PX_MM, base), t, font=f, fill=color, anchor="ls")
                        xx += f.getlength(t) / PX_MM
                yy += lh
            yy += entre

    def _vineta(self, x, y, lh, num):
        cy = y + lh / 2 + 0.3
        if num is None:
            d = 2.8
            self.elipse(x + 2.2, cy - d / 2, d, d, relleno=AZUL)
        else:
            self.insignia(x + 3.6, cy, num, d=7.4, pt=15, halo=False)

    def insignia(self, cx, cy, num, d=9.0, pt=18, halo=True):
        if halo:
            self.elipse(cx - d / 2 - HALO, cy - d / 2 - HALO, d + 2 * HALO, d + 2 * HALO, relleno=BLANCO)
        self.elipse(cx - d / 2, cy - d / 2, d, d, relleno=ROJO)
        self.d.text((cx * PX_MM, cy * PX_MM), str(num), font=fuente(pt, True), fill=BLANCO, anchor="mm")

    def imagen(self, ruta, x, y, w, h, descr=None, nombre=None):
        with Image.open(ruta) as im:
            if im.mode in ("RGBA", "LA", "P"):
                im = im.convert("RGBA")
                fondo = Image.new("RGB", im.size, BLANCO)
                fondo.paste(im, mask=im.split()[-1])
                im = fondo
            else:
                im = im.convert("RGB")
            im = im.resize((max(1, self.px(w)), max(1, self.px(h))), Image.LANCZOS)
        self.im.paste(im, (self.px(x), self.px(y)))

    def marco(self, x, y, w, h):
        """Recuadro rojo: el borde va POR FUERA, para no tapar lo que marca."""
        # Pillow dibuja el borde hacia ADENTRO de la caja: se agranda la caja y el borde termina justo en la marca.
        # Primero uno blanco un poco mas ancho (se ve aunque la app sea oscura), encima el rojo.
        for extra, color in ((GROSOR_MARCA + HALO, BLANCO), (GROSOR_MARCA, ROJO)):
            self.d.rectangle([self.px(x - extra), self.px(y - extra), self.px(x + w + extra), self.px(y + h + extra)],
                             outline=color, width=max(1, self.px(extra)))

    def flecha(self, x1, y1, x2, y2):
        ang = math.atan2(y2 - y1, x2 - x1)
        largo_punta, ancho_punta, grosor = 7.5, 6.5, GROSOR_MARCA + 0.3
        bx, by = x2 - largo_punta * math.cos(ang), y2 - largo_punta * math.sin(ang)
        nx, ny = -math.sin(ang), math.cos(ang)
        for extra, color in ((HALO, BLANCO), (0.0, ROJO)):
            a = ancho_punta / 2 + extra
            punta = [(x2 + extra * 1.6 * math.cos(ang), y2 + extra * 1.6 * math.sin(ang)),
                     (bx - extra * math.cos(ang) + nx * a, by - extra * math.sin(ang) + ny * a),
                     (bx - extra * math.cos(ang) - nx * a, by - extra * math.sin(ang) - ny * a)]
            self.d.line([(x1 * PX_MM, y1 * PX_MM), (bx * PX_MM, by * PX_MM)], fill=color,
                        width=self.px(grosor + 2 * extra))
            self.d.polygon([(px_ * PX_MM, py_ * PX_MM) for px_, py_ in punta], fill=color)


# ================================================================== lienzo PowerPoint (python-pptx)
NS_DECORATIVO = "http://schemas.microsoft.com/office/drawing/2017/decorative"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
ESTILO_SIN_LINEAS = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"      # estilo de tabla "sin estilo, sin cuadricula"


def _rgb(c):
    return RGBColor(*c)


def _hex(c):
    return "%02X%02X%02X" % tuple(c)


def _cnvpr(sh):
    """El elemento con el nombre y el texto alternativo de cualquier forma."""
    return sh._element.xpath("./*[1]/p:cNvPr")[0]


def _decorativo(sh):
    """Lo marca 'decorativo' (igual que el boton de PowerPoint): el lector de pantalla no lo nombra."""
    ext = etree.SubElement(etree.SubElement(_cnvpr(sh), qn("a:extLst")), qn("a:ext"),
                           uri="{C183D7F6-B498-43B3-948B-1728B52AA6E4}")
    etree.SubElement(ext, "{%s}decorative" % NS_DECORATIVO, nsmap={"adec": NS_DECORATIVO}, val="1")


def _formato(r, pt, color, negrita):
    r.font.name = FUENTE_PPT
    r.font.size = Pt(pt)
    r.font.bold = bool(negrita)
    r.font.color.rgb = _rgb(color)
    r.font.language_id = IDIOMA_ID


def _mm(v):
    return str(int(Mm(v)))


class LienzoPPT:
    """Dibuja con formas nativas de PowerPoint: lo que se ve en la diapositiva se puede editar."""

    def __init__(self, prs, layout):
        self.s = prs.slides.add_slide(layout)
        self._cont = self.s.shapes          # donde caen las formas nuevas: la diapositiva o un grupo

    # ---- formas
    def _forma(self, tipo, x, y, w, h, relleno, borde, grosor, punteado=False, deco=True, nombre=None, adentro=True):
        # En PowerPoint la linea va CENTRADA sobre el contorno y en el PDF va hacia ADENTRO: se achica media
        # linea para que el borde ocupe el mismo lugar (los marcos rojos, que van por fuera, lo piden aparte).
        if adentro and borde and grosor > 0:
            x, y, w, h = x + grosor / 2, y + grosor / 2, w - grosor, h - grosor
        sh = self._cont.add_shape(tipo, Mm(x), Mm(y), Mm(w), Mm(h))
        sh.shadow.inherit = False
        if relleno:
            sh.fill.solid()
            sh.fill.fore_color.rgb = _rgb(relleno)
        else:
            sh.fill.background()
        if borde and grosor > 0:
            sh.line.color.rgb = _rgb(borde)
            sh.line.width = Mm(grosor)
            if punteado:
                sh.line.dash_style = MSO_LINE.DASH
        else:
            sh.line.fill.background()
        if nombre:
            sh.name = nombre
        if deco:
            _decorativo(sh)
        return sh

    def rect(self, x, y, w, h, relleno=None, borde=None, grosor=0.0, punteado=False):
        return self._forma(MSO_SHAPE.RECTANGLE, x, y, w, h, relleno, borde, grosor, punteado)

    def elipse(self, x, y, w, h, relleno=None, borde=None, grosor=0.0):
        return self._forma(MSO_SHAPE.OVAL, x, y, w, h, relleno, borde, grosor)

    @contextlib.contextmanager
    def grupo(self, nombre):
        """Junta en un grupo todo lo que se dibuje adentro: la captura y sus marcas se mueven juntas."""
        g = self._cont.add_group_shape()
        g.name = nombre
        antes, self._cont = self._cont, g.shapes
        try:
            yield
        finally:
            self._cont = antes

    # ---- titulo: es el titulo de verdad de la diapositiva (vista de esquema, lector de pantalla)
    def banda(self):
        """La franja azul vive en la plantilla (layout), no en cada diapositiva."""

    def titulo(self, x, y, w, h, texto, pt, color):
        ph = self.s.shapes.title
        ph.left, ph.top, ph.width, ph.height = Mm(x), Mm(y + ajuste_ppt(pt, 1.0)), Mm(w), Mm(h)
        poner_texto_placeholder(ph, [texto], pt, color, True, "izq", interlinea=1.0)

    # ---- texto
    def texto(self, x, y, w, h, parrafos, pt, color=TINTA, negrita=False, alinear="izq",
              anclar="arriba", interlinea=INTERLINEA, entre=0.0, vinetas=False, deco=False, nombre=None):
        # Un poco mas ancho que en el PDF (HOLGURA_PPT): si PowerPoint mide la letra apenas distinto en otra
        # compu, un renglon que entra justo no se corta en otro lado. Se reparte segun la alineacion.
        x -= {"izq": 0.0, "centro": HOLGURA_PPT / 2, "der": HOLGURA_PPT}[alinear]
        tb = self._cont.add_textbox(Mm(x), Mm(y + ajuste_ppt(pt, interlinea)), Mm(w + HOLGURA_PPT), Mm(h))
        tb.name = nombre or "Texto"
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"arriba": MSO_ANCHOR.TOP, "medio": MSO_ANCHOR.MIDDLE, "abajo": MSO_ANCHOR.BOTTOM}[anclar]
        for i, par in enumerate(parrafos):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"izq": PP_ALIGN.LEFT, "centro": PP_ALIGN.CENTER, "der": PP_ALIGN.RIGHT}[alinear]
            p.line_spacing = Pt(interlineado_ppt(pt, interlinea))
            if entre and i < len(parrafos) - 1:
                # el espacio entre parrafos se redondea para que el paso de un parrafo al otro quede como en el PDF
                p.space_after = Pt(max(0, round(pt * interlinea + entre / MM_PT - interlineado_ppt(pt, interlinea))))
            num, cuerpo = separar_numero(par) if vinetas else (None, par)
            if vinetas:
                self._vineta(p, num)
            for t, neg in trozos(cuerpo):
                r = p.add_run()
                r.text = t
                _formato(r, pt, color, neg or negrita)
        if deco:
            _decorativo(tb)
        return tb

    @staticmethod
    def _vineta(p, num):
        """Punto azul, o circulo rojo con el numero adentro (numeracion automatica de PowerPoint,
        'circleNumWdBlackPlain': los numeros siguen siendo una lista numerada para el lector de pantalla)."""
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", _mm(SANGRIA))
        # el punto arranca un poco adentro (centrado donde lo pone el PDF); el circulo arranca al ras
        pPr.set("indent", str(-int(Mm(SANGRIA - (2.0 if num is None else 0.0)))))

        def hijo(tag, **atr):
            return etree.SubElement(pPr, qn(tag), **atr)

        clr = hijo("a:buClr")
        etree.SubElement(clr, qn("a:srgbClr"), val=_hex(AZUL if num is None else ROJO))
        hijo("a:buSzPct", val=str((VINETA_PUNTO_PCT if num is None else VINETA_NUMERO_PCT) * 1000))
        hijo("a:buFont", typeface="Segoe UI Symbol")
        if num is None:
            hijo("a:buChar", char="●")
        else:
            hijo("a:buAutoNum", type="circleNumWdBlackPlain", startAt=str(num))

    # ---- marcas sobre las capturas
    def insignia(self, cx, cy, num, d=9.0, pt=18, halo=True):
        if halo:
            self._forma(MSO_SHAPE.OVAL, cx - d / 2 - HALO, cy - d / 2 - HALO, d + 2 * HALO, d + 2 * HALO,
                        BLANCO, None, 0.0)
        sh = self._forma(MSO_SHAPE.OVAL, cx - d / 2, cy - d / 2, d, d, ROJO, None, 0.0, deco=False,
                         nombre="Numero %s" % num)
        tf = sh.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.word_wrap = False
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        _formato(r, pt, BLANCO, True)
        r.text = str(num)

    def imagen(self, ruta, x, y, w, h, descr=None, nombre=None):
        pic = self._cont.add_picture(str(ruta), Mm(x), Mm(y), Mm(w), Mm(h))
        if nombre:
            pic.name = nombre
        if descr:
            _cnvpr(pic).set("descr", descr)
        else:
            _decorativo(pic)
        return pic

    def marco(self, x, y, w, h):
        """El rojo va POR FUERA de lo marcado (de x-g a x); el borde blanco, pegado por fuera del rojo."""
        g = GROSOR_MARCA
        bl = HALO + 0.2                       # un poco mas ancho, por debajo del rojo: sin rendija entre los dos
        d = g + HALO - bl / 2                 # del borde marcado al centro de la linea blanca
        self._forma(MSO_SHAPE.RECTANGLE, x - d, y - d, w + 2 * d, h + 2 * d, None, BLANCO, bl,
                    nombre="Recuadro rojo (borde blanco)", adentro=False)
        self._forma(MSO_SHAPE.RECTANGLE, x - g / 2, y - g / 2, w + g, h + g, None, ROJO, g,
                    nombre="Recuadro rojo", adentro=False)

    def flecha(self, x1, y1, x2, y2):
        for color, grosor, nombre in ((BLANCO, GROSOR_MARCA + 0.3 + 2 * HALO, "Flecha roja (borde blanco)"),
                                      (ROJO, GROSOR_MARCA + 0.3, "Flecha roja")):
            c = self._cont.add_connector(MSO_CONNECTOR.STRAIGHT, Mm(x1), Mm(y1), Mm(x2), Mm(y2))
            c.name = nombre
            c.shadow.inherit = False
            c.line.color.rgb = _rgb(color)
            c.line.width = Mm(grosor)
            ln = c.line._get_or_add_ln()
            ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
            _decorativo(c)

    # ---- tabla: una tabla de PowerPoint de verdad (se edita celda por celda, y el lector de pantalla la entiende)
    def tabla(self, x, y, anchos, altos, filas, descr=None):
        gf = self.s.shapes.add_table(len(filas), len(anchos), Mm(x), Mm(y), Mm(sum(anchos)), Mm(sum(altos)))
        gf.name = "Tabla"
        if descr:
            _cnvpr(gf).set("descr", descr)
        tbl = gf.table
        tbl.first_row = True                  # la primera fila es el encabezado
        tbl.horz_banding = False
        sid = tbl._tbl.tblPr.find(qn("a:tableStyleId"))
        if sid is None:
            sid = etree.SubElement(tbl._tbl.tblPr, qn("a:tableStyleId"))
        sid.text = ESTILO_SIN_LINEAS          # todo (relleno, bordes, letra) se define celda por celda
        for j, a in enumerate(anchos):
            tbl.columns[j].width = Mm(a)
        dy = ajuste_ppt(TABLA_PT, 1.22)
        reng_pdf = TABLA_PT * 1.22 * MM_PT
        reng_ppt = interlineado_ppt(TABLA_PT, 1.22) * MM_PT       # PowerPoint redondea el interlineado al punto
        for i, (fila, alto) in enumerate(zip(filas, altos)):
            tbl.rows[i].height = Mm(alto)
            lineas = max(1, round((alto - 2 * PAD_TABLA) / reng_pdf))
            aire = max(0.5, alto - lineas * reng_ppt)             # lo que sobra arriba y abajo para que la fila no crezca
            for j, c in enumerate(fila):
                celda = tbl.cell(i, j)
                tf = celda.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.line_spacing = Pt(interlineado_ppt(TABLA_PT, 1.22))
                r = p.add_run()
                _formato(r, TABLA_PT, c["color"], c["negrita"])
                r.text = c["texto"]
                tcPr = celda._tc.get_or_add_tcPr()
                tcPr.set("marL", _mm(PAD_TABLA))
                tcPr.set("marR", _mm(PAD_TABLA - HOLGURA_PPT))        # mas lugar para el texto, como en HOLGURA_PPT
                tcPr.set("marT", _mm(max(0.0, aire / 2 + dy)))    # el texto, centrado, baja dy: queda donde lo pone el PDF
                tcPr.set("marB", _mm(max(0.0, aire / 2 - dy)))
                tcPr.set("anchor", "ctr")
                for lado in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):                 # primero los bordes...
                    ln = etree.SubElement(tcPr, qn(lado), w=_mm(0.25), cap="flat", cmpd="sng", algn="ctr")
                    etree.SubElement(etree.SubElement(ln, qn("a:solidFill")), qn("a:srgbClr"), val=_hex(GRIS_BORDE))
                    etree.SubElement(ln, qn("a:prstDash"), val="solid")
                etree.SubElement(etree.SubElement(tcPr, qn("a:solidFill")), qn("a:srgbClr"), val=_hex(c["fondo"]))  # ...despues el relleno
        return gf


def poner_texto_placeholder(ph, parrafos, pt, color, negrita, alinear, interlinea=1.0):
    """Escribe en un cuadro de titulo/subtitulo de la plantilla (la posicion y la letra las trae la plantilla)."""
    tf = ph.text_frame
    for i, par in enumerate(parrafos):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"izq": PP_ALIGN.LEFT, "centro": PP_ALIGN.CENTER, "der": PP_ALIGN.RIGHT}[alinear]
        p.line_spacing = Pt(interlineado_ppt(pt, interlinea))
        for t, neg in trozos(par):
            r = p.add_run()
            _formato(r, pt, color, neg or negrita)
            r.text = t


# ------------------------------------------------------------------ la plantilla del PowerPoint
PORTADA_BLOQUE_Y = 92.0                 # donde empieza el bloque azul oscuro de la portada
PORTADA_TITULO_PT = 60
PORTADA_SUBTITULO_PT = 34
PORTADA_DATOS_PT = 22
CIERRE_PASO_PT = 28
CIERRE_AVISO_PT = 26


def _forma_de_plantilla(spTree, ident, nombre, x, y, w, h, color):
    """Rectangulo liso en la plantilla (layout), debajo de todo. Va marcado como decorativo."""
    sp = etree.fromstring(
        '<p:sp xmlns:p="%s" xmlns:a="%s" xmlns:adec="%s"><p:nvSpPr>'
        '<p:cNvPr id="%d" name="%s"><a:extLst><a:ext uri="{C183D7F6-B498-43B3-948B-1728B52AA6E4}">'
        '<adec:decorative val="1"/></a:ext></a:extLst></p:cNvPr><p:cNvSpPr/><p:nvPr userDrawn="1"/></p:nvSpPr>'
        '<p:spPr><a:xfrm><a:off x="%s" y="%s"/><a:ext cx="%s" cy="%s"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="%s"/></a:solidFill>'
        '<a:ln><a:noFill/></a:ln></p:spPr></p:sp>'
        % (NS_P, NS_A, NS_DECORATIVO, ident, nombre, _mm(x), _mm(y), _mm(w), _mm(h), _hex(color)))
    spTree.insert(2, sp)                  # despues de nvGrpSpPr y grpSpPr: queda al fondo


def _estilo_de_placeholder(ph, pt, color, negrita, alinear, anclar):
    """Letra y alineacion del cuadro de la plantilla: lo que escriba alguien despues sale igual."""
    tf = ph.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"arriba": MSO_ANCHOR.TOP, "medio": MSO_ANCHOR.MIDDLE, "abajo": MSO_ANCHOR.BOTTOM}[anclar]
    txBody = ph._element.txBody
    lst = txBody.find(qn("a:lstStyle"))
    if lst is None:
        lst = etree.Element(qn("a:lstStyle"))
        txBody.find(qn("a:bodyPr")).addnext(lst)
    for hijo in list(lst):
        lst.remove(hijo)
    nivel = etree.SubElement(lst, qn("a:lvl1pPr"), marL="0", indent="0", algn={"izq": "l", "centro": "ctr"}[alinear])
    etree.SubElement(nivel, qn("a:buNone"))
    d = etree.SubElement(nivel, qn("a:defRPr"), sz=str(int(pt * 100)), b="1" if negrita else "0", lang=IDIOMA)
    etree.SubElement(etree.SubElement(d, qn("a:solidFill")), qn("a:srgbClr"), val=_hex(color))
    etree.SubElement(d, qn("a:latin"), typeface=FUENTE_PPT)


def preparar_plantilla(prs, ancho_orig, alto_orig):
    """La plantilla de python-pptx es 4:3 y de Calibri. Se deja a medida: A4 apaisado, letra Segoe UI y solo
    las dos plantillas del manual (portada y pagina con franja azul). Devuelve (layout_portada, layout_pagina)."""
    sx, sy = prs.slide_width / ancho_orig, prs.slide_height / alto_orig
    maestra = prs.slide_master
    for sh in maestra.shapes:
        sh.left, sh.top = int(sh.left * sx), int(sh.top * sy)
        sh.width, sh.height = int(sh.width * sx), int(sh.height * sy)
    # letra del tema: lo que se agregue despues en PowerPoint sale en Segoe UI
    tema = maestra.part.part_related_by(
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme")
    tema._blob = re.sub(rb'(<a:latin typeface=")Calibri(")', rb"\1" + FUENTE_PPT.encode() + rb"\2", tema.blob)
    tema._blob = tema._blob.replace(b'name="Office Theme"', b'name="Manual Barack"')
    prs._element.find(qn("p:sldSz")).attrib.pop("type", None)    # la plantilla decia 4:3; 297 x 210 mm es "personalizado" para PowerPoint
    portada = next(l for l in prs.slide_layouts if l.name == "Title Slide")
    pagina = next(l for l in prs.slide_layouts if l.name == "Title Only")
    for lay in list(prs.slide_layouts):
        if lay.name not in ("Title Slide", "Title Only"):
            prs.slide_layouts.remove(lay)
    for lay, nombre in ((portada, "Manual - portada"), (pagina, "Manual - pagina")):
        lay._element.cSld.set("name", nombre)
        for ph in list(lay.placeholders):           # fecha, pie y numero de diapositiva no se usan
            if ph.placeholder_format.type in (PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER):
                ph._element.getparent().remove(ph._element)
        for sh in lay.shapes:
            if not sh.is_placeholder:
                sh._element.getparent().remove(sh._element)

    # pagina: franja azul arriba + el titulo adentro de la franja
    _forma_de_plantilla(pagina.shapes._spTree, 101, "Franja del titulo", 0, 0, HOJA_W, BANDA_H, AZUL_OSC)
    d = 17.0
    titulo = next(p for p in pagina.placeholders if p.placeholder_format.type == PP_PLACEHOLDER.TITLE)
    titulo.left, titulo.top = Mm(MARGEN + d + 6), Mm(ajuste_ppt(TITULO_PT, 1.0))
    titulo.width, titulo.height = Mm(ANCHO - d - 6), Mm(BANDA_H)
    _estilo_de_placeholder(titulo, TITULO_PT, BLANCO, True, "izq", "medio")

    # portada: bloque azul oscuro abajo; titulo y subtitulo adentro
    # (cada una se mete al fondo de todo: la que va encima se mete primero)
    _forma_de_plantilla(portada.shapes._spTree, 102, "Raya de color", MARGEN + 4, PORTADA_BLOQUE_Y + 12, 46, 2.4, AZUL)
    _forma_de_plantilla(portada.shapes._spTree, 101, "Bloque azul", 0, PORTADA_BLOQUE_Y, HOJA_W,
                        HOJA_H - PORTADA_BLOQUE_Y, AZUL_OSC)
    ph_titulo = next(p for p in portada.placeholders if p.placeholder_format.type == PP_PLACEHOLDER.CENTER_TITLE)
    ph_sub = next(p for p in portada.placeholders if p.placeholder_format.type == PP_PLACEHOLDER.SUBTITLE)
    ph_titulo.left, ph_titulo.top = Mm(MARGEN + 4), Mm(PORTADA_BLOQUE_Y + 20 + ajuste_ppt(PORTADA_TITULO_PT, 1.0))
    ph_titulo.width, ph_titulo.height = Mm(ANCHO - 8), Mm(26)
    _estilo_de_placeholder(ph_titulo, PORTADA_TITULO_PT, BLANCO, True, "izq", "abajo")
    ph_sub.left, ph_sub.top = Mm(MARGEN + 4), Mm(PORTADA_BLOQUE_Y + 50 + ajuste_ppt(PORTADA_SUBTITULO_PT, 1.0))
    ph_sub.width, ph_sub.height = Mm(ANCHO - 8), Mm(16)
    _estilo_de_placeholder(ph_sub, PORTADA_SUBTITULO_PT, CELESTE, False, "izq", "arriba")
    return portada, pagina


def poner_transicion(slide):
    """Fundido corto, el mismo en todas las diapositivas. Sin animaciones por elemento."""
    t = etree.fromstring('<p:transition xmlns:p="%s" spd="%s"><p:fade/></p:transition>' % (NS_P, TRANSICION))
    sld = slide._element
    anterior = sld.find(qn("p:clrMapOvr"))
    if anterior is None:
        anterior = sld.find(qn("p:cSld"))
    anterior.addnext(t)


def poner_propiedades_app(prs, n_notas):
    """docProps/app.xml de la plantilla dice '4:3', 0 diapositivas y 'Macintosh': se deja con lo que es."""
    xml = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
           '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
           'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
           '<TotalTime>0</TotalTime><Application>Microsoft Office PowerPoint</Application>'
           '<PresentationFormat>Personalizado</PresentationFormat><Slides>%d</Slides><Notes>%d</Notes>'
           '<HiddenSlides>0</HiddenSlides><ScaleCrop>false</ScaleCrop><LinksUpToDate>false</LinksUpToDate>'
           '<SharedDoc>false</SharedDoc><HyperlinksChanged>false</HyperlinksChanged><AppVersion>16.0000</AppVersion>'
           '</Properties>' % (len(prs.slides), n_notas))
    for parte in prs.part.package.iter_parts():
        if str(parte.partname) == "/docProps/app.xml":
            parte._blob = xml.encode("utf-8")


def poner_notas(slide, texto):
    tf = slide.notes_slide.notes_text_frame
    tf.text = texto
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.language_id = IDIOMA_ID


# ================================================================== la pagina (igual para los dos lienzos)
def ruta_captura(cap, dir_capturas):
    return Path(dir_capturas) / Path(cap["archivo"]).name


def numero_captura(cap):
    return Path(cap["archivo"]).name[:2]


def encajar(iw, ih, x, y, w, h, arriba=False):
    """Mete una imagen de iw x ih adentro de la caja sin deformarla, centrada (o pegada arriba)."""
    k = min(w / iw, h / ih)
    nw, nh = iw * k, ih * k
    return x + (w - nw) / 2, (y if arriba else y + (h - nh) / 2), nw, nh


SIN_AVISOS = False     # --sin-avisos: la copia para mostrar no lleva la franja "a confirmar"


def parrafos_chip(pag):
    if SIN_AVISOS:
        return []
    grupos = {"app": [], "plan": []}
    for it in pag.get("a_confirmar", []):
        grupos[it["donde"]].append(it["que"])
    out = []
    if grupos["app"]:
        out.append("**A confirmar en la app:** " + " · ".join(grupos["app"]) + ".")
    if grupos["plan"]:
        out.append("**A confirmar antes de entregar:** " + " · ".join(grupos["plan"]) + ".")
    return out


def poner_chip(l, pag, y_fin):
    """La franja ambar "a confirmar" al pie. Devuelve hasta donde llega ahora el texto de arriba."""
    chip = parrafos_chip(pag)
    if not chip:
        return y_fin
    alto_chip, n_chip = alto_parrafos(chip, CHIP_PT, ANCHO - 6, interlinea=1.25)
    if n_chip > 4:
        raise ErrorDeContenido("pagina %s: lo que hay 'a confirmar' ocupa mas de 4 renglones: acortalo"
                               % pag.get("numero", pag.get("titulo", "?")))
    alto_chip += 3.0
    l.rect(MARGEN, CHIP_FONDO - alto_chip, ANCHO, alto_chip, relleno=AMBAR_FONDO)
    l.rect(MARGEN, CHIP_FONDO - alto_chip, 1.2, alto_chip, relleno=AMBAR)
    l.texto(MARGEN + 4, CHIP_FONDO - alto_chip + 1.5, ANCHO - 6, alto_chip - 3, chip, CHIP_PT, AMBAR, interlinea=1.25)
    return CHIP_FONDO - alto_chip - 4.5


def poner_captura(l, cap, caja, dir_capturas, lugar_rotulo=False, arriba=False):
    """lugar_rotulo: deja el lugar del rotulo aunque esta captura no lo tenga (para que dos capturas
    de la misma hoja arranquen a la misma altura). arriba: pega la imagen arriba en vez de centrarla."""
    x, y, w, h = caja
    if cap.get("rotulo"):
        l.texto(x, y, w, ROTULO_H, [cap["rotulo"]], ROTULO_PT, AZUL_OSC, negrita=True, interlinea=1.15)
    if cap.get("rotulo") or lugar_rotulo:
        y, h = y + ROTULO_H, h - ROTULO_H
    ruta = ruta_captura(cap, dir_capturas)
    if not ruta.exists():
        l.rect(x, y, w, h, relleno=GRIS_FONDO, borde=GRIS, grosor=0.6, punteado=True)
        chico = w < 150
        partes = ["**FALTA LA CAPTURA %s**" % numero_captura(cap), cap["que_se_ve"], cap["archivo"]]
        pts = (20 if chico else 28, 14 if chico else 18, 11 if chico else 13)
        altos = [alto_parrafos([p], pt, w - 16, interlinea=1.25)[0] for p, pt in zip(partes, pts)]
        yy = y + (h - sum(altos) - 8) / 2
        for p, pt, a, col in zip(partes, pts, altos, (GRIS_FALTA, GRIS_FALTA, GRIS)):
            l.texto(x + 8, yy, w - 16, a, [p], pt, col, alinear="centro", interlinea=1.25)
            yy += a + 4
        return
    with Image.open(ruta) as im:
        iw, ih = im.size
    ix, iy, nw, nh = encajar(iw, ih, x, y, w, h, arriba=arriba)
    with l.grupo("Captura %s con sus marcas" % numero_captura(cap)):
        l.imagen(ruta, ix, iy, nw, nh, descr=cap["que_se_ve"], nombre="Captura %s" % numero_captura(cap))
        l.rect(ix, iy, nw, nh, borde=GRIS_BORDE, grosor=0.3)
        for m in cap.get("marcas", []):
            if m["tipo"] == "recuadro":
                mx, my = ix + nw * m["x"] / 100, iy + nh * m["y"] / 100
                mw, mh = nw * m["ancho"] / 100, nh * m["alto"] / 100
                l.marco(mx, my, mw, mh)
                if m.get("numero") is not None:
                    # el numero va chico y corrido hacia afuera de la esquina: sobre la esquina tapaba lo marcado
                    # "numero_en": "izquierda" o "derecha" lo saca al costado del recuadro, a media altura
                    donde = m.get("numero_en")
                    if donde == "izquierda":
                        l.insignia(mx - 4.4, my + mh / 2, m["numero"], d=6.4, pt=13)
                    elif donde == "derecha":
                        l.insignia(mx + mw + 4.4, my + mh / 2, m["numero"], d=6.4, pt=13)
                    else:
                        l.insignia(mx - 2.0, my - 2.0, m["numero"], d=6.4, pt=13)
            elif m["tipo"] == "flecha":
                l.flecha(ix + nw * m["desde"][0] / 100, iy + nh * m["desde"][1] / 100,
                         ix + nw * m["hasta"][0] / 100, iy + nh * m["hasta"][1] / 100)


def poner_tabla(l, tabla, caja):
    x, y, w, h = caja
    anchos = [w * f for f in (0.50, 0.27, 0.23)]
    filas = [tabla["encabezado"]] + tabla["filas"]
    altos, celdas = [], []
    for i, fila in enumerate(filas):
        resaltada = (i - 1) == tabla.get("resaltar", -1)
        gruesa = i == 0 or resaltada
        n = max(len(partir(c, TABLA_PT, a - 2 * PAD_TABLA, negrita=gruesa)) for c, a in zip(fila, anchos))
        altos.append(n * TABLA_PT * 1.22 * MM_PT + 2 * PAD_TABLA)
        fondo = AZUL_OSC if i == 0 else (AMBAR_FONDO if resaltada else (BLANCO if i % 2 else CELESTE))
        celdas.append([dict(texto=c, fondo=fondo, color=BLANCO if i == 0 else TINTA, negrita=gruesa) for c in fila])
    yy = y + max(0.0, (h - sum(altos)) / 2)
    l.tabla(x, yy, anchos, altos, celdas,
            descr="Tabla con las columnas: " + ", ".join(tabla["encabezado"]) + ".")


def poner_tarjetas(l, tarjetas, caja):
    x, y, w, h = caja
    cols, hueco = 4, 4.0
    filas = math.ceil(len(tarjetas) / cols)
    tw, th = (w - hueco * (cols - 1)) / cols, (h - hueco * (filas - 1)) / filas
    pad = 4.0
    for i, t in enumerate(tarjetas):
        tx, ty = x + (i % cols) * (tw + hueco), y + (i // cols) * (th + hueco)
        alto_p, _ = alto_parrafos([t["pregunta"]], TARJETA_PT, tw - 2 * pad)
        libre = th - 13.5 - 9.5
        if alto_p > libre + 0.01:
            raise ErrorDeContenido("la pregunta de la tarjeta de %s no entra: acortala" % t["area"])
        with l.grupo("Tarjeta %s" % t["area"]):
            l.rect(tx, ty, tw, th, relleno=BLANCO, borde=GRIS_BORDE, grosor=0.3)
            l.rect(tx, ty, tw, 1.6, relleno=AZUL)
            l.texto(tx + pad, ty + 5.0, tw - 2 * pad, 7.0, [t["area"].upper()], 14, AZUL, negrita=True,
                    interlinea=1.15)
            l.texto(tx + pad, ty + 13.5, tw - 2 * pad, libre, [t["pregunta"]], TARJETA_PT, TINTA)
            l.texto(tx + pad, ty + th - 9.0, tw - 2 * pad, 6.5, ["Documento: **%s**" % t["documento"]], 13, GRIS,
                    interlinea=1.15)


def dibujar_pagina(l, pag, n, total, manual, dir_capturas):
    # franja del titulo, con el numero de pagina
    l.banda()
    d = 17.0
    l.elipse(MARGEN, (BANDA_H - d) / 2, d, d, relleno=BLANCO)
    l.texto(MARGEN, (BANDA_H - d) / 2, d, d, [str(n)], 28, AZUL_OSC, negrita=True, alinear="centro",
            anclar="medio", interlinea=1.0, deco=True, nombre="Numero de pagina")
    l.titulo(MARGEN + d + 6, 0, ANCHO - d - 6, BANDA_H, pag["titulo"], TITULO_PT, BLANCO)

    # de abajo hacia arriba: pie, franja "a confirmar", texto
    l.texto(MARGEN, PIE_Y, ANCHO * 0.7, 5, ["%s · %s" % (manual["pie"], manual["version"])], 10, GRIS, interlinea=1.2)
    l.texto(MARGEN + ANCHO * 0.7, PIE_Y, ANCHO * 0.3, 5, ["Página %d de %d" % (n, total)], 10, GRIS,
            alinear="der", interlinea=1.2)
    y_fin = poner_chip(l, pag, TEXTO_FONDO)
    alto_texto, n_reng = alto_parrafos(pag["texto"], TEXTO_PT, ANCHO, entre=ENTRE_PARRAFOS, vinetas=True)
    if n_reng > MAX_RENGLONES:
        raise ErrorDeContenido("pagina %d (%s): el texto ocupa %d renglones impresos y el maximo es %d: acortalo"
                               % (pag["numero"], pag["titulo"], n_reng, MAX_RENGLONES))
    y_texto = y_fin - alto_texto
    l.texto(MARGEN, y_texto, ANCHO, alto_texto + 1.0, pag["texto"], TEXTO_PT, TINTA, entre=ENTRE_PARRAFOS, vinetas=True)

    # la imagen: todo lo que queda
    area = (MARGEN, Y0, ANCHO, y_texto - AIRE_IMAGEN_TEXTO - Y0)
    if area[3] < ALTO_MINIMO_IMAGEN:
        raise ErrorDeContenido("pagina %d: a la imagen le quedan %.0f mm de alto (minimo %.0f): hay demasiado texto"
                               % (pag["numero"], area[3], ALTO_MINIMO_IMAGEN))
    if pag.get("tarjetas"):
        poner_tarjetas(l, pag["tarjetas"], area)
        return
    caps = [c for c in pag.get("capturas", [])
            if ruta_captura(c, dir_capturas).exists() or not c.get("opcional")]
    x, y, w, h = area
    if pag.get("tabla"):
        w1 = (w - SEP) * 0.56
        poner_captura(l, caps[0], (x, y, w1, h), dir_capturas)
        poner_tabla(l, pag["tabla"], (x + w1 + SEP, y, w - SEP - w1, h))
    elif len(caps) == 2:
        w1 = (w - SEP) * 0.60
        rot = any(c.get("rotulo") for c in caps)
        poner_captura(l, caps[0], (x, y, w1, h), dir_capturas, lugar_rotulo=rot, arriba=True)
        poner_captura(l, caps[1], (x + w1 + SEP, y, w - SEP - w1, h), dir_capturas, lugar_rotulo=rot, arriba=True)
    else:
        poner_captura(l, caps[0], area, dir_capturas)


# ================================================================== leer y controlar contenido.json
def reemplazar(obj, variables):
    if isinstance(obj, str):
        for k, v in variables.items():
            obj = obj.replace("{%s}" % k, v)
        return obj
    if isinstance(obj, list):
        return [reemplazar(o, variables) for o in obj]
    if isinstance(obj, dict):
        return {k: reemplazar(v, variables) for k, v in obj.items()}
    return obj


def cargar(ruta):
    datos = json.loads(Path(ruta).read_text(encoding="utf-8"))
    for clave in ("paginas", "portada", "cierre"):
        if clave in datos:
            datos[clave] = reemplazar(datos[clave], datos.get("variables", {}))
    return datos


def controlar(datos):
    """Lo que se puede controlar sin dibujar. Devuelve la lista de errores (vacia = bien)."""
    errores = []
    activas = [p for p in datos["paginas"] if p.get("activa", True)]
    if not activas:
        errores.append("no hay ninguna pagina activa")
    if len(activas) > MAX_PAGINAS:
        errores.append("hay %d paginas y el maximo es %d" % (len(activas), MAX_PAGINAS))
    vistos = set()
    for p in datos["paginas"]:
        donde = "pagina %s (%s)" % (p.get("numero"), p.get("titulo"))
        titulo = p.get("titulo", "")
        if not titulo or len(titulo.split()) > MAX_PALABRAS_TITULO:
            errores.append("%s: el titulo tiene que tener de 1 a %d palabras" % (donde, MAX_PALABRAS_TITULO))
        elif ancho_txt(titulo, TITULO_PT, True) > ANCHO - 23:
            errores.append("%s: el titulo no entra en un renglon" % donde)
        texto = p.get("texto", [])
        if not 2 <= len(texto) <= 4:
            errores.append("%s: el texto tiene que tener de 2 a 4 renglones (tiene %d)" % (donde, len(texto)))
        numeros = [separar_numero(t)[0] for t in texto if separar_numero(t)[0] is not None]
        if numeros != list(range(1, len(numeros) + 1)):
            errores.append("%s: los renglones numerados tienen que ir 1, 2, 3 seguidos" % donde)
        for t in texto + [c.get("que_se_ve", "") for c in p.get("capturas", [])]:
            if t.count("**") % 2:
                errores.append("%s: hay un ** sin cerrar en: %s" % (donde, t))
            if "{" in t or "}" in t:
                errores.append("%s: quedo una llave sin reemplazar en: %s" % (donde, t))
        caps = p.get("capturas", [])
        if not p.get("tarjetas") and not 1 <= len(caps) <= 2:
            errores.append("%s: tiene que tener 1 o 2 capturas (o tarjetas)" % donde)
        if p.get("tabla") and len(caps) != 1:
            errores.append("%s: una pagina con tabla lleva una sola captura" % donde)
        if caps and caps[0].get("opcional"):
            errores.append("%s: la captura grande no puede ser opcional" % donde)
        for c in caps:
            arch = c.get("archivo", "")
            if not re.fullmatch(r"capturas/\d\d-[a-z0-9-]+\.png", arch):
                errores.append("%s: el archivo '%s' no tiene la forma capturas/NN-nombre.png" % (donde, arch))
            if arch in vistos:
                errores.append("%s: la captura '%s' esta repetida" % (donde, arch))
            vistos.add(arch)
            if not c.get("que_se_ve"):
                errores.append("%s: a la captura '%s' le falta 'que_se_ve'" % (donde, arch))
            for m in c.get("marcas", []):
                if m.get("tipo") == "recuadro":
                    vals = [m.get("x"), m.get("y"), m.get("ancho"), m.get("alto")]
                    bien = all(isinstance(v, (int, float)) for v in vals) and m["ancho"] > 0 and m["alto"] > 0 \
                        and 0 <= m["x"] and 0 <= m["y"] and m["x"] + m["ancho"] <= 100 and m["y"] + m["alto"] <= 100
                elif m.get("tipo") == "flecha":
                    pts = list(m.get("desde", [])) + list(m.get("hasta", []))
                    bien = len(pts) == 4 and all(isinstance(v, (int, float)) and 0 <= v <= 100 for v in pts)
                else:
                    bien = False
                if not bien:
                    errores.append("%s: una marca de '%s' esta mal (recuadro o flecha, en porcentajes de 0 a 100): %s"
                                   % (donde, arch, json.dumps(m, ensure_ascii=False)))
        for it in p.get("a_confirmar", []):
            if it.get("donde") not in ("app", "plan") or not it.get("que"):
                errores.append("%s: cada 'a_confirmar' lleva donde ('app' o 'plan') y que" % donde)
    errores += controlar_portada_y_cierre(datos)
    return errores


def controlar_portada_y_cierre(datos):
    """La portada y el cierre salen solo en el PowerPoint, pero tienen que estar bien escritos igual."""
    errores = []
    por, cie = datos.get("portada") or {}, datos.get("cierre") or {}
    if not por.get("titulo") or not por.get("subtitulo"):
        errores.append("portada: faltan el titulo y el subtitulo")
    elif ancho_txt(por["titulo"], PORTADA_TITULO_PT, True) > ANCHO - 8:
        errores.append("portada: el titulo no entra en un renglon")
    pasos = cie.get("pasos", [])
    if not cie.get("titulo") or not cie.get("aviso") or not 2 <= len(pasos) <= 4:
        errores.append("cierre: faltan el titulo, el aviso o los pasos (de 2 a 4)")
    for t in [por.get("titulo", ""), por.get("subtitulo", ""), cie.get("titulo", ""), cie.get("aviso", "")] + pasos:
        if t.count("**") % 2:
            errores.append("portada/cierre: hay un ** sin cerrar en: %s" % t)
        if "{" in t or "}" in t:
            errores.append("portada/cierre: quedo una llave sin reemplazar en: %s" % t)
    for it in cie.get("a_confirmar", []):
        if it.get("donde") not in ("app", "plan") or not it.get("que"):
            errores.append("cierre: cada 'a_confirmar' lleva donde ('app' o 'plan') y que")
    return errores


def revisar_capturas(datos, dir_capturas):
    """Que capturas faltan y que marcas siguen en posicion estimada."""
    faltan, opcionales, estimadas = [], [], []
    n = 0
    for p in datos["paginas"]:
        if not p.get("activa", True):
            continue
        n += 1
        for c in p.get("capturas", []):
            existe = ruta_captura(c, dir_capturas).exists()
            fila = (numero_captura(c), c["archivo"], n, c["que_se_ve"])
            if not existe:
                (opcionales if c.get("opcional") else faltan).append(fila)
            elif any(not m.get("posicion_confirmada") for m in c.get("marcas", [])):
                estimadas.append(fila)
    return faltan, opcionales, estimadas


# ================================================================== armar
def armar_pdf(datos, dir_capturas):
    activas = [p for p in datos["paginas"] if p.get("activa", True)]
    hojas = []
    for n, pag in enumerate(activas, 1):
        l = LienzoPDF()
        dibujar_pagina(l, pag, n, len(activas), datos["manual"], dir_capturas)
        hojas.append(l.im)
    return hojas


MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre",
         "noviembre", "diciembre"]


def fecha_larga(dia=None):
    dia = dia or datetime.date.today()
    return "%d de %s de %d" % (dia.day, MESES[dia.month - 1], dia.year)


def diapositiva_portada(prs, layout, datos, hoy=None):
    """Solo en el PowerPoint: titulo, 'Manual de arranque', el logo, version y fecha."""
    por, man = datos["portada"], datos["manual"]
    l = LienzoPPT(prs, layout)
    poner_texto_placeholder(l.s.shapes.title, [por["titulo"]], PORTADA_TITULO_PT, BLANCO, True, "izq")
    sub = next(p for p in l.s.placeholders if p.placeholder_format.type == PP_PLACEHOLDER.SUBTITLE)
    poner_texto_placeholder(sub, [por["subtitulo"]], PORTADA_SUBTITULO_PT, CELESTE, False, "izq")
    if LOGO.exists():
        with Image.open(LOGO) as im:
            iw, ih = im.size
        lw = 78.0
        l.imagen(LOGO, MARGEN + 4, 22.0, lw, lw * ih / iw, descr="Logo de Barack Mercosul", nombre="Logo")
    l.texto(MARGEN + 4, PORTADA_BLOQUE_Y + 82, ANCHO - 8, 22,
            ["**Versión:** %s" % man["version"], "**Fecha:** %s" % fecha_larga(hoy)], PORTADA_DATOS_PT, BLANCO,
            interlinea=1.3, entre=1.0, nombre="Version y fecha")
    return l.s


def diapositiva_cierre(prs, layout, datos):
    """Solo en el PowerPoint: los tres pasos, y a quien avisar."""
    cie, man = datos["cierre"], datos["manual"]
    l = LienzoPPT(prs, layout)
    l.titulo(MARGEN, 0, ANCHO, BANDA_H, cie["titulo"], TITULO_PT, BLANCO)
    pasos = cie["pasos"]
    hueco, y, alto = 8.0, 38.0, 96.0
    ancho = (ANCHO - hueco * (len(pasos) - 1)) / len(pasos)
    for i, paso in enumerate(pasos, 1):
        x = MARGEN + (i - 1) * (ancho + hueco)
        if len(partir(paso, CIERRE_PASO_PT, ancho - 12)) > 3:
            raise ErrorDeContenido("cierre: el paso %d no entra en la tarjeta: acortalo" % i)
        with l.grupo("Paso %d" % i):
            l.rect(x, y, ancho, alto, relleno=BLANCO, borde=GRIS_BORDE, grosor=0.3)
            l.rect(x, y, ancho, 1.6, relleno=AZUL)
            l.insignia(x + ancho / 2, y + 25, i, d=26, pt=40, halo=False)
            l.texto(x + 6, y + 46, ancho - 12, alto - 52, [paso], CIERRE_PASO_PT, TINTA, alinear="centro",
                    anclar="medio", interlinea=1.2)
    ya, ah = y + alto + 10, 26.0
    l.rect(MARGEN, ya, ANCHO, ah, relleno=CELESTE)
    l.rect(MARGEN, ya, 1.6, ah, relleno=AZUL)
    l.texto(MARGEN + 8, ya, ANCHO - 12, ah, [cie["aviso"]], CIERRE_AVISO_PT, TINTA, anclar="medio", interlinea=1.2)
    poner_chip(l, cie, ya + ah)
    l.texto(MARGEN, PIE_Y, ANCHO * 0.7, 5, ["%s · %s" % (man["pie"], man["version"])], 10, GRIS, interlinea=1.2)
    return l.s


def armar_pptx(datos, dir_capturas, con_notas=False, hoy=None):
    """El PowerPoint: portada + las paginas del manual (iguales al PDF) + cierre."""
    activas = [p for p in datos["paginas"] if p.get("activa", True)]
    prs = Presentation()
    ancho0, alto0 = prs.slide_width, prs.slide_height
    prs.slide_width, prs.slide_height = Mm(HOJA_W), Mm(HOJA_H)
    portada, pagina = preparar_plantilla(prs, ancho0, alto0)
    cp = prs.core_properties
    cp.title = datos["manual"]["archivo"]
    cp.author = cp.last_modified_by = "Ingeniería - Barack Mercosul"
    cp.subject = cp.keywords = cp.comments = cp.category = ""
    cp.language = IDIOMA
    cp.created = cp.modified = datetime.datetime.now()
    cp.revision = 1

    s = diapositiva_portada(prs, portada, datos, hoy)
    if con_notas:
        poner_notas(s, "Portada. Después siguen las %d páginas del manual, en orden." % len(activas))
    for n, pag in enumerate(activas, 1):
        l = LienzoPPT(prs, pagina)
        dibujar_pagina(l, pag, n, len(activas), datos["manual"], dir_capturas)
        if con_notas and pag.get("error_que_evita"):
            poner_notas(l.s, "Para qué está esta página: evita " + pag["error_que_evita"][0].lower()
                        + pag["error_que_evita"][1:])
    s = diapositiva_cierre(prs, pagina, datos)
    if con_notas:
        poner_notas(s, "Cierre. Repasar los tres pasos y a quién avisar si algo no anda.")
    for diapositiva in prs.slides:
        poner_transicion(diapositiva)
    poner_propiedades_app(prs, sum(1 for d in prs.slides if d.has_notes_slide))
    return prs


def guardar(hojas, prs, dir_salida, nombre):
    dir_salida.mkdir(parents=True, exist_ok=True)
    pdf, pptx = dir_salida / (nombre + ".pdf"), dir_salida / (nombre + ".pptx")
    try:
        hojas[0].save(pdf, "PDF", save_all=True, append_images=hojas[1:], resolution=DPI, quality=95,
                      subsampling=0, title=nombre, author="Ingeniería - Barack Mercosul")
        prs.save(str(pptx))
    except PermissionError as e:
        cual = Path(e.filename).name if e.filename else nombre
        raise ErrorDeContenido("no pude guardar '%s': cerralo si esta abierto y volve a correr" % cual)
    return pdf, pptx


def hacer_grilla(datos, dir_capturas):
    """Cada captura con una cuadricula de 10 en 10 y sus marcas actuales, para leer la posicion real."""
    destino = Path(dir_capturas) / "_grilla"
    destino.mkdir(exist_ok=True)
    hechas = []
    for p in datos["paginas"]:
        for c in p.get("capturas", []):
            ruta = ruta_captura(c, dir_capturas)
            if not ruta.exists():
                continue
            im = Image.open(ruta).convert("RGB")
            w, h = im.size
            d = ImageDraw.Draw(im, "RGBA")
            f = ImageFont.truetype(str(DIR_FUENTES / "segoeuib.ttf"), max(14, w // 80))
            for k in range(0, 101, 5):
                fuerte = k % 10 == 0
                col = (0, 120, 255, 200 if fuerte else 90)
                d.line([(w * k / 100, 0), (w * k / 100, h)], fill=col, width=2 if fuerte else 1)
                d.line([(0, h * k / 100), (w, h * k / 100)], fill=col, width=2 if fuerte else 1)
                if fuerte and k < 100:
                    for pos in ((w * k / 100 + 4, 2), (2, h * k / 100 + 2)):
                        caja = d.textbbox(pos, str(k), font=f)
                        d.rectangle(caja, fill=(255, 255, 255, 220))
                        d.text(pos, str(k), font=f, fill=(0, 70, 200, 255))
            for m in c.get("marcas", []):
                if m["tipo"] == "recuadro":
                    d.rectangle([w * m["x"] / 100, h * m["y"] / 100, w * (m["x"] + m["ancho"]) / 100,
                                 h * (m["y"] + m["alto"]) / 100], outline=ROJO + (255,), width=max(4, w // 300))
                else:
                    d.line([(w * m["desde"][0] / 100, h * m["desde"][1] / 100),
                            (w * m["hasta"][0] / 100, h * m["hasta"][1] / 100)], fill=ROJO + (255,), width=max(4, w // 300))
            im.save(destino / ruta.name)
            hechas.append(ruta.name)
    return destino, hechas


def main(argv=None):
    for flujo in (sys.stdout, sys.stderr):
        try:
            flujo.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass
    ap = argparse.ArgumentParser(description="Arma el manual de arranque de Claude en Barack (PDF y PowerPoint).")
    ap.add_argument("--contenido", default=str(AQUI / "contenido.json"), help="otro contenido.json (para pruebas)")
    ap.add_argument("--capturas", default=str(AQUI / "capturas"), help="otra carpeta de capturas (para pruebas)")
    ap.add_argument("--salida", default=str(SALIDA), help="otra carpeta de salida (para pruebas)")
    ap.add_argument("--grilla", action="store_true", help="solo deja las capturas con cuadricula en capturas/_grilla/")
    ap.add_argument("--notas", action="store_true",
                    help="pone 'para que esta la pagina' en las notas del PowerPoint (con --sin-avisos van siempre)")
    ap.add_argument("--sin-avisos", action="store_true",
                    help="copia para mostrar: sin la franja 'a confirmar' (el archivo sale con ' - para mostrar' en el nombre)")
    a = ap.parse_args(argv)
    global SIN_AVISOS
    SIN_AVISOS = a.sin_avisos

    try:
        datos = cargar(a.contenido)
    except (OSError, ValueError, KeyError) as e:
        print("ERROR: no pude leer %s: %s" % (a.contenido, e))
        return 2
    errores = controlar(datos)
    if errores:
        print("ERROR: contenido.json tiene %d cosa(s) para corregir:" % len(errores))
        for e in errores:
            print("  - " + e)
        return 2

    if a.grilla:
        destino, hechas = hacer_grilla(datos, a.capturas)
        print("Cuadricula hecha para %d captura(s) en %s" % (len(hechas), destino))
        for h in hechas:
            print("  " + h)
        if not hechas:
            print("  (todavia no hay ninguna captura en la carpeta)")
        return 0

    try:
        hojas = armar_pdf(datos, a.capturas)
        prs = armar_pptx(datos, a.capturas, con_notas=a.notas or a.sin_avisos)
        pdf, pptx = guardar(hojas, prs, Path(a.salida),
                            datos["manual"]["archivo"] + (" - para mostrar" if a.sin_avisos else ""))
    except ErrorDeContenido as e:
        print("ERROR: %s" % e)
        return 2

    faltan, opcionales, estimadas = revisar_capturas(datos, a.capturas)
    print("Manual armado: %d paginas" % len(hojas))
    print("  PDF : %s" % pdf)
    print("  PPTX: %s" % pptx)
    pendientes = [(p["numero"], it) for p in datos["paginas"] if p.get("activa", True) for it in p.get("a_confirmar", [])]
    pendientes += [("cierre", it) for it in datos.get("cierre", {}).get("a_confirmar", [])]
    if pendientes or datos.get("a_confirmar_general"):
        print("\nA CONFIRMAR (sale impreso en la franja ambar de cada pagina hasta que se borre de contenido.json):")
        for num, it in pendientes:
            print("  pag. %2s  [%s]  %s" % (num, "en la app" if it["donde"] == "app" else "antes de entregar", it["que"]))
        for g in datos.get("a_confirmar_general", []):
            print("  general  %s" % g)
    if opcionales:
        print("\nCAPTURAS OPCIONALES QUE NO ESTAN (la pagina sale sin ellas):")
        for num, arch, pag, que in opcionales:
            print("  %s  %s  (pag. %d)  %s" % (num, arch, pag, que))
    if estimadas:
        print("\nCAPTURAS PUESTAS CON LA MARCA ROJA TODAVIA EN POSICION ESTIMADA (%d):" % len(estimadas))
        for num, arch, pag, _ in estimadas:
            print("  %s  %s  (pag. %d)" % (num, arch, pag))
        print("  -> python generar_manual.py --grilla, leer la posicion real en capturas/_grilla/, corregirla en")
        print("     contenido.json y poner \"posicion_confirmada\": true")
    if faltan:
        print("\nFALTAN %d CAPTURAS (el manual salio con un cartel gris en su lugar):" % len(faltan))
        for num, arch, pag, que in sorted(faltan):
            print("  %s  %s  (pag. %d)  %s" % (num, arch, pag, que))
    if faltan or estimadas:
        print("\nEL MANUAL NO ESTA TERMINADO: no entregarlo asi.")
        return 1
    print("\nManual completo.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
