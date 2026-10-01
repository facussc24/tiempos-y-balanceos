# -*- coding: utf-8 -*-
"""svlib_lamina.py — piezas para armar la presentacion (python-pptx, Python del sistema).

El formato es el que le gusto a Fak el 01/10/2026 (*"me encanta el power point"*): 16:9, fondo
blanco, titulo azul Barack, logo arriba a la derecha, una idea por hoja, el NUMERO grande en
una caja de color y la imagen grande. Sin notas explicativas ni marca de IA.

  m = Mazo(logo)                           # 13,33 x 7,5 pulgadas
  s = m.base('Titulo', 'subtitulo')        # hoja con titulo, logo y linea
  m.imagen(s, 'foto.png', x, y, w_max, h_max)
  m.caja(s, x, y, w, h, AZUL, 'Lo que se ve tapizado', '0,10 m²')
  m.texto(s, x, y, w, h, 'texto', 14, GRIS)
  m.tabla(s, filas, x, y, w, anchos)
  m.barra(s, x, y, w, h, [(0.14, OSC, 'Queda en la pieza'), (0.09, GRIS_BARRA, 'Recorte')])
  m.guardar('salida.pptx');  exportar('salida.pptx', 'qa', 'salida.pdf')

Los numeros se escriben con `es()` (coma decimal) y se CALCULAN en el script que arma el mazo,
con asserts de que las partes suman el total: la frase que resume numeros la arma el codigo.
"""
import os
import subprocess

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

AZUL_B = RGBColor(0x1B, 0x4F, 0x8F)       # azul Barack (titulos)
AZUL = RGBColor(0x2A, 0x7F, 0xD4)         # lo que se ve
NARANJA = RGBColor(0xFF, 0x8C, 0x1A)      # el borde que se dobla
OSC = RGBColor(0x2B, 0x2B, 0x2B)
GRIS = RGBColor(0x55, 0x55, 0x55)
GRIS_BARRA = RGBColor(0x9A, 0x9A, 0x9A)
ROJO = RGBColor(0xD6, 0x27, 0x28)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
CLARO = RGBColor(0xF2, 0xF4, 0xF7)
FUENTE = 'Segoe UI'
LOGO = (r'C:\Users\FacundoS-PC\BARACK ARGENTINA SRL\Ingeniería y Proyecto - General'
        r'\INGENIERIA BARACK (NUNCA BORRAR)\barack_logo.png')
AQUI = os.path.dirname(os.path.abspath(__file__))
IZQ, CEN = PP_ALIGN.LEFT, PP_ALIGN.CENTER
MEDIO = MSO_ANCHOR.MIDDLE


def es(x, dec):
    """0.14 -> '0,14'"""
    return ('%.*f' % (dec, x)).replace('.', ',')


def miles(x):
    """1401.7 -> '1.402'"""
    s = '%d' % round(x)
    return s[:-3] + '.' + s[-3:] if len(s) > 3 else s


class Mazo(object):
    def __init__(self, logo=LOGO, autor='Ingeniería Barack', titulo=''):
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333); self.prs.slide_height = Inches(7.5)
        self.logo = logo if logo and os.path.exists(logo) else None
        self.prs.core_properties.author = autor
        self.prs.core_properties.title = titulo

    def texto(self, s, x, y, w, h, t, size=14, color=OSC, bold=False, align=IZQ, anchor=MSO_ANCHOR.TOP):
        """t puede ser un texto o una lista de renglones (texto, tamaño, color, negrita)."""
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.02)
        renglones = t if isinstance(t, list) else [(t, size, color, bold)]
        for i, (tx, sz, col, b) in enumerate(renglones):
            par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            par.alignment = align
            run = par.add_run(); run.text = tx
            run.font.size = Pt(sz); run.font.bold = b; run.font.color.rgb = col; run.font.name = FUENTE
        return tb

    def caja(self, s, x, y, w, h, relleno, etiqueta, valor, color_txt=BLANCO, sz_et=16, sz_val=40):
        """Caja de color con una etiqueta chica y el NUMERO grande."""
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.adjustments[0] = 0.08
        sh.fill.solid(); sh.fill.fore_color.rgb = relleno; sh.line.fill.background(); sh.shadow.inherit = False
        self.texto(s, x + 0.15, y + 0.12, w - 0.3, 0.5, etiqueta, sz_et, color_txt, False, CEN)
        self.texto(s, x + 0.15, y + 0.58, w - 0.3, h - 0.65, valor, sz_val, color_txt, True, CEN, MEDIO)

    def imagen(self, s, path, x, y, w_max, h_max):
        """La imagen entera, sin deformar, centrada en el hueco."""
        im = Image.open(path); r = im.width / im.height
        w = w_max; h = w / r
        if h > h_max:
            h = h_max; w = h * r
        s.shapes.add_picture(path, Inches(x + (w_max - w) / 2), Inches(y + (h_max - h) / 2), Inches(w), Inches(h))

    def base(self, titulo, sub=None):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.texto(s, 0.5, 0.32, 10.5, 0.7, titulo, 30, AZUL_B, True)
        if sub:
            self.texto(s, 0.5, 0.98, 10.5, 0.4, sub, 16, GRIS)
        if self.logo:
            s.shapes.add_picture(self.logo, Inches(11.55), Inches(0.33), Inches(1.3))
        ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1.45), Inches(12.35), Emu(19050))
        ln.fill.solid(); ln.fill.fore_color.rgb = AZUL_B; ln.line.fill.background()
        return s

    def tabla(self, s, filas, x=0.5, y=1.85, w=12.35, anchos=(7.35, 2.5, 2.5), alto_fila=0.62, negrita_hasta=1):
        """filas[0] es el encabezado. La primera columna va a la izquierda, el resto centrado."""
        tb = s.shapes.add_table(len(filas), len(filas[0]), Inches(x), Inches(y), Inches(w), Inches(alto_fila * len(filas))).table
        for j, a in enumerate(anchos):
            tb.columns[j].width = Inches(a)
        for i, f in enumerate(filas):
            for j, v in enumerate(f):
                c = tb.cell(i, j); c.text = ''
                par = c.text_frame.paragraphs[0]; par.alignment = IZQ if j == 0 else CEN
                run = par.add_run(); run.text = v; run.font.name = FUENTE
                run.font.size = Pt(18 if i else 16); run.font.bold = (i <= negrita_hasta)
                run.font.color.rgb = BLANCO if i == 0 else OSC
                c.vertical_anchor = MEDIO
                c.fill.solid(); c.fill.fore_color.rgb = AZUL_B if i == 0 else (CLARO if i % 2 else BLANCO)
        return tb

    def barra(self, s, x, y, w, h, tramos, size=18):
        """Barra partida en proporcion: tramos = [(valor, color, texto), ...]."""
        total = float(sum(t[0] for t in tramos)); xx = x
        for v, col, t in tramos:
            ww = w * v / total
            sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(xx), Inches(y), Inches(ww), Inches(h))
            sh.fill.solid(); sh.fill.fore_color.rgb = col; sh.line.fill.background(); sh.shadow.inherit = False
            self.texto(s, xx, y, ww, h, t, size, BLANCO, True, CEN, MEDIO)
            xx += ww

    def pie(self, s, renglones, y=6.6):
        self.texto(s, 0.5, y, 12.35, 0.9, [(t, 13, GRIS, False) for t in renglones], 13)

    def guardar(self, path):
        self.prs.save(path)
        return path


def exportar(pptx, carpeta_png, pdf=''):
    """PNG de cada hoja (para mirarlas antes de entregar) y PDF. Necesita PowerPoint.
    Comprueba que salio una imagen NUEVA por hoja: mirar las hojas de la corrida anterior
    creyendo que son las de ahora es verificar la orden y no el resultado."""
    import glob
    import re
    import time
    os.makedirs(carpeta_png, exist_ok=True)
    t0 = time.time()
    args = ['powershell', '-NoProfile', '-ExecutionPolicy', 'Bypass', '-File', os.path.join(AQUI, 'sv_exportar_pptx.ps1'),
            os.path.abspath(pptx), os.path.abspath(carpeta_png)]
    if pdf:
        args.append(os.path.abspath(pdf))
    r = subprocess.run(args, capture_output=True, text=True)
    m = re.search(r'hojas=(\d+)', r.stdout)
    nuevas = [f for f in glob.glob(os.path.join(carpeta_png, '*.PNG'))
              if re.match(r'(Diapositiva|Slide)\d+\.png$', os.path.basename(f), re.I) and os.path.getmtime(f) >= t0 - 2]
    if r.returncode != 0 or not m or int(m.group(1)) < 1 or len(nuevas) < int(m.group(1)):
        raise SystemExit('no se pudo exportar el PowerPoint (hojas=%s, imagenes nuevas=%d):\n%s\n%s'
                         % (m.group(1) if m else '?', len(nuevas), r.stdout[-600:], r.stderr[-600:]))
    if pdf and not (os.path.exists(pdf) and os.path.getmtime(pdf) >= t0 - 2):
        raise SystemExit('PowerPoint no dejo el PDF: %s' % pdf)
    return r.stdout.strip()
