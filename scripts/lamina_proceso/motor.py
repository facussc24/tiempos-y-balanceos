# -*- coding: utf-8 -*-
"""Motor de dibujo de las laminas de proceso (A3 apaisado, PowerPoint editable).

Trae las piezas con las que se arma una lamina: caja, texto, flecha, foto (recortada al marco o
entera sobre fondo), tarjeta de paso (numero, titulo, fotos con rotulo, recuadro de uso por pieza,
pastillas de "que se junta aca") y banda de color por componente.

Lo importa el armado de cada lamina (ver insert_patagonia.py):
    python scripts/lamina_proceso/insert_patagonia.py <salida.pptx> <carpeta de datos>
La carpeta de datos trae fotos.json (que foto va en cada paso), usos.json y las fotos.
Como se hace una lamina entera: skill `lamina-de-proceso`.
"""
import io
import json
import os
import sys

from lxml import etree
from PIL import Image, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Mm, Pt

AQUI = os.path.dirname(os.path.abspath(__file__))
SALIDA = sys.argv[1] if len(sys.argv) > 1 else "lamina.pptx"
# carpeta de DATOS de la lamina (fotos.json, usos.json y las fotos): 2do argumento, o la carpeta actual
DATOS = os.path.abspath(sys.argv[2]) if len(sys.argv) > 2 else os.getcwd()
FOTOS = json.load(open(os.path.join(DATOS, "fotos.json"), encoding="utf-8"))

# ---------------------------------------------------------------- paleta
AZUL = "004E98"       # azul del logo de Barack
AZUL_OSC = "0B2545"
AZUL_CLARO = "E6EEF7"
GRIS_TXT = "5A6672"
GRIS_BORDE = "D5DBE2"
FONDO = "F4F6F9"
BLANCO = "FFFFFF"
PIZARRA = "4A5A6A"
PIZARRA_CLARO = "E9EDF1"
AMBAR = "B87A00"
AMBAR_CLARO = "FBF1D9"
FUENTE = "Segoe UI"
FUENTE_B = "Segoe UI Semibold"

LINEAS = {
    "V": dict(color=AZUL, claro=AZUL_CLARO, nombre="VINILO"),
    "S": dict(color=PIZARRA, claro=PIZARRA_CLARO, nombre="SUSTRATO PLÁSTICO"),
    "E": dict(color=AMBAR, claro=AMBAR_CLARO, nombre="ESPUMA", tinta="7A5200"),
    "A": dict(color=AZUL_OSC, claro="DDE3EC", nombre="ARMADO Y TAPIZADO"),
}

# ---------------------------------------------------------------- pagina
W, H = 420.0, 297.0
prs = Presentation()
prs.slide_width, prs.slide_height = Mm(W), Mm(H)
slide = prs.slides.add_slide(prs.slide_layouts[6])
sp = slide.shapes


def rgb(h):
    return RGBColor.from_string(h)


def sombra(shape, blur=5.0, dist=1.6, alpha=20):
    spPr = shape._element.spPr
    for viejo in spPr.findall(qn("a:effectLst")):
        spPr.remove(viejo)
    ef = etree.SubElement(spPr, qn("a:effectLst"))
    sh = etree.SubElement(ef, qn("a:outerShdw"), blurRad=str(int(Mm(blur))), dist=str(int(Mm(dist))),
                          dir="5400000", algn="t", rotWithShape="0")
    c = etree.SubElement(sh, qn("a:srgbClr"), val=AZUL_OSC)
    etree.SubElement(c, qn("a:alpha"), val=str(alpha * 1000))


def sin_sombra(shape):
    spPr = shape._element.spPr
    if spPr.find(qn("a:effectLst")) is None:
        etree.SubElement(spPr, qn("a:effectLst"))


def caja(x, y, w, h, relleno=BLANCO, borde=None, radio=None, grosor=0.75, nombre=None):
    forma = MSO_SHAPE.ROUNDED_RECTANGLE if radio else MSO_SHAPE.RECTANGLE
    s = sp.add_shape(forma, Mm(x), Mm(y), Mm(w), Mm(h))
    if radio:
        s.adjustments[0] = min(0.5, radio / min(w, h))
    if relleno is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(relleno)
    if borde:
        s.line.color.rgb = rgb(borde)
        s.line.width = Pt(grosor)
    else:
        s.line.fill.background()
    sin_sombra(s)
    if nombre:
        s.name = nombre
    return s


def texto(x, y, w, h, t, tam=10, color=AZUL_OSC, negrita=False, alin="l", anc="m", fuente=None,
          interlinea=None, espaciado=None):
    tb = sp.add_textbox(Mm(x), Mm(y), Mm(w), Mm(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anc]
    lineas = t if isinstance(t, list) else [t]
    for i, ln in enumerate(lineas):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[alin]
        if interlinea:
            p.line_spacing = interlinea
        r = p.add_run()
        r.text = ln
        f = r.font
        f.size = Pt(tam)
        f.bold = negrita
        f.name = fuente or (FUENTE_B if negrita else FUENTE)
        f.color.rgb = rgb(color)
        if espaciado:
            r._r.get_or_add_rPr().set("spc", str(espaciado))
    return tb


def flecha(puntos, color=AZUL, grosor=2.6, punta=True):
    """Polilinea con punta de flecha al final. puntos en mm."""
    x0, y0 = puntos[0]
    fb = sp.build_freeform(Mm(x0), Mm(y0), scale=1.0)
    fb.add_line_segments([(Mm(x), Mm(y)) for x, y in puntos[1:]], close=False)
    s = fb.convert_to_shape()
    s.fill.background()
    s.line.color.rgb = rgb(color)
    s.line.width = Pt(grosor)
    ln = s._element.spPr.find(qn("a:ln"))
    ln.set("cap", "rnd")
    for hijo in ln.findall(qn("a:round")):
        ln.remove(hijo)
    etree.SubElement(ln, qn("a:round"))
    if punta:
        etree.SubElement(ln, qn("a:tailEnd"), type="triangle", w="lg", len="med")
    sin_sombra(s)
    return s


# ---------------------------------------------------------------- fotos
def _abrir(ruta):
    im = Image.open(ruta)
    im = ImageOps.exif_transpose(im)
    return im


def foto(ruta, x, y, w, h, modo="cubrir", radio=1.6, fondo=BLANCO, rotar=0, borde=GRIS_BORDE, foco=(0.5, 0.5),
         margen=1.5, rotulo_vacio=None, recorte=None, corte_abajo=0.0):
    """Pone una foto en el marco (x,y,w,h en mm).
    modo 'cubrir': recorta para llenar el marco. modo 'entera': la muestra completa sobre el fondo."""
    marco = caja(x, y, w, h, relleno=fondo, borde=borde, radio=radio, grosor=0.5)
    if not ruta or not os.path.exists(ruta):
        marco.line.color.rgb = rgb(GRIS_BORDE)
        marco.line.dash_style = 4  # guiones
        marco.fill.solid()
        marco.fill.fore_color.rgb = rgb("EEF1F5")
        if rotulo_vacio:
            texto(x + 2, y, w - 4, h, rotulo_vacio, tam=7.5, color="8A96A3", alin="c")
        return None
    im = _abrir(ruta)
    if recorte:  # [x0, y0, x1, y1] en fracciones de la foto
        im = im.crop((int(recorte[0] * im.width), int(recorte[1] * im.height),
                      int(recorte[2] * im.width), int(recorte[3] * im.height)))
    if im.mode == "RGBA":  # recorte con transparencia: se le saca el margen vacio
        caja_alfa = im.split()[3].point(lambda v: 255 if v > 40 else 0).getbbox()
        if caja_alfa:
            g = int(0.02 * max(im.size))
            im = im.crop((max(0, caja_alfa[0] - g), max(0, caja_alfa[1] - g),
                          min(im.width, caja_alfa[2] + g), min(im.height, caja_alfa[3] + 2 * g)))
    if corte_abajo:  # se muestra solo la parte de arriba de la pieza; el corte queda en el borde del marco
        im = im.crop((0, 0, im.width, int(im.height * (1 - corte_abajo))))
    if rotar:
        im = im.rotate(rotar, expand=True)
    DPI = 300
    if modo == "cubrir":
        pw, ph = int(w / 25.4 * DPI), int(h / 25.4 * DPI)
        esc = max(pw / im.width, ph / im.height)
        if esc < 1:  # solo se achica; nunca se agranda aca
            im = im.resize((max(1, round(im.width * esc)), max(1, round(im.height * esc))), Image.LANCZOS)
        # recorte al aspecto del marco
        asp = w / h
        if im.width / im.height > asp:
            nw = round(im.height * asp)
            x0 = round((im.width - nw) * foco[0])
            im = im.crop((x0, 0, x0 + nw, im.height))
        else:
            nh = round(im.width / asp)
            y0 = round((im.height - nh) * foco[1])
            im = im.crop((0, y0, im.width, y0 + nh))
        px, py, pw_mm, ph_mm = x, y, w, h
    else:
        aw, ah = w - 2 * margen, h - 2 * margen
        esc = min(aw / im.width, ah / im.height)
        pw_mm, ph_mm = im.width * esc, im.height * esc
        px, py = x + (w - pw_mm) / 2, y + (h - ph_mm) / 2
        if corte_abajo:
            esc = min(aw / im.width, (h - margen) / im.height)
            pw_mm, ph_mm = im.width * esc, im.height * esc
            px, py = x + (w - pw_mm) / 2, y + h - ph_mm
    buf = io.BytesIO()
    if im.mode in ("RGBA", "LA", "P"):
        im.convert("RGBA").save(buf, "PNG")
    else:
        im.convert("RGB").save(buf, "JPEG", quality=95, subsampling=0)
    buf.seek(0)
    pic = sp.add_picture(buf, Mm(px), Mm(py), Mm(pw_mm), Mm(ph_mm))
    if modo == "cubrir":
        pic.auto_shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
        geom = pic._element.spPr.find(qn("a:prstGeom"))
        av = geom.find(qn("a:avLst"))
        if av is None:
            av = etree.SubElement(geom, qn("a:avLst"))
        etree.SubElement(av, qn("a:gd"), name="adj", fmla="val %d" % int(min(0.5, radio / min(w, h)) * 100000))
        if borde:
            pic.line.color.rgb = rgb(borde)
            pic.line.width = Pt(0.5)
    elif im.mode == "RGBA" and FOTOS.get("_sombra_piezas"):
        sombra(pic, blur=2.2, dist=0.9, alpha=32)
    return pic


def F(paso, clave):
    r = FOTOS.get(str(paso), {}).get(clave)
    if r and not os.path.isabs(r):
        r = os.path.join(DATOS, r)
    return r


# ---------------------------------------------------------------- tarjeta
def tarjeta(n, x, y, w, h, titulo, linea, fotos, uso=None, reserva=0.0, entra=None):
    """fotos: lista de filas; cada fila es lista de dicts {clave, rotulo, modo, peso, rotar, vacio}."""
    L = LINEAS[linea]
    c = caja(x, y, w, h, relleno=BLANCO, radio=2.6, nombre="Paso %s" % n)
    sombra(c)
    # encabezado
    eh = 11.5
    d = 8.4
    circ = sp.add_shape(MSO_SHAPE.OVAL, Mm(x + 2.6), Mm(y + (eh - d) / 2 + 0.6), Mm(d), Mm(d))
    circ.fill.solid()
    circ.fill.fore_color.rgb = rgb(L["color"])
    circ.line.fill.background()
    sin_sombra(circ)
    tf = circ.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(n)
    r.font.size, r.font.bold, r.font.name = Pt(13), True, FUENTE_B
    r.font.color.rgb = rgb(BLANCO)
    texto(x + 2.6 + d + 2.2, y + 0.6, w - d - 7.4, eh, titulo, tam=10.5, color=AZUL_OSC, negrita=True,
          interlinea=0.9)
    # cuerpo
    pad = 2.6
    bx, by = x + pad, y + eh + 1.4
    bw, bh = w - 2 * pad, h - eh - 1.4 - pad
    if uso:
        uh = 5.0 + 4.3 * len(uso)
        uy = y + h - pad - uh
        caja(bx, uy, bw, uh, relleno=L["claro"], radio=1.4)
        caja(bx, uy, 1.1, uh, relleno=L["color"])
        texto(bx + 2.8, uy + 0.7, bw - 3.8, 3.8, "USO POR PIEZA", tam=7.5, color=L.get("tinta", L["color"]),
              negrita=True, espaciado=80)
        for k, ln in enumerate(uso):
            texto(bx + 2.8, uy + 4.5 + 4.3 * k, bw - 3.8, 4.3, ln, tam=9, color=AZUL_OSC)
        bh -= uh + 1.8
    bh -= reserva
    if entra:  # pastillas con los componentes que se juntan en este paso: [(linea, texto), ...]
        px_ = bx
        for k, (ln_, tx_) in enumerate(entra):
            if k or entra[0][1].startswith("+"):
                pass
            pw_ = len(tx_) * 1.55 + 5.5
            pil = caja(px_, by, pw_, 5.4, relleno=LINEAS[ln_]["color"], radio=2.7)
            texto(px_, by, pw_, 5.4, tx_, tam=7.6, color=BLANCO, negrita=True, alin="c")
            px_ += pw_ + 1.6
            if k < len(entra) - 1:
                texto(px_, by, 3.2, 5.4, "+", tam=10, color=AZUL_OSC, negrita=True, alin="c")
                px_ += 3.2 + 1.6
        by += 5.4 + 2.0
        bh -= 5.4 + 2.0
    rot_h = 4.0
    gap = 1.8
    pesos = [f[0].get("alto", 1.0) for f in fotos]
    util = bh - gap * (len(fotos) - 1)
    cy = by
    for fila, peso in zip(fotos, pesos):
        fh = util * peso / sum(pesos)
        anchos = [f.get("peso", 1.0) for f in fila]
        utilw = bw - gap * (len(fila) - 1)
        cx = bx
        for f, a in zip(fila, anchos):
            fw = utilw * a / sum(anchos)
            con_rotulo = bool(f.get("rotulo"))
            alto_foto = fh - (rot_h if con_rotulo else 0)
            foto(F(n, f["clave"]), cx, cy, fw, alto_foto, modo=f.get("modo", "cubrir"),
                 rotar=f.get("rotar", 0), foco=f.get("foco", (0.5, 0.5)),
                 fondo=f.get("fondo", "F2F4F7" if f.get("modo") == "entera" else BLANCO),
                 borde=f.get("borde", None if f.get("modo") == "entera" else GRIS_BORDE),
                 rotulo_vacio=f.get("vacio"), recorte=f.get("recorte"), corte_abajo=f.get("corte_abajo", 0.0),
                 margen=f.get("margen", 1.5))
            if con_rotulo:
                texto(cx, cy + alto_foto + 0.2, fw, rot_h - 0.2, f["rotulo"], tam=7.2, color=GRIS_TXT, alin="c")
            cx += fw + gap
        cy += fh + gap
    return c


def banda(x, y, w, h, linea, rotulo_w=None):
    L = LINEAS[linea]
    b = caja(x, y, w, h, relleno=L["claro"], radio=4.0)
    # pestaña con el nombre del componente
    tw = rotulo_w or (len(L["nombre"]) * 2.35 + 9)
    t = caja(x + 4, y - 3.2, tw, 6.4, relleno=L["color"], radio=3.2)
    texto(x + 4, y - 3.2, tw, 6.4, L["nombre"], tam=8.5, color=BLANCO, negrita=True, alin="c", espaciado=120)
    return b


