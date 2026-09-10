# -*- coding: utf-8 -*-
"""Mapa de shapes por INDICE de una slide del deck. Uso: py mapear.py <archivo> <idx> [idx...]"""
import io, sys, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pptx import Presentation
from pptx.util import Emu

p = Presentation(sys.argv[1])
for arg in sys.argv[2:]:
    i = int(arg)
    s = p.slides[i]
    print("\n===== SLIDE %d =====" % i)
    for j, sh in enumerate(s.shapes):
        t = ""
        if sh.has_text_frame:
            t = sh.text_frame.text.replace("\n", " / ")[:90]
        elif sh.has_table:
            t = "TABLA %dx%d" % (len(sh.table.rows), len(sh.table.columns))
        elif sh.shape_type == 13:
            t = "<PICTURE>"
        pos = ""
        try:
            pos = "x=%.1f y=%.1f w=%.1f h=%.1f" % (Emu(sh.left).cm, Emu(sh.top).cm,
                                                  Emu(sh.width).cm, Emu(sh.height).cm)
        except Exception:
            pos = "(sin pos)"
        print("  [%2d] id=%-4s %-22s %-34s %s" % (j, sh.shape_id, str(sh.shape_type).split(" ")[0][:22], pos, t))
