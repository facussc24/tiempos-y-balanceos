# -*- coding: utf-8 -*-
"""Agrega la jornada del Dia 6 (08/09/2026) al deck de TryOut IMG, en ES y EN.

La jornada NO se dibuja: se clonan slides del propio deck de Carlos y se les cambia el
texto (memoria informe_tryout_extender_deck_carlos). Se trabaja por INDICE de shape, no
por shape_id: los ids difieren entre ES y EN pero el orden del spTree es identico, asi
que los dos decks salen paralelos por construccion.

Las 9 laminas nuevas entran en el indice 45, ANTES del PLAN DE TRIAL de Carlos (45-46)
y de la lamina de GRACIAS (47): Carlos pidio que esas queden al final. 48 -> 57 slides.

  A  clon de 41  portada Dia 6
  B  clon de 22  resumen de actividades + chips de estado
  C  clon de 23  secuencia de trials T3 a T7
  D  clon de 23  secuencia de trials T8 a T14
  E  clon de 25  estado por cavidad al inicio (2 fotos verticales)
  F  clon de 25  cavidad 4, retiro del suplemento (2 fotos verticales)
  G  clon de 24  tiempos de vacio antes/despues (2 fotos horizontales)
  H  clon de 44  situacion y proximos pasos
  I  clon de 39  plan de accion vigente (tabla: 13 filas heredadas + 3 nuevas)

Contra el Dia 5 hay dos trampas nuevas, y las dos se comen texto de otra jornada EN
SILENCIO si no se controlan:
  · los paneles de titulo+cuerpo son UN parrafo con N runs separados por salto de linea.
    Si el contenido trae menos pares que runs, los runs sobrantes conservan el texto del
    Dia 3. Por eso `set_pairs_exacto` EXIGE que la cantidad coincida.
  · lo mismo con los parrafos de las laminas de situacion (`set_lineas_exacto`).
"""
import copy, io, os, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from datetime import datetime
from pptx import Presentation
from pptx.util import Cm, Emu
from pptx.oxml.ns import qn
from clonlib import dup, mover, set_txt
from contenido_dia6 import ES, EN

MAT = (r"Y:\BARACK\CALIDAD\DOCUMENTACION SGC\PPAP CLIENTES\NOVAX\Tapizadas puerta"
       r"\28- Corrida de Produccion\01- TryOut\T0 IMG BARACK\T0 Dia 5 02-09-2026")
ENT = (r"Y:\BARACK\CALIDAD\DOCUMENTACION SGC\PPAP CLIENTES\NOVAX\Tapizadas puerta"
       r"\28- Corrida de Produccion\01- TryOut\T0 IMG BARACK\T0 Dia 6 08-09-2026")
IMG = os.path.join(os.environ.get("TRYOUT_IMG", ""), "") or None
if IMG is None:
    IMG = (r"C:\Users\FacundoS-PC\AppData\Local\Temp\claude\C--Dev-BarackMercosul"
           r"\6e8e78c2-e21b-47a8-aef7-7eabf7675977\scratchpad\tryout\img")

F_CAV4 = os.path.join(IMG, "cav4_hundimiento.jpg")      # 900x1600
F_CAV3 = os.path.join(IMG, "cav3_marcas_vacio.jpg")     # 900x1600
F_SUPL = os.path.join(IMG, "suplemento_cav4.jpg")       # 900x1600
F_ARRU = os.path.join(IMG, "arruga_cav4.jpg")           # 900x1600
F_ANTE = os.path.join(IMG, "hmi_antes.jpg")             # 770x529 — la captura de 5,3 s
F_AHOR = os.path.join(IMG, "hmi_ahora.jpg")             # 740x515 — la captura de 11,0 s

# ---- indices de shape (identicos en ES y EN; se verifican con AFIRMA antes de escribir)
P_KICK, P_TIT, P_SUB, P_DESC, P_FECHA, P_EQUIPOS = 2, 3, 4, 5, 8, 10
R_TIT, R_PIE, R_ROT = 3, 4, 6
R_ITEMS = [(10, 11), (14, 15), (18, 19), (22, 23), (26, 27)]
R_PANEL_TIT = 29
R_CHIPS = [(31, 33, 34, 32), (36, 38, 39, 37), (41, 43, 44, 42),
           (46, 48, 49, 47), (51, 53, 54, 52)]           # barra, pastilla, texto, label
T_PANEL_TIT, T_PANEL = 29, 30                             # laminas de trials (clon de 23)
I_TIT, I_PIE = 3, 4
I_MARCO_V, I_FOTO_V, I_MARCO_H, I_FOTO_H = 6, 7, 8, 9
I_CAP, I_CHIP_IZQ, I_CHIP_DER, I_PANEL = 10, 12, 14, 16
S_TIT, S_PIE, S_ROT, S_BULLETS, S_ROT2 = 3, 4, 6, 7, 8
S_CAJAS = [(9, 10), (12, 13), (15, 16)]
S_PANEL_TIT, S_PANEL = 18, 20
A_TIT, A_PIE, A_ROT, A_TABLA, A_CIERRE = 3, 4, 6, 7, 9


# --------------------------------------------------------------------- helpers
def sh(slide, idx):
    return slide.shapes[idx]


def AFIRMA(cond, msg):
    if not cond:
        raise AssertionError(msg)


def texto(shape):
    return shape.text_frame.text.strip()


def set_pairs_exacto(shape, vals, donde):
    """Panel titulo+cuerpo: UN parrafo con N runs separados por <a:br>.
    Si sobran runs queda texto de otra jornada; si sobran vals, se pierde contenido."""
    p = shape.text_frame.paragraphs[0]
    runs = p.runs
    AFIRMA(len(runs) == len(vals),
           "%s: el panel tiene %d runs y el contenido trae %d" % (donde, len(runs), len(vals)))
    for r, v in zip(runs, vals):
        r.text = v


def set_lineas_exacto(shape, lineas, donde):
    """Escribe una linea por parrafo. Los parrafos que sobran se BORRAN enteros (si no,
    queda una vineta con texto de otra jornada). Mas lineas que parrafos es un error."""
    paras = list(shape.text_frame.paragraphs)
    AFIRMA(len(lineas) <= len(paras),
           "%s: el shape tiene %d parrafos y el contenido trae %d lineas"
           % (donde, len(paras), len(lineas)))
    for par, txt in zip(paras, lineas):
        if par.runs:
            par.runs[0].text = txt
            for r in par.runs[1:]:
                r.text = ""
        else:
            par.text = txt
    for par in paras[len(lineas):]:
        par._p.getparent().remove(par._p)


def fill_de(shape):
    for tag in ("a:solidFill", "a:gradFill"):
        el = shape._element.spPr.find(qn(tag))
        if el is not None:
            return copy.deepcopy(el)
    return None


def poner_fill(dst, fill):
    if fill is None:
        return
    sp = dst._element.spPr
    for tag in ("a:solidFill", "a:noFill", "a:gradFill", "a:pattFill", "a:blipFill"):
        el = sp.find(qn(tag))
        if el is not None:
            sp.remove(el)
    nuevo = copy.deepcopy(fill)
    ref = sp.find(qn("a:prstGeom"))
    (ref.addnext(nuevo) if ref is not None else sp.insert(0, nuevo))


def reemplazar_imagen(slide, pic, ruta):
    _, rId = slide.part.get_or_add_image_part(ruta)
    pic._element.blipFill.blip.set(qn("r:embed"), rId)


def caja(shape, x, y, w, h=None):
    shape.left, shape.top, shape.width = Cm(x), Cm(y), Cm(w)
    if h is not None:
        shape.height = Cm(h)


def clonar_shape(slide, shape):
    """Duplica un shape dentro de la misma slide (para los rotulos ANTES/AHORA).
    OJO: al insertar shapes cambian los indices de todos los que vienen despues, asi que
    lo que se vaya a tocar por indice se resuelve ANTES de llamar a esta funcion."""
    nuevo = copy.deepcopy(shape._element)
    shape._element.addnext(nuevo)
    for s2 in slide.shapes:
        if s2._element is nuevo:
            return s2
    raise AssertionError("el clon no aparecio en el spTree")


# ------------------------------------------------------------- las 9 laminas
def portada(s, c):
    AFIRMA("TRYOUT IMG" in texto(sh(s, P_TIT)), "portada: indice de titulo")
    set_txt(sh(s, P_KICK), c["kicker"])
    set_txt(sh(s, P_TIT), c["tit_portada"])
    set_txt(sh(s, P_SUB), c["sub"])
    set_txt(sh(s, P_DESC), c["desc"])
    set_txt(sh(s, P_FECHA), c["fecha"])
    set_txt(sh(s, P_EQUIPOS), c["equipos"])


def resumen(s, c):
    AFIRMA(texto(sh(s, R_PANEL_TIT)), "resumen: indice del titulo del panel")
    set_txt(sh(s, R_TIT), c["tit_resumen"])
    set_txt(sh(s, R_PIE), c["pie"])
    set_txt(sh(s, R_ROT), c["rotulo"])
    set_txt(sh(s, R_PANEL_TIT), c["tit_panel_resumen"])
    for (i_t, i_c), (tit, cue) in zip(R_ITEMS, c["items"]):
        set_txt(sh(s, i_t), tit)
        set_txt(sh(s, i_c), cue)
    # El color de estado se COPIA del chip original, pero la foto de los fills se saca
    # ANTES de repintar ninguno: si se copia de un chip que el loop ya piso, se propaga
    # el color equivocado (paso el 03/09: tres chips verdes diciendo EN PROCESO).
    paleta = {}
    for b, p, t, _l in R_CHIPS:
        paleta.setdefault(texto(sh(s, t)), (fill_de(sh(s, b)), fill_de(sh(s, p))))
    for (b, p, t, l), (etq, est) in zip(R_CHIPS, c["chips"]):
        AFIRMA(est in paleta, "chip %r: el estado %r no existe en el deck de Carlos" % (etq, est))
        set_txt(sh(s, l), etq)
        set_txt(sh(s, t), est)
        f_barra, f_past = paleta[est]
        poner_fill(sh(s, b), f_barra)
        poner_fill(sh(s, p), f_past)


def trials(s, c, tit, items, tit_panel, panel, tag):
    set_txt(sh(s, R_TIT), tit)
    set_txt(sh(s, R_PIE), c["pie"])
    set_txt(sh(s, R_ROT), c["rotulo"])
    for (i_t, i_c), (t, cue) in zip(R_ITEMS, items):
        set_txt(sh(s, i_t), t)
        set_txt(sh(s, i_c), cue)
    set_txt(sh(s, T_PANEL_TIT), tit_panel)
    set_pairs_exacto(sh(s, T_PANEL), panel, tag)


def dos_verticales(s, c, tit, chip_i, chip_d, cap, panel, f_izq, f_der, tag):
    """Clon de la lamina de issue (una foto vertical + una horizontal). Las dos fotos
    del Dia 6 son verticales 900x1600: se reusa el marco horizontal como segundo slot
    vertical y se recentra el par en la columna de fotos (x 1,4 a 20,06)."""
    set_txt(sh(s, I_TIT), tit)
    set_txt(sh(s, I_PIE), c["pie"])
    set_txt(sh(s, I_CHIP_IZQ), chip_i)
    set_txt(sh(s, I_CHIP_DER), chip_d)
    caja(sh(s, I_MARCO_V), 3.31, 3.71, 6.92, 11.94)
    caja(sh(s, I_FOTO_V), 3.54, 3.94, 6.46, 11.48)
    caja(sh(s, I_MARCO_H), 11.23, 3.71, 6.92, 11.94)
    caja(sh(s, I_FOTO_H), 11.46, 3.94, 6.46, 11.48)
    reemplazar_imagen(s, sh(s, I_FOTO_V), f_izq)
    reemplazar_imagen(s, sh(s, I_FOTO_H), f_der)
    caja(sh(s, I_CAP), 1.40, 15.95, 18.66, 1.27)
    set_txt(sh(s, I_CAP), cap)
    set_pairs_exacto(sh(s, I_PANEL), panel, tag)


def dos_horizontales(s, c, tag):
    """Clon de la lamina de parametros. Las dos capturas de HMI van lado a lado, a la
    misma ALTURA (5,82 cm) para que se comparen fila contra fila; cada una conserva su
    proporcion (ANTES 770x529, AHORA 740x515), por eso los anchos son levemente distintos.
    La pantalla va como la foto real, con el rotulo ARRIBA: nada la tapa."""
    set_txt(sh(s, I_TIT), c["tit_tiempos"])
    set_txt(sh(s, I_PIE), c["pie"])
    set_txt(sh(s, I_CHIP_IZQ), c["chip_izq_tiempos"])
    set_txt(sh(s, I_CHIP_DER), c["chip_der_tiempos"])
    caja(sh(s, I_MARCO_V), 1.40, 6.20, 8.93, 6.28)
    caja(sh(s, I_FOTO_V), 1.63, 6.43, 8.47, 5.82)
    caja(sh(s, I_MARCO_H), 11.23, 6.20, 8.82, 6.28)
    caja(sh(s, I_FOTO_H), 11.46, 6.43, 8.36, 5.82)
    reemplazar_imagen(s, sh(s, I_FOTO_V), F_ANTE)
    reemplazar_imagen(s, sh(s, I_FOTO_H), F_AHOR)
    cap, panel = sh(s, I_CAP), sh(s, I_PANEL)     # se resuelven ANTES de clonar
    caja(cap, 1.40, 12.85, 18.65, 1.27)
    set_txt(cap, c["cap_tiempos"])
    set_pairs_exacto(panel, c["panel_tiempos"], tag)
    # rotulos de cada captura: se clonan del epigrafe para heredar su formato
    lbl2 = clonar_shape(s, cap)
    lbl1 = clonar_shape(s, cap)
    caja(lbl1, 1.40, 5.20, 8.93, 0.85)
    caja(lbl2, 11.23, 5.20, 8.82, 0.85)
    set_txt(lbl1, c["lbl_antes"])
    set_txt(lbl2, c["lbl_ahora"])


def situacion(s, c, tag):
    set_txt(sh(s, S_TIT), c["tit_situ"])
    set_txt(sh(s, S_PIE), c["pie"])
    set_txt(sh(s, S_ROT), c["rotulo_situ"])
    set_lineas_exacto(sh(s, S_BULLETS), c["bullets"], tag + "/bullets")
    set_txt(sh(s, S_ROT2), c["rotulo2"])
    for (i_f, i_t), lineas in zip(S_CAJAS, c["cajas"]):
        set_lineas_exacto(sh(s, i_t), lineas, tag + "/caja")
    # las tres cajas son verificaciones pendientes: la 2 es la neutra del deck original
    fill_neutro = fill_de(sh(s, S_CAJAS[1][0]))
    for i_f, _ in (S_CAJAS[0], S_CAJAS[2]):
        poner_fill(sh(s, i_f), fill_neutro)
    set_txt(sh(s, S_PANEL_TIT), c["tit_panel_situ"])
    set_lineas_exacto(sh(s, S_PANEL), c["verif"], tag + "/verif")


def plan(s, c, tag):
    """El plan heredado se arrastra FILA POR FILA: las 13 del 31/08 quedan intactas y las
    3 del 08/09 se agregan al final (memoria plan_heredado_se_arrastra_fila_por_fila)."""
    set_txt(sh(s, A_TIT), c["tit_plan"])
    set_txt(sh(s, A_PIE), c["pie"])
    set_txt(sh(s, A_ROT), c["rotulo_plan"])
    set_txt(sh(s, A_CIERRE), c["cierre_plan"])
    forma = sh(s, A_TABLA)
    AFIRMA(forma.has_table, "%s: el shape %d no es la tabla" % (tag, A_TABLA))
    tbl = forma.table
    n0 = len(tbl.rows)
    AFIRMA(n0 == 14, "%s: la tabla heredada tiene %d filas, se esperaban 14" % (tag, n0))
    heredadas = [[c2.text for c2 in r.cells] for r in tbl.rows]
    ultima = tbl._tbl.tr_lst[-1]
    for _ in c["filas_nuevas"]:
        tbl._tbl.append(copy.deepcopy(ultima))
    for k, fila in enumerate(c["filas_nuevas"]):
        celdas = tbl.rows[n0 + k].cells
        set_txt(celdas[0], str(n0 + k))              # Nº: 14, 15, 16
        for j, v in enumerate(fila, start=1):
            set_txt(celdas[j], v)
    # nada de lo heredado se movio
    for i in range(n0):
        AFIRMA([c2.text for c2 in tbl.rows[i].cells] == heredadas[i],
               "%s: se toco la fila heredada %d" % (tag, i))
    # 17 filas no entran con la altura original (0,9 cm): se comprimen las de cuerpo.
    # La altura de fila es un MINIMO — PowerPoint la agranda sola si el texto envuelve.
    for r in list(tbl.rows)[1:]:
        r.height = Cm(0.72)
    forma.height = Cm(1.0 + 0.72 * (len(tbl.rows) - 1))


# ------------------------------------------------------------------ construir
def construir(c, tag):
    prs = Presentation(os.path.join(MAT, c["src"]))
    n0 = len(prs.slides._sldIdLst)
    AFIRMA(n0 == 48, "%s: el deck de origen tiene %d slides, se esperaban 48" % (tag, n0))

    A = dup(prs, 41)
    B = dup(prs, 22)
    C = dup(prs, 23)
    D = dup(prs, 23)
    E = dup(prs, 25)
    F = dup(prs, 25)
    G = dup(prs, 24)
    H = dup(prs, 44)
    I = dup(prs, 39)

    portada(A, c)
    resumen(B, c)
    trials(C, c, c["tit_trials_a"], c["items_a"], c["tit_panel_a"], c["panel_a"], tag + "/T3-T7")
    trials(D, c, c["tit_trials_b"], c["items_b"], c["tit_panel_b"], c["panel_b"], tag + "/T8-T14")
    dos_verticales(E, c, c["tit_defectos"], c["chip_izq_def"], c["chip_der_def"],
                   c["cap_def"], c["panel_def"], F_CAV4, F_CAV3, tag + "/defectos")
    dos_verticales(F, c, c["tit_supl"], c["chip_izq_supl"], c["chip_der_supl"],
                   c["cap_supl"], c["panel_supl"], F_SUPL, F_ARRU, tag + "/suplemento")
    dos_horizontales(G, c, tag + "/tiempos")
    situacion(H, c, tag + "/situacion")
    plan(I, c, tag + "/plan")

    # las 9 entran en el 45: el PLAN DE TRIAL de Carlos y el GRACIAS quedan al final
    for k in range(9):
        mover(prs, n0 + k, 45 + k)

    cambios = 0
    for i, s in enumerate(prs.slides):
        for x in s.shapes:
            if not x.has_text_frame or x.top is None or x.left is None:
                continue
            t = texto(x)
            if t.isdigit() and Emu(x.top).cm > 17.5 and Emu(x.left).cm > 30 and t != str(i + 1):
                set_txt(x, str(i + 1))
                cambios += 1

    cp = prs.core_properties
    cp.author = "Barack Mercosul"
    cp.last_modified_by = "Facundo Santoro"
    cp.title = "Informe Tecnico de TryOut IMG - Proyecto Patagonia"
    cp.modified = datetime(2026, 9, 10, 12, 0, 0)

    os.makedirs(ENT, exist_ok=True)
    destino = os.path.join(ENT, c["out"])
    prs.save(destino)
    print("%-28s %d slides, %d pies renumerados" %
          (c["out"], len(prs.slides._sldIdLst), cambios))
    return destino


if __name__ == "__main__":
    for cfg, tg in ((ES, "ES"), (EN, "EN")):
        construir(cfg, tg)
