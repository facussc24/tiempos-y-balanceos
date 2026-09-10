# -*- coding: utf-8 -*-
"""Clonado de slides dentro del MISMO pptx, con remapeo de relaciones.
Metodo verificado el 29/08 (memoria informe_tryout_extender_deck_carlos)."""
import copy, re
from pptx.oxml.ns import qn

def dup(prs, idx):
    """Clona la slide idx al final del deck. Devuelve la slide nueva."""
    src = prs.slides[idx]
    dest = prs.slides.add_slide(src.slide_layout)
    # vaciar el spTree que trae del layout
    for sh in list(dest.shapes):
        sh._element.getparent().remove(sh._element)
    # copiar los shapes
    for sh in src.shapes:
        dest.shapes._spTree.append(copy.deepcopy(sh._element))
    # copiar el fondo
    bg = src._element.find(qn('p:cSld')).find(qn('p:bg'))
    if bg is not None:
        cSld = dest._element.find(qn('p:cSld'))
        cSld.insert(0, copy.deepcopy(bg))
    # remapear TODA relacion r:embed / r:id / r:link
    # OJO: la notesSlide NO se copia. Una hoja de notas pertenece a UNA sola slide y
    # tiene relacion de vuelta hacia ella; si dos slides apuntan a la misma, el archivo
    # queda invalido y PowerPoint no lo abre (python-pptx si lo abre: no alcanza como control).
    NOTES = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/notesSlide'
    old2new = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype == NOTES:
            continue
        if rel.is_external:
            old2new[rid] = dest.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            old2new[rid] = dest.part.relate_to(rel.target_part, rel.reltype)
    for attr in ('embed', 'id', 'link'):
        q = qn('r:' + attr)
        for el in dest._element.iter():
            v = el.get(q)
            if v is not None and v in old2new:
                el.set(q, old2new[v])
    return dest

def mover(prs, desde, hasta):
    """Mueve la slide del indice `desde` al indice `hasta`."""
    sldIdLst = prs.slides._sldIdLst
    ids = list(sldIdLst)
    el = ids[desde]
    sldIdLst.remove(el)
    sldIdLst.insert(hasta, el)

def byid(slide, sid):
    for sh in slide.shapes:
        if sh.shape_id == sid:
            return sh
    raise KeyError(f"shape_id {sid} no esta en la slide")

def set_pairs(shape, pairs):
    """Reescribe un panel titulo+cuerpo respetando los RUNS (no los parrafos).
    pairs = [(texto_run0, texto_run1), ...] uno por parrafo con 2 runs,
    o [(texto,)] para parrafos de un solo run. None = no tocar ese run."""
    paras = shape.text_frame.paragraphs
    for p, vals in zip(paras, pairs):
        runs = p.runs
        for r, v in zip(runs, vals):
            if v is not None:
                r.text = v

def set_txt(shape, texto):
    """Reemplaza el texto conservando el formato del primer run de cada parrafo."""
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    if not p0.runs:
        tf.text = texto
        return
    p0.runs[0].text = texto
    for r in p0.runs[1:]:
        r.text = ""
    for p in tf.paragraphs[1:]:
        for r in p.runs:
            r.text = ""

def dump_runs(shape):
    out = []
    for i, p in enumerate(shape.text_frame.paragraphs):
        out.append([(j, r.text) for j, r in enumerate(p.runs)])
    return out
