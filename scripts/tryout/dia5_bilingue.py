# -*- coding: utf-8 -*-
"""Agrega la jornada del Dia 5 (02/09/2026) al deck de TryOut IMG, en ES y EN.

La jornada NO se dibuja: se clonan slides del propio deck de Carlos y se les cambia
el texto (memoria informe_tryout_extender_deck_carlos). Se trabaja por INDICE de
shape dentro de la slide, no por shape_id, porque los ids difieren entre ES y EN
pero el orden del spTree es el mismo -> los dos decks salen paralelos por construccion.
"""
import copy, io, os, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pptx import Presentation
from datetime import datetime
from pptx.util import Cm, Emu, Pt
from pptx.oxml.ns import qn
from clonlib import dup, mover, set_pairs, set_txt

BASE = r"C:\Users\FacundoS-PC\OneDrive - BARACK ARGENTINA SRL\Desktop\Informe TryOut IMG - Dia 5 (02-09)"
MAT = os.path.join(BASE, "01- Material")
ENT = os.path.join(BASE, "04- Entregable")
FOTO = os.path.join(MAT, r"fotos\_agujero_vacuum_recorte.jpg")

# indices de shape dentro de cada slide clonada (identicos en ES y EN)
P_TIT, P_SUB, P_DESC, P_FECHA = 5, 6, 7, 10
R_TIT, R_PIE = 3, 4
R_ITEMS = [(10, 11), (14, 15), (18, 19), (22, 23), (26, 27)]
R_CHIPS = [(31, 33, 34, 32), (36, 38, 39, 37), (41, 43, 44, 42),
           (46, 48, 49, 47), (51, 53, 54, 52)]          # barra, pastilla, texto, label
I_TIT, I_PIE, I_MARCO_V, I_FOTO_V = 3, 4, 6, 7
I_MARCO_H, I_FOTO_H, I_CAP, I_KICK, I_FECHA, I_PANEL = 8, 9, 10, 12, 14, 16
S_TIT, S_PIE, S_BULLETS, S_ROT2 = 3, 4, 7, 8
S_CAJAS = [(9, 10), (12, 13), (15, 16)]                  # fondo, texto
S_PANEL_TIT, S_PANEL = 18, 20


def sh(slide, idx):
    return slide.shapes[idx]


def fill_de(shape):
    """Foto del relleno de un shape, para poder copiarlo despues de haberlo pisado."""
    for tag in ("a:solidFill", "a:gradFill"):
        el = shape._element.spPr.find(qn(tag))
        if el is not None:
            return copy.deepcopy(el)
    return None


def poner_fill(dst, fill):
    """Aplica un relleno guardado con fill_de(). None = no toca nada."""
    if fill is None:
        return
    sp = dst._element.spPr
    for tag in ("a:solidFill", "a:noFill", "a:gradFill", "a:pattFill", "a:blipFill"):
        el = sp.find(qn(tag))
        if el is not None:
            sp.remove(el)
    nuevo = copy.deepcopy(fill)
    ref = sp.find(qn("a:prstGeom"))
    (ref.addnext(nuevo) if ref is not None else sp.insert(0, nuevo))


def copiar_fill(dst, src):
    poner_fill(dst, fill_de(src))


def reemplazar_imagen(slide, pic, ruta):
    _, rId = slide.part.get_or_add_image_part(ruta)
    pic._element.blipFill.blip.set(qn("r:embed"), rId)


def set_lineas(shape, lineas):
    """Escribe `lineas` en los parrafos del shape. Si sobran parrafos originales
    (paso el 03/09: 6 parrafos de origen contra 5 lineas nuevas), el sobrante NO
    se deja con su texto viejo — se BORRA el parrafo entero, para que no quede
    una vineta vacia ni, peor, una frase de otra jornada colandose sin tocar."""
    paras = list(shape.text_frame.paragraphs)
    for par, txt in zip(paras, lineas):
        if par.runs:
            par.runs[0].text = txt
            for r in par.runs[1:]:
                r.text = ""
    for par in paras[len(lineas):]:
        par._p.getparent().remove(par._p)


from contenido_dia5 import ES, EN


def construir(cfg):
    prs = Presentation(os.path.join(MAT, cfg["src"]))
    n0 = len(prs.slides._sldIdLst)
    port, resu, inte, situ = (dup(prs, 30), dup(prs, 31), dup(prs, 32), dup(prs, 37))

    # ---- portada
    set_txt(sh(port, P_TIT), cfg["tit_portada"])
    set_txt(sh(port, P_SUB), cfg["sub"])
    set_txt(sh(port, P_DESC), cfg["desc"])
    set_txt(sh(port, P_FECHA), cfg["fecha"])
    # La unica foto del dia es la del agujero, y va en su propia lamina. La que traia
    # esta portada es la del molde delantero de los Dias 3 y 4: reusarla seria mostrarle
    # al cliente la misma imagen por tercera vez, y ademas de la jornada que no es.
    for idx in (3, 2):                        # el PIC y su marco, en ese orden
        e = sh(port, idx)._element
        e.getparent().remove(e)
    for s in list(port.shapes)[12:]:          # letras A/B/C/D de cavidades (solo ES)
        s._element.getparent().remove(s._element)
    # sin foto a la izquierda, el bloque de texto se corre para no dejar media lamina vacia
    NUEVO_X = {2: (3.30, 15.75), 3: (3.13, 16.76), 4: (3.23, 16.26), 5: (3.23, 17.50),
               6: (3.30, 20.00), 7: (3.30, 6.48), 8: (3.30, 6.48),
               9: (10.16, 8.38), 10: (10.16, 8.38)}
    for idx, (x, w) in NUEVO_X.items():
        f = sh(port, idx)
        f.left, f.width = Cm(x), Cm(w)

    # ---- resumen
    set_txt(sh(resu, R_TIT), cfg["tit_resumen"])
    set_txt(sh(resu, R_PIE), cfg["pie"])
    for (i_t, i_c), (tit, cue) in zip(R_ITEMS, cfg["items"]):
        set_txt(sh(resu, i_t), tit)
        set_txt(sh(resu, i_c), cue)
    # Los colores de estado se COPIAN de los chips originales, pero hay que sacarles
    # una foto ANTES de repintar nada: si se copia de un chip que el propio loop ya
    # modifico, se propaga el color equivocado (paso el 03/09: tres chips en verde
    # diciendo EN PROCESO / EN OPTIMIZACION). Lo caza el render, no el codigo.
    paleta = {}                               # estado original -> (fill barra, fill pastilla)
    for b, p, t, _l in R_CHIPS:
        est_orig = sh(resu, t).text_frame.text.strip()
        paleta.setdefault(est_orig, (fill_de(sh(resu, b)), fill_de(sh(resu, p))))
    for (b, p, t, l), (etq, est) in zip(R_CHIPS, cfg["chips"]):
        set_txt(sh(resu, l), etq)
        set_txt(sh(resu, t), est)
        if est in paleta:
            f_barra, f_past = paleta[est]
            poner_fill(sh(resu, b), f_barra)
            poner_fill(sh(resu, p), f_past)

    # ---- intervencion: hay UNA sola foto
    set_txt(sh(inte, I_TIT), cfg["tit_inter"])
    set_txt(sh(inte, I_PIE), cfg["pie"])
    set_txt(sh(inte, I_KICK), cfg["kicker"])
    set_txt(sh(inte, I_FECHA), cfg["fecha_inter"])
    marco, foto = sh(inte, I_MARCO_H), sh(inte, I_FOTO_H)
    cap = sh(inte, I_CAP)
    panel = sh(inte, I_PANEL)
    for s in (sh(inte, I_FOTO_V), sh(inte, I_MARCO_V)):
        s._element.getparent().remove(s._element)
    marco.left, marco.top, marco.width, marco.height = Cm(2.60), Cm(3.71), Cm(15.10), Cm(11.55)
    foto.left, foto.top, foto.width, foto.height = Cm(2.83), Cm(3.94), Cm(14.64), Cm(11.09)
    reemplazar_imagen(inte, foto, FOTO)
    cap.left, cap.top, cap.width = Cm(2.60), Cm(15.55), Cm(15.10)
    set_txt(cap, cfg["caption"])
    # circulo sobre el agujero: en el encuadre lo que domina es la banda mecanizada
    # brillante, no el agujero, que es chico y oscuro. Sin marca se mira otra cosa.
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.dml.color import RGBColor
    marca = inte.shapes.add_shape(MSO_SHAPE.OVAL, Cm(9.60), Cm(8.55), Cm(2.30), Cm(2.30))
    marca.fill.background()
    marca.line.color.rgb = RGBColor(0xC0, 0x6E, 0x2E)
    marca.line.width = Pt(2.5)
    marca.shadow.inherit = False
    set_pairs(panel, [cfg["panel"]])

    # ---- situacion
    set_txt(sh(situ, S_TIT), cfg["tit_situ"])
    set_txt(sh(situ, S_PIE), cfg["pie"])
    set_lineas(sh(situ, S_BULLETS), cfg["bullets"])
    set_txt(sh(situ, S_ROT2), cfg["rotulo2"])
    for (i_f, i_t), lineas in zip(S_CAJAS, cfg["cajas"]):
        set_lineas(sh(situ, i_t), lineas)
    # la caja 1 venia VERDE ("en produccion"): aca las tres son verificaciones pendientes.
    # Se saca la foto del fill ANTES de tocar ninguna, por lo mismo que los chips.
    fill_neutro = fill_de(sh(situ, S_CAJAS[1][0]))
    for i_f, _ in (S_CAJAS[0], S_CAJAS[2]):
        poner_fill(sh(situ, i_f), fill_neutro)
    set_txt(sh(situ, S_PANEL_TIT), cfg["panel_tit"])
    set_lineas(sh(situ, S_PANEL), cfg["verif"])

    # ---- reubicar y renumerar
    for k in range(4):
        mover(prs, n0 + k, 41 + k)
    cambios = 0
    for i, s in enumerate(prs.slides):
        for x in s.shapes:
            if not x.has_text_frame or x.top is None or x.left is None:
                continue
            t = x.text_frame.text.strip()
            if t.isdigit() and Emu(x.top).cm > 17.5 and Emu(x.left).cm > 30 and t != str(i + 1):
                set_txt(x, str(i + 1))
                cambios += 1
    # Propiedades: el EN salia con last_modified_by="python-pptx" y title="PowerPoint
    # Presentation". Eso se ve en Archivo > Informacion y delata como se genero.
    cp = prs.core_properties
    cp.author = "Barack Mercosul"
    cp.last_modified_by = "Facundo Santoro"
    cp.title = "Informe Tecnico de TryOut IMG - Proyecto Patagonia"
    cp.modified = datetime(2026, 9, 3, 12, 0, 0)
    destino = os.path.join(ENT, cfg["out"])
    prs.save(destino)
    print("%-26s %d slides, %d pies renumerados" % (cfg["out"], len(prs.slides._sldIdLst), cambios))
    return destino


for cfg in (ES, EN):
    construir(cfg)
