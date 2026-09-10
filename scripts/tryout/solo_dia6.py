# -*- coding: utf-8 -*-
"""Del acumulado saca el pptx de la jornada sola (las 9 laminas del Dia 6)."""
import io, os, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation
from pptx.util import Emu
from clonlib import set_txt

ENT = (r"Y:\BARACK\CALIDAD\DOCUMENTACION SGC\PPAP CLIENTES\NOVAX\Tapizadas puerta"
       r"\28- Corrida de Produccion\01- TryOut\T0 IMG BARACK\T0 Dia 6 08-09-2026")
SRC = os.path.join(ENT, "TryOut_IMG_Dia1a6.pptx")
OUT = os.path.join(ENT, "TryOut_IMG_Dia6.pptx")

QUEDAN = set(range(45, 54))          # portada + resumen + 2 trials + 3 fotos + situacion + plan

prs = Presentation(SRC)
assert len(prs.slides._sldIdLst) == 57, "el acumulado no tiene 57 slides"
lst = prs.slides._sldIdLst
for i in reversed(range(len(lst))):
    if i not in QUEDAN:
        el = lst[i]
        rId = el.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        prs.part.drop_rel(rId)
        lst.remove(el)

# renumerar el pie: el numero es indice + 1, y la portada no lleva
cambios = 0
for i, s in enumerate(prs.slides):
    for sh in s.shapes:
        if not sh.has_text_frame or sh.top is None or sh.left is None:
            continue
        t = sh.text_frame.text.strip()
        if t.isdigit() and Emu(sh.top).cm > 17.5 and Emu(sh.left).cm > 30 and t != str(i + 1):
            set_txt(sh, str(i + 1))
            cambios += 1

prs.save(OUT)
print("guardado: %s  (%d slides, %d pies renumerados)"
      % (OUT, len(prs.slides._sldIdLst), cambios))
