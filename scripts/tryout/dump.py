import sys
from pptx import Presentation
from pptx.util import Emu

path = sys.argv[1]
lo = int(sys.argv[2]) if len(sys.argv) > 2 else 0
hi = int(sys.argv[3]) if len(sys.argv) > 3 else 10**9
prs = Presentation(path)
print("FILE:", path)
print("SLIDES:", len(prs.slides), "SIZE:", prs.slide_width, prs.slide_height)


def emit(sh, indent="  "):
    try:
        geom = "L=%.2f T=%.2f W=%.2f H=%.2f" % (
            Emu(sh.left).cm, Emu(sh.top).cm, Emu(sh.width).cm, Emu(sh.height).cm)
    except Exception:
        geom = "geom=?"
    print("%s[%s] name=%r %s" % (indent, sh.shape_type, sh.name, geom))
    if sh.shape_type == 6:
        for sub in sh.shapes:
            emit(sub, indent + "    ")
        return
    if sh.has_text_frame:
        for p in sh.text_frame.paragraphs:
            t = "".join(r.text for r in p.runs)
            if t.strip():
                sizes, colors, fonts, bolds = set(), set(), set(), set()
                for r in p.runs:
                    if r.font.size:
                        sizes.add(r.font.size.pt)
                    try:
                        if r.font.color and r.font.color.type is not None:
                            colors.add(str(r.font.color.rgb))
                    except Exception:
                        pass
                    if r.font.name:
                        fonts.add(r.font.name)
                    bolds.add(r.font.bold)
                print("%s    TXT: %r sz=%s col=%s f=%s b=%s" % (
                    indent, t, sorted(sizes), sorted(colors), sorted(fonts), sorted(bolds, key=str)))
    if getattr(sh, "has_table", False) and sh.has_table:
        for ri, row in enumerate(sh.table.rows):
            print("%s    ROW%d: %r" % (indent, ri, [c.text for c in row.cells]))


for i, s in enumerate(prs.slides):
    if i < lo or i > hi:
        continue
    print("=" * 80)
    print("--- SLIDE idx=%d (page %d) layout=%s" % (i, i + 1, s.slide_layout.name))
    for sh in s.shapes:
        emit(sh)
