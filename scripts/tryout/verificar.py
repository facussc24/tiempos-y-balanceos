# -*- coding: utf-8 -*-
"""Controles del entregable. Cada uno puede dar ROJO."""
import io, os, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pptx import Presentation

BASE = r"C:\Users\FacundoS-PC\OneDrive - BARACK ARGENTINA SRL\Desktop\Informe TryOut IMG - Dia 5 (02-09)"
MAT, ENT = os.path.join(BASE, "01- Material"), os.path.join(BASE, "04- Entregable")

fallas = []


def chequear(nombre, ok, detalle=""):
    print("  %-52s %s %s" % (nombre, "OK  " if ok else "ROJO", detalle))
    if not ok:
        fallas.append(nombre)


def textos(s):
    return [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]


def fotos(s):
    return sum(1 for sh in s.shapes if sh.shape_type == 13)


print("\n=== 1. el bloque heredado no se toco (slides 0..40) ===")
for base_f, new_f in (("TryOut_IMG_Dia1a4.pptx", "TryOut_IMG_Dia1a5.pptx"),
                      ("TryOut_IMG_Day1to4_EN.pptx", "TryOut_IMG_Day1to5_EN.pptx")):
    b, n = Presentation(os.path.join(MAT, base_f)), Presentation(os.path.join(ENT, new_f))
    dif = sum(1 for i in range(41) if textos(b.slides[i]) != textos(n.slides[i]))
    chequear(new_f + " vs original", dif == 0, "%d diferencias" % dif)

print("\n=== 2. las slides de Carlos siguen al final ===")
for f in ("TryOut_IMG_Dia1a5.pptx", "TryOut_IMG_Day1to5_EN.pptx"):
    p = Presentation(os.path.join(ENT, f))
    ult = " | ".join(t.strip()[:30] for t in textos(p.slides[47]) if t.strip())
    chequear(f + " -> ultima slide", "GRACIAS" in ult.upper() or "THANK" in ult.upper(), ult[:40])

print("\n=== 3. paridad estructural ES / EN (48 slides) ===")
es = Presentation(os.path.join(ENT, "TryOut_IMG_Dia1a5.pptx"))
en = Presentation(os.path.join(ENT, "TryOut_IMG_Day1to5_EN.pptx"))
chequear("misma cantidad de slides", len(es.slides._sldIdLst) == len(en.slides._sldIdLst))
# lo preexistente se mide contra los originales, no se da por bueno a ojo
oe = Presentation(os.path.join(MAT, "TryOut_IMG_Dia1a4.pptx"))
oi = Presentation(os.path.join(MAT, "TryOut_IMG_Day1to4_EN.pptx"))
previo = {i for i in range(44)
          if len(oe.slides[i].shapes) != len(oi.slides[i].shapes)
          or fotos(oe.slides[i]) != fotos(oi.slides[i])}
malos = [i for i, (a, b) in enumerate(zip(es.slides, en.slides))
         if len(a.shapes) != len(b.shapes) or fotos(a) != fotos(b)]
nuevos = [i for i in malos if i not in previo]
chequear("mismos shapes y fotos por slide (sin lo heredado)", not nuevos,
         "nuevas que difieren: %s" % (nuevos or "ninguna"))
if previo:
    print("  %-52s %s %s" % ("(preexistente en los decks de Carlos)", "nota",
                             "slides %s: la portada ES lleva las letras de cavidades, la EN no"
                             % sorted(previo)))

print("\n=== 4. la jornada nueva esta completa y en su lugar ===")
ESPERADO_ES = ["TRYOUT IMG \u2014 D\u00cdA 5", "Resumen de actividades \u2014 D\u00eda 5",
               "agujeros de vac\u00edo en el extremo", "al cierre del 02/09"]
ESPERADO_EN = ["TRYOUT IMG \u2014 DAY 5", "Activity summary \u2014 Day 5",
               "vacuum holes at the end", "at the close of 02/09"]
for p, esp, tag in ((es, ESPERADO_ES, "ES"), (en, ESPERADO_EN, "EN")):
    for k, frag in enumerate(esp):
        cuerpo = " ".join(textos(p.slides[41 + k]))
        chequear("%s slide %d contiene %r" % (tag, 41 + k, frag[:34]), frag in cuerpo)

print("\n=== 5. pie y numeracion de la jornada nueva ===")
for p, dia, tag in ((es, "D\u00eda 5", "ES"), (en, "Day 5", "EN")):
    for k in (1, 2, 3):                       # la portada no lleva pie numerado
        cuerpo = textos(p.slides[41 + k])
        chequear("%s slide %d: pie dice %s y 02/09/2026" % (tag, 41 + k, dia),
                 any(dia in t and "02/09/2026" in t for t in cuerpo))
        chequear("%s slide %d: numero de pagina = %d" % (tag, 41 + k, 42 + k),
                 any(t.strip() == str(42 + k) for t in cuerpo))

print("\n=== 6. control GEMELO: el control sabe dar rojo ===")
# si comparo el deck nuevo contra si mismo desplazado, el control 1 TIENE que fallar
b = Presentation(os.path.join(MAT, "TryOut_IMG_Dia1a4.pptx"))
n = Presentation(os.path.join(ENT, "TryOut_IMG_Dia1a5.pptx"))
dif_gemelo = sum(1 for i in range(41) if textos(b.slides[i]) != textos(n.slides[i + 1]))
chequear("gemelo: comparar contra el deck corrido da distinto", dif_gemelo > 0,
         "%d diferencias (tiene que ser > 0)" % dif_gemelo)

print("\n=== 7. los archivos del entregable existen y pesan ===")
for f in ("TryOut_IMG_Dia1a5.pptx", "TryOut_IMG_Dia5.pptx", "TryOut_IMG_Day1to5_EN.pptx",
          "_vista previa - TryOut_IMG_Dia1a5.pdf", "_vista previa - TryOut_IMG_Dia5.pdf"):
    ruta = os.path.join(ENT, f)
    ok = os.path.exists(ruta) and os.path.getsize(ruta) > 100_000
    chequear(f, ok, "%d KB" % (os.path.getsize(ruta) // 1024 if os.path.exists(ruta) else 0))

print("\n=== 8bis. el COLOR de cada chip dice lo mismo que su TEXTO ===")
# El 03/09 tres chips salieron VERDES diciendo EN PROCESO / EN OPTIMIZACION, porque el
# fill se copiaba de un chip que el mismo loop ya habia repintado. Ningun control lo vio:
# solo el render. Criterio: el color de un estado tiene que ser el que ese estado tiene
# en el deck ORIGINAL de Carlos.
from pptx.oxml.ns import qn as _qn

def color_pastilla(shape):
    sf = shape._element.spPr.find(_qn("a:solidFill"))
    if sf is None:
        return None
    c = sf.find(_qn("a:srgbClr"))
    if c is not None:
        return c.get("val")
    c = sf.find(_qn("a:schemeClr"))
    return "scheme:" + c.get("val") if c is not None else "?"

CHIP_TXT, CHIP_PAST = [34, 39, 44, 49, 54], [33, 38, 43, 48, 53]
for base_f, new_f, tag in (("TryOut_IMG_Dia1a4.pptx", "TryOut_IMG_Dia1a5.pptx", "ES"),
                           ("TryOut_IMG_Day1to4_EN.pptx", "TryOut_IMG_Day1to5_EN.pptx", "EN")):
    ref = {}                                   # estado -> color, segun el deck de Carlos
    o = Presentation(os.path.join(MAT, base_f))
    for sl in (31, 22, 4):                     # las laminas de resumen ya existentes
        for it, ip in zip(CHIP_TXT, CHIP_PAST):
            try:
                ref.setdefault(o.slides[sl].shapes[it].text_frame.text.strip(),
                               color_pastilla(o.slides[sl].shapes[ip]))
            except Exception:
                pass
    n = Presentation(os.path.join(ENT, new_f))
    for it, ip in zip(CHIP_TXT, CHIP_PAST):
        est = n.slides[42].shapes[it].text_frame.text.strip()
        got, want = color_pastilla(n.slides[42].shapes[ip]), ref.get(est)
        chequear("%s chip %r -> color del estado" % (tag, est[:22]), want is None or got == want,
                 "tiene %s, el deck usa %s" % (got, want))

print("\n=== 8. el JUEZ del pptx es PowerPoint, no python-pptx ===")
# python-pptx abre archivos que PowerPoint rechaza: el 03/09 el deck EN quedo invalido
# (dos slides compartiendo la misma notesSlide) y los 7 controles de arriba daban verde.
import shutil
import win32com.client
app = win32com.client.Dispatch("PowerPoint.Application")
for f in ("TryOut_IMG_Dia1a5.pptx", "TryOut_IMG_Dia5.pptx", "TryOut_IMG_Day1to5_EN.pptx"):
    try:
        pres = app.Presentations.Open(os.path.join(ENT, f), WithWindow=False)
        n = pres.Slides.Count
        pres.Close()
        chequear("PowerPoint abre " + f, True, "%d slides" % n)
    except Exception as e:
        chequear("PowerPoint abre " + f, False, str(e)[:70])

# gemelo del control 8: un pptx corrompido a proposito TIENE que dar rojo.
# el gemelo queda en el scratchpad, no se borra (guard de borrado en lote).
gemelo = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_gemelo_roto.pptx")
shutil.copy(os.path.join(ENT, "TryOut_IMG_Dia5.pptx"), gemelo)
with open(gemelo, "r+b") as fh:
    fh.seek(600)
    fh.write(b"\x00" * 400)
try:
    pres = app.Presentations.Open(gemelo, WithWindow=False)
    pres.Close()
    abrio = True
except Exception:
    abrio = False
chequear("gemelo: PowerPoint RECHAZA un pptx corrompido", not abrio)
app.Quit()

print("\n" + ("=" * 62))
print("RESUMEN: %d control(es) en ROJO" % len(fallas))
for f in fallas:
    print("   - " + f)
sys.exit(1 if fallas else 0)
