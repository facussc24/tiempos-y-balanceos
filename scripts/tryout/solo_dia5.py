# -*- coding: utf-8 -*-
"""Del acumulado saca el pptx de la jornada sola (portada + 3 laminas del Dia 5)."""
import io, os, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pptx import Presentation
from pptx.util import Emu
from clonlib import set_txt

BASE = r"C:\Users\FacundoS-PC\OneDrive - BARACK ARGENTINA SRL\Desktop\Informe TryOut IMG - Dia 5 (02-09)\04- Entregable"
SRC = os.path.join(BASE, "TryOut_IMG_Dia1a5.pptx")
OUT = os.path.join(BASE, "TryOut_IMG_Dia5.pptx")

QUEDAN = {41, 42, 43, 44}          # portada + resumen + intervencion + situacion

prs = Presentation(SRC)
lst = prs.slides._sldIdLst
for i in reversed(range(len(lst))):
    if i not in QUEDAN:
        el = lst[i]
        rId = el.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        prs.part.drop_rel(rId)
        lst.remove(el)

# renumerar el pie: el numero es indice + 1, y la portada no lleva
cambios = 0
for i, s in enumerate(prs.slides):
    for sh in s.shapes:
        if not sh.has_text_frame or sh.top is None or sh.left is None:
            continue
        t = sh.text_frame.text.strip()
        if t.isdigit() and Emu(sh.top).cm > 17.5 and Emu(sh.left).cm > 30:
            if t != str(i + 1):
                set_txt(sh, str(i + 1))
                cambios += 1

prs.save(OUT)
print("guardado: %s  (%d slides, %d pies renumerados)"
      % (OUT, len(prs.slides._sldIdLst), cambios))
