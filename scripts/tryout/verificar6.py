# -*- coding: utf-8 -*-
"""Controles del entregable del Dia 6. Cada uno puede dar ROJO.

Los que existen porque ya fallaron una vez:
  · 6 y 10 son GEMELOS: le dan de comer al control algo que TIENE que rechazar. Un control
    que no se vio fallar por su motivo no esta probado.
  · 9 abre los archivos con PowerPoint: python-pptx abrio sin chistar un deck que PowerPoint
    rechazaba (03/09, dos slides compartiendo la misma notesSlide) con 7 controles en verde.
  · 8bis mira el COLOR de cada chip contra el estado que ese color tiene en el deck de Carlos.
  · 8ter cuenta el arrastre del plan de accion FILA POR FILA.
  · 4bis busca texto de la jornada CLONADA que se haya quedado pegado en la nueva.
  · 4ter clava cual captura de HMI va en cada slot: ese par salio INVERTIDO el
    10/09 y ningun control de texto lo marcaba (el texto existia y era coherente).
"""
import io, os, shutil, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pptx import Presentation
from pptx.oxml.ns import qn

MAT = (r"Y:\BARACK\CALIDAD\DOCUMENTACION SGC\PPAP CLIENTES\NOVAX\Tapizadas puerta"
       r"\28- Corrida de Produccion\01- TryOut\T0 IMG BARACK\T0 Dia 5 02-09-2026")
ENT = (r"Y:\BARACK\CALIDAD\DOCUMENTACION SGC\PPAP CLIENTES\NOVAX\Tapizadas puerta"
       r"\28- Corrida de Produccion\01- TryOut\T0 IMG BARACK\T0 Dia 6 08-09-2026")
NUEVAS = range(45, 54)                       # las 9 laminas del Dia 6, 0-based
fallas = []


def chequear(nombre, ok, detalle=""):
    print("  %-56s %s %s" % (nombre, "OK  " if ok else "ROJO", detalle))
    if not ok:
        fallas.append(nombre)


def textos(s):
    return [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]


def fotos(s):
    return sum(1 for sh in s.shapes if sh.shape_type == 13)


print("\n=== 1. el bloque heredado no se toco (slides 0..44) ===")
for base_f, new_f in (("TryOut_IMG_Dia1a5.pptx", "TryOut_IMG_Dia1a6.pptx"),
                      ("TryOut_IMG_Day1to5_EN.pptx", "TryOut_IMG_Day1to6_EN.pptx")):
    b, n = Presentation(os.path.join(MAT, base_f)), Presentation(os.path.join(ENT, new_f))
    dif = [i for i in range(45) if textos(b.slides[i]) != textos(n.slides[i])]
    chequear(new_f + " vs original (texto)", not dif, "diferencias en %s" % (dif or "ninguna"))
    dif2 = [i for i in range(45) if len(b.slides[i].shapes) != len(n.slides[i].shapes)
            or fotos(b.slides[i]) != fotos(n.slides[i])]
    chequear(new_f + " vs original (shapes y fotos)", not dif2, "diferencias en %s" % (dif2 or "ninguna"))

print("\n=== 2. lo de Carlos sigue al final ===")
for f in ("TryOut_IMG_Dia1a6.pptx", "TryOut_IMG_Day1to6_EN.pptx"):
    p = Presentation(os.path.join(ENT, f))
    chequear(f + " -> 57 slides", len(p.slides._sldIdLst) == 57, str(len(p.slides._sldIdLst)))
    ult = " | ".join(t.strip()[:30] for t in textos(p.slides[56]) if t.strip())
    chequear(f + " -> ultima slide", "GRACIAS" in ult.upper() or "THANK" in ult.upper(), ult[:40])
    plan_c = " ".join(textos(p.slides[55]))
    chequear(f + " -> PLAN DE TRIAL de Carlos en la 56", "Trial" in plan_c, plan_c.strip()[:44])

print("\n=== 3. paridad estructural ES / EN ===")
es = Presentation(os.path.join(ENT, "TryOut_IMG_Dia1a6.pptx"))
en = Presentation(os.path.join(ENT, "TryOut_IMG_Day1to6_EN.pptx"))
chequear("misma cantidad de slides", len(es.slides._sldIdLst) == len(en.slides._sldIdLst))
oe = Presentation(os.path.join(MAT, "TryOut_IMG_Dia1a5.pptx"))
oi = Presentation(os.path.join(MAT, "TryOut_IMG_Day1to5_EN.pptx"))
previo = {i for i in range(48)
          if len(oe.slides[i].shapes) != len(oi.slides[i].shapes)
          or fotos(oe.slides[i]) != fotos(oi.slides[i])}
malos = [i for i, (a, b) in enumerate(zip(es.slides, en.slides))
         if len(a.shapes) != len(b.shapes) or fotos(a) != fotos(b)]
nuevos = [i for i in malos if i not in previo and i not in range(45, 48)]
# las heredadas que ya diferian en el deck de Carlos se corren 9 lugares al insertar
corridas = [i for i in malos if (i - 9) in previo and i >= 54]
chequear("mismos shapes y fotos por slide (sin lo heredado)",
         not [i for i in nuevos if i not in corridas],
         "difieren: %s" % ([i for i in nuevos if i not in corridas] or "ninguna"))

print("\n=== 4. la jornada nueva esta completa y en su lugar ===")
ESPERADO_ES = ["TRYOUT IMG — DÍA 6", "Resumen de actividades — Día 6",
               "Secuencia de trials — T3 a T7", "Secuencia de trials — T8 a T14",
               "Estado por cavidad al inicio de la jornada — T3",
               "Cavidad 4 — retiro del suplemento",
               "Tiempos de vacío del molde inferior — antes y después",
               "al cierre del 08/09", "acciones técnicas al 08/09"]
ESPERADO_EN = ["TRYOUT IMG — DAY 6", "Activity summary — Day 6",
               "Trial sequence — T3 to T7", "Trial sequence — T8 to T14",
               "Status per cavity at the start of the day — T3",
               "Cavity 4 — removal of the shim",
               "Lower mold vacuum times — before and after",
               "at the close of 08/09", "technical actions as of 08/09"]
for p, esp, tag in ((es, ESPERADO_ES, "ES"), (en, ESPERADO_EN, "EN")):
    for k, frag in enumerate(esp):
        cuerpo = " ".join(textos(p.slides[45 + k]))
        chequear("%s slide %d contiene %r" % (tag, 45 + k, frag[:36]), frag in cuerpo)

print("\n=== 4bis. no quedo texto de la jornada que se clono ===")
PROHIBIDO_ES = ["Marcación de cavidades", "top roll trasero", "Ajustes aplicados",
                "Secuencia de vacío", "Bastidor y producción", "Estirado del vinilo",
                "Relación con el Día 2", "trimming tool", "Haartz",
                "ajustes de herramental", "al 31/08", "Día 3", "Día 5", "02/09", "29/08"]
PROHIBIDO_EN = ["Cavity marking", "rear top roll", "Applied adjustments",
                "Vacuum sequence", "Frame and production", "Vinyl stretching",
                "Relation to Day 2", "trimming tool", "Haartz",
                "tooling and process", "as of 31/08", "Day 3", "Day 5", "02/09", "29/08"]
for p, prohibido, tag in ((es, PROHIBIDO_ES, "ES"), (en, PROHIBIDO_EN, "EN")):
    cuerpo = " ".join(" ".join(textos(p.slides[i])) for i in NUEVAS)
    for frag in prohibido:
        chequear("%s: la jornada nueva NO dice %r" % (tag, frag[:30]), frag not in cuerpo)

print("\n=== 4ter. la lamina de tiempos: cada captura en su lugar ===")
# El 10/09 este par salio INVERTIDO: las dos capturas se mandaron a las 16:02 y yo las ordene
# por la hora de envio, pero Carlos las identifico a las 16:03 y el ANTES es el de 5,3 s
# (mandado 7 s DESPUES). Lo confirman la pantalla de las 16:12:48 y el audio de las 16:17.
# Las dos capturas tienen tamano distinto: el tamano prueba CUAL foto quedo en cada slot.
G = 51                                        # indice 0-based de la lamina de tiempos
ANTES_PX, AHORA_PX = (770, 529), (740, 515)   # 5,3 s  /  11,0 s
FLECHA = " s  \u2192  "
for p, marca_a, marca_b, val_a, val_b, tag in (
        (es, "ANTES", "AHORA", "5,3", "11,0", "ES"),
        (en, "BEFORE", "AFTER", "5.3", "11.0", "EN")):
    s = p.slides[G]
    pics = sorted((sh for sh in s.shapes                 # el logo tambien es una imagen
                   if sh.shape_type == 13 and sh.top > 4 * 360000), key=lambda sh: sh.left)
    chequear("%s tiempos: hay 2 capturas de HMI" % tag, len(pics) == 2, "%d" % len(pics))
    if len(pics) == 2:
        chequear("%s tiempos: a la IZQUIERDA va la captura de %s s" % (tag, val_a),
                 pics[0].image.size == ANTES_PX, "%s" % (pics[0].image.size,))
        chequear("%s tiempos: a la DERECHA va la captura de %s s" % (tag, val_b),
                 pics[1].image.size == AHORA_PX, "%s" % (pics[1].image.size,))
    rot = [x.strip() for x in textos(s) if x.strip().startswith((marca_a, marca_b))]
    chequear("%s tiempos: el rotulo %s dice %s s" % (tag, marca_a, val_a),
             any(x.startswith(marca_a) and val_a in x for x in rot), " | ".join(rot))
    chequear("%s tiempos: el rotulo %s dice %s s" % (tag, marca_b, val_b),
             any(x.startswith(marca_b) and val_b in x for x in rot))
    cuerpo = " ".join(textos(s))
    chequear("%s tiempos: el panel dice %s s -> %s s" % (tag, val_a, val_b),
             (val_a + FLECHA + val_b + " s") in cuerpo)
    chequear("%s tiempos: el panel NO dice %s s -> %s s" % (tag, val_b, val_a),
             (val_b + FLECHA + val_a + " s") not in cuerpo)

# gemelo de 4ter: con las dos capturas intercambiadas el control TIENE que dar rojo
tmp_g = os.path.join(os.environ.get("TEMP", "."), "_gemelo_tiempos.pptx")
pg = Presentation(os.path.join(ENT, "TryOut_IMG_Dia1a6.pptx"))
pics_g = sorted((sh for sh in pg.slides[G].shapes
                 if sh.shape_type == 13 and sh.top > 4 * 360000), key=lambda sh: sh.left)
blips = [sh._element.find(qn("p:blipFill")).find(qn("a:blip")) for sh in pics_g]
r0, r1 = (b.get(qn("r:embed")) for b in blips)
blips[0].set(qn("r:embed"), r1)
blips[1].set(qn("r:embed"), r0)
pg.save(tmp_g)
pics_g = sorted((sh for sh in Presentation(tmp_g).slides[G].shapes
                 if sh.shape_type == 13 and sh.top > 4 * 360000), key=lambda sh: sh.left)
chequear("gemelo 4ter: con las capturas cambiadas el control da ROJO",
         pics_g[0].image.size != ANTES_PX, "izquierda quedo en %s" % (pics_g[0].image.size,))
os.remove(tmp_g)

print("\n=== 5. pie y numeracion de la jornada nueva ===")
for p, dia, tag in ((es, "Día 6", "ES"), (en, "Day 6", "EN")):
    for i in NUEVAS:
        if i == 45:
            continue                          # la portada no lleva pie numerado
        cuerpo = textos(p.slides[i])
        chequear("%s slide %d: pie dice %s y 08/09/2026" % (tag, i, dia),
                 any(dia in t and "08/09/2026" in t for t in cuerpo))
        chequear("%s slide %d: numero de pagina = %d" % (tag, i, i + 1),
                 any(t.strip() == str(i + 1) for t in cuerpo))

print("\n=== 6. control GEMELO: el control 1 sabe dar rojo ===")
b = Presentation(os.path.join(MAT, "TryOut_IMG_Dia1a5.pptx"))
n = Presentation(os.path.join(ENT, "TryOut_IMG_Dia1a6.pptx"))
dif_gemelo = sum(1 for i in range(45) if textos(b.slides[i]) != textos(n.slides[i + 1]))
chequear("gemelo: comparar contra el deck corrido da distinto", dif_gemelo > 0,
         "%d diferencias (tiene que ser > 0)" % dif_gemelo)

print("\n=== 7. los archivos del entregable existen y pesan ===")
for f in ("TryOut_IMG_Dia1a6.pptx", "TryOut_IMG_Dia6.pptx", "TryOut_IMG_Day1to6_EN.pptx",
          "_vista previa - TryOut_IMG_Dia1a6.pdf", "_vista previa - TryOut_IMG_Dia6.pdf",
          "_vista previa - TryOut_IMG_Day1to6_EN.pdf"):
    ruta = os.path.join(ENT, f)
    ok = os.path.exists(ruta) and os.path.getsize(ruta) > 100_000
    chequear(f, ok, "%d KB" % (os.path.getsize(ruta) // 1024 if os.path.exists(ruta) else 0))

print("\n=== 8bis. el COLOR de cada chip dice lo mismo que su TEXTO ===")


def color_pastilla(shape):
    sf = shape._element.spPr.find(qn("a:solidFill"))
    if sf is None:
        return None
    c = sf.find(qn("a:srgbClr"))
    if c is not None:
        return c.get("val")
    c = sf.find(qn("a:schemeClr"))
    return "scheme:" + c.get("val") if c is not None else "?"


CHIP_TXT, CHIP_PAST = [34, 39, 44, 49, 54], [33, 38, 43, 48, 53]
for base_f, new_f, tag in (("TryOut_IMG_Dia1a5.pptx", "TryOut_IMG_Dia1a6.pptx", "ES"),
                           ("TryOut_IMG_Day1to5_EN.pptx", "TryOut_IMG_Day1to6_EN.pptx", "EN")):
    ref = {}
    o = Presentation(os.path.join(MAT, base_f))
    for sl in (42, 31, 22, 4):                 # las laminas de resumen ya existentes
        for it, ip in zip(CHIP_TXT, CHIP_PAST):
            try:
                ref.setdefault(o.slides[sl].shapes[it].text_frame.text.strip(),
                               color_pastilla(o.slides[sl].shapes[ip]))
            except Exception:
                pass
    nn = Presentation(os.path.join(ENT, new_f))
    for it, ip in zip(CHIP_TXT, CHIP_PAST):
        est = nn.slides[46].shapes[it].text_frame.text.strip()
        got, want = color_pastilla(nn.slides[46].shapes[ip]), ref.get(est)
        chequear("%s chip %r -> color del estado" % (tag, est[:22]),
                 want is not None and got == want, "tiene %s, el deck usa %s" % (got, want))

print("\n=== 8ter. el plan de accion se arrastro fila por fila ===")
for base_f, new_f, tag in (("TryOut_IMG_Dia1a5.pptx", "TryOut_IMG_Dia1a6.pptx", "ES"),
                           ("TryOut_IMG_Day1to5_EN.pptx", "TryOut_IMG_Day1to6_EN.pptx", "EN")):
    o = Presentation(os.path.join(MAT, base_f)).slides[39].shapes[7].table
    nn = Presentation(os.path.join(ENT, new_f)).slides[53].shapes[7].table
    viejas = [[c.text for c in r.cells] for r in o.rows]
    nuevas = [[c.text for c in r.cells] for r in nn.rows]
    chequear("%s: la tabla pasa de %d a 17 filas" % (tag, len(viejas)), len(nuevas) == 17,
             "%d filas" % len(nuevas))
    chequear("%s: las 14 filas heredadas estan intactas" % tag, nuevas[:14] == viejas)
    chequear("%s: las 3 filas nuevas estan numeradas 14, 15, 16" % tag,
             [f[0] for f in nuevas[14:]] == ["14", "15", "16"],
             str([f[0] for f in nuevas[14:]]))

print("\n=== 9. el JUEZ del pptx es PowerPoint, no python-pptx ===")
import win32com.client
app = win32com.client.Dispatch("PowerPoint.Application")
for f in ("TryOut_IMG_Dia1a6.pptx", "TryOut_IMG_Dia6.pptx", "TryOut_IMG_Day1to6_EN.pptx"):
    try:
        pres = app.Presentations.Open(os.path.join(ENT, f), WithWindow=False)
        cant = pres.Slides.Count
        pres.Close()
        chequear("PowerPoint abre " + f, True, "%d slides" % cant)
    except Exception as e:
        chequear("PowerPoint abre " + f, False, str(e)[:70])

print("\n=== 10. control GEMELO: un pptx roto TIENE que ser rechazado ===")
gemelo = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_gemelo_roto.pptx")
shutil.copy(os.path.join(ENT, "TryOut_IMG_Dia6.pptx"), gemelo)
with open(gemelo, "r+b") as fh:
    fh.seek(600)
    fh.write(b"\x00" * 400)
try:
    pres = app.Presentations.Open(gemelo, WithWindow=False)
    pres.Close()
    abrio = True
except Exception:
    abrio = False
chequear("gemelo: PowerPoint RECHAZA un pptx corrompido", not abrio)
app.Quit()

print("\n" + ("=" * 70))
print("RESUMEN: %d control(es) en ROJO" % len(fallas))
for f in fallas:
    print("   - " + f)
sys.exit(1 if fallas else 0)
