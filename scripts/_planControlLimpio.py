"""
_planControlLimpio.py -- el plan de control que se entrega, mirado como lo miro Fak el 15/10/2026.

POR QUE EXISTE: LECCIONES 15/10 ("si se rearma entero va limpio: sin amarillo ni notas, una caracteristica
por fila, la D sola"; Fak: "todo amarillo con demasiadas notas... no me gusta", "¿por que no solo D?") y la
skill `planes-de-control` §5, puntos 3 a 6. El generador de los APC (`scripts/apc/_pcApcPatagonia.py`) ya
lo hacia por dentro; este chequeo mira el ARCHIVO que se entrega, lo haya armado quien lo haya armado.

  python scripts/_planControlLimpio.py <plan.xls|plan.xlsx> [--rearmado] [--definitivo]
  python scripts/_planControlLimpio.py --selftest

Siempre (cualquier plan de control, tambien las correcciones de pocas celdas):
  D_TLD          ROJO   la columna CLASIF. dice D/TLD o TLD: en el plan de control la marca critica es la D
                        sola, en rojo (§5 punto 5; el AMFE sigue escribiendo D/TLD)
  VARIAS_MARCAS  ROJO   dos marcas o mas en la misma celda de CLASIF.: una caracteristica por fila (§5 punto 4:
                        con varias en la misma celda se lee "SC SC SC")
  SIGLA_RARA     AVISO  una marca que no es D, SC ni CC (`caracteristicas-especiales.md` punto 5: no se adivina)
  MEDIDA_VISUAL  ROJO   la especificacion es una medida (numero con unidad) y el control es solo visual, tactil o
                        una prueba funcional: se mide con un instrumento con codigo (§3.3 bis). No cuenta la medida
                        que es parte del nombre de un insumo con su codigo ("bolsa PP 600 x 500 mm (427VAR002TAP01)")
  CRITERIO_AMBIGUO ROJO control visual o funcional sin muestra patron/limite, calibre ni documento contra el que se
                        compara, y con una especificacion sin ningun numero ni codigo ("Bien colocado", "Abre y cierra")
  SIN_TOLERANCIA AVISO  una medida sin tolerancia (sin ±, rango, minimo ni maximo)
  Los tres salen del mail de Agustin Aguayo (Calidad, 22/10/2026) sobre el plan del APB trasero: "el control de la
  cinta de manipulacion indica una metrica y dice que se hace visual, ademas de que falta las tolerancias" y "los
  controles de costura, cierre, portavasos y el funcionamiento (que son visuales) utiliza terminos muy ambiguos".
  Calibrados contra su captura: marco 3.0 a 7.0 de la operacion 70 y dejo pasar 1.0, 2.0 y 8.0 a 12.0.
Con --rearmado (el plan se rearmo entero, §5 punto 3):
  AMARILLO       ROJO   celda de la tabla (debajo del encabezado CLASIF.) con relleno amarillo
  NOTA           ROJO   comentario o nota en una celda de la tabla
Con --definitivo (implica --rearmado; §5 punto 6, Fak 15/10: "un plan de control verdadero sin esas cosas"):
  TBD            ROJO   TBD o CONFLICTO adentro de una celda de la tabla

Lo de arriba del encabezado (titulos, revisiones, firmas) no se mira: la plantilla de Calidad trae banners
amarillos. Sale con 1 si hay un ROJO, 0 si no, 2 si no se pudo leer el archivo o no tiene la columna CLASIF.
Solo lee: no escribe nada.
"""
import os
import re
import sys
import tempfile

MARCAS_VALIDAS = {'D', 'SC', 'CC'}
AMARILLOS = {'FFFF00'}                      # el amarillo puro con que se marcan las celdas cambiadas
RE_TLD = re.compile(r'\bTLD\b', re.I)
RE_TBD = re.compile(r'\bTBD\b|\bCONFLICTO\b', re.I)
# una medida: numero seguido de su unidad ("75 mm", "4 ± 1 mm", "250 °C", "140 gr/m2", "7,00 s", "65 %")
RE_MEDIDA = re.compile(r'\d\s*(mm|cm|m|°\s*C|º\s*C|N|kN|g|gr|kg|s|seg|min|bar|%|g/m2|gr/m2|kg/min|mm/min|g/s)(?![A-Za-z0-9])')
RE_TOLERANCIA = re.compile(r'±|\+/?-|\d\s*(a|-)\s*\d|[<>≤≥]|m[aá]x|m[ií]n|entre|toler', re.I)
# el nombre de un insumo con su codigo: letras y numeros mezclados (427VAR002TAP01, ET-SATO-50X20, CARTSEP03)
RE_CODIGO = re.compile(r'\b(?=[A-Z0-9.\-]*[A-Z]{2})(?=[A-Z0-9.\-]*\d{2})[A-Z0-9][A-Z0-9.\-]{5,}\b')
# control solo con los sentidos: visual, tactil, conteo, prueba funcional a mano
RE_SENTIDOS = re.compile(r'visual|t[aá]ctil|funcional|conteo|manual', re.I)
# lo que vuelve objetivo un control visual: una referencia fisica o escrita contra la que se compara, o un instrumento
RE_REFERENCIA = re.compile(r'muestra|patr[oó]n|l[ií]mite|cat[aá]logo|biblia|calibre|regla|cinta|\bMC\s*\d|balanza|pantalla|'
                           r'display|contra |dispositivo|plantilla|mylar|term[oó]metro|pir[oó]metro|dinam[oó]metro|sistema|'
                           r'galga|torqu|programa|l[aá]ser|l[ií]neas? |certificado|'
                           r'seg[uú]n (la |el )?(lista|orden|planilla|plano|ficha|hoja|especificaci)', re.I)


def _rgb6(valor):
    """'FFFFFF00' / '00FFFF00' / 'FFFF00' -> 'FFFF00'; cualquier otra cosa -> ''."""
    v = str(valor or '').upper()
    return v[-6:] if re.fullmatch(r'[0-9A-F]{6}([0-9A-F]{2})?', v) or re.fullmatch(r'[0-9A-F]{8}', v) else ''


def leer(ruta):
    """Devuelve una lista de hojas: {nombre, celdas: {(r, c): texto}, amarillas: set((r, c)), notas: set((r, c))}."""
    hojas = []
    if ruta.lower().endswith('.xls'):
        import xlrd  # solo para el formato viejo de Calidad; el CI no lo necesita
        libro = xlrd.open_workbook(ruta, formatting_info=True)
        for s in libro.sheets():
            celdas, amarillas = {}, set()
            for r in range(s.nrows):
                for c in range(s.ncols):
                    v = s.cell_value(r, c)
                    if str(v).strip():
                        celdas[(r, c)] = str(v)
                    xf = libro.xf_list[s.cell_xf_index(r, c)]
                    color = libro.colour_map.get(xf.background.pattern_colour_index)
                    if xf.background.fill_pattern and color == (255, 255, 0):
                        amarillas.add((r, c))
            hojas.append({'nombre': s.name, 'celdas': celdas, 'amarillas': amarillas, 'notas': set(s.cell_note_map)})
        return hojas
    import openpyxl
    libro = openpyxl.load_workbook(ruta)
    for s in libro.worksheets:
        celdas, amarillas, notas = {}, set(), set()
        for fila in s.iter_rows():
            for cel in fila:
                r, c = cel.row - 1, cel.column - 1
                if cel.value is not None and str(cel.value).strip():
                    celdas[(r, c)] = str(cel.value)
                relleno = cel.fill
                if relleno is not None and relleno.fill_type == 'solid' and _rgb6(relleno.fgColor.rgb) in AMARILLOS:
                    amarillas.add((r, c))
                if cel.comment is not None:
                    notas.add((r, c))
        hojas.append({'nombre': s.title, 'celdas': celdas, 'amarillas': amarillas, 'notas': notas})
    return hojas


def encabezado(celdas):
    """(fila, columna) del primer encabezado CLASIF., o None."""
    candidatos = sorted(k for k, v in celdas.items() if 'CLASIF' in v.upper())
    return candidatos[0] if candidatos else None


def marcas(texto):
    """'SC\\nSC' -> ['SC', 'SC']; 'D/TLD' -> ['D/TLD']."""
    return [t for t in re.split(r'[\s,;]+', texto.strip().upper()) if t]


def columnas_metodo(celdas, fila0, col):
    """(columna de ESPECIFICACION, columna de CALIBRE) por el texto del encabezado; si no esta, las del formulario
    I-AC-005.1 (una y dos a la derecha de CLASIF.)."""
    esp = ins = None
    for (r, c), v in celdas.items():
        if fila0 <= r <= fila0 + 2 and c > col:
            u = v.upper()
            if esp is None and 'ESPECIFICAC' in u:
                esp = c
            if ins is None and ('CALIBRE' in u or 'TECNICA' in u or 'TÉCNICA' in u):
                ins = c
    return (esp if esp is not None else col + 1), (ins if ins is not None else col + 2)


def revisar_criterio(esp, ins, proceso=False):
    """Codigos de las tres reglas del mail de Agustin Aguayo (22/10/2026) para una fila: especificacion y calibre.
    CRITERIO_AMBIGUO mira solo las caracteristicas de PRODUCTO (lo que se acepta o rechaza de la pieza): una de proceso
    ("Plato elevado antes de iniciar el corte") es un estado de la maquina, no un criterio de aceptacion."""
    esp, ins = (esp or '').strip(), (ins or '').strip()
    if not esp or not ins or RE_TBD.search(esp):
        return []
    sentidos = bool(RE_SENTIDOS.search(ins)) and not RE_REFERENCIA.search(ins)
    medida = bool(RE_MEDIDA.search(esp)) and not RE_CODIGO.search(esp)
    salida = []
    if medida and sentidos:
        salida.append(('ROJO', 'MEDIDA_VISUAL'))
    if medida and not RE_TOLERANCIA.search(esp):
        salida.append(('AVISO', 'SIN_TOLERANCIA'))
    if sentidos and not proceso and not re.search(r'\d', esp) and not RE_REFERENCIA.search(esp):
        salida.append(('ROJO', 'CRITERIO_AMBIGUO'))
    return salida


def revisar(hojas, rearmado=False, definitivo=False):
    """Lista de hallazgos (nivel, codigo, hoja, celda, evidencia). Con ninguna hoja con CLASIF., un ERROR."""
    rearmado = rearmado or definitivo
    hallazgos, con_tabla = [], 0
    for h in hojas:
        enc = encabezado(h['celdas'])
        if not enc:
            continue
        con_tabla += 1
        fila0, col = enc
        nombre = h['nombre'].strip()
        celda = lambda r, c: f'{nombre}!{chr(65 + c) if c < 26 else c}{r + 1}'
        c_esp, c_ins = columnas_metodo(h['celdas'], fila0, col)
        for r in sorted({r for (r, _c) in h['celdas'] if r > fila0 + 1}):
            esp, ins = h['celdas'].get((r, c_esp), ''), h['celdas'].get((r, c_ins), '')
            # en el formulario I-AC-005.1 PROCESO va justo a la izquierda de CLASIF. y PRODUCTO a la izquierda de PROCESO
            proceso = bool(h['celdas'].get((r, col - 1), '').strip()) and not h['celdas'].get((r, col - 2), '').strip()
            for nivel, cod in revisar_criterio(esp, ins, proceso):
                hallazgos.append((nivel, cod, nombre, celda(r, c_esp), f'{esp} | {ins}'.replace('\n', ' / ')[:90]))
        for (r, c), v in sorted(h['celdas'].items()):
            if r <= fila0:
                continue
            if c == col and 'CLASIF' not in v.upper():
                ms = marcas(v)
                if RE_TLD.search(v):
                    hallazgos.append(('ROJO', 'D_TLD', nombre, celda(r, c), v.replace('\n', ' / ')))
                if len(ms) > 1:
                    hallazgos.append(('ROJO', 'VARIAS_MARCAS', nombre, celda(r, c), v.replace('\n', ' / ')))
                for m in ms:
                    if m not in MARCAS_VALIDAS and not RE_TLD.search(m):
                        hallazgos.append(('AVISO', 'SIGLA_RARA', nombre, celda(r, c), m))
            if definitivo and RE_TBD.search(v):
                hallazgos.append(('ROJO', 'TBD', nombre, celda(r, c), v.replace('\n', ' / ')[:80]))
        if rearmado:
            for (r, c) in sorted(h['amarillas']):
                if r > fila0:
                    hallazgos.append(('ROJO', 'AMARILLO', nombre, celda(r, c), h['celdas'].get((r, c), '')[:60]))
            for (r, c) in sorted(h['notas']):
                if r > fila0:
                    hallazgos.append(('ROJO', 'NOTA', nombre, celda(r, c), h['celdas'].get((r, c), '')[:60]))
    if not con_tabla:
        hallazgos.append(('ERROR', 'SIN_CLASIF', '-', '-', 'ninguna hoja tiene el encabezado CLASIF.: no parece un plan de control I-AC-005.1'))
    return hallazgos


def codigo_salida(hallazgos):
    if any(h[0] == 'ERROR' for h in hallazgos):
        return 2
    return 1 if any(h[0] == 'ROJO' for h in hallazgos) else 0


def main(argv):
    if '--selftest' in argv:
        return selftest()
    rutas = [a for a in argv if not a.startswith('--')]
    if len(rutas) != 1:
        print(__doc__)
        return 2
    try:
        hojas = leer(rutas[0])
    except Exception as e:  # el archivo no abre: no se juzga
        print(f'No pude leer {rutas[0]}: {e}')
        return 2
    hallazgos = revisar(hojas, rearmado='--rearmado' in argv, definitivo='--definitivo' in argv)
    modo = 'definitivo' if '--definitivo' in argv else 'rearmado' if '--rearmado' in argv else 'correccion'
    for nivel, cod, hoja, cel, ev in hallazgos[:200]:
        print(f'{nivel:5} {cod:13} {cel:22} {ev}')
    rojos = sum(1 for h in hallazgos if h[0] == 'ROJO')
    print(f'plan de control ({modo}): {rojos} ROJO · {sum(1 for h in hallazgos if h[0] == "AVISO")} AVISO')
    return codigo_salida(hallazgos)


# ---------------------------------------------------------------- selftest (sin archivos reales: el CI no los tiene)

def _plan(dir_, nombre, filas, amarillas=(), notas=(), banner_amarillo=False, con_clasif=True, metodos=()):
    """metodos: filas (producto, proceso, especificacion, calibre) que van debajo de `filas`."""
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import PatternFill
    amarillo = PatternFill('solid', fgColor='FFFF00')
    wb = openpyxl.Workbook()
    s = wb.active
    s.title = 'PC CENTRAL '
    s['A2'] = 'PLAN DE CONTROL'
    if banner_amarillo:
        s['A2'].fill = amarillo
    s.cell(row=5, column=7, value='CLASIF. CARAC.' if con_clasif else 'METODOS')
    for i, (clasif, espec) in enumerate(filas):
        s.cell(row=7 + i, column=5, value='Caracteristica')
        if clasif:
            s.cell(row=7 + i, column=7, value=clasif)
        if espec:
            s.cell(row=7 + i, column=8, value=espec)
    for i, (prod, proc, espec, ins) in enumerate(metodos, start=len(filas)):
        s.cell(row=7 + i, column=5 if prod else 6, value=prod or proc)
        s.cell(row=7 + i, column=8, value=espec)
        s.cell(row=7 + i, column=9, value=ins)
    for i in amarillas:
        s.cell(row=7 + i, column=8).fill = amarillo
    for i in notas:
        s.cell(row=7 + i, column=8).comment = Comment('fuente + ANTES', 'F.Santoro')
    ruta = os.path.join(dir_, nombre + '.xlsx')
    wb.save(ruta)
    return ruta


def selftest():
    casos = []  # (nombre, ruta, flags, codigo esperado, codigos que tienen que aparecer)
    with tempfile.TemporaryDirectory() as d:
        limpio = _plan(d, 'limpio', [('D', '< 100 mm/min'), ('SC', '4 ± 1 mm'), ('', 'Visual')], banner_amarillo=True)
        casos += [
            ('limpio, correccion', limpio, [], 0, []),
            ('limpio con banner amarillo de la plantilla, definitivo', limpio, ['--definitivo'], 0, []),
            ('D/TLD en CLASIF.', _plan(d, 'dtld', [('D/TLD', '< 100 mm/min')]), [], 1, ['D_TLD']),
            ('dos SC en una celda', _plan(d, 'varias', [('SC\nSC', '4 ± 1 mm')]), [], 1, ['VARIAS_MARCAS']),
            ('sigla W', _plan(d, 'rara', [('W', '4 ± 1 mm')]), [], 0, ['SIGLA_RARA']),
        ]
        amar = _plan(d, 'amarillo', [('SC', '4 ± 1 mm')], amarillas=[0])
        nota = _plan(d, 'nota', [('SC', '4 ± 1 mm')], notas=[0])
        tbd = _plan(d, 'tbd', [('SC', 'TBD')])
        casos += [
            ('amarillo en la tabla, correccion de pocas celdas', amar, [], 0, []),
            ('amarillo en la tabla, rearmado', amar, ['--rearmado'], 1, ['AMARILLO']),
            ('nota en la tabla, correccion de pocas celdas', nota, [], 0, []),
            ('nota en la tabla, rearmado', nota, ['--rearmado'], 1, ['NOTA']),
            ('TBD, rearmado (version de trabajo)', tbd, ['--rearmado'], 0, []),
            ('TBD, definitivo', tbd, ['--definitivo'], 1, ['TBD']),
            ('sin columna CLASIF.', _plan(d, 'sin', [('SC', 'x')], con_clasif=False), [], 2, ['SIN_CLASIF']),
        ]
        # las 12 filas de la operacion 70 del plan del APB trasero, tal cual las marco Agustin Aguayo (Calidad, 22/10/2026)
        rojas = [('Costuras', '', 'Sin costuras abiertas', 'Visual'),
                 ('Cierre', '', 'Bien colocado, sin aberturas', 'Visual'),
                 ('Cinta de manipulación', '', 'Presente; 75 mm afuera de la pieza', 'Visual'),
                 ('Portavasos y marco', '', 'Bien colocados', 'Visual'),
                 ('Funcionamiento del portavasos', '', 'Abre y cierra', 'Prueba funcional manual')]
        blancas = [('Deformaciones', '', 'Sin deformaciones', 'Visual / muestra patrón'),
                   ('Marcas en el vinilo', '', 'Sin marcas', 'Visual / muestra patrón'),
                   ('Dimensional', '', 'Según calibre de forma', 'Calibre de forma APB TAPA PATAGONIA MC505'),
                   ('Identificación de la pieza', '', 'Etiqueta ET-SATO-50X20 con código 2HC.885.081 RL1', 'Visual'),
                   ('Embolsado', '', 'Pieza en bolsa PP 600 x 500 mm (427VAR002TAP01)', 'Visual'),
                   ('Rebaba de espuma visible', '', 'Sin rebaba visible', 'Visual con iluminación dirigida / muestra límite'),
                   ('', 'Carga en el sistema ARB', 'Cantidad OK / cantidad NOK', 'Sistema ARB')]
        casos += [
            ('Aguayo: las 7 filas que dejo en blanco', _plan(d, 'ag_blancas', [], metodos=blancas), [], 0, []),
            ('Aguayo: costura sin costuras abiertas, visual', _plan(d, 'ag_cost', [], metodos=rojas[:1]), [], 1, ['CRITERIO_AMBIGUO']),
            ('Aguayo: cierre bien colocado, visual', _plan(d, 'ag_cie', [], metodos=rojas[1:2]), [], 1, ['CRITERIO_AMBIGUO']),
            ('Aguayo: cinta 75 mm visual sin tolerancia', _plan(d, 'ag_cin', [], metodos=rojas[2:3]), [], 1,
             ['MEDIDA_VISUAL', 'SIN_TOLERANCIA']),
            ('Aguayo: portavasos bien colocados, visual', _plan(d, 'ag_por', [], metodos=rojas[3:4]), [], 1, ['CRITERIO_AMBIGUO']),
            ('Aguayo: funcionamiento abre y cierra', _plan(d, 'ag_fun', [], metodos=rojas[4:5]), [], 1, ['CRITERIO_AMBIGUO']),
            ('cinta medida con regla y tolerancia', _plan(d, 'cin_ok', [], metodos=[
                ('Cinta de manipulación: largo afuera', '', '75 ± 2 mm', 'Regla MC 403')]), [], 0, []),
            ('referencia escrita en el control (lineas de la mesa, laser)', _plan(d, 'refs', [], metodos=[
                ('Alineación del vinilo', '', 'Coincidente con las líneas rojas de la mesa', 'Visual'),
                ('Sobrante', '', 'Máximo 2,5 cm por lado', 'Visual y láser de la máquina')]), [], 0, []),
            ('caracteristica de proceso visual (estado de la maquina)', _plan(d, 'proc', [], metodos=[
                ('', 'Plato elevado', 'Plato elevado antes de iniciar el corte', 'Visual')]), [], 0, []),
        ]
        malos = 0
        for nombre, ruta, flags, esperado, codigos in casos:
            hall = revisar(leer(ruta), rearmado='--rearmado' in flags, definitivo='--definitivo' in flags)
            sale = codigo_salida(hall)
            vistos = {h[1] for h in hall}
            ok = sale == esperado and set(codigos) <= vistos and (codigos or not [h for h in hall if h[0] != 'AVISO'])
            malos += not ok
            print(f'  {"ok" if ok else "MAL"}  {nombre}  esperado={esperado} salio={sale} codigos={sorted(vistos)}')
    print('selftest planControlLimpio: todo verde' if not malos else f'selftest planControlLimpio: {malos} MAL')
    return 1 if malos else 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1:]))
