# -*- coding: utf-8 -*-
"""Convierte el indice de la portada en un MAPA por etapas, en el pptx que ya existe.

Por que: leidas en fila, las 17 hojas tienen una logica clara — se prepara, se produce, se
reacciona si algo interrumpe, y se termina. Pero la portada las lista planas, una debajo de
otra, y el que llega nuevo no distingue lo que hace SIEMPRE de lo que hace SOLO SI PASA ALGO.
Agrupandolas, el documento se entiende de un vistazo.

Toca UNA caja de texto de UNA lamina. Antes: resguardo. Despues: se verifica que el texto de
las otras 17 laminas no se movio.

  portada_mapa.py --pptx <ruta> [--dry-run]
"""
import argparse, os, shutil, sys
from datetime import datetime
from pptx import Presentation
from pptx.util import Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import hoja_pptx as HP

ETAPAS = [
    ("PREPARAR Y ARRANCAR", "en este orden, una vez por turno", [
        ("20.1",  "Reconocimiento de la máquina y riesgos"),
        ("20.2",  "Puesta en marcha del fusor de adhesivo"),
        ("20.3",  "Encendido y acceso al HMI"),
        ("20.4",  "Carga de parámetros de producto y temperatura"),
        ("20.5",  "Calentamiento y espera"),
        ("20.6",  "Montaje del rollo en el desbobinador"),
        ("20.7",  "Enhebrado del material"),
        ("20.8",  "Centrado y tensión de la banda"),
        ("20.9",  "Arranque y alineación"),
    ]),
    ("PRODUCIR", "acá pasa el turno", [
        ("20.10", "Laminado — control durante la marcha"),
    ]),
    ("SI PASA ALGO", "sólo cuando ocurre", [
        ("20.11", "Empalme del material"),
        ("20.12", "Cambio de rollo por alarma de fin de material"),
        ("20.13", "Destrabe del material trabado"),
    ]),
    ("TERMINAR", "al cerrar el turno", [
        ("20.14", "Corte de la plancha"),
        ("20.15", "Parada de la máquina"),
        ("20.16", "Apertura de rodillos y colocación de la bandeja"),
        ("20.17", "Limpieza de rodillos por pasos"),
    ]),
]


def texto_completo(prs):
    out = []
    for s in prs.slides:
        cajas = sorted([(round(sh.top / 360000., 1), round(sh.left / 360000., 1),
                         sh.text_frame.text.strip())
                        for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()])
        out.append("\n".join(t for _, _, t in cajas))
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pptx", required=True)
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    prs = Presentation(a.pptx)
    antes = texto_completo(prs)
    port = prs.slides[0]

    # la caja del indice: la mas alta de la mitad derecha de la portada
    cand = [sh for sh in port.shapes
            if sh.has_text_frame and sh.left / 360000.0 > 13
            and sh.height / 360000.0 > 4 and "20.1" in sh.text_frame.text]
    if len(cand) != 1:
        sys.exit("esperaba 1 caja de indice en la portada y encontre %d" % len(cand))
    caja = cand[0]

    print("--- PLAN ---")
    print("portada, caja del indice de %.1f x %.1f cm"
          % (caja.width / 360000.0, caja.height / 360000.0))
    print("  antes: lista plana de %d renglones"
          % len([l for l in caja.text_frame.text.split("\n") if l.strip()]))
    print("  queda: %d etapas" % len(ETAPAS))
    for tit, sub, hojas in ETAPAS:
        print("     %-22s %-34s %d hojas" % (tit, "(" + sub + ")", len(hojas)))
    if a.dry_run:
        print("DRY-RUN: no se toco nada.")
        return

    res = os.path.join(os.path.dirname(a.pptx),
                       "_resguardo portada %s.pptx" % datetime.now().strftime("%Y-%m-%d %H%M"))
    shutil.copy2(a.pptx, res)
    print("\nresguardo:", os.path.basename(res))

    tf = caja.text_frame
    for p in list(tf.paragraphs)[1:]:
        p._element.getparent().remove(p._element)
    tf.paragraphs[0].clear()

    primero = True
    for tit, sub, hojas in ETAPAS:
        p = tf.paragraphs[0] if primero else tf.add_paragraph()
        primero = False
        p.space_before = Pt(0 if p is tf.paragraphs[0] else 6)
        p.space_after = Pt(1)
        r = p.add_run(); r.text = tit + "   "
        r.font.size = Pt(8.5); r.font.bold = True
        r.font.name = "Calibri"; r.font.color.rgb = HP.AZUL
        r = p.add_run(); r.text = sub
        r.font.size = Pt(7); r.font.italic = True
        r.font.name = "Calibri"; r.font.color.rgb = HP.RGBColor(0x70, 0x76, 0x84)
        for num, nom in hojas:
            q = tf.add_paragraph()
            q.space_after = Pt(0)
            r = q.add_run(); r.text = "     %s   " % num
            r.font.size = Pt(7.5); r.font.bold = True
            r.font.name = "Calibri"; r.font.color.rgb = HP.AZUL
            r = q.add_run(); r.text = nom
            r.font.size = Pt(7.5); r.font.name = "Calibri"; r.font.color.rgb = HP.NEGRO

    prs.save(a.pptx)

    despues = texto_completo(Presentation(a.pptx))
    intrusas = [k + 1 for k, (x, y) in enumerate(zip(antes, despues)) if x != y and k != 0]
    if intrusas:
        sys.exit("ABORTAR: cambio texto en %s" % intrusas)
    print("las otras 17 laminas: texto intacto.")


main()
