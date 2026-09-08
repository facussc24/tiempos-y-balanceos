# -*- coding: utf-8 -*-
"""Diff del texto COMPLETO de cada lamina: el de Fak contra el del generador."""
import glob, difflib
from pptx import Presentation

def todo(ruta):
    prs = Presentation(ruta)
    out = []
    for s in prs.slides:
        cajas = sorted([(round(sh.top/360000.,1), round(sh.left/360000.,1),
                         sh.text_frame.text.strip())
                        for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()])
        out.append("\n".join(t for _, _, t in cajas))
    return out

f = glob.glob("C:/Users/FacundoS-PC/BARACK ARGENTINA SRL/Ingenier*/INGENIERIA BARACK (NUNCA BORRAR)/1- GENERAL/INSTRUCTIVOS/INSTRUCCIONES OPERATIVAS/HOTMELT/HOJAS DE PROCESO*.pptx")[0]
a, b = todo(f), todo("_mio.pptx")
print("laminas: Fak %d / generador %d" % (len(a), len(b)))
dif = 0
for i, (x, y) in enumerate(zip(b, a)):
    if x != y:
        dif += 1
        print("\n=== LAMINA %d ===" % (i + 1))
        for l in difflib.unified_diff(x.split("\n"), y.split("\n"),
                                      "generador", "Fak", lineterm="", n=0):
            if not l.startswith(("---", "+++", "@@")):
                print(l)
print("\nlaminas con diferencia de texto:", dif)
