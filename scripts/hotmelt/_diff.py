# -*- coding: utf-8 -*-
"""Compara el texto del deck de Fak contra el que sale del generador, hoja por hoja."""
import glob, re, difflib
from pptx import Presentation

def pasos(ruta):
    prs = Presentation(ruta)
    d = {}
    for s in prs.slides:
        op = None
        txt = None
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            t = sh.text_frame.text.strip()
            if re.fullmatch(r"20\.\d+", t):
                op = t
            if sh.left / 360000.0 > 16 and 5 < sh.top / 360000.0 < 15:
                txt = t
        if op:
            d[op] = txt or ""
    return d

f = glob.glob("C:/Users/FacundoS-PC/BARACK ARGENTINA SRL/Ingenier*/INGENIERIA BARACK (NUNCA BORRAR)/1- GENERAL/INSTRUCTIVOS/INSTRUCCIONES OPERATIVAS/HOTMELT/HOJAS DE PROCESO*.pptx")[0]
a, b = pasos(f), pasos("_mio.pptx")
for op in sorted(a, key=lambda x: float(x[3:])):
    if a[op] != b.get(op, ""):
        print("\n" + "#" * 70)
        print("#  " + op)
        print("#" * 70)
        for l in difflib.unified_diff(b.get(op, "").split("\n"), a[op].split("\n"),
                                      "generador", "el de Fak", lineterm="", n=0):
            if l.startswith(("---", "+++", "@@")):
                continue
            print(l)
