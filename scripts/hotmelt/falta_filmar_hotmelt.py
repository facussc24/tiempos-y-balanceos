# -*- coding: utf-8 -*-
"""La lista de lo que falta filmar o preguntar para cerrar las hojas de la MAQUINA HOTMELT.

Por que existe: los videos existentes cubren principalmente pasadas parciales de material y
pantallas del panel. Las maniobras criticas de puesta en marcha, manejo del eje neumático,
corte sobre mesa y limpieza de fin de turno no quedaron filmadas o se hicieron de modo
no estándar (ej. corte en el piso).

Sale un PDF A4 apaisado corporativo, una página por bloque.

    py -3 scripts/hotmelt/falta_filmar_hotmelt.py
"""
import os
import sys
import shutil

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Cm, Pt

AZUL  = RGBColor(0x44, 0x54, 0x6A)
AZUL2 = RGBColor(0x1F, 0x49, 0x7D)
BLANCO= RGBColor(0xFF, 0xFF, 0xFF)
NEGRO = RGBColor(0x00, 0x00, 0x00)
GRIS  = RGBColor(0xF2, 0xF2, 0xF2)
ROJO  = RGBColor(0xC0, 0x30, 0x30)

W, H = 29.7, 21.0
M = 1.1

BASE = os.path.dirname(os.path.abspath(__file__))
SALIDA_PPTX = os.path.join(BASE, "QUE FALTA FILMAR - MAQUINA HOTMELT.pptx")
ESCRITORIO = (r"C:\Users\FacundoS-PC\OneDrive - BARACK ARGENTINA SRL\Desktop"
              r"\Hojas de proceso maquina HOTMELT - desde los videos")

ENCABEZADO = ("Barridos los videos existentes, las transcripciones de planta y las observaciones "
              "de Fak al 28/09/2026. Lo que está en esta lista no aparece filmado con estándar "
              "de trabajo o falta registrar. Con esto se cierran al 100% las hojas de proceso.")

TOMAS = [
    ("1 · ENERGIZACIÓN Y PUESTA EN MARCHA (OP 20.1) — falta el arranque físico", [
        ("Giro del interruptor general principal del tablero a posición ON",
         "Los videos arrancan con la máquina ya energizada. Falta el gesto de encendido general."),
        ("Manómetro y válvula de entrada de aire comprimido a 6 bar",
         "No está filmada la verificación de presión de línea antes de operar."),
        ("Rearme de pulsadores de parada de emergencia en frío",
         "Verificación visual de que todas las paradas perimetrales estén desbloqueadas antes de arrancar."),
    ]),
    ("2 · FUSOR DE ADHESIVO PUR Y DOSIFICACIÓN (OP 20.2) — tanque y boquillas", [
        ("Carga física de panes o pastillas de adhesivo dentro del tanque del fusor",
         "Solo está filmado el tanque cerrado o la tapa. Falta el nivel y estado de carga de materia prima."),
        ("Inspección de las dos boquillas dosificadoras de adhesivo sobre los rodillos",
         "Verificación en frío y apertura de paso de adhesivo hacia los caños calefaccionados."),
        ("Procedimiento ante boquilla obstruida",
         "Criterio de actuación del operario sin intervenir en caliente con elementos metálicos."),
    ]),
    ("3 · MONTAJE Y CAMBIO DE ROLLO (OP 20.6 / 20.12) — eje neumático expansible", [
        ("Inserción del eje neumático en el centro del buje de cartón del vinilo",
         "No está filmado el momento en que se introduce el eje dentro del rollo nuevo."),
        ("Insuflado de aire a presión con pistola en la válvula del eje expansible",
         "El gesto clave que traba las zapatas contra el cartón no quedó registrado en video."),
        ("Alce y montaje del rollo entre dos operarios sobre las cunas del desbobinador",
         "La maniobra ergonómica de montaje y bloqueo de la perilla ON/OFF del soporte."),
        ("Despresurización del eje al retirar el buje de cartón vacío al agotarse la bobina",
         "Accionamiento de la válvula de alivio para liberar el tubo de cartón."),
    ]),
    ("4 · CORTE REGLAMENTARIO DE LA PLANCHA (OP 20.14) — en mesa con guantes anticorte", [
        ("Fraccionamiento de vinilo con cutter de seguridad sobre la mesa plana de trabajo",
         "Los videos viejos mostraban corte improvisado en el piso. Debe filmarse sobre mesa limpia."),
        ("Uso obligatorio de guantes anticorte durante el manipuleo de herramienta de filo",
         "El estándar de seguridad exige guantes certificados puestos durante la maniobra."),
        ("Corte de esquinas con bordes redondeados",
         "Evitar cortes rectos en ángulo vivo que provocan desgarros en el conformado posterior."),
    ]),
    ("5 · APAGADO, BANDEJA Y LIMPIEZA EN CALIENTE (OP 20.15 a 20.17) — fin de turno", [
        ("Puesta en modo Standby (espera a 150 °C) en el panel HMI ante paradas mayores a 30 min",
         "Para prevenir la degradación térmica del adhesivo PUR dentro de mangueras y rodillos."),
        ("Apertura neumática de rodillos y calce de la bandeja recolectora de goteo",
         "Inserción de la bandeja metálica sobre los topes para proteger la bancada."),
        ("Aplicación de parafina / cera sólida de limpieza sobre rodillos girando lento",
         "La secuencia de limpieza en caliente paso a paso utilizando raspador no abrasivo."),
    ]),
]

COMO = ("Como filmar: celular en horizontal (16:9), foco fijo en la mano y el comando, 5 a 10 "
        "segundos por toma. Si hay cartel o display, terminar con 3 segundos quietos de cerca.")

def _caja(slide, x, y, w, h, relleno=None, borde=None, ancho=Pt(1)):
    sh = slide.shapes.add_shape(1, Cm(x), Cm(y), Cm(w), Cm(h))
    if relleno is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = relleno
    if borde is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = borde
        sh.line.width = ancho
    sh.shadow.inherit = False
    return sh

def _txt(sh, texto, size=11, bold=False, color=NEGRO, align=PP_ALIGN.CENTER):
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Cm(0.2)
    tf.margin_top = tf.margin_bottom = Cm(0.04)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = texto
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = "Calibri"
    r.font.color.rgb = color
    return sh

def armar():
    prs = Presentation()
    prs.slide_width = Cm(W)
    prs.slide_height = Cm(H)
    blank = prs.slide_layouts[6]
    n = 0

    for i, (bloque, tomas) in enumerate(TOMAS):
        slide = prs.slides.add_slide(blank)
        _caja(slide, M, M, W - 2 * M, H - 2 * M, BLANCO, AZUL, Pt(1.5))
        _txt(_caja(slide, M, M, W - 2 * M, 1.2, AZUL, AZUL),
             f"QUÉ FALTA FILMAR · MÁQUINA LAMINADORA HOTMELT  —  página {i + 1} de {len(TOMAS)}",
             size=16, bold=True, color=BLANCO)
        _txt(_caja(slide, M, M + 1.25, W - 2 * M, 0.9, GRIS, None),
             ENCABEZADO, size=9.5, bold=False, color=AZUL2)

        _txt(_caja(slide, M, M + 2.25, W - 2 * M, 0.75, AZUL2, AZUL2),
             bloque.upper(), size=12, bold=True, color=BLANCO)

        y = M + 3.10
        total_h = H - M - 1.2 - y
        alto_fila = total_h / len(tomas)

        for que, porque in tomas:
            n += 1
            _txt(_caja(slide, M, y, 1.0, alto_fila, GRIS, AZUL), str(n),
                 size=12, bold=True, align=PP_ALIGN.CENTER)
            sh = _caja(slide, M + 1.0, y, W - 2 * M - 1.0, alto_fila, None, AZUL)
            tf = sh.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = Cm(0.25)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            r = p.add_run()
            r.text = que + "\n"
            r.font.size = Pt(11)
            r.font.bold = True
            r.font.name = "Calibri"
            r.font.color.rgb = NEGRO

            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.LEFT
            r2 = p2.add_run()
            r2.text = porque
            r2.font.size = Pt(9.5)
            r2.font.name = "Calibri"
            r2.font.color.rgb = AZUL2
            y += alto_fila

        _txt(_caja(slide, M, H - M - 1.0, W - 2 * M, 1.0, GRIS, ROJO, Pt(1.25)),
             COMO, size=10, bold=False, color=NEGRO)

    prs.save(SALIDA_PPTX)
    print(f"[OK] {SALIDA_PPTX}  ({len(TOMAS)} páginas, {n} tomas)")
    return SALIDA_PPTX

def a_pdf(pptx_path):
    import win32com.client
    pdf = os.path.splitext(pptx_path)[0] + ".pdf"
    ppt = win32com.client.Dispatch("PowerPoint.Application")
    habia_abiertas = ppt.Presentations.Count
    try:
        pres = ppt.Presentations.Open(os.path.abspath(pptx_path), WithWindow=False)
        pres.SaveAs(os.path.abspath(pdf), 32)
        pres.Close()
    finally:
        if habia_abiertas == 0 and ppt.Presentations.Count == 0:
            ppt.Quit()
    print(f"[OK] {pdf}")
    return pdf

if __name__ == "__main__":
    p = armar()
    try:
        pdf = a_pdf(p)
        if os.path.isdir(ESCRITORIO):
            dst = os.path.join(ESCRITORIO, os.path.basename(pdf))
            shutil.copy2(pdf, dst)
            print(f"[OK] Copiado al Escritorio: {dst}")
    except Exception as e:
        print(f"[AVISO] No se pudo exportar PDF automático: {e}")
