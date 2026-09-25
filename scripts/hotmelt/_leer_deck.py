# -*- coding: utf-8 -*-
"""Vuelca el deck como lo leeria un operario: hoja por hoja, en orden, solo lo que se lee."""
import glob, io
from pptx import Presentation
p = glob.glob("Y:/BARACK/CALIDAD/DOCUMENTACION SGC/HOJAS DE OPERACIONES/1- CLIENTES/NOVAX/Tapizadas puerta/TOP ROLL/HOJAS DE PROCESO*.pptx")
prs = Presentation(p[0])
out = []
for i, s in enumerate(prs.slides):
    cajas = [(sh.top / 360000.0, sh.left / 360000.0, sh.text_frame.text.strip())
             for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    cajas.sort()
    out.append("\n" + "=" * 78)
    out.append("LAMINA %d" % (i + 1))
    out.append("=" * 78)
    for _, _, t in cajas:
        out.append(t)
io.open("_deck_lectura.txt", "w", encoding="utf-8").write("\n".join(out))
print("lineas:", len(out))
