# -*- coding: utf-8 -*-
"""Aplica los criterios de imagen a las 17 hojas, SOBRE EL ARCHIVO DE FAK.

No regenera el deck: Fak lo viene editando desde el 03/09 y su texto manda. Toca UNA cosa
por lamina —el bloque de IMAGENES— y despues verifica que el texto de las 18 quedo igual
byte a byte. Si cambio uno solo, aborta y deja el resguardo.

El PLAN de cada hoja es una lista de lo que va a quedar, en orden, donde:
    int    = la imagen que la lamina YA tiene en esa posicion (se conserva tal cual)
    "..."  = un archivo nuevo que la reemplaza o se agrega
y el segundo valor es cual de esas manda: la que el paso obliga a mirar o a leer.
Lo que no aparece en la lista, sale de la hoja — y abajo esta escrito por que.

    aplicar_criterios.py [--dry-run]
"""
import argparse
import io
import os
import re
import shutil
import sys
from datetime import datetime

from pptx import Presentation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
# hoja_pptx agrega al path los scripts del skill: de ahi sale el gate CANONICO.
# Hasta el 08/09/2026 esto importaba `_hojaProcesoCheck`, una copia local que
# reimplementaba sus propios umbrales y no conocia el modo `secuencia` (0 menciones
# contra 16 en el canonico). Dos gates con numeros propios es exactamente lo que la
# regla `hojas-proceso.md` prohibe: "una sola fuente".
import hoja_pptx as HP                  # noqa: E402
import hoja_proceso_check as CK         # noqa: E402

RUTA = ("C:/Users/FacundoS-PC/BARACK ARGENTINA SRL/Ingeniería y Proyecto - General/"
        "INGENIERIA BARACK (NUNCA BORRAR)/1- GENERAL/INSTRUCTIVOS/INSTRUCCIONES OPERATIVAS/"
        "HOTMELT/HOJAS DE PROCESO - MAQUINA HOTMELT - Rev.A.pptx")

OP = "_pantalla_operacion_%s.png"

PLAN = {
    # hoja      lo que queda, en orden                              principal, tenia
    "20.1":  ([0, 1, 2],                                                    0, 3),
    # el celular del traductor tapaba media foto y era la imagen mas grande de la hoja
    "20.2":  (["_pantalla_fusor.png", "_foto_fusor_unidad.jpg", 1],         0, 3),
    "20.3":  ([0, 1, 2],                                                    0, 3),
    # dos pantallas en un bloque no entran legibles: se fusionaron en una, a ancho completo
    "20.4":  (["_pantalla_parametros.png"],                                 0, 2),
    "20.5":  ([OP % "calentamiento_ok", 1],                                 0, 2),
    "20.6":  ([1, 0, 2],                                                    0, 3),
    "20.7":  ([0, 1, 2],                                                    0, 3),
    "20.8":  ([0, 1, 2],                                                    0, 3),
    "20.9":  ([OP % "automatico", 1, 2],                                    0, 3),
    # sale la foto del HMI en chino con la anotacion: ilegible, y su dato ya vive en la 20.9
    "20.10": ([1, 2],                                                       0, 3),
    "20.11": ([1, 0, 2],                                                    0, 3),
    # sale una foto de piso vacio con un zapato: no muestra nada de lo que el paso pide
    "20.12": ([1, 0],                                                       0, 3),
    "20.13": ([1, 0, 2],                                                    0, 3),
    "20.14": ([0, 1],                                                       0, 2),
    # 2do lugar: la pantalla de señales de seguridad, que es el paso 4 de esta hoja
    "20.15": ([OP % "detenido", 1, 0],                                      0, 3),
    "20.16": ([1, 0, 2],                                                    0, 3),
    "20.17": (["_pantalla_limpieza.png", 0],                                0, 2),
}


def texto(prs):
    return ["\n".join(t for _, _, t in sorted(
        (round(sh.top / 360000., 1), round(sh.left / 360000., 1), sh.text_frame.text.strip())
        for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()))
        for s in prs.slides]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    prs = Presentation(RUTA)
    antes = texto(prs)
    tmp = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_conservadas")
    os.makedirs(tmp, exist_ok=True)

    acciones = []
    for s in prs.slides:
        op = None
        for sh in s.shapes:
            if sh.has_text_frame and re.fullmatch(r"20\.\d+", sh.text_frame.text.strip()):
                op = sh.text_frame.text.strip()
                break
        if op not in PLAN:
            continue
        lista, principal, esperadas = PLAN[op]
        viejas = sorted([sh for sh in s.shapes if CK.es_contenido(sh)],
                        key=lambda q: (round(q.top / 360000., 1), round(q.left / 360000., 1)))
        # Este script NO es idempotente: los indices del PLAN apuntan a las imagenes que la
        # lamina tiene ANTES de tocarla; corriendolo dos veces apuntan a otra cosa. Por eso
        # cada hoja declara cuantas espera encontrar, y si no coinciden aborta antes de
        # escribir nada, en vez de dejar el deck a medio arreglar.
        if len(viejas) != esperadas:
            sys.exit("ABORTAR: %s tiene %d imagenes y el plan esperaba %d. ¿Ya se aplico? "
                     "Restaurar del resguardo antes de volver a correrlo."
                     % (op, len(viejas), esperadas))

        rutas = []
        for tok in lista:
            if isinstance(tok, int):
                if tok >= len(viejas):
                    sys.exit("ABORTAR: %s pide su imagen [%d] y tiene %d" % (op, tok, len(viejas)))
                blob = viejas[tok].image.blob
                ext = viejas[tok].image.ext or "png"
                d = os.path.join(tmp, "%s_%d.%s" % (op.replace(".", "_"), tok, ext))
                io.open(d, "wb").write(blob)
                rutas.append(d)
            else:
                if not os.path.exists(tok):
                    sys.exit("ABORTAR: falta %s (lo pide %s)" % (tok, op))
                rutas.append(tok)

        sacadas = len(viejas) - sum(1 for t in lista if isinstance(t, int))
        acciones.append((op, len(viejas), len(rutas), sacadas, principal,
                         [os.path.basename(r) for r in rutas]))
        if a.dry_run:
            continue

        for sh in viejas:
            sh._element.getparent().remove(sh._element)
        HP.bloque_imagenes_solo_fotos(s, rutas, principal=principal)

    print("--- PLAN ---")
    for op, v, n, sac, pr, nombres in acciones:
        print("  %-6s %d -> %d imagenes%s   principal: %s"
              % (op, v, n, ("  (saca %d)" % sac) if sac else "          ", nombres[pr]))
    if a.dry_run:
        print("\nDRY-RUN: no se toco nada.")
        return

    res = os.path.join(os.path.dirname(RUTA),
                       "_resguardo criterios %s.pptx" % datetime.now().strftime("%Y-%m-%d %H%M"))
    shutil.copy2(RUTA, res)
    prs.save(RUTA)

    despues = texto(Presentation(RUTA))
    intrusas = [k + 1 for k, (x, y) in enumerate(zip(antes, despues)) if x != y]
    if intrusas:
        shutil.copy2(res, RUTA)
        sys.exit("ABORTAR: se movio texto en las laminas %s. Se restauro el archivo." % intrusas)
    print("\nresguardo: %s" % os.path.basename(res))
    print("texto de las 18 laminas: intacto.")


main()
