# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Cm
pr = Presentation("_deck.pptx")
for i, s in enumerate(pr.slides):
    if i == 0: continue
    fotos = [sh for sh in s.shapes if sh.shape_type == 13
             and sh.left/360000 < 17.5 and sh.top/360000 > 5]
    op = ""
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().startswith("20."):
            op = sh.text_frame.text.strip(); break
    dims = ", ".join(f"{sh.width/360000:.1f}x{sh.height/360000:.1f}" for sh in fotos)
    area = sum(sh.width/360000 * sh.height/360000 for sh in fotos)
    print(f"lam{i+1:>2} {op:<6} n={len(fotos)}  {dims:<40} ocupa {area/(16.2*9.3)*100:.0f}% del bloque")
