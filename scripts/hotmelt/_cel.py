# -*- coding: utf-8 -*-
from pptx import Presentation
pr = Presentation("_deck.pptx")
s = pr.slides[10]      # lamina 11 = 20.10
for sh in s.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip():
        t = sh.text_frame.text.strip()
        if sh.top/360000 > 14.5 and sh.top/360000 < 18.4:
            r = sh.text_frame.paragraphs[0].runs
            sz = r[0].font.size.pt if r and r[0].font.size else None
            print(f"y={sh.top/360000:5.2f} h={sh.height/360000:4.2f} w={sh.width/360000:5.2f} size={sz}  {t[:45]}")
