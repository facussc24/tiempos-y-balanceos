# -*- coding: utf-8 -*-
"""Cambia UNA imagen de UNA lamina del pptx que ya existe, por su hash. Nada mas."""
import argparse, hashlib, os, shutil, sys
from datetime import datetime
from pptx import Presentation
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import hoja_pptx as HP

ap = argparse.ArgumentParser()
ap.add_argument("--pptx", required=True)
ap.add_argument("--hoja", required=True)
ap.add_argument("--vieja", help="imagen que hay que sacar (ruta)")
ap.add_argument("--por-ext", help="o: la unica imagen con esta extension (png = pantalla preparada)")
ap.add_argument("--nueva", required=True)
ap.add_argument("--dry-run", action="store_true")
a = ap.parse_args()

h_vieja = hashlib.md5(open(a.vieja, "rb").read()).hexdigest() if a.vieja else None
prs = Presentation(a.pptx)

def es_contenido(sh):
    if sh.shape_type != 13: return False
    x, y = sh.left/360000., sh.top/360000.
    if y < 4.5: return False
    if x > 17.0 and sh.width/360000. < 2.5: return False
    return True

for i, s in enumerate(prs.slides):
    if not any(sh.has_text_frame and sh.text_frame.text.strip() == a.hoja for sh in s.shapes):
        continue
    fotos = sorted([sh for sh in s.shapes if es_contenido(sh)],
                   key=lambda q: (round(q.top/360000.,1), round(q.left/360000.,1)))
    tmp = os.path.join(os.environ["TEMP"], "cambiar1"); os.makedirs(tmp, exist_ok=True)
    rutas, hallada = [], False
    for k, sh in enumerate(fotos):
        if h_vieja:
            coincide = hashlib.md5(sh.image.blob).hexdigest() == h_vieja
        else:
            coincide = sh.image.ext.lower() == a.por_ext.lower()
        if coincide:
            rutas.append(a.nueva); hallada = True
        else:
            d = os.path.join(tmp, "f%d.%s" % (k, sh.image.ext))
            open(d, "wb").write(sh.image.blob); rutas.append(d)
    print("lamina %d, hoja %s: %d fotos, la vieja %s" %
          (i+1, a.hoja, len(fotos), "ENCONTRADA" if hallada else "NO ESTA"))
    for r in rutas: print("   " + os.path.basename(r))
    if not hallada: sys.exit("no encontre la imagen a cambiar")
    if a.dry_run: print("DRY-RUN"); sys.exit(0)
    shutil.copy2(a.pptx, os.path.join(os.path.dirname(a.pptx),
        "_resguardo img %s.pptx" % datetime.now().strftime("%Y-%m-%d %H%M%S")))
    for sh in fotos: sh._element.getparent().remove(sh._element)
    HP.bloque_imagenes_solo_fotos(s, rutas)
    prs.save(a.pptx); print("guardado.")
    sys.exit(0)
sys.exit("no encontre la hoja " + a.hoja)
