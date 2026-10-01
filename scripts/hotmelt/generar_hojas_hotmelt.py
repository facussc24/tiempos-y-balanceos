# -*- coding: utf-8 -*-
"""
Generador oficial de HOJAS DE OPERACIONES para la MÁQUINA LAMINADORA HOTMELT KINGPOWER
Proyecto: Top Roll Patagonia / VW (OP 20 del Flujograma 155).
Formulario SGC oficial: I-IN-002.4-R01 (A4 apaisado 29.7 x 21.0 cm).

Estructura Oficial Dividida por Partes (Criterio Fak 28/09/2026):
  1. PARTE 1: OPERACIÓN ESTÁNDAR / PRODUCCIÓN (12 láminas: 20.1 a 20.10, 20.14, 20.15)
  2. PARTE 2: CONTINGENCIAS / SI PASA ALGO (3 láminas: 20.11 a 20.13)
  3. PARTE 3: MANTENIMIENTO OPERATIVO Y LIMPIEZA (2 láminas: 20.16, 20.17)
  4. DECK COMPLETO UNIFICADO (17 operaciones: 20.1 a 20.17)

Criterios de Fak y SGC cumplidos estrictamente:
  - 100% fotos reales de planta y fotogramas (0% dibujos vectoriales artificiales, 0% IA).
  - Eliminación de la foto 3 en 20.5 (los pedazos de adhesivo sólido tirados a mano sobre el rodillo eran de una prueba informal; el adhesivo fluye por cañería térmica interna).
  - Fotos del fusor en 20.2 rectificadas y limpias (se elimina la foto del celular en la mano y cinta de embalar).
  - Pantallas de HMI en 20.4 con recorte exacto del LCD activo (sin marco plástico gris para maximizar legibilidad).
  - Plan de Reacción Ante No Conforme específico e individualizado por cada operación (sin copy-paste genérico).
  - Lenguaje directo y claro de planta ("Meter el eje expansible adentro del tubo de cartón", "con la cara texturada hacia arriba", "guantes anticorte").
  - Clarificación en 20.7: el enhebrado en serpentina es sólo arranque inicial/rotura; el cambio de rollo normal es por empalme (20.11).
  - Eliminación del término "potenciómetro de guía" en 20.9.
  - Sizing balanceado en 20.15 con la pantalla de seguridades rectificada grande.
"""
import os
import sys
import shutil
import functools

# Gate canónico de redacción de la skill y hojalib
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                '..', '..', '.claude', 'skills',
                                'hojas-de-proceso', 'scripts'))
try:
    from redaccion import gate_redaccion
except ImportError:
    gate_redaccion = None

try:
    import hojalib as HL
except ImportError:
    HL = None

from PIL import Image, ImageFont
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

try:
    from datos_privados import CLAVE_HMI
except ImportError:            # el repo es publico: sin este archivo va TBD
    CLAVE_HMI = "TBD"

# ─── PALETA CORPORATIVA SGC ──────────────────────────────────────────────────
AZUL        = RGBColor(0x44, 0x54, 0x6A)  # Azul oscuro cabeceras (#44546A)
AZUL2       = RGBColor(0x44, 0x72, 0xC4)  # Azul banda seguridad (#4472C4)
AZUL_TEXTO  = RGBColor(0x1F, 0x49, 0x7D)  # Azul énfasis texto (#1F497D)
BLANCO      = RGBColor(0xFF, 0xFF, 0xFF)
NEGRO       = RGBColor(0x00, 0x00, 0x00)
GRISF       = RGBColor(0xF2, 0xF2, 0xF2)  # Gris fondo celdas secundarias

# ─── GEOMETRÍA OFICIAL FORMULARIO I-IN-002.4-R01 (cm) ────────────────────────
W, H = 29.7, 21.0
M = 0.70
X0, X1 = M, W - M                        # 0.70 .. 29.00 (ancho útil 28.30 cm)
HDR_Y, HDR_H = M, 4.00                   # 0.70 .. 4.70 cm

BODY_Y, BODY_H = 4.90, 9.90              # 4.90 .. 14.80 cm
IMG_W = 16.20
IMG_X = X0
DSC_X = X0 + IMG_W + 0.25                # 17.15 cm
DSC_W = X1 - DSC_X                       # 11.85 cm

CIC_Y, CIC_H = 15.00, 3.10               # 15.00 .. 18.10 cm
EPP_W = 6.40
CIC_W = X1 - X0 - EPP_W - 0.25           # 21.65 cm
EPP_X = X0 + CIC_W + 0.25                # 22.60 cm

PLN_Y = 18.30                            # 18.30 .. 20.30 cm
PLN_H = H - M - PLN_Y                    # 2.00 cm

# ─── RUTAS DE ACTIVOS LOCALES ────────────────────────────────────────────────
AQUI = os.path.dirname(os.path.abspath(__file__))
FOTOS_DIR = r"C:\Users\FacundoS-PC\OneDrive - BARACK ARGENTINA SRL\Desktop\Hojas de proceso maquina HOTMELT - desde los videos\_trabajo\fotos_hoja"
CELULAR_DIR = r"C:\Users\FacundoS-PC\OneDrive - BARACK ARGENTINA SRL\Desktop\Hojas de proceso maquina HOTMELT - desde los videos\Fotos del celular de Fak"
DESKTOP_DIR = r"C:\Users\FacundoS-PC\OneDrive - BARACK ARGENTINA SRL\Desktop\Hojas de proceso maquina HOTMELT - desde los videos\_trabajo"
DESKTOP_ROOT = r"C:\Users\FacundoS-PC\OneDrive - BARACK ARGENTINA SRL\Desktop\Hojas de proceso maquina HOTMELT - desde los videos"
SGC_DIR = r"Y:\BARACK\CALIDAD\DOCUMENTACION SGC\HOJAS DE OPERACIONES\1- CLIENTES\NOVAX\Tapizadas puerta\TOP ROLL"
EPP_DIR = os.path.join(AQUI, "epp")

LOGO_BARACK = r"C:\Users\FacundoS-PC\BARACK ARGENTINA SRL\Ingeniería y Proyecto - General\INGENIERIA BARACK (NUNCA BORRAR)\barack_logo.png"

# Iconos EPP
ICO_ROPA     = os.path.join(EPP_DIR, "ico_13756.png")
ICO_CALZADO  = os.path.join(EPP_DIR, "ico_4449.png")
ICO_GUANTES  = os.path.join(EPP_DIR, "ico_11789.png")
ICO_ANTEOJOS = os.path.join(EPP_DIR, "ico_16034.png")
ICO_BARBIJO  = os.path.join(EPP_DIR, "ico_barbijo.png")

EPP_HOTMELT = [ICO_ROPA, ICO_CALZADO, ICO_GUANTES, ICO_BARBIJO, ICO_ANTEOJOS]

# Activos especiales
FOTO_SEGURIDAD = os.path.join(AQUI, "_pantalla_seguridad.png")
FOTO_PANTALLA_FUSOR = os.path.join(AQUI, "_pantalla_fusor.png")

# ─── TIPOGRAFÍA Y MEDICIÓN REAL ──────────────────────────────────────────────
_TTF = {
    "Calibri": r"C:\Windows\Fonts\calibri.ttf",
    "Calibri-b": r"C:\Windows\Fonts\calibrib.ttf",
    "Arial": r"C:\Windows\Fonts\arial.ttf",
    "Arial-b": r"C:\Windows\Fonts\arialbd.ttf"
}
PT_CM = 0.03527777

@functools.lru_cache(maxsize=512)
def _fuente(nombre, bold, px):
    ruta = _TTF.get(nombre + ("-b" if bold else ""), _TTF["Calibri"])
    return ImageFont.truetype(ruta, max(int(px), 4))

def _ancho_cm(texto, size, fuente, bold):
    f = _fuente(fuente, bold, round(size * 96 / 72))
    return f.getlength(texto) / 96 * 2.54

def _achicar(texto, w_cm, h_cm, base, margen=0.06, minimo=5.5, fuente="Calibri", bold=False):
    if not texto:
        return base
    util_w = max(w_cm - 2 * margen - 0.24, 0.4)
    util_h = max(h_cm - 0.10, 0.2)
    palabras = str(texto).split()
    s = base
    while s >= minimo:
        lh = 1.22 * s * PT_CM
        lineas, actual = 1, ""
        cabe = True
        for p in palabras:
            if _ancho_cm(p, s, fuente, bold) > util_w:
                cabe = False
                break
            probar = (actual + " " + p) if actual else p
            if _ancho_cm(probar, s, fuente, bold) <= util_w:
                actual = probar
            else:
                lineas += 1
                actual = p
        if cabe and lineas * lh <= util_h:
            return s
        s -= 0.5
    return minimo

def _lineas_wrap(texto, size, w_cm, bold=False, fuente="Calibri"):
    lineas, actual = 1, ""
    for p in texto.split():
        probar = (actual + " " + p) if actual else p
        if _ancho_cm(probar, size, fuente, bold) <= w_cm:
            actual = probar
        else:
            lineas += 1
            actual = p
    return lineas

# ─── PRIMITIVAS GRÁFICAS ─────────────────────────────────────────────────────
def _caja(slide, x, y, w, h, relleno=None, borde=NEGRO, ancho=Pt(1)):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
    if relleno is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = relleno
    if borde is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = borde
        sh.line.width = ancho
    sh.shadow.inherit = False
    sh.text_frame.word_wrap = True
    return sh

def _txt(sh, texto, size=11, bold=False, color=NEGRO, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, fuente="Calibri", margen=0.06):
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Cm(margen)
    tf.margin_top = tf.margin_bottom = Cm(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = texto
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = fuente
    return sh

def _celda(slide, x, y, w, h, texto="", relleno=BLANCO, borde=NEGRO,
           ancho=Pt(1), ajustar=True, **kw):
    sh = _caja(slide, x, y, w, h, relleno, borde, ancho)
    if texto != "":
        if ajustar:
            kw["size"] = _achicar(texto, w, h, kw.get("size", 11),
                                  kw.get("margen", 0.06),
                                  fuente=kw.get("fuente", "Calibri"),
                                  bold=kw.get("bold", False))
        _txt(sh, texto, **kw)
    return sh

def _banda(slide, x, y, w, h, texto, size=11, azul=AZUL):
    return _celda(slide, x, y, w, h, texto, relleno=azul, size=size, bold=True, color=BLANCO)

# ─── 1. CAJETÍN OFICIAL (I-IN-002.4-R01) ─────────────────────────────────────
C_OP, C_DEN, C_CLI, C_PUE = 3.40, 10.60, 4.20, 3.00
C_LAB, C_VAL = 2.80, 4.30
C_MOD = C_CLI + C_PUE

def cajetin(slide, d, logo=None):
    top_h = 1.60
    fila = (HDR_H - top_h) / 4

    lw, rw = 4.30, C_LAB + C_VAL
    _caja(slide, X0, HDR_Y, lw, top_h, BLANCO, borde=NEGRO, ancho=Pt(1))
    if logo and os.path.exists(logo):
        im = Image.open(logo)
        ar = im.width / im.height
        ih = min(top_h - 0.30, (lw - 0.50) / ar)
        iw = ih * ar
        slide.shapes.add_picture(logo, Cm(X0 + (lw - iw) / 2),
                                 Cm(HDR_Y + (top_h - ih) / 2), Cm(iw), Cm(ih))

    _celda(slide, X0 + lw, HDR_Y, X1 - X0 - lw - rw, top_h,
           d.get("titulo_hoja", "HOJA DE OPERACIONES"), size=24, bold=True)
    _celda(slide, X1 - rw, HDR_Y, rw, top_h * 0.42,
           f"Form: {d.get('form', 'I-IN-002.4-R01')}", size=10.5, bold=True)
    _celda(slide, X1 - rw, HDR_Y + top_h * 0.42, rw, top_h * 0.58,
           d.get("ho", "HO-TBD"), size=20, bold=True)

    y = HDR_Y + top_h
    izq = [
        [("N° DE OPERACIÓN", C_OP), ("DENOMINACION DE LA OPERACIÓN", C_DEN),
         ("MODELO O VEHICULO", C_MOD)],
        [(d.get("op", ""), C_OP), (d.get("denominacion", ""), C_DEN),
         (d.get("modelo", ""), C_MOD)],
        [("SECTOR", C_OP), ("COD. DE PIEZA / DESCRIPCION", C_DEN),
         ("CLIENTE", C_CLI), ("N° PUESTO", C_PUE)],
        [(d.get("sector", ""), C_OP), (d.get("pieza", ""), C_DEN),
         (d.get("cliente", ""), C_CLI), (d.get("puesto", "-"), C_PUE)],
    ]
    der = [("REALIZO:", d.get("realizo", "")), ("APROBO:", d.get("aprobo", "")),
           ("FECHA:", d.get("fecha", "")), ("REV.", d.get("rev", "A"))]

    for i, fila_datos in enumerate(izq):
        etiqueta = (i % 2 == 0)
        yy = y + i * fila
        x = X0
        for texto, an in fila_datos:
            _celda(slide, x, yy, an, fila, str(texto),
                   relleno=(AZUL if etiqueta else BLANCO),
                   color=(BLANCO if etiqueta else NEGRO),
                   size=(8 if etiqueta else 9.5),
                   bold=(not etiqueta))
            x += an
        lab, val = der[i]
        _celda(slide, x, yy, C_LAB, fila, lab, relleno=AZUL, color=BLANCO,
               size=8.5, align=PP_ALIGN.LEFT, margen=0.12)
        _celda(slide, x + C_LAB, yy, C_VAL, fila, str(val),
               size=9.5, bold=True)

# ─── 2. BLOQUE DE IMÁGENES MULTI-FOTO SECUENCIAL ─────────────────────────────
def _badge_numero(slide, x, y, n, diam=0.86):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Cm(x), Cm(y), Cm(diam), Cm(diam))
    sh.fill.solid()
    sh.fill.fore_color.rgb = AZUL
    sh.line.color.rgb = BLANCO
    sh.line.width = Pt(1.25)
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(n)
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.name = "Calibri"
    r.font.color.rgb = BLANCO
    return sh

def _pie_foto(slide, x, y, w, h, texto):
    _celda(slide, x, y, w, h, texto, relleno=AZUL, borde=AZUL, color=BLANCO,
           size=8.5, bold=True, align=PP_ALIGN.CENTER)

MAX_FOTOS = 4

def bloque_imagenes(slide, imagenes, pies=None, numerar=True, principal=None):
    _banda(slide, IMG_X, BODY_Y, IMG_W, 0.60, "IMÁGENES", size=12)
    y0 = BODY_Y + 0.60
    h = BODY_H - 0.60
    _caja(slide, IMG_X, y0, IMG_W, h, BLANCO, borde=NEGRO, ancho=Pt(1))

    faltan = [i for i in imagenes if not os.path.exists(i)]
    if faltan:
        raise SystemExit("FOTOS QUE NO EXISTEN:\n  " + "\n  ".join(faltan))
    if not imagenes:
        return
    n = len(imagenes)
    if n > MAX_FOTOS:
        raise SystemExit(f"{n} fotos en una hoja: el tope es {MAX_FOTOS}.")

    pad = 0.15
    pie_h = 0.52
    W_util = IMG_W - 2 * pad
    H_util = h - 2 * pad

    ars = [Image.open(p).width / Image.open(p).height for p in imagenes]

    if principal is not None and HL is not None:
        layout = HL.layout_principal(ars, W_util, H_util - pie_h, principal)
        cajas = [None] * n
        otras = [k for k in range(n) if k != principal]
        pr_box = layout[principal]
        if pr_box[1] == 0.0 and len(otras) > 1 and layout[otras[0]][1] > 0.5:
            h_top = pr_box[3]
            h_bottom = (H_util - pie_h) - h_top - 0.16
            slot_w = W_util / len(otras)
            cajas[principal] = (pr_box[0], pr_box[1], pr_box[2], pr_box[3],
                                pr_box[0], pr_box[1] + pr_box[3] + 0.04, pr_box[2], pie_h)
            for idx_otra, k in enumerate(otras):
                iw, ih = layout[k][2], layout[k][3]
                slot_x = idx_otra * slot_w
                px = slot_x + (slot_w - iw) / 2
                py = h_top + 0.16 + (h_bottom - ih) / 2
                pw = min(slot_w - 0.20, max(iw, 4.5))
                p_x = slot_x + (slot_w - pw) / 2
                cajas[k] = (px, py, iw, ih, p_x, py + ih + 0.04, pw, pie_h)
        else:
            for k, (lx, ly, lw, lh) in enumerate(layout):
                cajas[k] = (lx, ly, lw, lh, lx, ly + lh + 0.04, lw, pie_h)
    else:
        if all(a < 1.0 for a in ars):
            cols = n
            filas = 1
        elif all(a >= 1.2 for a in ars):
            cols = 2 if n > 1 else 1
            filas = (n + cols - 1) // cols
        else:
            lo = HL.layout_orientacion(ars, W_util, H_util - pie_h) if HL else None
            if lo and min(w * hh for _, _, w, hh in lo) >= 25.0:
                cajas = []
                for lx, ly, lw, lh in lo:
                    cajas.append((lx, ly, lw, lh, lx, ly + lh + 0.04, lw, pie_h))
                cols = None
            else:
                cols = 2 if n > 1 else 1
                filas = (n + cols - 1) // cols

        if cols is not None:
            cw = W_util / cols
            ch = H_util / filas
            ch_foto = ch - pie_h - 0.06
            cajas = []
            for k, ar in enumerate(ars):
                fila = k // cols
                col = k % cols
                items_fila = min(cols, n - fila * cols)
                start_x = (W_util - items_fila * cw) / 2
                cell_x = start_x + col * cw
                cell_y = fila * ch
                ih = min(ch_foto, (cw - 0.20) / ar)
                iw = ih * ar
                fx = cell_x + (cw - iw) / 2
                fy = cell_y + (ch_foto - ih) / 2
                p_y = fy + ih + 0.04
                p_w = min(cw - 0.20, max(iw, 4.5))
                p_x = cell_x + (cw - p_w) / 2
                cajas.append((fx, fy, iw, ih, p_x, p_y, p_w, pie_h))

    for k, (fx, fy, fw, fh, p_x, p_y, p_w, p_h) in enumerate(cajas):
        px = IMG_X + pad + fx
        py = y0 + pad + fy
        ruta = imagenes[k]
        slide.shapes.add_picture(ruta, Cm(px), Cm(py), Cm(fw), Cm(fh))
        if numerar:
            _badge_numero(slide, px + 0.10, py + 0.10, k + 1)
        texto_pie = pies[k] if pies and k < len(pies) else ""
        if texto_pie:
            pie_abs_x = IMG_X + pad + p_x
            pie_abs_y = y0 + pad + p_y
            _pie_foto(slide, pie_abs_x, pie_abs_y, p_w, p_h, texto_pie)

# ─── 3. BLOQUE DE DESCRIPCIÓN DE LA OPERACIÓN ────────────────────────────────
def bloque_pasos(slide, pasos, nota=None, parametros=None):
    _banda(slide, DSC_X, BODY_Y, DSC_W, 0.60, "DESCRIPCION DE LA OPERACIÓN", size=12)
    y = BODY_Y + 0.60
    h = BODY_H - 0.60
    sh = _caja(slide, DSC_X, y, DSC_W, h, BLANCO, borde=NEGRO, ancho=Pt(1))
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = Cm(0.25)
    tf.margin_top = Cm(0.20)

    util_h = h - 0.70
    util_w = DSC_W - 2 * 0.25 - 0.85
    size = 12.0
    while size > 7.0:
        lh = 1.22 * size * PT_CM
        sep = (6 if size >= 10 else 4) * PT_CM
        alto = 0
        for t in pasos:
            alto += _lineas_wrap(t, size, util_w, False) * lh + sep
        for k, v in (parametros or []):
            alto += 1.22 * max(size - 0.5, 8.0) * PT_CM + 2 * PT_CM
        if nota:
            sn = max(size - 1, 7.5)
            alto += (_lineas_wrap(nota, sn, DSC_W - 2 * 0.25, True) * 1.22 * sn * PT_CM + 5 * PT_CM)
        if alto <= util_h:
            break
        size -= 0.5

    for i, texto in enumerate(pasos):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(4.0 if size >= 10.0 else 2.5)
        
        r = p.add_run()
        r.text = f"{i+1}.  "
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.name = "Calibri"
        r.font.color.rgb = NEGRO

        r2 = p.add_run()
        r2.text = texto
        r2.font.size = Pt(size)
        r2.font.name = "Calibri"
        r2.font.color.rgb = NEGRO

    for k, v in (parametros or []):
        p = tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(2.0)
        r = p.add_run()
        r.text = "▸ " + k + ": "
        r.font.size = Pt(max(size - 0.5, 8.0))
        r.font.name = "Calibri"
        r.font.color.rgb = NEGRO
        r2 = p.add_run()
        r2.text = str(v)
        r2.font.size = Pt(max(size - 0.5, 8.0))
        r2.font.bold = True
        r2.font.name = "Calibri"
        r2.font.color.rgb = AZUL_TEXTO

    if nota:
        p = tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(3.0)
        r = p.add_run()
        r.text = "NOTA: " + nota
        r.font.size = Pt(max(size - 1, 7.5))
        r.font.bold = True
        r.font.name = "Calibri"
        r.font.color.rgb = AZUL_TEXTO

# ─── 4. BLOQUE CICLO DE CONTROL (VACIADO SISTEMÁTICO LISTO PARA CALIDAD) ─────
COLS_CIC = [("Características a controlar", 8.0), ("Método de control", 5.6),
            ("Resp.", 2.7), ("Frec.", 2.9), ("Registro", 3.4)]

def bloque_ciclo(slide, filas=None):
    _banda(slide, X0, CIC_Y, CIC_W, 0.55, "CICLO DE CONTROL", size=12)
    y = CIC_Y + 0.55
    total = sum(w for _, w in COLS_CIC)
    anchos = [w / total * CIC_W for _, w in COLS_CIC]
    hh = 0.48
    x = X0
    for (lab, _), an in zip(COLS_CIC, anchos):
        _celda(slide, x, y, an, hh, lab, relleno=AZUL, color=BLANCO, size=8.5, bold=True)
        x += an
    y += hh
    
    filas = filas or [("", "", "", "", ""), ("", "", "", "", "")]
    fh = (CIC_Y + CIC_H - y) / len(filas)
    for f in filas:
        x = X0
        for val, an in zip(f, anchos):
            _celda(slide, x, y, an, fh, str(val), size=8.5)
            x += an
        y += fh

# ─── 5. ELEMENTOS DE SEGURIDAD (EPP) ─────────────────────────────────────────
def bloque_epp(slide, iconos, refs=("OP - Operador de Producción",)):
    _banda(slide, EPP_X, CIC_Y, EPP_W, 0.55, "ELEMENTOS DE SEGURIDAD", size=9, azul=AZUL2)
    y = CIC_Y + 0.55
    h = CIC_H - 0.55 - 0.44 * len(refs)
    _caja(slide, EPP_X, y, EPP_W, h, BLANCO, borde=NEGRO, ancho=Pt(1))
    
    iconos_ok = [ic for ic in (iconos or []) if os.path.exists(ic)]
    if iconos_ok:
        n = len(iconos_ok)
        cw = (EPP_W - 0.20) / n
        s = min(cw - 0.10, h - 0.16)
        for k, ic in enumerate(iconos_ok):
            cx = EPP_X + 0.10 + k * cw + (cw - s) / 2
            slide.shapes.add_picture(ic, Cm(cx), Cm(y + (h - s) / 2), Cm(s), Cm(s))

    yr = y + h
    for r in refs:
        _celda(slide, EPP_X, yr, EPP_W, 0.44, f"Referencia: {r}", size=7,
               align=PP_ALIGN.LEFT, margen=0.12)
        yr += 0.44

# ─── 6. PLAN DE REACCIÓN ANTE NO CONFORME (ESPECÍFICO POR OPERACIÓN) ─────────
FIJAS = ["DETENGA LA OPERACIÓN",
         "NOTIFIQUE DE INMEDIATO A SU LIDER O SUPERVISOR",
         "ESPERE LA DEFINICION DEL LIDER O SUPERVISOR"]

def bloque_plan(slide, disparador, acciones=None):
    _banda(slide, X0, PLN_Y, X1 - X0, 0.48, "PLAN DE REACCION ANTE NO CONFORME", size=11)
    y = PLN_Y + 0.48
    hh = PLN_H - 0.48
    izq = 15.20

    _celda(slide, X0, y, izq, hh * 0.28, disparador, size=8.0, bold=True,
           fuente="Arial", relleno=GRISF, align=PP_ALIGN.LEFT, margen=0.15)
    fy = y + hh * 0.28
    fh = (hh * 0.72) / 3
    for t in FIJAS:
        _celda(slide, X0, fy, izq, fh, t, size=8.5, bold=True, fuente="Arial",
               align=PP_ALIGN.LEFT, margen=0.15)
        fy += fh

    der = X1 - X0 - izq
    sh = _caja(slide, X0 + izq, y, der, hh, BLANCO, borde=NEGRO, ancho=Pt(1))
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = Cm(0.20)
    tf.margin_top = Cm(0.10)

    acciones = acciones or [
        "1. Segregar e identificar el material afectado en el cajón de scrap.",
        "2. Dar aviso según procedimiento P-09/I.",
        "3. No reiniciar la producción sin autorización del Líder o Supervisor."
    ]
    for i, t in enumerate(acciones):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(1.5)
        r = p.add_run()
        r.text = t
        r.font.size = Pt(8.0)
        r.font.name = "Arial"
        r.font.color.rgb = NEGRO

# ─── 7. PORTADA CORPORATIVA ADAPTABLE ────────────────────────────────────────
def portada(prs, d, logo=None, foto=None, ops_indice=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _caja(slide, X0, M, X1 - X0, H - 2 * M, BLANCO, borde=NEGRO, ancho=Pt(1.5))

    cab_h = 3.20
    _caja(slide, X0, M, X1 - X0, cab_h, BLANCO, borde=AZUL, ancho=Pt(1.5))
    _caja(slide, X0, M + cab_h - 0.08, X1 - X0, 0.08, AZUL, borde=AZUL)

    lw = 5.20
    if logo and os.path.exists(logo):
        im = Image.open(logo)
        ar = im.width / im.height
        ih = min(cab_h - 0.60, (lw - 0.80) / ar)
        iw = ih * ar
        slide.shapes.add_picture(logo, Cm(X0 + (lw - iw) / 2),
                                 Cm(M + (cab_h - ih) / 2), Cm(iw), Cm(ih))

    tx = X0 + lw + 0.20
    tw = X1 - X0 - lw - 0.40
    _celda(slide, tx, M + 0.40, tw, 1.40,
           d.get("titulo", "HOJAS DE PROCESO — MÁQUINA HOTMELT"),
           size=24, bold=True, color=AZUL, relleno=BLANCO, borde=None, align=PP_ALIGN.LEFT)
    _celda(slide, tx, M + 1.80, tw, 1.00,
           d.get("subtitulo", "Laminadora de vinilo con adhesivo hot melt · OP 20 del FLUJOGRAMA 155 TOP ROLL PATAGONIA"),
           size=11.5, bold=False, color=AZUL2, relleno=BLANCO, borde=None, align=PP_ALIGN.LEFT)

    y_body = M + cab_h + 0.35
    fw = 13.80
    bh = H - M - y_body - 0.10
    if foto and os.path.exists(foto):
        im = Image.open(foto)
        ar = im.width / im.height
        iw = fw - 0.20
        ih = iw / ar
        if ih > bh - 0.20:
            ih = bh - 0.20
            iw = ih * ar
        bh_real = min(bh, ih + 0.20)
        top = y_body + (bh - bh_real) / 2
        _caja(slide, X0 + 0.10, top, fw, bh_real, BLANCO, borde=AZUL, ancho=Pt(1))
        slide.shapes.add_picture(foto, Cm(X0 + 0.10 + (fw - iw) / 2),
                                 Cm(top + (bh_real - ih) / 2), Cm(iw), Cm(ih))
    else:
        _caja(slide, X0 + 0.10, y_body, fw, bh, BLANCO, borde=AZUL, ancho=Pt(1))

    xd = X0 + fw + 0.50
    wd = X1 - xd - 0.10
    filas = [
        ("Documento SGC", d.get("ho", "HO-TBD")),
        ("Formulario Oficial", d.get("form", "I-IN-002.4-R01")),
        ("Operación Flujograma", d.get("op_flujo", "20 — ADHESIVADO HOT MELT")),
        ("Cliente / Modelo", d.get("cliente_modelo", "VW / PATAGONIA")),
        ("Pieza / Conjunto", d.get("pieza", "TOP ROLL PATAGONIA — N 216 / N 256 / N 285 / N 315")),
        ("Máquina / Celda", d.get("maquina", "Laminadora hot melt KINGPOWER + fusor de adhesivo")),
        ("Elaboró / Aprobó", d.get("firmas", "F. Santoro / C. Baptista")),
        ("Fecha / Revisión", d.get("fecha_rev", "28/09/2026 · Rev. A")),
    ]
    hh = 0.65
    y_meta = y_body + 0.15
    for lab, val in filas:
        _celda(slide, xd, y_meta, 4.60, hh, lab, relleno=AZUL, color=BLANCO,
               size=8.5, bold=True, align=PP_ALIGN.LEFT, margen=0.15)
        _celda(slide, xd + 4.60, y_meta, wd - 4.60, hh, val, relleno=BLANCO,
               color=NEGRO, size=9, bold=True, align=PP_ALIGN.LEFT, margen=0.15)
        y_meta += hh + 0.08

    # Cuadro de Índice de Operaciones
    ops_indice = ops_indice if ops_indice is not None else [
        ("20.1", "Puesta en marcha y energización"),
        ("20.2", "Puesta en marcha del fusor de adhesivo"),
        ("20.3", "Acceso al HMI y reset de alarmas"),
        ("20.4", "Carga de parámetros de producto y temp."),
        ("20.5", "Calentamiento y espera de temperatura"),
        ("20.6", "Montaje del rollo en el desbobinador"),
        ("20.7", "Enhebrado del material en serpentina"),
        ("20.8", "Centrado y tensión de la banda"),
        ("20.9", "Arranque y alineación de borde"),
        ("20.10", "Laminado — control durante la marcha"),
        ("20.11", "Empalme del material con cinta adhesiva"),
        ("20.12", "Cambio de rollo por fin de material"),
        ("20.13", "Destrabe de material atascado"),
        ("20.14", "Corte de la plancha de vinilo"),
        ("20.15", "Parada y bloqueo de la máquina"),
        ("20.16", "Apertura de rodillos y bandeja"),
        ("20.17", "Limpieza de rodillos en caliente"),
    ]
    _banda(slide, xd, y_meta + 0.15, wd, 0.55, d.get("titulo_indice", "ÍNDICE DE OPERACIONES ESTÁNDAR"), size=10)
    y_ind = y_meta + 0.70
    hh_ind = (H - M - y_ind - 0.20) / len(ops_indice)
    hh_ind = min(hh_ind, 0.62)
    for num, txt in ops_indice:
        _celda(slide, xd, y_ind, 1.40, hh_ind, f"OP {num}", size=7.5, bold=True)
        _celda(slide, xd + 1.40, y_ind, wd - 1.40, hh_ind, txt, size=7.5,
               bold=False, align=PP_ALIGN.LEFT, margen=0.12)
        y_ind += hh_ind

# ─── 8. HOJA COMPLETA ────────────────────────────────────────────────────────
def hoja(prs, d, logo=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _caja(slide, X0, M, X1 - X0, H - 2 * M, BLANCO, borde=NEGRO, ancho=Pt(1.5))
    cajetin(slide, d, logo)
    bloque_imagenes(slide, d.get("imagenes", []), d.get("pies"), numerar=True,
                    principal=d.get("principal"))
    bloque_pasos(slide, d.get("pasos", []), nota=d.get("nota"), parametros=d.get("parametros"))
    bloque_ciclo(slide, d.get("ciclo", []))
    bloque_epp(slide, d.get("epp", EPP_HOTMELT))
    bloque_plan(slide, d.get("disparador", 'SI DETECTA "PRODUCTO" O "PROCESO" NO CONFORME'),
                d.get("acciones"))
    return slide

# ─── 9. ESPECIFICACIÓN CANÓNICA DE HOJAS (17 LÁMINAS) ────────────────────────
def P(nombre):
    """Resuelve la ruta absoluta de una foto en fotos_hoja o ruta directa."""
    if os.path.isabs(nombre) and os.path.exists(nombre):
        return nombre
    p = os.path.join(FOTOS_DIR, nombre)
    if os.path.exists(p):
        return p
    p_jpg = os.path.join(FOTOS_DIR, nombre + ".jpg")
    if os.path.exists(p_jpg):
        return p_jpg
    p_png = os.path.join(FOTOS_DIR, nombre + ".png")
    if os.path.exists(p_png):
        return p_png
    p_local = os.path.join(AQUI, nombre)
    if os.path.exists(p_local):
        return p_local
    p_local_png = os.path.join(AQUI, nombre + ".png")
    if os.path.exists(p_local_png):
        return p_local_png
    p_local_jpg = os.path.join(AQUI, nombre + ".jpg")
    if os.path.exists(p_local_jpg):
        return p_local_jpg
    return p_jpg

HOJAS_HOTMELT = [
    {
        "op": "20.1", "etapa": "PREPARAR Y ARRANCAR",
        "denominacion": "PUESTA EN MARCHA Y ENERGIZACIÓN",
        "imagenes": [P("h01_a_llave_general"), P("h01_a_desbobinador"), P("h01_b_rodillos_reja")],
        "pies": ["Interruptor general principal", "Sector del desbobinador", "Zona de rodillos de aplicación"],
        "pasos": [
            "Girar a posición ON el interruptor general principal de la máquina.",
            "Verificar la presión del suministro de aire comprimido en el manómetro de entrada (6 bar).",
            "Inspeccionar visualmente que los rodillos no tengan restos metálicos, suciedad ni adhesivo endurecido.",
            "Verificar que los botones de parada de emergencia estén rearmados y despejados."
        ],
        "nota": "Los rodillos trabajan entre 150 y 185 °C. Operar con guantes térmicos, ropa de trabajo ajustada y sin cordones ni elementos colgantes que impliquen riesgo de atrapamiento. Trabajar con barbijo para protección contra vapores de adhesivo.",
        "disparador": "SI LA MÁQUINA NO ENCIENDE O FALTA PRESIÓN NEUMÁTICA (< 6 BAR)",
        "acciones": [
            "1. Verificar que el interruptor general y las paradas de emergencia estén liberadas.",
            "2. Verificar el manómetro de la línea general de aire comprimido.",
            "3. Dar aviso al Líder de Producción y a Mantenimiento; no puentear seguridades."
        ]
    },
    {
        "op": "20.2", "etapa": "PREPARAR Y ARRANCAR",
        "principal": 2,
        "denominacion": "PUESTA EN MARCHA DEL FUSOR DE ADHESIVO",
        "imagenes": [P("_foto_fusor_unidad"), P("h02_e_tanque_adhesivo"), FOTO_PANTALLA_FUSOR],
        "pies": ["Unidad fusor de adhesivo PUR", "Tanque de carga de adhesivo", "Pantalla de control del fusor"],
        "pasos": [
            "Encender el fusor de adhesivo accionando el interruptor rojo general.",
            "Verificar visualmente que el tanque contenga carga suficiente de adhesivo sólido.",
            "Habilitar el calentamiento del tanque, mangueras calefaccionadas y pistolas dosificadoras.",
            "Verificar en la pantalla del fusor la consigna de temperatura a 160 °C en los 5 circuitos.",
            "Esperar a que el adhesivo alcance el estado fundido completo antes de habilitar bombeo.",
            "Inspeccionar que las boquillas dosificadoras no presenten obstrucciones sólidas.",
            "Dar aviso de inmediato al líder si una boquilla queda obstruida; no intervenir en caliente."
        ],
        "nota": "El fusor alimenta el adhesivo por mangueras térmicas internas hacia dos caños de aplicación sobre el rodillo. La temperatura de trabajo es de 160 °C. El adhesivo fundido quema y se adhiere a la piel: utilizar guantes térmicos y protección ocular.",
        "disparador": "SI EL FUSOR NO CALIENTA, INDICA ALARMA O UNA BOQUILLA ESTÁ TAPADA",
        "acciones": [
            "1. No intentar desobstruir mangueras ni pistolas presurizadas en caliente.",
            "2. Verificar nivel de adhesivo sólido en el tanque de carga.",
            "3. Dar aviso inmediato a Mantenimiento y al Líder de Producción."
        ]
    },
    {
        "op": "20.3", "etapa": "PREPARAR Y ARRANCAR",
        "denominacion": "ENCENDIDO Y ACCESO AL HMI",
        "imagenes": [P("h03_b_panel"), P("h03_a_login_hmi"), P("h03_c_operario_panel")],
        "pies": ["Tablero de control principal", "Pantalla de acceso al HMI", "Puesto de mando del operador"],
        "pasos": [
            "Verificar el arranque completo del sistema operativo del panel HMI.",
            "Ingresar la contraseña de operador " + CLAVE_HMI + " en la pantalla de acceso del HMI.",
            "Pulsar el botón amarillo RESET en caso de existir fallas activas retenidas.",
            "Verificar en la pantalla principal que no queden alarmas activas en el sistema."
        ],
        "nota": "Las alarmas comunes bloquean el arranque: puertas abiertas, botones de parada accionados o fallas en ejes de tracción. Si el panel inicia en idioma chino, seleccionar la opción English en la esquina inferior.",
        "disparador": "SI LA PANTALLA HMI NO RESPONDE O PERSISTE UNA ALARMA ACTIVA TRAS EL RESET",
        "acciones": [
            "1. No insistir reiteradamente con el pulsador de arranque.",
            "2. Tomar nota del código de alarma o mensaje de falla visualizado en el HMI.",
            "3. Notificar al Líder de Producción y esperar asistencia técnica."
        ]
    },
    {
        "op": "20.4", "etapa": "PREPARAR Y ARRANCAR",
        "denominacion": "CARGA DE PARÁMETROS DE PRODUCTO Y TEMPERATURA",
        "imagenes": [P("h04_a_pantalla_operacion"), P("h04_b_parametros_temp")],
        "pies": ["Pantalla de recetas y producto", "Pantalla de control de temperaturas"],
        "pasos": [
            "Ingresar al menú de Selección de Receta en el panel HMI.",
            "Seleccionar la Receta 1 correspondiente a Top Roll Patagonia.",
            "Verificar la velocidad de línea programada en 3,0 m/min.",
            "Verificar la velocidad del rodillo dosificador programada en 0,040 m/min.",
            "Verificar el espesor de producto y la separación de luz entre rodillos (0,250 mm).",
            "Verificar en la pantalla de temperaturas la consigna de 185 °C para ambos rodillos.",
            "Verificar la temperatura de habilitación mínima de 150 °C (interlock de seguridad)."
        ],
        "parametros": [
            ("Velocidad de línea", "3,0 m/min"),
            ("Velocidad dosificador", "0,040 m/min"),
            ("Luz entre rodillos", "0,250 mm"),
            ("Temperatura consigna", "185 °C"),
            ("Temperatura interlock", "150 °C")
        ],
        "nota": "Las cuatro piezas de Top Roll (N 216, N 256, N 285 y N 315) comparten la misma receta con luz de 0,250 mm. Los rodillos poseen un interlock de seguridad: no arrancan a girar si la temperatura real está por debajo de 150 °C.",
        "disparador": "SI LA RECETA NO CARGA O LOS PARÁMETROS DIFIEREN DE LA HOJA DE PROCESO",
        "acciones": [
            "1. Prohibido modificar recetas o separación de rodillos sin autorización de Calidad.",
            "2. Recargar Receta 1 (Top Roll) y verificar velocidad (3,0 m/min) y luz (0,250 mm).",
            "3. Dar aviso al Líder de Producción para validar la configuración antes de arrancar."
        ]
    },
    {
        "op": "20.5", "etapa": "PREPARAR Y ARRANCAR",
        "denominacion": "CALENTAMIENTO Y ESPERA DE TEMPERATURA",
        "imagenes": [P("h05_a_calentando"), P("h05_b_calentamiento_ok")],
        "pies": ["Calentamiento de rodillos activo", "Temperatura de régimen alcanzada"],
        "pasos": [
            "Esperar a que los rodillos alcancen la temperatura de régimen (185 °C).",
            "Verificar en la pantalla del HMI el indicador de calentamiento completado.",
            "Verificar que no existan alarmas por desviación de temperatura (banda ±10 °C).",
            "Habilitar el giro lento de rodillos únicamente cuando se supere la cota de 150 °C.",
            "Verificar que el adhesivo comience a fluir automáticamente por las mangueras internas."
        ],
        "nota": "Desde máquina fría, el precalentamiento demora hasta 60 minutos. El adhesivo se alimenta automáticamente desde el fusor por mangueras térmicas internas hacia los rodillos (no colocar adhesivo manual). Los rodillos no arrancan a girar por debajo de 150 °C.",
        "disparador": "SI LOS RODILLOS NO LLEGAN A 185 °C O HAY DESVIACIÓN MAYOR A ±10 °C",
        "acciones": [
            "1. No intentar forzar el giro de rodillos si la temperatura es inferior a 150 °C.",
            "2. Verificar en el HMI el estado de las resistencias de calentamiento.",
            "3. Dar aviso al Líder de Producción y Mantenimiento si no alcanza régimen en 60 min."
        ]
    },
    {
        "op": "20.6", "etapa": "PREPARAR Y ARRANCAR",
        "denominacion": "MONTAJE DEL ROLLO EN EL DESBOBINADOR",
        "imagenes": [P("rollo_2_montaje"), P("h06_d_mandril_aire"), P("rollo_3_onoff")],
        "pies": ["Posicionamiento del rollo", "Inserción de eje expansible", "Bloqueo neumático en soporte"],
        "pasos": [
            "Verificar que el código y lote de la bobina coincidan con la orden de producción.",
            "Fijar la punta suelta del vinilo con cinta adhesiva antes de mover el rollo.",
            "Meter el eje expansible adentro del tubo de cartón del rollo.",
            "Centrar el rollo sobre el eje usando como guía la regla graduada.",
            "Cargar aire con la pistola en la válvula del eje hasta que quede bien apretado.",
            "Montar el rollo entre dos operarios sobre las cunas del desbobinador.",
            "Girar la perilla de soporte a posición ON para trabarlo."
        ],
        "nota": "Orientar el rollo de modo que el vinilo salga hacia adelante, con la cara texturada hacia arriba y desenrollando directo hacia los rodillos.",
        "disparador": "SI EL EJE NEUMÁTICO PIERDE PRESIÓN O EL ROLLO QUEDA FLOJO / DESCENTRADO",
        "acciones": [
            "1. Detener la maniobra y no operar con el rollo mal amarrado al eje expansible.",
            "2. Revisar la válvula de inflado del eje y verificar presión de aire con el manómetro.",
            "3. Reajustar el centrado con la regla graduada antes de habilitar el soporte."
        ]
    },
    {
        "op": "20.7", "etapa": "PREPARAR Y ARRANCAR",
        "denominacion": "ENHEBRADO DEL MATERIAL EN SERPENTINA",
        "imagenes": [P("h07_b_primer_rodillo"), P("h07_f_serpentina"), P("h07_d_sensor_negro")],
        "pies": ["Ingreso por rodillo frontal", "Paso por serpentina y tensor", "Sensor óptico de guiado inferior"],
        "pasos": [
            "Poner el selector de tensión del desbobinador en modo MANUAL.",
            "Tirar de la punta del vinilo y pasarla por el primer rodillo frontal.",
            "Pasar la lámina respetando el recorrido en serpentina a través de los rodillos guía y el tensor.",
            "Pasar el vinilo a través de la barra de alineación inferior y el sensor óptico.",
            "Colocar la punta de la lámina en la boca de entrada de los rodillos calefaccionados."
        ],
        "nota": "ATENCIÓN: El enhebrado completo en serpentina se realiza únicamente al iniciar la máquina desde cero o si se cortó el vinilo. En producción normal, el cambio de rollo se hace por EMPALME rápido con cinta adhesiva (ver OP 20.11), sin volver a pasar todo el circuito.",
        "disparador": "SI EL VINILO SE CRUZA, SE PLIEGA O SE ENGANCHA AL PASAR POR LOS RODILLOS",
        "acciones": [
            "1. Mantener los rodillos detenidos y la reja de resguardo abierta durante el enhebrado.",
            "2. Cortar el tramo deformado de vinilo con cutter de seguridad de punta redondeada.",
            "3. Reenhebrar la serpentina asegurando que la cara texturada quede hacia arriba."
        ]
    },
    {
        "op": "20.8", "etapa": "PREPARAR Y ARRANCAR",
        "denominacion": "CENTRADO Y TENSIÓN DE LA BANDA",
        "imagenes": [P("h08_a_centrado"), P("h08_b_sensor_borde"), P("h08_c_tablero_tension")],
        "pies": ["Centrado en barra de medición", "Ajuste de sensor óptico de borde", "Controlador digital de tensión"],
        "pasos": [
            "Alinear el borde del vinilo con la marca milimetrada de la barra guía.",
            "Posicionar el sensor óptico enfocado al borde de la lámina.",
            "Verificar en la pantalla de tensión la consigna de trabajo (20 kg).",
            "Girar el selector de tensión de MANUAL a AUTOMÁTICO.",
            "Verificar que el vinilo avance parejo, estirado y sin arrugas."
        ],
        "nota": "Una tensión baja causa arrugas en la lámina; una tensión excesiva estira el vinilo y modifica las medidas de la pieza.",
        "disparador": "SI LA BANDA AVANZA CON ARRUGAS, ONDULACIONES O DESVÍO LATERAL",
        "acciones": [
            "1. Pasar el control de tensión a MANUAL y verificar alineación contra la regla guía.",
            "2. Limpiar el lente del sensor óptico de borde si pierde la referencia de lámina.",
            "3. Reajustar la tensión a 20 kg y conmutar a AUTOMÁTICO bajo supervisión del Líder."
        ]
    },
    {
        "op": "20.9", "etapa": "PREPARAR Y ARRANCAR",
        "principal": 0,
        "denominacion": "ARRANQUE Y ALINEACIÓN DE BORDE",
        "imagenes": [P("h09_a_hmi_arranque"), P("h09_b_reset"), P("h09_b_rodillo_gira")],
        "pies": ["Selección de modo automático", "Pulsador de rearme RESET", "Avance y arrastre de rodillos"],
        "pasos": [
            "Verificar que la reja de resguardo de los rodillos esté cerrada y bloqueada.",
            "Pulsar el botón de arranque de ciclo en la botonera de mando.",
            "Verificar en la pantalla del HMI el cambio a estado AUTOMÁTICO.",
            "Verificar el arrastre sincronizado del vinilo a través de los rodillos.",
            "Controlar visualmente que el borde del vinilo mantenga la marca de alineación.",
            "Detener la marcha y corregir la posición en la guía si la lámina se corre de costado."
        ],
        "nota": "Probar preventivamente el pulsador de parada de emergencia en vacío antes de iniciar el laminado de producción.",
        "disparador": "SI AL ARRANCAR SE DISPARA UNA ALARMA O EL MATERIAL DERIVA DEL CENTRO",
        "acciones": [
            "1. Detener la marcha desde la botonera de mando de la celda.",
            "2. Corregir la alineación lateral de la banda antes de habilitar el avance continuo.",
            "3. Notificar al Líder si el sistema automático de guiado no compensa el desvío."
        ]
    },
    {
        "op": "20.10", "etapa": "PRODUCIR",
        "denominacion": "LAMINADO — CONTROL DURANTE LA MARCHA",
        "imagenes": [P("h10_b_material_saliendo"), P("h10_c_enfriamiento")],
        "pies": ["Salida de vinilo hacia la mesa", "Zona de rodillos de enfriamiento y sopladores"],
        "pasos": [
            "Inspeccionar de forma continua que el vinilo salga plano y libre de arrugas.",
            "Verificar que la capa de adhesivo hot melt sea uniforme, continua y sin grumos.",
            "Controlar que la superficie del vinilo no presente marcas por sobrecalentamiento.",
            "Verificar periódicamente la temperatura de proceso en los displays de los rodillos.",
            "Verificar el funcionamiento del soplador de aire de enfriamiento a la salida."
        ],
        "nota": "No tocar bajo ninguna circunstancia el material ni los rodillos en la zona de aplicación mientras la máquina esté en movimiento. Ante anomalías, detener la línea desde el pulsador.",
        "disparador": "SI APARECEN ARRUGAS, BURBUJAS O LA CAPA DE ADHESIVO ES DISCONTINUA",
        "acciones": [
            "1. Detener inmediatamente el avance de línea desde el pulsador de parada.",
            "2. Segregar el material no conforme e identificarlo con tarjeta roja de SCRAP.",
            "3. Dar aviso inmediato al Líder de Producción y a Control de Calidad."
        ]
    },
    {
        "op": "20.11", "etapa": "SI PASA ALGO",
        "denominacion": "EMPALME DEL MATERIAL CON CINTA ADHESIVA",
        "imagenes": [P("h11_a_corte_diagonal"), P("h11_c_marca_referencia"), P("h11_b_cinta_blanca")],
        "pies": ["Corte recto de extremo", "Alineación contra barra patrón", "Fijación con cinta adhesiva"],
        "pasos": [
            "Detener el avance de la máquina accionando el pulsador de parada de ciclo.",
            "Cortar el extremo final de la bobina agotada en ángulo recto sobre la barra de corte.",
            "Marcar la posición de empalme sobre la barra guía de referencia.",
            "Presentar el extremo inicial de la nueva bobina a tope contra el corte anterior.",
            "Aplicar cinta adhesiva de alta resistencia cubriendo toda la línea de unión.",
            "Presionar firmemente la cinta adhesiva para asegurar una unión plana sin bordes levantados.",
            "Rearmar los seguros de máquina y reiniciar la marcha a velocidad reducida."
        ],
        "nota": "Esta maniobra se realiza en producción cada vez que se termina una bobina para empalmar la nueva sin tener que reenhebrar toda la serpentina. Marcar el tramo con cinta adhesiva con marcador rojo: ese sector debe ser descartado como scrap y nunca puede ingresar al termoformado.",
        "disparador": "SI EL EMPALME QUEDA DESALINEADO, CON ESPACIO ABIERTO O CINTA ADHESIVA FLOJA",
        "acciones": [
            "1. Detener el arrastre y no permitir que un empalme defectuoso ingrese a rodillos.",
            "2. Retirar la cinta adhesiva, reajustar los bordes a tope recto y volver a aplicar cinta adhesiva limpia.",
            "3. Marcar con fibrón indeleble rojo la zona empalmada para su posterior descarte en mesa."
        ]
    },
    {
        "op": "20.12", "etapa": "SI PASA ALGO",
        "denominacion": "CAMBIO DE ROLLO POR FIN DE MATERIAL",
        "imagenes": [P("h12_a_pulsador_rojo"), P("h12_c_volante_mandril"), P("h12_e_rollo_nuevo")],
        "pies": ["Pulsador de parada en desbobinador", "Despresurización de eje expansible", "Montaje de bobina nueva"],
        "pasos": [
            "Verificar la detención automática de la máquina al activarse la alarma de fin de bobina.",
            "Pulsar la parada de seguridad del desbobinador.",
            "Despresurizar el eje expansible accionando la válvula de alivio.",
            "Retirar el buje de cartón vacío de los soportes mecánicos.",
            "Montar el nuevo rollo de vinilo verificando código de materia prima y sentido de giro.",
            "Cargar aire con la pistola en la válvula del eje hasta que quede bien apretado.",
            "Ejecutar el procedimiento de empalme de vinilo según la hoja 20.11."
        ],
        "nota": "Manipular los rollos pesados con la ayuda de un segundo operario o dispositivo de elevación para prevenir sobreesfuerzos lumbares.",
        "disparador": "SI EL EJE NO DESPRESURIZA O EL BUJE DE CARTÓN QUEDA TRABADO",
        "acciones": [
            "1. Verificar corte de aire y accionar manualmente la válvula de alivio del eje.",
            "2. No golpear el buje ni el eje con herramientas de acero para evitar dañarlo.",
            "3. Solicitar ayuda de un segundo operario o dar aviso a Mantenimiento."
        ]
    },
    {
        "op": "20.13", "etapa": "SI PASA ALGO",
        "denominacion": "DESTRABE DE MATERIAL ATASCADO",
        "imagenes": [P("h13_c_maquina_detenida"), P("h13_b_material_trabado"), P("h13_a_vinilo_mal_pasado")],
        "pies": ["Parada de máquina ante atasco", "Inspección de material trabado", "Verificación de recorrido por rodillos"],
        "pasos": [
            "Golpear el pulsador de parada de emergencia ante un atasco o enrollamiento en rodillos.",
            "Accionar el comando neumático de apertura y separación de rodillos.",
            "Cortar con herramienta de filo protegida el tramo de vinilo plegado o dañado.",
            "Retirar los restos de vinilo y adhesivo adheridos utilizando espátula de latón o teflón.",
            "Rearmar el botón de parada de emergencia y presionar RESET en el panel HMI.",
            "Reenhebrar el material a través de la serpentina antes de reanudar el ciclo."
        ],
        "nota": "No utilizar elementos cortantes de acero templado sobre la superficie de los rodillos recubiertos de silicona o cromo: rayaduras en los rodillos transfieren defectos a toda la producción.",
        "disparador": "SI EL VINILO SE ENROLLA EN LOS RODILLOS O SE PRODUCE UN ATASCO SEVERO",
        "acciones": [
            "1. Presionar de inmediato la Parada de Emergencia y accionar apertura neumática de rodillos.",
            "2. Utilizar únicamente espátula de latón o teflón; terminantemente prohibido usar acero.",
            "3. Segregar todo el material dañado e inspeccionar rodillos antes de reiniciar ciclo."
        ]
    },
    {
        "op": "20.14", "etapa": "TERMINAR",
        "denominacion": "CORTE DE LA PLANCHA DE VINILO",
        "imagenes": [P("h14_a_corte_cuchillo"), P("h14_b_guia_corte")],
        "pies": ["Corte sobre mesa de apoyo", "Guiado de corte perimetral"],
        "pasos": [
            "Posicionar la lámina laminada plana sobre la mesa de trabajo asignada.",
            "Utilizar guantes anticorte durante toda la maniobra.",
            "Realizar el corte de fraccionamiento con cutter de seguridad siguiendo la línea guía.",
            "Redondear los vértices del corte para evitar desgarros durante la posterior manipulación.",
            "Depositar las láminas cortadas en el contenedor intermedio correspondiente."
        ],
        "nota": "Está terminantemente prohibido cortar material directamente sobre el piso. El corte debe realizarse sobre mesa de trabajo plana y limpia para no contaminar el adhesivo con polvo o partículas.",
        "disparador": "SI LA PLANCHA PRESENTA REBABAS, CORTE DESVIADO O MEDIDAS FUERA DE TOLERANCIA",
        "acciones": [
            "1. Segregar las planchas cortadas fuera de medida en el contenedor de scrap.",
            "2. Reemplazar la hoja del cutter de seguridad si presenta desgaste o pérdida de filo.",
            "3. Cortar siempre sobre la mesa plana designada; prohibido cortar sobre el piso."
        ]
    },
    {
        "op": "20.15", "etapa": "TERMINAR",
        "principal": 0,
        "denominacion": "PARADA Y BLOQUEO DE LA MÁQUINA",
        "imagenes": [FOTO_SEGURIDAD],
        "pies": ["Pantalla de señales de seguridad (IO) — estado de actuadores y paradas"],
        "pasos": [
            "Pulsar el botón de parada de ciclo para detener el arrastre de material.",
            "Desactivar la alimentación de adhesivo desde el panel del fusor.",
            "Seleccionar en el HMI el modo de espera (standby a 150 °C) o enfriamiento a 50 °C.",
            "Verificar en la pantalla de seguridades que todos los actuadores queden en reposo.",
            "Girar el interruptor general a OFF únicamente cuando los rodillos bajen de temperatura."
        ],
        "nota": "Si la máquina se detiene por más de 30 minutos, se debe pasar a modo espera para no degradar ni quemar el adhesivo alojado en el circuito de distribución.",
        "disparador": "SI UN ACTUADOR NO QUEDA EN REPOSO O PERSISTE TEMPERATURA PELIGROSA",
        "acciones": [
            "1. Verificar que el fusor haya cortado el flujo de adhesivo a las mangueras.",
            "2. No cortar la energía general mientras los rodillos permanezcan a más de 50 °C.",
            "3. Reportar anomalías al Líder de Producción antes de retirarse del puesto."
        ]
    },
    {
        "op": "20.16", "etapa": "TERMINAR",
        "denominacion": "APERTURA DE RODILLOS Y BANDEJA",
        "imagenes": [P("h16_a_ejes_rotacion"), P("h16_b_plato_pegamento"), P("h16_c_plato_manchas")],
        "pies": ["Apertura y elevación de rodillos", "Colocación de bandeja recolectora", "Inspección y limpieza de adhesivo"],
        "pasos": [
            "Verificar que la rotación de los rodillos esté completamente detenida.",
            "Accionar el mando de apertura neumática hasta la separación máxima de los rodillos.",
            "Insertar la bandeja metálica recolectora debajo de la zona de goteo de los rodillos.",
            "Verificar que la bandeja quede correctamente apoyada sobre los topes mecánicos.",
            "Inspeccionar que no caiga adhesivo sobre las guías o partes motrices de la máquina."
        ],
        "nota": "La bandeja de goteo protege la estructura y el piso de la celda ante escurrimientos durante la limpieza o los períodos de parada prolongada.",
        "disparador": "SI LOS RODILLOS NO ABREN A TOPE O LA BANDEJA NO CALZA EN SUS SOPORTES",
        "acciones": [
            "1. Verificar la presión de aire de la línea del cilindro de apertura neumática.",
            "2. Limpiar restos de adhesivo endurecido en las guías de apoyo de la bandeja.",
            "3. Dar aviso a Mantenimiento si la separación mecánica no se completa."
        ]
    },
    {
        "op": "20.17", "etapa": "TERMINAR",
        "denominacion": "LIMPIEZA DE RODILLOS EN CALIENTE",
        "imagenes": [P("h17_a_pantalla_limpieza"), P("h17_c_rodillos_pegamento")],
        "pies": ["Menú de limpieza en panel HMI", "Película de adhesivo en rodillos"],
        "pasos": [
            "Ingresar al menú de Limpieza en la pantalla del panel HMI.",
            "Verificar que la temperatura de rodillos se encuentre entre 150 y 185 °C.",
            "Aplicar parafina industrial de limpieza sobre la superficie de los rodillos.",
            "Habilitar el ciclo de rotación lenta para disolver y arrastrar restos de adhesivo degradado.",
            "Retirar los residuos desprendidos empleando raspador no abrasivo de teflón.",
            "Drenar el residuo hacia la bandeja de recolección inferior.",
            "Verificar que los rodillos queden limpios antes del enfriamiento final del equipo."
        ],
        "nota": "La limpieza debe efectuarse en caliente. Usar guantes térmicos para alta temperatura, protección ocular y barbijo para vapores orgánicos. Prohibido usar cuchillas de acero que rayen los rodillos.",
        "disparador": "SI QUEDAN RESTOS DE ADHESIVO QUEMADO O SE OBSERVA UNA RAYADURA EN RODILLOS",
        "acciones": [
            "1. Reaplicar parafina sólida y repetir rotación lenta sin aumentar la presión de raspado.",
            "2. Prohibido usar cuchillos, lijas o espátulas duras que rayen el recubrimiento.",
            "3. Dar aviso inmediato al Líder y Mantenimiento si se detecta un rodillo rayado."
        ]
    }
]

# ─── 10. AGRUPACIONES POR PARTES (ESTRUCTURA DECK POR BLOQUES) ───────────────
OPS_MAP = {h["op"]: h for h in HOJAS_HOTMELT}

PARTE_1_OPS = ["20.1", "20.2", "20.3", "20.4", "20.5", "20.6", "20.7", "20.8", "20.9", "20.10", "20.14", "20.15"]
PARTE_2_OPS = ["20.11", "20.12", "20.13"]
PARTE_3_OPS = ["20.16", "20.17"]

HOJAS_PARTE_1 = [OPS_MAP[op] for op in PARTE_1_OPS]
HOJAS_PARTE_2 = [OPS_MAP[op] for op in PARTE_2_OPS]
HOJAS_PARTE_3 = [OPS_MAP[op] for op in PARTE_3_OPS]

# ─── 11. EXPORTADOR COM A PDF ────────────────────────────────────────────────
def exportar_pdf(pptx_path):
    """Exporta un archivo PPTX a PDF usando Microsoft PowerPoint COM."""
    import win32com.client
    pdf_path = os.path.splitext(pptx_path)[0] + ".pdf"
    ppt = win32com.client.Dispatch("PowerPoint.Application")
    habia_abiertas = ppt.Presentations.Count
    try:
        pres = ppt.Presentations.Open(os.path.abspath(pptx_path), WithWindow=False)
        pres.SaveAs(os.path.abspath(pdf_path), 32) # 32 = ppSaveAsPDF
        pres.Close()
    finally:
        if habia_abiertas == 0 and ppt.Presentations.Count == 0:
            ppt.Quit()
    print(f"     -> PDF generado: {pdf_path}")
    return pdf_path

# ─── 12. GENERACIÓN PRINCIPAL ────────────────────────────────────────────────
def generar_deck(lista_hojas, subtitulo_portada, ruta_pptx, titulo_indice="ÍNDICE DE OPERACIONES ESTÁNDAR"):
    os.makedirs(os.path.dirname(ruta_pptx), exist_ok=True)
    prs = Presentation()
    prs.slide_width = Cm(W)
    prs.slide_height = Cm(H)

    foto_portada = P("cand_30") if os.path.exists(P("cand_30")) else P("h01_a_desbobinador")
    indice_ops = [(h["op"], h["denominacion"]) for h in lista_hojas]

    datos_portada = {
        "titulo": "HOJAS DE PROCESO — MÁQUINA HOTMELT",
        "subtitulo": subtitulo_portada,
        "ho": "HO-TBD",
        "form": "I-IN-002.4-R01",
        "op_flujo": "20 — ADHESIVADO HOT MELT",
        "cliente_modelo": "VW / PATAGONIA",
        "pieza": "TOP ROLL PATAGONIA — N 216 / N 256 / N 285 / N 315",
        "maquina": "Laminadora hot melt KINGPOWER + fusor de adhesivo",
        "firmas": "F. Santoro / C. Baptista",
        "fecha_rev": "28/09/2026 · Rev. A",
        "titulo_indice": titulo_indice
    }
    portada(prs, datos_portada, logo=LOGO_BARACK, foto=foto_portada, ops_indice=indice_ops)

    cajetin_base = {
        "titulo_hoja": "HOJA DE OPERACIONES - PRELIMINAR",
        "ho": "HO-TBD",
        "form": "I-IN-002.4-R01",
        "modelo": "PATAGONIA",
        "cliente": "VW",
        "sector": "ADHESIVADO HOT MELT",
        "pieza": "TOP ROLL PATAGONIA — N 216 / N 256 / N 285 / N 315",
        "puesto": "-",
        "realizo": "F. Santoro",
        "aprobo": "C. Baptista",
        "fecha": "28/09/2026",
        "rev": "A"
    }

    for h in lista_hojas:
        d = dict(cajetin_base)
        d.update(h)
        hoja(prs, d, logo=LOGO_BARACK)

    prs.save(ruta_pptx)
    print(f"\n[OK] Presentación generada: {ruta_pptx} ({len(lista_hojas)} hojas)")
    return ruta_pptx

def compilar_todos():
    print("=" * 80)
    print("COMPILANDO HOJAS DE PROCESO — MÁQUINA HOTMELT (SISTEMA OFICIAL SGC)")
    print("=" * 80)

    # 1. Gate de redacción canónico sobre todas las 17 hojas
    if gate_redaccion:
        print("\nVerificando gate de redacción canónico sobre las 17 hojas...")
        for h in HOJAS_HOTMELT:
            gate_redaccion(h)
        print("[PASS] Gate de redacción: 17 hojas en REGLA (0 errores).\n")

    decks = [
        # (lista_hojas, subtitulo, nombre_base, titulo_indice)
        (HOJAS_PARTE_1,
         "PARTE 1: OPERACIÓN ESTÁNDAR Y PRODUCCIÓN · OP 20 DEL FLUJOGRAMA 155",
         "HOJAS DE PROCESO - HOTMELT - 1. OPERACION ESTANDAR",
         "ÍNDICE: OPERACIÓN ESTÁNDAR"),
        (HOJAS_PARTE_2,
         "PARTE 2: CONTINGENCIAS Y RESOLUCIÓN DE DESVÍOS · OP 20 DEL FLUJOGRAMA 155",
         "HOJAS DE PROCESO - HOTMELT - 2. CONTINGENCIAS",
         "ÍNDICE: CONTINGENCIAS"),
        (HOJAS_PARTE_3,
         "PARTE 3: MANTENIMIENTO OPERATIVO Y LIMPIEZA · OP 20 DEL FLUJOGRAMA 155",
         "HOJAS DE PROCESO - HOTMELT - 3. LIMPIEZA Y MANTENIMIENTO",
         "ÍNDICE: MANTENIMIENTO Y LIMPIEZA"),
        (HOJAS_HOTMELT,
         "MANUAL OPERATIVO COMPLETO (17 OPERACIONES) · OP 20 DEL FLUJOGRAMA 155",
         "HOJAS DE PROCESO - MAQUINA HOTMELT",
         "ÍNDICE GENERAL DE FABRICACIÓN"),
    ]

    generados_pdf = []

    for hojas_sub, sub_portada, nombre_base, tit_ind in decks:
        pptx_path = os.path.join(DESKTOP_DIR, nombre_base + ".pptx")
        generar_deck(hojas_sub, sub_portada, pptx_path, tit_ind)
        try:
            pdf_path = exportar_pdf(pptx_path)
            generados_pdf.append(pdf_path)
            # Replicar PDF al directorio raíz de Desktop
            desktop_root_pdf = os.path.join(DESKTOP_ROOT, os.path.basename(pdf_path))
            shutil.copy2(pdf_path, desktop_root_pdf)
            print(f"     -> Replicado a Escritorio raíz: {desktop_root_pdf}")
            # Este generador (formato A4 viejo) YA NO publica en el servidor: desde el 01/10/2026
            # las hojas vigentes son las de `generar_hojas_v3.py` (HO 992 y HO 993) y los decks
            # que salian de aca quedaron en `HO 992 - TOP ROLL\OBSOLETO\`. Con la copia activa,
            # una corrida vieja volvia a dejar hojas obsoletas al lado de las vigentes.
            if False and os.path.exists(SGC_DIR):
                sgc_dest_pptx = os.path.join(SGC_DIR, os.path.basename(pptx_path))
                sgc_dest_pdf = os.path.join(SGC_DIR, os.path.basename(pdf_path))
                try:
                    shutil.copy2(pptx_path, sgc_dest_pptx)
                    shutil.copy2(pdf_path, sgc_dest_pdf)
                    print(f"     -> Replicado a SGC Y: {sgc_dest_pdf}")
                except Exception as e_sgc:
                    print(f"     [AVISO SGC] No se pudo copiar a Y: ({e_sgc})")
        except Exception as e:
            print(f"     [AVISO COM] No se pudo exportar PDF automático para {nombre_base}: {e}")

    print("\n" + "=" * 80)
    print("PROCESO COMPLETADO EXITOSAMENTE")
    print("=" * 80)
    for p in generados_pdf:
        print(f"  • {p}")

if __name__ == "__main__":
    compilar_todos()
