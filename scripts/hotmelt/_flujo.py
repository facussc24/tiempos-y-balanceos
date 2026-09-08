# -*- coding: utf-8 -*-
"""Extrae SOLO la secuencia: numero, titulo y pasos de cada hoja, para leerla en fila."""
import glob, io, re
from pptx import Presentation
p = glob.glob("C:/Users/FacundoS-PC/BARACK ARGENTINA SRL/Ingenier*/INGENIERIA BARACK (NUNCA BORRAR)/1- GENERAL/INSTRUCTIVOS/INSTRUCCIONES OPERATIVAS/HOTMELT/HOJAS DE PROCESO*.pptx")
prs = Presentation(p[0])
out = []
for i, s in enumerate(prs.slides):
    if i == 0:
        continue
    op = tit = None
    pasos = None
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if re.fullmatch(r"20\.\d+", t):
            op = t
        elif sh.top / 360000.0 < 4.5 and t.isupper() and len(t) > 12 and "HOJA DE" not in t \
             and "OPERAC" not in t.split("\n")[0][:12]:
            if tit is None and t not in ("ADHESIVADO HOT MELT", "PATAGONIA"):
                tit = t
        if sh.left / 360000.0 > 16 and sh.top / 360000.0 > 5 and sh.top / 360000.0 < 15:
            pasos = t
    out.append("\n%s   %s" % (op, tit))
    if pasos:
        out.append(pasos)
io.open("_flujo.txt", "w", encoding="utf-8").write("\n".join(out))
print("hojas:", sum(1 for l in out if l.startswith("\n20.")))
