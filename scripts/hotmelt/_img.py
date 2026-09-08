import sys, openpyxl, warnings
warnings.filterwarnings("ignore")
from openpyxl.utils import get_column_letter
for p in sys.argv[1:]:
    wb=openpyxl.load_workbook(p); tag=p.split("/")[-1][:6]
    print("###",tag)
    for sn in wb.sheetnames[:4]:
        ws=wb[sn]
        print(" hoja",sn,"imgs:",len(ws._images))
        for im in ws._images:
            try:
                a=im.anchor._from
                print(f"   col{get_column_letter(a.col+1)} row{a.row+1}  w={im.width} h={im.height}")
            except Exception as e: print("   ?",e)
        print("  K8/CLIENTE:",ws['K8'].value)
