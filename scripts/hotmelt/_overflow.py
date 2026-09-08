# -*- coding: utf-8 -*-
"""Chequea que NINGUN texto del deck se salga de su caja, midiendo con la fuente real."""
import sys
from pptx import Presentation
import hoja_pptx as HP
pr = Presentation("_deck.pptx")
malos = 0
for i, s in enumerate(pr.slides):
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        tf = sh.text_frame
        if not tf.text.strip():
            continue
        w = sh.width / 360000.0
        h = sh.height / 360000.0
        alto = 0.0
        ok = True
        for p in tf.paragraphs:
            texto = "".join(r.text for r in p.runs)
            if not texto.strip():
                continue
            # el run 0 de un paso es el numero "1.  " y va en NEGRITA: medir todo
            # el parrafo en negrita infla el ancho y da un falso positivo.
            r0 = max(p.runs, key=lambda r: len(r.text))
            size = r0.font.size.pt if r0.font.size else 11
            bold = bool(r0.font.bold)
            ml = tf.margin_left / 360000.0 if tf.margin_left else 0
            util = w - 2 * ml - 0.24
            if util < 0.3:
                continue
            lin = HP._lineas_wrap(texto, size, util, bold)
            alto += lin * 1.22 * size * HP.PT_CM
            if p.space_after: alto += p.space_after.pt * HP.PT_CM
            if p.space_before: alto += p.space_before.pt * HP.PT_CM
        mt = (tf.margin_top / 360000.0 if tf.margin_top else 0)
        if alto > h - mt + 0.02:
            malos += 1
            print(f"lam{i+1:>2}  caja {w:.2f}x{h:.2f}  texto pide {alto:.2f} cm  ->  {tf.text.strip()[:60]!r}")
print("cajas con texto que no entra:", malos)
sys.exit(1 if malos else 0)
