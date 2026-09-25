# -*- coding: utf-8 -*-
import glob, os, io
from pptx import Presentation
p = glob.glob("Y:/BARACK/CALIDAD/DOCUMENTACION SGC/HOJAS DE OPERACIONES/1- CLIENTES/NOVAX/Tapizadas puerta/TOP ROLL/*.pptx")
f = p[0]
print("existe:", os.path.exists(f), os.path.getsize(f))
with open(f, "rb") as fh:
    data = io.BytesIO(fh.read())
pr = Presentation(data)
print("slides:", len(pr.slides))
for sh in pr.slides[17].shapes:
    if sh.has_text_frame and sh.text_frame.text.strip():
        t = sh.text_frame.text.strip()
        if "limpieza" in t.lower():
            print("---"); print(t[:900].encode("ascii","replace").decode())
