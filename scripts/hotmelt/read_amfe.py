import openpyxl
wb = openpyxl.load_workbook(r"C:\Users\FACUND~1\AppData\Local\Temp\claude\C--Dev-BarackMercosul\c66b0fd1-ca90-4cc1-86f1-5c8410cdd456\scratchpad\docs_fuente\AMFE de Proceso - TOP ROLL.xlsx", data_only=True)
print(wb.sheetnames)
ws = wb["AMFE"] if "AMFE" in wb.sheetnames else wb.active
# print header rows
for r in range(1, 6):
    vals = [ws.cell(row=r, column=c).value for c in range(1, 20)]
    print(r, vals)
print("----- rows 44-55 -----")
for r in range(44, 56):
    vals = [ws.cell(row=r, column=c).value for c in range(1, 25)]
    print(r, vals)
