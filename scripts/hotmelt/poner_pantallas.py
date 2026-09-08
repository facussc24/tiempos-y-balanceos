# -*- coding: utf-8 -*-
"""Cambia las FOTOS de pantalla por las pantallas PREPARADAS, en el PPTX que ya existe.

Regla dura: **no se toca una sola letra**. El deck lo edita Fak desde el 03/09 y sus cambios
mandan. Este script solo saca y pone IMAGENES, y al final compara el texto de las 18 laminas
contra el volcado previo: si cambio una coma, aborta.

Segunda regla, aprendida en la primera corrida: **el pptx de Fak es la verdad, no el spec.**
En la hoja 20.5 el spec decia 3 fotos y el archivo tenia 2 — el saco una. Por eso:
  · la foto a cambiar se identifica por el MD5 de su contenido, no por su posicion;
  · las fotos que quedan se sacan del propio pptx (blob embebido), no de la carpeta.
Asi el resultado no depende de que el spec siga sincronizado con lo que el tiene.

  poner_pantallas.py --pptx <ruta> [--dry-run]
"""
import argparse, hashlib, os, shutil, sys, tempfile
from datetime import datetime
from pptx import Presentation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import hoja_pptx as HP          # se reusa el mismo empaquetado que dibujo el deck

AQUI = os.path.dirname(os.path.abspath(__file__))
FOTOS = (r"C:\Users\FacundoS-PC\OneDrive - BARACK ARGENTINA SRL\Desktop"
         r"\Hojas de proceso maquina HOTMELT - desde los videos\_trabajo\fotos_hoja")

# hoja -> [(foto original que se saca, pantalla preparada que entra), ...]
CAMBIOS = {
    "20.2":  [("h02_b_panel_glsc2.jpg",      "_pantalla_fusor.png")],
    "20.4":  [("h04_a_pantalla_operacion.jpg", "_pantalla_receta.png"),
              ("h04_b_parametros_temp.jpg",  "_pantalla_rodillos.png")],
    "20.5":  [("h05_a_calentando.jpg",       "_pantalla_operacion_calentamiento_ok.png")],
    "20.9":  [("h09_a_hmi_arranque.jpg",     "_pantalla_operacion_automatico.png")],
    "20.15": [("h15_a_stop.jpg",             "_pantalla_operacion_detenido.png")],
    "20.17": [("h17_a_pantalla_limpieza.jpg", "_pantalla_limpieza.png")],
}


def md5(b):
    return hashlib.md5(b).hexdigest()


def es_contenido(sh):
    """Logo e iconos de EPP no son fotos de contenido: se reconocen por donde estan."""
    if sh.shape_type != 13:
        return False
    x, y = sh.left / 360000.0, sh.top / 360000.0
    if y < 4.5:                       # cajetin (logo)
        return False
    if x > 17.0 and sh.width / 360000.0 < 2.5:   # banda de EPP
        return False
    return True


def texto_de(prs):
    return ["\n".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
            for s in prs.slides]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pptx", required=True)
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    # huella de cada foto original que hay que cambiar
    huella = {}
    for op, cambios in CAMBIOS.items():
        for viejo, nuevo in cambios:
            p = os.path.join(FOTOS, viejo)
            if not os.path.exists(p):
                sys.exit("no encuentro la foto original " + viejo)
            n = os.path.join(AQUI, nuevo)
            if not os.path.exists(n):
                sys.exit("no encuentro la pantalla " + nuevo)
            huella[md5(open(p, "rb").read())] = (op, viejo, n)

    prs = Presentation(a.pptx)
    antes = texto_de(prs)
    tmp = tempfile.mkdtemp(prefix="pptxfotos_")
    plan, encontradas = [], set()

    for i, s in enumerate(prs.slides):
        op = None
        for sh in s.shapes:
            t = sh.text_frame.text.strip() if sh.has_text_frame else ""
            if t in CAMBIOS:
                op = t
                break
        if not op:
            continue
        fotos = sorted([sh for sh in s.shapes if es_contenido(sh)],
                       key=lambda q: (round(q.top / 360000.0, 1),
                                      round(q.left / 360000.0, 1)))
        rutas = []
        for k, sh in enumerate(fotos):
            h = md5(sh.image.blob)
            if h in huella and huella[h][0] == op:
                _, viejo, nueva = huella[h]
                rutas.append(("CAMBIA", viejo, nueva))
                encontradas.add(h)
            else:                       # queda: se guarda tal cual esta en el pptx
                d = os.path.join(tmp, "s%02d_%d.%s" % (i, k, sh.image.ext))
                open(d, "wb").write(sh.image.blob)
                rutas.append(("queda", "(la que tenías)", d))
        plan.append((i, s, op, fotos, rutas))

    print("--- PLAN ---")
    for i, s, op, fotos, rutas in plan:
        print("lamina %-2d  hoja %-6s  %d fotos" % (i + 1, op, len(fotos)))
        for est, viejo, nueva in rutas:
            print("     %-6s %-34s %s" % (est, viejo,
                                          os.path.basename(nueva) if est == "CAMBIA" else ""))
    sin_hallar = [v for k, v in huella.items() if k not in encontradas]
    if sin_hallar:
        print("\n🔴 NO ENCONTRE en el pptx estas fotos que habia que cambiar:")
        for op, viejo, _ in sin_hallar:
            print("   hoja %s: %s" % (op, viejo))
        sys.exit("pará y mirá: puede que Fak ya las haya cambiado.")
    if a.dry_run:
        print("\nDRY-RUN: no se toco nada.")
        return

    res = os.path.join(os.path.dirname(a.pptx),
                       "_resguardo %s.pptx" % datetime.now().strftime("%Y-%m-%d %H%M"))
    shutil.copy2(a.pptx, res)
    print("\nresguardo:", os.path.basename(res))

    for i, s, op, fotos, rutas in plan:
        for sh in fotos:
            sh._element.getparent().remove(sh._element)
        HP.bloque_imagenes_solo_fotos(s, [r for _, _, r in rutas])
        print("lamina %-2d (%s): %d imagenes" % (i + 1, op, len(rutas)))

    prs.save(a.pptx)

    despues = texto_de(Presentation(a.pptx))
    if antes != despues:
        for k, (x, y) in enumerate(zip(antes, despues)):
            if x != y:
                print("🔴 CAMBIO DE TEXTO en la lamina", k + 1)
        sys.exit("ABORTAR: cambio texto. Restaurá el resguardo.")
    print("texto de las %d laminas: IDENTICO al de antes." % len(antes))


main()
