# -*- coding: utf-8 -*-
import io
from pptx import Presentation
pr = Presentation("_deck.pptx")
out = []
for i, s in enumerate(pr.slides):
    out.append(f"\n{'='*70}\nLAMINA {i+1}\n{'='*70}")
    for sh in sorted(s.shapes, key=lambda x: (round(x.top/360000,1), round(x.left/360000,1))):
        if sh.has_text_frame and sh.text_frame.text.strip():
            out.append(sh.text_frame.text.strip())
        elif sh.shape_type == 13:
            out.append(f"[IMAGEN {sh.width/360000:.1f}x{sh.height/360000:.1f} cm en x={sh.left/360000:.1f} y={sh.top/360000:.1f}]")
io.open("_deck.txt","w",encoding="utf-8").write("\n".join(out))
print("lineas:", len(out))
