import sys, openpyxl, warnings, collections, re
from openpyxl.utils import column_index_from_string
warnings.filterwarnings("ignore")
resp=collections.Counter(); frec=collections.Counter(); reg=collections.Counter(); meth=collections.Counter()
extra=collections.Counter()
for p in sys.argv[1:]:
    wb = openpyxl.load_workbook(p, data_only=True)
    tag=p.split("/")[-1][:6]
    for sn in wb.sheetnames:
        ws=wb[sn]
        hdr=None
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value,str) and c.value.strip().startswith("Características a controlar"):
                    hdr=c
                if isinstance(c.value,str):
                    t=c.value.strip()
                    if re.search(r'Segregar|P-09|BORRADOR|pendiente|a definir|RC\b', t, re.I):
                        extra[(tag,t[:90])]+=1
        if not hdr: continue
        r=hdr.row; ci=hdr.column
        # locate columns of headers in same row
        cols={}
        for c in ws[r]:
            if isinstance(c.value,str) and c.value.strip() in ("Método de control","Resp.","Frec.","Registro"):
                cols[c.value.strip()]=c.column
        for rr in range(r+1, r+20):
            a=ws.cell(rr,ci).value
            if a in (None,""): continue
            if isinstance(a,str) and a.strip().startswith(("PLAN DE","SI DETECTA","DETENGA","NOTIFIQUE","ESPERE")): continue
            if "Resp." in cols: resp[(tag,str(ws.cell(rr,cols["Resp."]).value))]+=1
            if "Frec." in cols: frec[(tag,str(ws.cell(rr,cols["Frec."]).value))]+=1
            if "Registro" in cols: reg[(tag,str(ws.cell(rr,cols["Registro"]).value))]+=1
            if "Método de control" in cols: meth[(tag,str(ws.cell(rr,cols["Método de control"]).value)[:60])]+=1
def show(name,ctr):
    print("="*60); print(name)
    for k,v in sorted(ctr.items(), key=lambda x:(-x[1],x[0])): print(f"  [{v:3}] {k}")
show("RESP.",resp); show("FREC.",frec); show("REGISTRO",reg)
show("METODO (top)",meth)
show("Busqueda Segregar/P-09/BORRADOR/RC",extra)
