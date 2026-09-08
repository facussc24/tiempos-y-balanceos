# -*- coding: utf-8 -*-
"""Mete una pantalla redibujada en UNA lamina del PPTX que ya existe, sin regenerarlo.

Existe porque desde el 03/09 el deck lo edita Fak a mano: correr `armar_deck.py` encima
le borraria los cambios. Esto abre el archivo que hay, toca UNA lamina y guarda.

  insertar_pantalla.py --pptx <ruta> --hoja 20.4 --imagen _pantalla_rodillos.png
                       [--reemplazar-fotos] [--dry-run]

Por defecto NO reemplaza nada: agrega la imagen ocupando el bloque IMAGENES y deja las
fotos que ya estaban debajo. Con --reemplazar-fotos saca las fotos de contenido de esa
lamina (nunca el logo ni los iconos de EPP, que se reconocen por tamano y posicion).

SIEMPRE hace copia de resguardo antes de escribir.
"""
import argparse, os, shutil, sys
from datetime import datetime
from pptx import Presentation
from pptx.util import Cm

# geometria del bloque IMAGENES, la misma que usa hoja_pptx.py
IMG_X, IMG_Y = 0.70, 5.50
IMG_W, IMG_H = 16.20, 9.30


def es_foto_de_contenido(sh):
    """Logo (arriba) e iconos de EPP (derecha abajo, chicos) NO son fotos de contenido."""
    if sh.shape_type != 13:
        return False
    x, y = sh.left / 360000.0, sh.top / 360000.0
    w, h = sh.width / 360000.0, sh.height / 360000.0
    if y < 4.5:                      # cajetin: ahi esta el logo
        return False
    if x > 17.0 and w < 2.0:         # banda de EPP
        return False
    return True


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--pptx", required=True)
    p.add_argument("--hoja", required=True, help='numero de operacion, ej "20.4"')
    p.add_argument("--imagen", required=True)
    p.add_argument("--reemplazar-fotos", action="store_true")
    p.add_argument("--dry-run", action="store_true")
    a = p.parse_args()

    if not os.path.exists(a.pptx):
        sys.exit("no existe el pptx: " + a.pptx)
    if not os.path.exists(a.imagen):
        sys.exit("no existe la imagen: " + a.imagen)

    prs = Presentation(a.pptx)
    objetivo = None
    for i, s in enumerate(prs.slides):
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip() == a.hoja:
                objetivo = (i, s)
                break
        if objetivo:
            break
    if not objetivo:
        sys.exit("no encontre la lamina de la hoja " + a.hoja)
    i, slide = objetivo

    fotos = [sh for sh in slide.shapes if es_foto_de_contenido(sh)]
    print("lamina %d  (hoja %s)" % (i + 1, a.hoja))
    print("  fotos de contenido que tiene ahora: %d" % len(fotos))
    for sh in fotos:
        print("     %.1f x %.1f cm  en x=%.1f y=%.1f"
              % (sh.width / 360000.0, sh.height / 360000.0,
                 sh.left / 360000.0, sh.top / 360000.0))
    print("  accion: %s" % ("REEMPLAZAR esas fotos por la pantalla"
                            if a.reemplazar_fotos else "AGREGAR la pantalla encima"))

    if a.dry_run:
        print("DRY-RUN: no se toco nada.")
        return

    # resguardo antes de escribir, siempre
    res = os.path.join(os.path.dirname(a.pptx) or ".",
                       "_resguardo %s %s" % (datetime.now().strftime("%Y-%m-%d %H%M"),
                                             os.path.basename(a.pptx)))
    shutil.copy2(a.pptx, res)
    print("  resguardo: " + os.path.basename(res))

    if a.reemplazar_fotos:
        for sh in fotos:
            sh._element.getparent().remove(sh._element)

    from PIL import Image
    im = Image.open(a.imagen)
    ar = im.width / im.height
    iw, ih = (IMG_W, IMG_W / ar) if IMG_W / ar <= IMG_H else (IMG_H * ar, IMG_H)
    x = IMG_X + (IMG_W - iw) / 2
    y = IMG_Y + (IMG_H - ih) / 2
    slide.shapes.add_picture(a.imagen, Cm(x), Cm(y), Cm(iw), Cm(ih))
    prs.save(a.pptx)
    print("  puesta: %.1f x %.1f cm.  guardado." % (iw, ih))


main()
