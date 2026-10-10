"""
firmaIA.py — ningun documento de Barack dice que lo hizo Claude o una IA.

Origen (Fak, 08/10/2026): el listado de hojas de proceso tenia "Claude" en la columna CREADO POR
de 14 filas (HO 972 a 984 y HO 118, cargadas en junio de 2026) y una pestaña oculta
"_CONTEXTO_CLAUDE" escrita "para el proximo Claude". *"Es un error gravisimo, no puede volver a
suceder nunca algo asi... en ningun tipo de documento"*.

Que revisa de un archivo, ademas de lo que se ve:
  - Excel / Word / PowerPoint (.xlsx .xlsm .docx .pptx y plantillas): TODAS las partes del zip —
    celdas, pestañas ocultas y sus nombres, nombres definidos, comentarios, notas del orador,
    cuadros de texto, encabezados y pies, y las propiedades del archivo (autor, ultimo en guardar,
    descripcion). El texto partido en varios "runs" se junta antes de buscar.
  - PDF: texto de cada pagina, anotaciones, marcadores, metadata y XMP.
  - Mail (.msg / .eml): asunto, cuerpo, nombres y contenido de los adjuntos.
  - Formatos viejos (.xls .doc .ppt) y binarios: los bytes, en latin-1 y en UTF-16.
  - Texto (.txt .csv .md .html .json .xml .dxf .plt .svg .rtf): el texto.
  - El NOMBRE del archivo y de sus carpetas (debajo de la raiz que se revisa).

Las palabras viven en `firmaIA.data.json` (una sola fuente: este modulo y el guardian de los hooks
la leen). Salida: lista de hallazgos con el lugar exacto (pestaña y celda, diapositiva, pagina).

Uso como libreria:
    from firmaIA import revisar_archivo, revisar_texto, exigir_sin_firma
    exigir_sin_firma([ruta1, ruta2])   # sale con 1 si alguno nombra a Claude o a una IA
CLI: `python scripts/_sinFirmaIA.py <archivos o carpetas> [--desde AAAA-MM-DD] [--json]`.
"""
from __future__ import annotations

import html
import io
import json
import os
import re
import stat
import sys
import tempfile
import zipfile
from dataclasses import dataclass, field, asdict

AQUI = os.path.dirname(os.path.abspath(__file__))
CANON_PATH = os.path.join(AQUI, 'firmaIA.data.json')

OOXML = {'.xlsx', '.xlsm', '.xltx', '.xltm', '.docx', '.docm', '.dotx', '.dotm',
         '.pptx', '.pptm', '.potx', '.potm', '.ppsx', '.vsdx'}
TEXTO = {'.txt', '.csv', '.md', '.html', '.htm', '.json', '.xml', '.dxf', '.plt', '.hpgl',
         '.svg', '.rtf', '.eml', '.tsv', '.ini', '.log'}
BINARIO_LEGADO = {'.xls', '.doc', '.ppt', '.xlsb', '.mpp', '.vsd', '.pub', '.dwg'}
IMAGEN = {'.png', '.jpg', '.jpeg', '.tif', '.tiff', '.bmp', '.gif', '.webp', '.emf', '.wmf'}
DOCUMENTO = OOXML | TEXTO | BINARIO_LEGADO | IMAGEN | {'.pdf', '.msg'}

# Atributos de OneDrive "solo en la nube": leerlo lo baja al disco.
_ATRIB_EN_LA_NUBE = 0x00400000 | 0x00040000 | 0x00001000  # RECALL_ON_DATA_ACCESS | RECALL_ON_OPEN | OFFLINE


@dataclass
class Hallazgo:
    archivo: str
    lugar: str          # 'pestaña INDICE, celda J74' · 'propiedades del archivo (autor)' · 'pagina 3'
    regla: str          # id del canon
    que: str            # explicacion corta
    texto: str          # el texto encontrado, recortado
    nivel: str = 'BLOQUEANTE'   # BLOQUEANTE | AVISO

    def renglon(self) -> str:
        return f'  [{self.nivel}] {self.archivo} — {self.lugar}: {self.que} → «{self.texto}»'


@dataclass
class _Canon:
    bloqueante: list = field(default_factory=list)   # (id, regex, que, solo_unido)
    aviso: list = field(default_factory=list)
    metadata_programa: set = field(default_factory=set)
    excluidas: list = field(default_factory=list)
    no_en_binario: set = field(default_factory=set)   # siglas cortas: en bytes comprimidos caen al azar


_CANON: _Canon | None = None


def canon() -> _Canon:
    global _CANON
    if _CANON is None:
        with open(CANON_PATH, encoding='utf-8') as f:
            d = json.load(f)
        c = _Canon()
        for nivel, destino in (('bloqueante', c.bloqueante), ('aviso', c.aviso)):
            for p in d.get(nivel, []):
                flags = 0 if p.get('distingue_mayusculas') else re.IGNORECASE
                destino.append((p['id'], re.compile(p['re'], flags), p['que'], bool(p.get('solo_unido'))))
                if p.get('no_en_binario'):
                    c.no_en_binario.add(p['id'])
        c.metadata_programa = {v.lower() for v in d['metadata_programa']['valores']}
        c.excluidas = [s.lower() for s in d['rutas_excluidas']['contiene']]
        _CANON = c
    return _CANON


def ruta_excluida(ruta: str) -> bool:
    r = os.path.abspath(ruta).lower() + os.sep
    return any(x in r for x in canon().excluidas)


def _recorte(t: str, m: re.Match, ancho: int = 45) -> str:
    a, b = max(0, m.start() - ancho), min(len(t), m.end() + ancho)
    s = re.sub(r'\s+', ' ', t[a:b]).strip()
    return ('…' if a else '') + s + ('…' if b < len(t) else '')


def revisar_texto(texto: str, archivo: str = '', lugar: str = 'texto', unido: str | None = None,
                  con_avisos: bool = True) -> list[Hallazgo]:
    """Busca las palabras del canon en un texto. `unido` = el mismo texto sin separadores entre
    pedazos (runs de Word/PowerPoint), para encontrar 'Cla'+'ude'."""
    out: list[Hallazgo] = []
    if not texto and not unido:
        return out
    for rid, rx, que, solo_unido in canon().bloqueante:
        fuente = unido if solo_unido else texto
        m = rx.search(fuente) if fuente else None
        if m:   # un hallazgo por regla y lugar alcanza para frenar
            out.append(Hallazgo(archivo, lugar, rid, que, _recorte(fuente, m)))
    if con_avisos:
        for rid, rx, que, _ in canon().aviso:
            m = rx.search(texto or '')
            if m:
                out.append(Hallazgo(archivo, lugar, rid, que, _recorte(texto, m), 'AVISO'))
    # 'claude-pegado' repite a 'claude' cuando el texto no estaba partido
    if any(h.regla == 'claude' for h in out):
        out = [h for h in out if h.regla != 'claude-pegado']
    return out


# ---------------------------------------------------------------- el logo de Barack (cola H9, 09/10/2026)
#
# Fak, 09/10/2026: "usaste un logo no oficial de barack, gravisimo". El oficial es UN archivo
# (VARIOS\Logo y color barack\barack_logo.png; copia en tools/flowchart/assets/). Lo que se mide de una imagen
# no es su hash (un programa que la vuelve a guardar lo cambia: 124 copias del oficial en exports/ dan distinto
# hash y la misma imagen) sino una HUELLA: sobre blanco, recortada al dibujo, achicada a 24x12 en color. Distancia
# = diferencia media por canal (0-255). Lo medido el 09/10 esta en el canon (`logo`), con los umbrales.

_RAIZ_REPO = os.path.dirname(os.path.dirname(AQUI))
_LOGOS: dict | None = None


def huella_imagen(crudo: bytes, max_lado: int = 5000, proporcion: tuple | None = None):
    """(huella_hex, (ancho, alto)) de una imagen, o None si no se puede leer, pasa `max_lado` o no tiene la `proporcion`
    (ancho/alto) pedida: las dos cosas se miran en la cabecera, antes de decodificar (una foto 4:3 no se abre)."""
    try:
        from PIL import Image
    except ImportError:
        return None
    try:
        im = Image.open(io.BytesIO(crudo))
        w, h = im.size                       # solo la cabecera: una foto grande se descarta sin decodificarla
        if max(w, h) > max_lado or not h or (proporcion and not (proporcion[0] <= w / h <= proporcion[1])):
            return None
        im = im.convert('RGBA')
        fondo = Image.new('RGBA', im.size, (255, 255, 255, 255))
        im = Image.alpha_composite(fondo, im).convert('RGB')
        caja = im.convert('L').point(lambda v: 255 if v < 245 else 0).getbbox()
        if caja:
            im = im.crop(caja)
        px = im.resize((24, 12), Image.LANCZOS).getdata()
        return ''.join(f'{c:02x}' for p in px for c in p), (w, h)
    except Exception:
        return None


def distancia_huellas(a: str, b: str) -> float:
    va, vb = bytes.fromhex(a), bytes.fromhex(b)
    return sum(abs(x - y) for x, y in zip(va, vb)) / max(1, len(va))


def _logos() -> dict:
    """Las referencias del canon, con la huella del oficial calculada de la copia del repo (y su hash verificado)."""
    global _LOGOS
    if _LOGOS is None:
        with open(CANON_PATH, encoding='utf-8') as f:
            L = json.load(f).get('logo') or {}
        ofi = L.get('oficial') or {}
        huella_ofi = None
        try:
            crudo = open(os.path.join(_RAIZ_REPO, ofi.get('copia_repo', '')), 'rb').read()
            import hashlib
            if hashlib.sha256(crudo).hexdigest() == ofi.get('sha256'):
                h = huella_imagen(crudo)
                huella_ofi = h[0] if h else None
        except OSError:
            pass
        _LOGOS = {'oficial': ofi, 'huella_oficial': huella_ofi, 'no_oficiales': L.get('no_oficiales') or [],
                  'umbral': L.get('umbral') or {}, 'proporcion': L.get('proporcion') or [1.5, 6.5]}
    return _LOGOS


def revisar_imagen(crudo: bytes, archivo: str, lugar: str, logo_bloquea: bool = False) -> list[Hallazgo]:
    """¿La imagen es el logo de Barack NO oficial, o se parece al oficial sin serlo? Por defecto es AVISO: se imprime y
    se manda igual, con el aviso (exigir_sin_firma lo muestra como OJO). Con `logo_bloquea` el no oficial FRENA.
    Por que no frena siempre (medido el 09/10/2026): el de letras finas esta ADENTRO del formulario de hoja de proceso de
    la casa (HO 21-9463 a 9475, termoformado, las de embalaje en PowerPoint); frenar ahi bloquearia corregir una HO de
    Calidad en su propio formulario. `--logo-bloquea` queda listo para prenderlo en el cierre del turno cuando Fak diga si
    el formulario tambien pasa al oficial. Solo mira imagenes con forma de logo (mas anchas que altas)."""
    L = _logos()
    if not crudo or len(crudo) > 8 * 1024 * 1024:
        return []
    import hashlib
    sha = hashlib.sha256(crudo).hexdigest()
    for n in L['no_oficiales']:
        if sha == n.get('sha256'):
            return [Hallazgo(archivo, lugar, 'logo-no-oficial', f'logo de Barack NO oficial ({n.get("como", "otro dibujo")}): va el oficial, {L["oficial"].get("ruta", "")}', 'mismo archivo', 'BLOQUEANTE' if logo_bloquea else 'AVISO')]
    if not L['huella_oficial']:
        _avisar_logo_apagado()
        return []
    h = huella_imagen(crudo, proporcion=tuple(L['proporcion']))
    if not h:
        return []
    huella, (w, hgt) = h
    ruta_ofi = L['oficial'].get('ruta', 'VARIOS\\Logo y color barack\\barack_logo.png')
    for n in L['no_oficiales']:
        d = distancia_huellas(huella, n.get('huella', '')) if n.get('huella') else 999
        if d <= float(L['umbral'].get('no_oficial', 5)):
            return [Hallazgo(archivo, lugar, 'logo-no-oficial',
                             f'logo de Barack NO oficial ({n.get("como", "otro dibujo")}): va el oficial, {ruta_ofi}',
                             f'imagen {w}x{hgt}, distancia {d:.1f}', 'BLOQUEANTE' if logo_bloquea else 'AVISO')]
    if L['huella_oficial']:
        d = distancia_huellas(huella, L['huella_oficial'])
        if float(L['umbral'].get('oficial', 10)) < d <= float(L['umbral'].get('parecido', 35)):
            return [Hallazgo(archivo, lugar, 'logo-parecido',
                             f'parece el logo de Barack pero no es el archivo oficial (recortado, re-guardado u otro color): va {ruta_ofi}',
                             f'imagen {w}x{hgt}, distancia {d:.1f}', 'AVISO')]
    return []


_EXT_IMAGEN_ZIP = ('.png', '.jpg', '.jpeg', '.gif', '.bmp', '.tif', '.tiff', '.webp')
_AVISADO = {'logo': False}


def _avisar_logo_apagado():
    """El chequeo del logo no puede correr (sin PIL, o la copia del oficial no esta o cambio): se dice UNA vez, no se calla."""
    if not _AVISADO['logo']:
        _AVISADO['logo'] = True
        print('  OJO: el chequeo del logo de Barack no corre en esta PC (falta Pillow o la copia oficial del repo no coincide): solo se compara por hash', file=sys.stderr)


# ---------------------------------------------------------------- OOXML (Excel, Word, PowerPoint)

def _xml_a_texto(xml: str) -> tuple[str, str]:
    """(texto con espacios entre etiquetas, texto pegado sin separadores)."""
    con_espacios = html.unescape(re.sub(r'<[^>]+>', ' ', xml))
    pegado = html.unescape(re.sub(r'<[^>]+>', '', xml))
    return con_espacios, pegado


def _mapa_pestanas(z: zipfile.ZipFile) -> dict[str, tuple[str, str]]:
    """xl/worksheets/sheetN.xml -> (nombre de la pestaña, estado visible/hidden/veryHidden)."""
    try:
        wb = z.read('xl/workbook.xml').decode('utf-8', 'replace')
        rels = z.read('xl/_rels/workbook.xml.rels').decode('utf-8', 'replace')
    except KeyError:
        return {}
    rid2target = {m.group(1): m.group(2) for m in re.finditer(r'<Relationship[^>]*Id="([^"]+)"[^>]*Target="([^"]+)"', rels)}
    rid2target.update({m.group(2): m.group(1) for m in re.finditer(r'<Relationship[^>]*Target="([^"]+)"[^>]*Id="([^"]+)"', rels)})
    out = {}
    for m in re.finditer(r'<sheet\b([^>]*)/?>', wb):
        attrs = m.group(1)
        nombre = re.search(r'name="([^"]*)"', attrs)
        rid = re.search(r'r:id="([^"]*)"', attrs)
        estado = re.search(r'state="([^"]*)"', attrs)
        if nombre and rid and rid.group(1) in rid2target:
            t = rid2target[rid.group(1)].lstrip('/')
            t = t if t.startswith('xl/') else 'xl/' + t
            out[t] = (html.unescape(nombre.group(1)), estado.group(1) if estado else 'visible')
    return out


def _lugar_ooxml(parte: str, pestanas: dict) -> str:
    p = parte.lower()
    if parte in pestanas:
        nombre, estado = pestanas[parte]
        oculta = '' if estado == 'visible' else ' (OCULTA)'
        return f'pestaña «{nombre}»{oculta}'
    if p == 'docprops/core.xml':
        return 'propiedades del archivo (autor, ultimo en guardar, descripcion)'
    if p == 'docprops/app.xml':
        return 'propiedades del archivo (aplicacion, empresa, nombres de pestañas)'
    if p == 'docprops/custom.xml':
        return 'propiedades personalizadas del archivo'
    if p == 'xl/workbook.xml':
        return 'libro (nombres de pestañas o nombres definidos)'
    if p == 'xl/sharedstrings.xml':
        return 'texto de las celdas'
    m = re.match(r'ppt/slides/slide(\d+)\.xml', p)
    if m:
        return f'diapositiva {m.group(1)}'
    m = re.match(r'ppt/notesslides/notesslide(\d+)\.xml', p)
    if m:
        return f'notas del orador ({m.group(1)})'
    if 'comment' in p:
        return f'comentarios ({parte})'
    if p.startswith('word/header') or p.startswith('word/footer'):
        return f'encabezado o pie ({parte})'
    if p == 'word/document.xml':
        return 'cuerpo del documento'
    if 'drawing' in p:
        return f'cuadro de texto o forma ({parte})'
    return parte


def _celdas_excel(ruta: str, hallazgos_texto: bool) -> list[tuple[str, str]]:
    """Ubica las celdas exactas (pestaña, celda, valor) que nombran a una IA."""
    if not hallazgos_texto:
        return []
    try:
        import openpyxl
        wb = openpyxl.load_workbook(ruta, data_only=False)
    except Exception:
        return []
    out = []
    for ws in wb.worksheets:
        estado = '' if ws.sheet_state == 'visible' else ' (OCULTA)'
        for fila in ws.iter_rows():
            for c in fila:
                if isinstance(c.value, str) and revisar_texto(c.value, con_avisos=False):
                    out.append((f'pestaña «{ws.title}»{estado}, celda {c.coordinate}', c.value))
            # comentarios de celda
        for fila in ws.iter_rows():
            for c in fila:
                if c.comment is not None and revisar_texto(c.comment.text or '', con_avisos=False):
                    out.append((f'pestaña «{ws.title}»{estado}, comentario en {c.coordinate}', c.comment.text))
    return out


def _metadata_core(xml: str, archivo: str) -> list[Hallazgo]:
    out = []
    for campo, etiqueta in (('dc:creator', 'autor'), ('cp:lastModifiedBy', 'ultimo en guardar')):
        m = re.search(rf'<{campo}\b[^>]*>(.*?)</{campo}>', xml, re.S)   # openpyxl escribe <dc:creator xmlns:dc=...>
        if m and html.unescape(m.group(1)).strip().lower() in canon().metadata_programa:
            out.append(Hallazgo(archivo, f'propiedades del archivo ({etiqueta})', 'metadata-programa',
                                'el archivo dice que lo armo un programa', html.unescape(m.group(1)).strip(), 'AVISO'))
    return out


def _revisar_ooxml(ruta: str, nombre: str, logo_bloquea: bool = False, con_avisos: bool = True) -> list[Hallazgo]:
    out: list[Hallazgo] = []
    with zipfile.ZipFile(ruta) as z:
        pestanas = _mapa_pestanas(z)
        partes_con_texto = False
        for parte in z.namelist():
            pl = parte.lower()
            if not (pl.endswith('.xml') or pl.endswith('.rels') or pl.endswith('.txt') or pl.endswith('.vml')):
                if '/media/' in pl and pl.endswith(_EXT_IMAGEN_ZIP):
                    if con_avisos or logo_bloquea:
                        out += revisar_imagen(z.read(parte), nombre, f'imagen ({parte})', logo_bloquea)
                    continue
                if pl.endswith('.bin') and 'vbaproject' in pl:
                    crudo = z.read(parte)
                    out += _revisar_bytes(crudo, nombre, f'macros ({parte})')
                continue
            xml = z.read(parte).decode('utf-8', 'replace')
            lugar = _lugar_ooxml(parte, pestanas)
            # 1) atributos y etiquetas (nombres de pestaña, nombres definidos, autor)
            hs = revisar_texto(html.unescape(xml), nombre, lugar, con_avisos=False)
            # 2) texto con los "runs" pegados
            esp, peg = _xml_a_texto(xml)
            hs += revisar_texto(esp, nombre, lugar, unido=peg, con_avisos=pl in ('xl/sharedstrings.xml', 'word/document.xml') or pl.startswith(('ppt/slides/', 'xl/worksheets/sheet')))
            vistos = set()
            for h in hs:
                k = (h.lugar, h.regla)
                if k not in vistos:
                    vistos.add(k)
                    out.append(h)
            # solo un BLOQUEANTE justifica cargar el libro entero con openpyxl (unos 4,5 s por HO de 17 MB): un aviso no
            if any(h.nivel == 'BLOQUEANTE' for h in hs) and (pl == 'xl/sharedstrings.xml' or pl.startswith('xl/worksheets/') or 'comment' in pl):
                partes_con_texto = True
            if pl == 'docprops/core.xml':
                out += _metadata_core(xml, nombre)
        # pestaña oculta: su nombre ya se revisa; ademas se dice que existe si nombra a una IA
    if partes_con_texto and os.path.splitext(ruta)[1].lower() in ('.xlsx', '.xlsm', '.xltx', '.xltm'):
        celdas = _celdas_excel(ruta, True)
        if celdas:
            # la celda exacta reemplaza al hallazgo generico de la misma pestaña
            hojas_con_celda = {l.split(', ')[0] for l, _ in celdas}
            out = [h for h in out if h.lugar != 'texto de las celdas' and h.lugar not in hojas_con_celda]
            for lugar, valor in celdas:
                out += revisar_texto(valor, nombre, lugar, con_avisos=False)[:1]
    return out


# ---------------------------------------------------------------- PDF

def _revisar_pdf(ruta: str, nombre: str, logo_bloquea: bool = False, con_avisos: bool = True) -> list[Hallazgo]:
    out: list[Hallazgo] = []
    try:
        import fitz  # PyMuPDF
    except ImportError:
        return _revisar_bytes(open(ruta, 'rb').read(), nombre, 'PDF (sin PyMuPDF: bytes)')
    try:
        doc = fitz.open(ruta)
    except Exception as e:
        return [Hallazgo(nombre, 'PDF', 'ilegible', f'no se pudo abrir ({e})', '', 'AVISO')]
    with doc:
        meta = doc.metadata or {}
        for k, v in meta.items():
            if isinstance(v, str) and v:
                out += revisar_texto(v, nombre, f'propiedades del PDF ({k})', con_avisos=False)
                if k in ('author', 'creator') and v.strip().lower() in canon().metadata_programa:
                    out.append(Hallazgo(nombre, f'propiedades del PDF ({k})', 'metadata-programa',
                                        'el archivo dice que lo armo un programa', v, 'AVISO'))
        try:
            xmp = doc.get_xml_metadata() or ''
            out += revisar_texto(xmp, nombre, 'propiedades del PDF (XMP)', con_avisos=False)
        except Exception:
            pass
        try:
            for nivel, titulo, pag in doc.get_toc():
                out += revisar_texto(titulo, nombre, f'marcador (pagina {pag})', con_avisos=False)
        except Exception:
            pass
        vistas = set()
        vistas_bytes = set()                 # el mismo logo guardado como objeto aparte en cada pagina se mira una vez
        mirar_logo = con_avisos or logo_bloquea
        lo, hi = _logos()['proporcion']
        for i, pag in enumerate(doc, start=1):
            t = pag.get_text() or ''
            out += revisar_texto(t, nombre, f'pagina {i}')
            # las imagenes con forma de logo (la proporcion se mira antes de extraerlas): cola H9
            for img in (pag.get_images(full=True) if mirar_logo else []):
                xref, ancho, alto = img[0], img[2], img[3]
                if xref in vistas or not alto or not (lo <= ancho / alto <= hi):
                    continue
                vistas.add(xref)
                try:
                    crudo = (doc.extract_image(xref) or {}).get('image')
                except Exception:
                    crudo = None
                if crudo:
                    import hashlib
                    hb = hashlib.sha1(crudo).hexdigest()
                    if hb in vistas_bytes:
                        continue
                    vistas_bytes.add(hb)
                    out += revisar_imagen(crudo, nombre, f'imagen en pagina {i}', logo_bloquea)
            for a in pag.annots() or []:
                info = a.info or {}
                for k in ('content', 'title', 'subject'):
                    if info.get(k):
                        out += revisar_texto(info[k], nombre, f'anotacion en pagina {i}', con_avisos=False)
            for w in pag.widgets() or []:
                if w.field_value and isinstance(w.field_value, str):
                    out += revisar_texto(w.field_value, nombre, f'campo de formulario en pagina {i}', con_avisos=False)
    return out


# ---------------------------------------------------------------- mail, binarios, texto

def _revisar_bytes(crudo: bytes, nombre: str, lugar: str) -> list[Hallazgo]:
    out: list[Hallazgo] = []
    vistos = set()
    for enc, txt in (('latin-1', crudo.decode('latin-1', 'replace')),
                     ('utf-16', crudo.decode('utf-16-le', 'replace')),
                     ('utf-16b', crudo[1:].decode('utf-16-le', 'replace'))):
        for h in revisar_texto(txt, nombre, lugar, con_avisos=False):
            if h.regla in canon().no_en_binario:
                continue
            if h.regla not in vistos:
                vistos.add(h.regla)
                out.append(h)
    return out


def _revisar_msg(ruta: str, nombre: str, logo_bloquea: bool = False, con_avisos: bool = True) -> list[Hallazgo]:
    out: list[Hallazgo] = []
    try:
        import extract_msg
        m = extract_msg.Message(ruta)
    except Exception:
        return _revisar_bytes(open(ruta, 'rb').read(), nombre, 'mail (bytes)')
    try:
        out += revisar_texto(m.subject or '', nombre, 'asunto del mail', con_avisos=False)
        out += revisar_texto(m.body or '', nombre, 'cuerpo del mail')
        for a in m.attachments:
            an = getattr(a, 'longFilename', None) or getattr(a, 'shortFilename', None) or 'adjunto'
            out += revisar_texto(an, nombre, 'nombre de un adjunto', con_avisos=False)
            datos = getattr(a, 'data', None)
            if isinstance(datos, (bytes, bytearray)) and os.path.splitext(an)[1].lower() in DOCUMENTO:
                with tempfile.TemporaryDirectory() as td:
                    p = os.path.join(td, os.path.basename(an))
                    with open(p, 'wb') as f:
                        f.write(datos)
                    for h in revisar_archivo(p, raiz=td, logo_bloquea=logo_bloquea, con_avisos=con_avisos):
                        h.archivo = f'{nombre} → adjunto {an}'
                        out.append(h)
    finally:
        try:
            m.close()
        except Exception:
            pass
    return out


def _revisar_texto_plano(ruta: str, nombre: str) -> list[Hallazgo]:
    crudo = open(ruta, 'rb').read()
    for enc in ('utf-8', 'utf-16', 'latin-1'):
        try:
            t = crudo.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        t = crudo.decode('latin-1', 'replace')
    return revisar_texto(t, nombre, 'contenido')


def en_la_nube(ruta: str) -> bool:
    try:
        st = os.stat(ruta)
        return bool(getattr(st, 'st_file_attributes', 0) & _ATRIB_EN_LA_NUBE)
    except OSError:
        return False


def revisar_archivo(ruta: str, raiz: str | None = None, max_mb: float = 60, logo_bloquea: bool = False, con_avisos: bool = True) -> list[Hallazgo]:
    """Hallazgos de UN archivo (contenido + nombre debajo de `raiz`)."""
    nombre = ruta
    ext = os.path.splitext(ruta)[1].lower()
    out: list[Hallazgo] = []
    # nombre del archivo y de las carpetas debajo de la raiz
    rel = os.path.relpath(ruta, raiz) if raiz else os.path.basename(ruta)
    out += revisar_texto(rel.replace('_', ' ').replace('-', ' '), nombre, 'nombre del archivo o de su carpeta', con_avisos=False)
    try:
        tam = os.path.getsize(ruta)
    except OSError as e:
        return out + [Hallazgo(nombre, 'archivo', 'ilegible', f'no se pudo leer ({e})', '', 'AVISO')]
    if tam > max_mb * 1024 * 1024:
        return out + [Hallazgo(nombre, 'archivo', 'grande', f'no se revisó: pesa {tam/1048576:.0f} MB', '', 'AVISO')]
    try:
        if ext in OOXML:
            try:
                out += _revisar_ooxml(ruta, nombre, logo_bloquea, con_avisos)
            except zipfile.BadZipFile:
                out += _revisar_bytes(open(ruta, 'rb').read(), nombre, 'contenido (no es un zip valido)')
        elif ext == '.pdf':
            out += _revisar_pdf(ruta, nombre, logo_bloquea, con_avisos)
        elif ext == '.msg':
            out += _revisar_msg(ruta, nombre, logo_bloquea, con_avisos)
        elif ext in TEXTO:
            out += _revisar_texto_plano(ruta, nombre)
        elif ext in BINARIO_LEGADO or ext in IMAGEN:
            crudo = open(ruta, 'rb').read()
            out += _revisar_bytes(crudo, nombre, 'contenido' if ext in BINARIO_LEGADO else 'datos de la imagen')
            if ext in IMAGEN and (con_avisos or logo_bloquea):
                out += revisar_imagen(crudo, nombre, 'la imagen', logo_bloquea)
    except PermissionError as e:
        out.append(Hallazgo(nombre, 'archivo', 'ilegible', f'abierto por otro programa ({e})', '', 'AVISO'))
    except Exception as e:  # un archivo raro no frena el barrido, pero se dice
        out.append(Hallazgo(nombre, 'archivo', 'ilegible', f'no se pudo revisar ({type(e).__name__}: {e})', '', 'AVISO'))
    return out


def recorrer(rutas: list[str], desde_ts: float | None = None, incluir_nube: bool = False,
             max_mb: float = 60):
    """Genera (ruta, raiz) de los documentos debajo de cada ruta. Respeta las rutas excluidas."""
    for r in rutas:
        r = os.path.abspath(r)
        if os.path.isfile(r):
            yield r, os.path.dirname(r)
            continue
        for d, subdirs, archivos in os.walk(r):
            if ruta_excluida(d):
                subdirs[:] = []
                continue
            for a in archivos:
                p = os.path.join(d, a)
                if a.startswith('~$') or os.path.splitext(a)[1].lower() not in DOCUMENTO:
                    continue
                if ruta_excluida(p):
                    continue
                if desde_ts is not None:
                    try:
                        if os.path.getmtime(p) < desde_ts:
                            continue
                    except OSError:
                        continue
                if not incluir_nube and en_la_nube(p):
                    yield p, r  # se informa como no revisado
                    continue
                yield p, r


def revisar(rutas: list[str], desde_ts: float | None = None, incluir_nube: bool = False,
            max_mb: float = 60, con_avisos: bool = True, logo_bloquea: bool = False):
    """(hallazgos, revisados, salteados_en_la_nube)."""
    hallazgos: list[Hallazgo] = []
    revisados = 0
    nube: list[str] = []
    for p, raiz in recorrer(rutas, desde_ts, incluir_nube, max_mb):
        if not incluir_nube and en_la_nube(p):
            nube.append(p)
            continue
        hs = revisar_archivo(p, raiz=raiz, max_mb=max_mb, logo_bloquea=logo_bloquea, con_avisos=con_avisos)
        revisados += 1
        hallazgos += [h for h in hs if con_avisos or h.nivel == 'BLOQUEANTE']
    return hallazgos, revisados, nube


def exigir_sin_firma(rutas: list[str], quien: str = '') -> None:
    """Para los generadores y los que mandan o imprimen: si algun archivo nombra a Claude o a una
    IA, imprime donde y sale con 1. No pregunta ni tiene escape: el texto se saca del archivo."""
    todos, _, _ = revisar(list(rutas), incluir_nube=True, con_avisos=True)
    hs = [h for h in todos if h.nivel == 'BLOQUEANTE']
    # el logo en un documento ajeno y las frases que delatan son AVISO: se imprime o se manda igual, pero se dice (H9, H11)
    for h in (x for x in todos if x.nivel == 'AVISO' and (x.regla.startswith('logo-') or x.regla in ('reproceso-delata', 'antes-despues'))):
        print(f'  OJO{(" (" + quien + ")") if quien else ""}: {h.archivo} — {h.lugar}: {h.que}', file=sys.stderr)
    if hs:
        print(f'\n✋ FIRMA DE IA EN EL DOCUMENTO{(" (" + quien + ")") if quien else ""} — no sale así '
              '(regla de Fak, 08/10/2026: ningún documento dice que lo hizo Claude o una IA):', file=sys.stderr)
        for h in hs:
            print(h.renglon(), file=sys.stderr)
        print('  Sacalo del archivo (o del generador) y volvé a correr.', file=sys.stderr)
        sys.exit(1)


_COMPLEMENTO_CLAUDE = re.compile(r'claude|WA200009404', re.IGNORECASE)


def quitar_complemento_claude(ruta: str, destino: str | None = None) -> list[str]:
    """Saca de un Excel / Word / PowerPoint la marca que deja el complemento "Claude para Excel /
    PowerPoint" (un panel de tareas guardado en `<xl|ppt|word>/webextensions/`, con la propiedad
    `claude.fileId` y la referencia de tienda WA200009404). Cirugia sobre el zip: el resto de las
    partes se copia byte a byte. Devuelve las partes que saco ([] si no habia nada).
    Si la marca esta enganchada desde otro lado que no sea el panel de tareas, NO toca nada y
    levanta RuntimeError (se arregla a mano)."""
    destino = destino or ruta
    with zipfile.ZipFile(ruta) as z:
        nombres = z.namelist()
        datos = {n: z.read(n) for n in nombres}
        infos = {i.filename: i for i in z.infolist()}
    claude_ext = [n for n in nombres if re.match(r'(xl|ppt|word)/webextensions/webextension\d+\.xml$', n)
                  and _COMPLEMENTO_CLAUDE.search(datos[n].decode('utf-8', 'replace'))]
    if not claude_ext:
        return []
    base = claude_ext[0].split('/')[0]
    dir_we = f'{base}/webextensions/'
    tp, tp_rels = dir_we + 'taskpanes.xml', dir_we + '_rels/taskpanes.xml.rels'
    # ¿alguien mas que el panel de tareas apunta a esas partes?
    for n in nombres:
        if n.endswith('.rels') and n != tp_rels:
            t = datos[n].decode('utf-8', 'replace')
            for ext in claude_ext:
                if os.path.basename(ext) in t and not n.startswith(dir_we + '_rels/webextension'):
                    raise RuntimeError(f'{ruta}: {ext} esta enganchada desde {n}; se arregla a mano')
    sacar = set(claude_ext)
    for ext in claude_ext:   # rels propios de la extension y lo que cuelga de ellos (capturas)
        r = dir_we + '_rels/' + os.path.basename(ext) + '.rels'
        if r in datos:
            sacar.add(r)
            for tgt in re.findall(r'Target="([^"]+)"', datos[r].decode('utf-8', 'replace')):
                p = os.path.normpath(os.path.join(dir_we, tgt)).replace('\\', '/')
                if p in datos:
                    sacar.add(p)
    nuevos = dict(datos)
    if tp in datos and tp_rels in datos:
        rels = datos[tp_rels].decode('utf-8')
        ids_claude = [m.group(1) for m in re.finditer(r'<Relationship\b[^>]*?Id="([^"]+)"[^>]*?/>', rels)
                      if any(os.path.basename(e) in m.group(0) for e in claude_ext)]
        tpx = datos[tp].decode('utf-8')
        for rid in ids_claude:
            tpx = re.sub(rf'<wetp:taskpane\b(?:(?!</wetp:taskpane>).)*?r:id="{rid}"(?:(?!</wetp:taskpane>).)*?</wetp:taskpane>', '', tpx, flags=re.S)
            rels = re.sub(rf'<Relationship\b[^>]*?Id="{rid}"[^>]*?/>', '', rels)
        if '<wetp:taskpane ' in tpx or '<wetp:taskpane>' in tpx:
            nuevos[tp], nuevos[tp_rels] = tpx.encode('utf-8'), rels.encode('utf-8')
        else:   # no queda ningun panel: se va el panel entero y su enganche en la raiz
            sacar |= {tp, tp_rels}
            raiz = datos['_rels/.rels'].decode('utf-8')
            raiz = re.sub(r'<Relationship\b[^>]*?Type="[^"]*webextensiontaskpanes"[^>]*?/>', '', raiz)
            nuevos['_rels/.rels'] = raiz.encode('utf-8')
    ct = datos['[Content_Types].xml'].decode('utf-8')
    for p in sacar:
        ct = re.sub(rf'<Override\b[^>]*?PartName="/{re.escape(p)}"[^>]*?/>', '', ct)
    nuevos['[Content_Types].xml'] = ct.encode('utf-8')
    tmp = destino + '.tmp-sinfirma'
    with zipfile.ZipFile(tmp, 'w') as zo:
        for n in nombres:
            if n in sacar:
                continue
            i = infos[n]
            zi = zipfile.ZipInfo(n, date_time=i.date_time)
            zi.compress_type, zi.external_attr = i.compress_type, i.external_attr
            zo.writestr(zi, nuevos[n])
    with zipfile.ZipFile(tmp) as zt:   # el resultado tiene que abrir y no tener la marca
        if zt.testzip() is not None or any(n in sacar for n in zt.namelist()):
            os.remove(tmp)
            raise RuntimeError(f'{ruta}: el archivo limpio no paso la verificacion; no se toco')
    os.replace(tmp, destino)
    return sorted(sacar)


AUTOR = 'Facundo Santoro'   # lo que escribe Excel en esta PC (docProps/core.xml del listado, 08/10/2026)


def limpiar_propiedades(ruta: str) -> list[str]:
    """Propiedades de un Excel / Word / PowerPoint que dejan ver que lo armo un programa o una IA:
    autor o ultimo en guardar = nombre de libreria ('python-pptx', 'openpyxl', 'Steve Canny',
    'Unknown') o una IA -> AUTOR; descripcion / palabras clave / asunto / titulo / categoria que
    dicen 'generated using python-pptx' o nombran una IA -> vacio. Solo toca docProps/core.xml
    (cirugia sobre el zip; el resto byte a byte). Devuelve lo que cambio."""
    with zipfile.ZipFile(ruta) as z:
        nombres = z.namelist()
        if 'docProps/core.xml' not in nombres:
            return []
        datos = {n: z.read(n) for n in nombres}
        infos = {i.filename: i for i in z.infolist()}
    xml = datos['docProps/core.xml'].decode('utf-8')
    cambios = []

    def _persona(m):
        tag, attrs, val = m.group(1), m.group(2), html.unescape(m.group(3)).strip()
        if val.lower() in canon().metadata_programa or revisar_texto(val, con_avisos=False):
            cambios.append(f'{tag}: «{val}» -> «{AUTOR}»')
            return f'<{tag}{attrs}>{AUTOR}</{tag}>'
        return m.group(0)

    def _texto(m):
        tag, attrs, val = m.group(1), m.group(2), html.unescape(m.group(3)).strip()
        if revisar_texto(val, con_avisos=False):
            cambios.append(f'{tag}: «{val}» -> vacio')
            return f'<{tag}{attrs}></{tag}>'
        return m.group(0)

    xml = re.sub(r'<(dc:creator|cp:lastModifiedBy)(\b[^>]*)>(.*?)</\1>', _persona, xml, flags=re.S)
    xml = re.sub(r'<(dc:description|cp:keywords|dc:subject|dc:title|cp:category)(\b[^>]*)>(.*?)</\1>', _texto, xml, flags=re.S)
    if not cambios:
        return []
    datos['docProps/core.xml'] = xml.encode('utf-8')
    tmp = ruta + '.tmp-sinfirma'
    with zipfile.ZipFile(tmp, 'w') as zo:
        for n in nombres:
            i = infos[n]
            zi = zipfile.ZipInfo(n, date_time=i.date_time)
            zi.compress_type, zi.external_attr = i.compress_type, i.external_attr
            zo.writestr(zi, datos[n])
    with zipfile.ZipFile(tmp) as zt:
        if zt.testzip() is not None:
            os.remove(tmp)
            raise RuntimeError(f'{ruta}: el archivo limpio no paso la verificacion; no se toco')
    os.replace(tmp, ruta)
    return cambios


def a_json(hallazgos: list[Hallazgo]) -> str:
    return json.dumps([asdict(h) for h in hallazgos], ensure_ascii=False, indent=1)
