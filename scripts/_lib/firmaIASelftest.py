"""
Selftest de firmaIA.py — cada forma de "lo hizo Claude" en ROJO, y el trabajo normal en VERDE.

Arma documentos de prueba en una carpeta temporal (Excel, PowerPoint, Word, PDF, CSV) con las
formas que encontramos el 08/10/2026 en el listado de hojas de proceso — la celda CREADO POR, la
pestaña oculta, la marca del complemento "Claude para Excel" — y las que puede dejar un generador
(propiedades del archivo, notas del orador, texto partido en dos "runs"). Un detector que no puede
dar verde esta tan roto como uno que no puede dar rojo.

    python scripts/_sinFirmaIA.py --selftest
"""
import os
import sys
import tempfile
import zipfile

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)
import firmaIA  # noqa: E402

WEBEXT = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<we:webextension '
          'xmlns:we="http://schemas.microsoft.com/office/webextensions/webextension/2010/11" '
          'id="{63D3F5EF-0625-4C39-A49D-EEB3A4D217E1}"><we:reference id="WA200009404" version="1.0.0.8" '
          'store="Omex" storeType="OMEX"/><we:properties><we:property name="claude.fileId" '
          'value="&quot;e1c6&quot;"/></we:properties><we:bindings/></we:webextension>')
TASKPANES = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<wetp:taskpanes '
             'xmlns:wetp="http://schemas.microsoft.com/office/webextensions/taskpanes/2010/11">'
             '<wetp:taskpane dockstate="right" visibility="0" width="525" row="1"><wetp:webextensionref '
             'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rId1"/>'
             '</wetp:taskpane></wetp:taskpanes>')
TP_RELS = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships '
           'xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" '
           'Type="http://schemas.microsoft.com/office/2011/relationships/webextension" Target="webextension1.xml"/></Relationships>')


def _xlsx(p, celdas=None, oculta=None, autor='Facundo Santoro'):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'INDICE HOJAS DE PROCESO'
    ws['J6'] = 'CREADO POR'
    for k, v in (celdas or {}).items():
        ws[k] = v
    if oculta:
        o = wb.create_sheet(oculta)
        o['A1'] = 'notas'
        o.sheet_state = 'hidden'
    wb.properties.creator = autor
    wb.properties.lastModifiedBy = autor
    wb.save(p)
    return p


def _con_complemento(p):
    """Le agrega a un xlsx la marca del complemento, igual que la del listado real."""
    with zipfile.ZipFile(p) as z:
        datos = {n: z.read(n) for n in z.namelist()}
    rels = datos['_rels/.rels'].decode()
    rels = rels.replace('</Relationships>', '<Relationship Id="rIdWE" Type="http://schemas.microsoft.com/office/2011/relationships/webextensiontaskpanes" Target="xl/webextensions/taskpanes.xml"/></Relationships>')
    ct = datos['[Content_Types].xml'].decode().replace('</Types>',
        '<Override PartName="/xl/webextensions/taskpanes.xml" ContentType="application/vnd.ms-office.webextensiontaskpanes+xml"/>'
        '<Override PartName="/xl/webextensions/webextension1.xml" ContentType="application/vnd.ms-office.webextension+xml"/></Types>')
    datos['_rels/.rels'], datos['[Content_Types].xml'] = rels.encode(), ct.encode()
    datos['xl/webextensions/taskpanes.xml'] = TASKPANES.encode()
    datos['xl/webextensions/_rels/taskpanes.xml.rels'] = TP_RELS.encode()
    datos['xl/webextensions/webextension1.xml'] = WEBEXT.encode()
    with zipfile.ZipFile(p, 'w', zipfile.ZIP_DEFLATED) as z:
        for n, d in datos.items():
            z.writestr(n, d)
    return p


def _pptx(p, texto='Colocar la pieza en el nido.', notas=None, partido=False, props_por_defecto=False):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tf = s.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame
    if partido:
        para = tf.paragraphs[0]
        for trozo in ('Revisado por Cla', 'ude el 08/10'):
            r = para.add_run()
            r.text = trozo
    else:
        tf.text = texto
    if notas:
        s.notes_slide.notes_text_frame.text = notas
    if not props_por_defecto:
        prs.core_properties.author = 'Facundo Santoro'
        prs.core_properties.last_modified_by = 'Facundo Santoro'
        prs.core_properties.comments = ''
        prs.core_properties.title = 'HO-990'
    prs.save(p)
    return p


def _docx(p, texto):
    from docx import Document
    d = Document()
    d.add_paragraph(texto)
    d.core_properties.author = 'Facundo Santoro'
    d.core_properties.last_modified_by = 'Facundo Santoro'
    d.core_properties.comments = ''
    d.save(p)
    return p


def _pdf(p, texto='HO-990 Rev.A', autor='Facundo Santoro'):
    import fitz
    doc = fitz.open()
    pag = doc.new_page()
    pag.insert_text((72, 72), texto)
    doc.set_metadata({'author': autor, 'creator': 'Microsoft Excel', 'producer': 'Microsoft Excel', 'title': 'HO'})
    doc.save(p)
    return p


def _imagen_png(dibujo, ancho=600, alto=200, gris=False):
    """Una imagen PNG en memoria: `dibujo(draw, ancho, alto)` pinta sobre blanco."""
    import io
    from PIL import Image, ImageDraw
    im = Image.new('RGB', (ancho, alto), 'white')
    dibujo(ImageDraw.Draw(im), ancho, alto)
    if gris:
        im = im.convert('L').convert('RGB')
    b = io.BytesIO()
    im.save(b, 'PNG')
    return b.getvalue()


def _letras_finas(d, w, h):
    """Un falso 'logo no oficial' para el CI: dos bandas de trazos finos y una raya en el medio."""
    for i in range(6):
        d.rectangle([40 + i * 90, 20, 50 + i * 90, h // 2 - 15], fill=(60, 60, 150))
        d.rectangle([20 + i * 95, h // 2 + 15, 30 + i * 95, h - 20], fill=(60, 60, 150))
    d.rectangle([10, h // 2 - 4, w - 10, h // 2 + 4], fill=(60, 60, 150))


def _pptx_con_imagen(p, crudo):
    import io
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(io.BytesIO(crudo), Inches(0.3), Inches(0.3), width=Inches(2))
    s.shapes.add_textbox(Inches(1), Inches(3), Inches(6), Inches(1)).text_frame.text = 'Colocar la pieza en el nido.'
    prs.core_properties.author = 'Facundo Santoro'
    prs.core_properties.last_modified_by = 'Facundo Santoro'
    prs.core_properties.comments = ''
    prs.core_properties.title = 'HO-990'
    prs.save(p)
    return p


def correr() -> int:
    fallas = 0
    n = 0

    def caso(nombre, ruta, debe_frenar, debe_decir=None, debe_avisar=None, logo_bloquea=False, sin_avisos=False):
        nonlocal fallas, n
        n += 1
        hs = firmaIA.revisar_archivo(ruta, raiz=os.path.dirname(ruta), logo_bloquea=logo_bloquea)
        bloq = [h for h in hs if h.nivel == 'BLOQUEANTE']
        ok = bool(bloq) == debe_frenar
        if ok and debe_decir:
            ok = any(debe_decir in (h.lugar + ' ' + h.texto + ' ' + h.regla) for h in bloq)
        if ok and debe_avisar:
            ok = any(debe_avisar == h.regla for h in hs if h.nivel == 'AVISO')
        if ok and sin_avisos:
            ok = not [h for h in hs if h.regla.startswith('logo-') or h.regla in ('reproceso-delata', 'antes-despues')]
        print(f'  {"ok  " if ok else "MAL "}  {nombre}')
        if not ok:
            fallas += 1
            for h in hs:
                print('        ', h.renglon())

    with tempfile.TemporaryDirectory() as td:
        j = lambda x: os.path.join(td, x)
        print('ROJOS — tiene que frenar')
        caso('08/10: CREADO POR = Claude en la celda J74', _xlsx(j('a.xlsx'), {'J74': 'Claude'}), True, 'celda J74')
        caso('08/10: pestaña oculta _CONTEXTO_CLAUDE', _xlsx(j('b.xlsx'), oculta='_CONTEXTO_CLAUDE'), True, '_CONTEXTO_CLAUDE')
        caso('08/10: texto "para el proximo Claude"', _xlsx(j('c.xlsx'), {'A1': 'CONTEXTO PARA EL PRÓXIMO CLAUDE'}), True, 'celda A1')
        caso('08/10: texto "segun el usuario"', _xlsx(j('c2.xlsx'), {'A3': 'REGLA NUEVA (segun el usuario): todas juntas'}), True, 'proximo-claude')
        caso('08/10: marca del complemento Claude para Excel', _con_complemento(_xlsx(j('d.xlsx'))), True, 'claude.fileId')
        caso('propiedades: autor Claude', _xlsx(j('e.xlsx'), autor='Claude'), True, 'propiedades')
        caso('PowerPoint: notas del orador "hecho con Claude"', _pptx(j('f.pptx'), notas='Hoja hecha con Claude'), True, 'notas')
        caso('PowerPoint: Claude partido en dos runs', _pptx(j('g.pptx'), partido=True), True, 'claude-pegado')
        caso('PowerPoint: propiedades por defecto de python-pptx', _pptx(j('h.pptx'), props_por_defecto=True), True, 'generado-por-libreria')
        caso('Word: "generado por IA"', _docx(j('i.docx'), 'Este informe fue generado por IA.'), True, 'hecho-con-ia')
        caso('Word: inteligencia artificial', _docx(j('i2.docx'), 'Preparado con inteligencia artificial'), True, 'inteligencia-artificial')
        caso('PDF: autor Claude en la metadata', _pdf(j('k.pdf'), autor='Claude'), True, 'propiedades del PDF')
        caso('PDF: "Generado con ChatGPT" en la pagina', _pdf(j('l.pdf'), texto='Generado con ChatGPT'), True, 'pagina 1')
        with open(j('m.csv'), 'w', encoding='utf-8') as f:
            f.write('HO;CREADO POR\n972;Claude\n')
        caso('CSV: Claude en una columna', j('m.csv'), True, 'Claude')
        with open(j('_CONTEXTO_CLAUDE.txt'), 'w', encoding='utf-8') as f:
            f.write('nada\n')
        caso('nombre de archivo con Claude', j('_CONTEXTO_CLAUDE.txt'), True, 'nombre del archivo')
        caso('Excel: Gemini en una celda', _xlsx(j('n.xlsx'), {'B2': 'Revisado con Gemini'}), True, 'gemini')

        print('VERDES — no tiene que frenar')
        caso('CREADO POR = F.Santoro', _xlsx(j('v1.xlsx'), {'J74': 'F.Santoro'}), False)
        caso('Jean-Claude es una persona', _xlsx(j('v2.xlsx'), {'B2': 'Contacto: Jean-Claude Dupont'}), False)
        caso('Claudio y Claudia no son Claude', _xlsx(j('v3.xlsx'), {'B2': 'Claudio Pérez', 'B3': 'Claudia Gómez'}), False)
        caso('autoría, geminis y copiloto no son IA', _docx(j('v4.docx'), 'La autoría del plano es del cliente; signo Géminis; copiloto del auto.'), False)
        caso('HO en PowerPoint con propiedades puestas', _pptx(j('v5.pptx')), False)
        caso('PDF de una HO', _pdf(j('v6.pdf')), False)
        caso('IA como codigo: aviso, no freno', _xlsx(j('v7.xlsx'), {'E9': '2HC.858.417 IA'}), False, debe_avisar='ia-suelta')

        print('FRASES QUE DELATAN (cola H11) — aviso, no freno')
        caso('09/10: «pedido del cliente de no poner reprocesos... el acta»', _docx(j('fd1.docx'), 'Pedido del cliente de no poner reprocesos. Que quede escrito en el acta.'), False, debe_avisar='reproceso-delata')
        caso('25/09: el antes y el despues de una correccion propia', _docx(j('fd2.docx'), 'El consumo antes decía 0,08 kg; queda 0,023 kg.'), False, debe_avisar='antes-despues')
        caso('«OP 80 REPROCESO DE COSTURA» en un flujograma: nada', _docx(j('fd3.docx'), 'OP 80 REPROCESO DE COSTURA'), False, sin_avisos=True)
        caso('nombres de operación reales (auditor 09/10): «Reproceso: eliminación de hilo sobrante», «...de arrugas en horno», «Reproceso de costura o scrap / Pieza rechazada por cliente»: nada',
             _docx(j('fd5.docx'), 'Reproceso: eliminación de hilo sobrante. Reproceso: eliminación de arrugas en horno. Reproceso de costura o scrap / Pieza rechazada por cliente. Sacabocado exacto, impacta en el acta compacta.'), False, sin_avisos=True)
        caso('«eliminamos los reprocesos a pedido del cliente»: avisa', _docx(j('fd6.docx'), 'Se eliminaron los reprocesos de la OP 32 a pedido del cliente.'), False, debe_avisar='reproceso-delata')
        caso('«Antes de coser, verificar la tensión»: nada', _docx(j('fd4.docx'), 'Antes de coser, verificar la tensión del hilo.'), False, sin_avisos=True)

        print('LOGO DE BARACK (cola H9) — el no oficial frena en un documento PROPIO y avisa en uno ajeno')
        import hashlib
        L = firmaIA._logos()
        guardado = list(L['no_oficiales'])
        falso = _imagen_png(_letras_finas)
        L['no_oficiales'] = guardado + [{'id': 'falso-ci', 'como': 'falso para el CI', 'sha256': hashlib.sha256(falso).hexdigest(),
                                         'huella': firmaIA.huella_imagen(falso)[0]}]
        try:
            caso('no oficial en un PowerPoint propio (--logo-bloquea): FRENA', _pptx_con_imagen(j('lg1.pptx'), falso), True, 'logo-no-oficial', logo_bloquea=True)
            caso('el mismo en un PowerPoint ajeno: aviso, no freno', _pptx_con_imagen(j('lg2.pptx'), falso), False, debe_avisar='logo-no-oficial')
            import io
            from PIL import Image
            Image.open(io.BytesIO(falso)).resize((900, 300), Image.LANCZOS).save(j('lg3.png'))   # el mismo, re-guardado mas grande
            caso('no oficial re-guardado más grande, suelto: frena con --logo-bloquea', j('lg3.png'), True, 'logo-no-oficial', logo_bloquea=True)
        finally:
            L['no_oficiales'] = guardado
        oficial = open(os.path.join(os.path.dirname(os.path.dirname(AQUI)), 'tools', 'flowchart', 'assets', 'barack_logo.png'), 'rb').read()
        caso('el logo OFICIAL en un PowerPoint: nada', _pptx_con_imagen(j('lg4.pptx'), oficial), False, sin_avisos=True, logo_bloquea=True)
        import io
        from PIL import Image
        gris = io.BytesIO(); Image.open(io.BytesIO(oficial)).convert('L').save(gris, 'PNG')
        caso('el oficial pasado a gris: aviso "parece el logo"', _pptx_con_imagen(j('lg5.pptx'), gris.getvalue()), False, debe_avisar='logo-parecido', logo_bloquea=True)
        foto = _imagen_png(lambda d, w, h: d.ellipse([50, 50, w - 50, h - 50], fill=(120, 90, 60)), ancho=800, alto=600)
        caso('una foto 4:3: ni se mira', _pptx_con_imagen(j('lg6.pptx'), foto), False, sin_avisos=True, logo_bloquea=True)
        real = os.path.join(os.path.expanduser('~'), 'BARACK ARGENTINA SRL', 'Ingeniería y Proyecto - General', 'LOGO BARACK.png')
        if os.path.exists(real):
            caso('el LOGO BARACK.png real de la raiz de la nube: frena con --logo-bloquea', real, True, 'logo-no-oficial', logo_bloquea=True)
        else:
            print('  --    (el LOGO BARACK.png real no esta en esta PC: ese caso se saltea)')

        print('ARREGLO — quitar_complemento_claude')
        p = _con_complemento(_xlsx(j('x.xlsx'), {'J74': 'F.Santoro'}))
        sacadas = firmaIA.quitar_complemento_claude(p)
        n += 1
        import openpyxl
        wb = openpyxl.load_workbook(p)
        limpio = not [h for h in firmaIA.revisar_archivo(p) if h.nivel == 'BLOQUEANTE']
        ok = len(sacadas) == 3 and limpio and wb.active['J74'].value == 'F.Santoro'
        print(f'  {"ok  " if ok else "MAL "}  saca las 3 partes, el libro abre y queda limpio ({sacadas})')
        fallas += 0 if ok else 1
        n += 1
        ok = firmaIA.quitar_complemento_claude(_xlsx(j('y.xlsx'))) == []
        print(f'  {"ok  " if ok else "MAL "}  un libro sin la marca no se toca')
        fallas += 0 if ok else 1

        print('ARREGLO — limpiar_propiedades')
        p = _pptx(j('z.pptx'), props_por_defecto=True)
        cambios = firmaIA.limpiar_propiedades(p)
        n += 1
        from pptx import Presentation
        cp = Presentation(p).core_properties
        ok = (not [h for h in firmaIA.revisar_archivo(p) if h.nivel == 'BLOQUEANTE'] and cp.comments == ''
              and cp.last_modified_by == firmaIA.AUTOR and len(Presentation(p).slides) == 1)
        print(f'  {"ok  " if ok else "MAL "}  saca "generated using python-pptx" y "Steve Canny"; el deck abre ({cambios})')
        fallas += 0 if ok else 1
        n += 1
        ok = firmaIA.limpiar_propiedades(_pptx(j('z2.pptx'))) == []
        print(f'  {"ok  " if ok else "MAL "}  propiedades de una persona no se tocan')
        fallas += 0 if ok else 1

    print(f'\nselftest firmaIA: {"todo verde" if not fallas else f"{fallas} MAL"} ({n} casos)')
    return 1 if fallas else 0


if __name__ == '__main__':
    sys.exit(correr())
