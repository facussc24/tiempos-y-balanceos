# -*- coding: utf-8 -*-
"""revisarCasos.py — los casos de respuesta conocida de scripts/revisar.py (tanda 16, pieza 2).

Son los 13 entregables que Fak devolvio por la forma (informe G de formatos, tabla 1.3; las filas 1 y 2 son el mismo
PowerPoint). Los archivos reales no estan en el repo (exports/ esta fuera de git y el repo es publico), asi que cada
caso es un GEMELO MINIMO que reproduce lo que Fak reclamo, armado aca con su cita al lado:
  - los 10 por CONTENIDO: la version rechazada da ROJO (en la regla que dice el caso) y la corregida da VERDE;
  - los 3 VISUALES: la version que quedo da VERDE (lo visual lo ve una persona; donde el reclamo tenia ademas una
    parte de contenido que el revisor si ve, la version rechazada tambien se prueba en ROJO).
Corre con `python scripts/revisar.py --selftest`; lo llama __tests__/scripts/revisar.test.mjs.
"""
import io
import os
import sys
import tempfile

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(os.path.dirname(AQUI))
LOGO = os.path.join(RAIZ, "tools", "flowchart", "assets", "barack_logo.png")
AUTOR = "Facundo Santoro"


# ---------------------------------------------------------------- armadores de gemelos
def foto(ruta, w=400, h=300):
    """Una «foto» gris (proporcion 4:3, lejos de la del logo para que el control de logo no la confunda)."""
    from PIL import Image, ImageDraw
    im = Image.new("RGB", (w, h), (120, 128, 136))
    dr = ImageDraw.Draw(im)
    for i in range(0, w, 20):
        dr.line([(i, 0), (w - i, h)], fill=(90 + i % 60, 100, 110), width=3)
    im.save(ruta)
    return ruta


def pptx(ruta, laminas, logo=False, carpeta=None):
    """laminas: [{"textos": [(texto, pt|None)], "foto": bool}]"""
    from pptx import Presentation
    from pptx.util import Cm, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(33.867), Cm(19.05)
    for lam in laminas:
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        if logo:
            sl.shapes.add_picture(LOGO, Cm(1), Cm(0.6), height=Cm(1.5))
        y = 2.5
        for texto, pt in lam["textos"]:
            tb = sl.shapes.add_textbox(Cm(1.5), Cm(y), Cm(30), Cm(1.6))
            run = tb.text_frame.paragraphs[0].add_run()
            run.text = texto
            run.font.name = "Calibri"
            if pt:
                run.font.size = Pt(pt)
            y += 1.8
        if lam.get("foto"):
            sl.shapes.add_picture(foto(os.path.join(carpeta, f"f{len(prs.slides)}.png")), Cm(18), Cm(5), width=Cm(12))
    cp = prs.core_properties
    cp.author, cp.last_modified_by, cp.comments, cp.title = AUTOR, AUTOR, "", ""
    prs.save(ruta)
    return ruta


def xlsx(ruta, hojas):
    """hojas: [{"nombre", "celdas": {"B2": valor | {"v": valor, "fondo": "FFFF00", "gris": True, "nota": "..."}}, "oculta"}]"""
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for h in hojas:
        ws = wb.create_sheet(h["nombre"])
        if h.get("oculta"):
            ws.sheet_state = "hidden"
        for ref, x in h["celdas"].items():
            x = x if isinstance(x, dict) else {"v": x}
            c = ws[ref]
            c.value = x.get("v")
            if x.get("fondo"):
                c.fill = PatternFill("solid", fgColor=x["fondo"])
            if x.get("gris"):
                c.font = Font(name="Calibri", italic=True, color="808080")
            if x.get("nota"):
                c.comment = Comment(x["nota"], "F.Santoro")
    wb.properties.creator = AUTOR
    wb.properties.lastModifiedBy = AUTOR
    wb.save(ruta)
    return ruta


def pdf(ruta, renglones):
    import fitz
    doc = fitz.open()
    p = doc.new_page()
    y = 60
    for t in renglones:
        p.insert_text((50, y), t, fontsize=11)
        y += 18
    doc.set_metadata({"author": AUTOR, "creator": "", "producer": ""})
    doc.save(ruta)
    return ruta


def texto(ruta, contenido):
    io.open(ruta, "w", encoding="utf-8").write(contenido)
    return ruta


CAJETIN = "HOJA DE OPERACIONES Form: I-IN-002.4-R01 HO-996 N° DE OPERACIÓN COD. DE PIEZA / DESCRIPCIÓN REALIZÓ: F.SANTORO"
TESTAMENTO = ("Estimado Carlos, buenas tardes. Por medio del presente queria comentarte que, en el marco de la revision "
              "integral que venimos llevando adelante sobre el plan de control del apoyacabezas del proyecto Patagonia, "
              "hemos procedido a realizar un analisis exhaustivo de cada una de las caracteristicas, contrastandolas con "
              "el AMFE vigente, el flujograma y las fichas de embalaje. Como resultado de dicho analisis, se identificaron "
              "diversas oportunidades de mejora que detallo a continuacion: en primer lugar, la frecuencia de control de "
              "la costura; en segundo lugar, el metodo de evaluacion de la espuma; en tercer lugar, la reaccion ante un "
              "desvio en el peso; y por ultimo, la identificacion de las caracteristicas especiales segun el criterio del "
              "manual. Quedo a entera disposicion para coordinar una reunion en la que podamos repasar cada punto en "
              "detalle, y agradezco de antemano la atencion prestada. Sin otro particular, saludo a usted muy atentamente.")


# ---------------------------------------------------------------- los casos
# Cada caso: (id, n° del informe, clase, cita, argumentos, armar_antes | None, regla que frena el antes, armar_despues)
def casos(d):
    j = lambda n: os.path.join(d, n)  # noqa: E731
    return [
        ("imds-gerente", "1-2", "visual", "09/10 PowerPoint del IMDS: «parece hecho con IA»; «usaste un logo no oficial de Barack, gravísimo»",
         ["--tipo", "leer"], None, None,
         lambda: pptx(j("imds_despues.pptx"), [{"textos": [("La declaración IMDS del Insert quedó lista para cargar", 24),
                                                        ("Captura del portal con el número de MDS", 12)], "foto": True}], logo=True, carpeta=d)),
        ("hoja-reunion", "3", "visual", "09/10 hoja para la reunión de AMFE: «que no parezca IA»",
         ["--tipo", "una-pagina"], None, None,
         lambda: pdf(j("reunion_despues.pdf"), ["Necesito que Calidad confirme la frecuencia de control de la costura",
                                                 "OP 30 costura: hoy se controla al inicio del turno", "OP 50 espumado: sin cambios"])),
        ("meta-ads", "4", "visual", "14/10 PDF de Meta Ads: «el qué haría está de más… no agregues cosas de más, el formato no me convence»",
         ["--tipo", "una-pagina"],
         lambda: pdf(j("meta_antes.pdf"), ["Campañas de Instagram de octubre: 3 activas", "Alcance total 12.400 personas",
                                            "Qué haría yo: subir el presupuesto de la campaña 2"]), "recomendaciones",
         lambda: pdf(j("meta_despues.pdf"), ["Campañas de Instagram de octubre: 3 activas", "Alcance total 12.400 personas"])),
        ("camino-inspeccion", "5", "contenido", "13/10 Carlos Baptista: «el formato está mal. Tenemos que hacer una hoja de proceso…»",
         ["--oficial", "--documento", "camino de inspeccion", "--tipo", "leer"],
         lambda: pptx(j("camino_antes.pptx"), [{"textos": [("Camino de inspección de grampas - Top Roll", None), ("x4 x3", None)], "foto": True}], carpeta=d),
         "formulario-oficial",
         lambda: pptx(j("camino_despues.pptx"), [{"textos": [(CAJETIN, None), ("50 UBICACION DE LAS GRAMPAS - TOP ROLL", None)], "foto": True}], carpeta=d)),
        ("plan-control-amarillo", "6", "contenido", "15/10 plan de control del apoyacabezas: «todo amarillo con demasiadas notas… ¿por qué no solo D?»",
         ["--oficial", "--documento", "plan de control"],
         lambda: xlsx(j("pc_antes.xlsx"), [{"nombre": "PC CENTRAL", "celdas": {
             "A1": "PLAN DE CONTROL I-AC-032", "A3": "Caracteristica", "B3": "Clase", "C3": "Metodo",
             "A4": {"v": "Costura", "fondo": "FFFF00", "nota": "cambiado por Ingenieria"}, "B4": {"v": "D / D", "fondo": "FFFF00"}, "C4": "Visual",
             "A6": "NOTA: se agregaron las caracteristicas del AMFE"}}]), "amarillo-y-antes",
         lambda: xlsx(j("pc_despues.xlsx"), [{"nombre": "PC CENTRAL", "celdas": {
             "A1": "PLAN DE CONTROL I-AC-032", "A3": "Caracteristica", "B3": "Clase", "C3": "Metodo",
             "A4": "Costura", "B4": "D", "C4": "Visual"}}])),
        ("cambio-de-molde", "7", "contenido", "25/09 hojas del cambio de molde de la IMG: «era extremadamente difícil de comprender»",
         ["--oficial", "--documento", "hoja de proceso", "--tipo", "leer"],
         lambda: pptx(j("molde_antes.pptx"), [
             {"textos": [(CAJETIN, None), ("30 CAMBIO DE MOLDE — MANEJO DE LA PANTALLA", None), ("Tocar el botón CAMBIO y esperar", None)], "foto": True},
             {"textos": [(CAJETIN, None), ("40 DESMONTAJE DEL MOLDE AUXILIAR", None)], "foto": True}], carpeta=d), "piezas-antes-que-pasos",
         lambda: pptx(j("molde_despues.pptx"), [
             {"textos": [(CAJETIN, None), ("10 RECONOCIMIENTO DEL MOLDE", None), ("Piezas que se mueven: placa superior y mordazas", None)], "foto": True},
             {"textos": [(CAJETIN, None), ("30 CAMBIO DE MOLDE — MANEJO DE LA PANTALLA", None), ("Tocar el botón CAMBIO y esperar", None)], "foto": True}], carpeta=d)),
        ("banda", "8", "contenido", "11/10 banda del tablero: «283 llamadas en curso???»; «renueva pusiste hora pero no día»",
         ["--tipo", "html"],
         lambda: texto(j("banda_antes.html"), "<html><body><div>Cupo de la API</div><div>283 llamadas · en curso</div><div>renueva 13:00</div></body></html>"),
         "contador-y-hora",
         lambda: texto(j("banda_despues.html"), "<html><body><div>Cupo de la semana: gastado 63 %, quedan 37 %</div><div>se renueva el viernes 23/10 a las 13:00</div></body></html>")),
        ("excel-folio", "9", "contenido", "03/09 Excel del folio: «sacá esas aclaraciones en gris… las flechitas… TÍTULO Y LOS DATOS»; B2 «requisito directo del dueño»",
         ["--tipo", "excel"],
         lambda: xlsx(j("folio_antes.xlsx"), [{"nombre": "Consumo", "celdas": {
             "A1": "Consumo del folio protector", "A3": "Pieza", "B3": "m2",
             "A4": "Tapa", "B4": 0.42, "A5": {"v": "el consumo sale del marker de agosto", "gris": True}, "C4": "→ ver plano"}}]), "titulo-en-b2",
         lambda: xlsx(j("folio_despues.xlsx"), [{"nombre": "Consumo", "celdas": {
             "B2": "Consumo del folio protector", "B4": "Pieza", "C4": "m2", "B5": "Tapa", "C5": 0.42}}])),
        ("notas-y-fuente", "10", "contenido", "24/07 tablas con bloque NOTAS y columna Fuente: «queda horrible, queda re que lo hizo Claude»",
         ["--tipo", "excel"],
         lambda: xlsx(j("tabla_antes.xlsx"), [{"nombre": "Tabla", "celdas": {
             "B2": "Pesos por pieza", "B4": "Pieza", "C4": "Peso g", "D4": "Fuente", "B5": "Bracket", "C5": 35, "D5": "plano rev C",
             "B7": "NOTAS:", "B8": "el peso del bracket sale del plano"}}]), "notas-de-mas",
         lambda: xlsx(j("tabla_despues.xlsx"), [{"nombre": "Tabla", "celdas": {
             "B2": "Pesos por pieza", "B4": "Pieza", "C4": "Peso g", "B5": "Bracket", "C5": 35}}])),
        ("txt-de-un-3d", "11", "contenido", "31/08 .txt y .html de un 3D: «los txt son al pedo… un PDF fácil de entender a prueba de boludos con fotos reales»",
         ["--tipo", "explicacion"],
         lambda: texto(j("dispositivo_antes.txt"), "El dispositivo apoya la pieza en tres puntos.\nLa pieza entra por arriba.\n"), "medio",
         lambda: pdf(j("dispositivo_despues.pdf"), ["El dispositivo apoya la pieza en tres puntos", "La pieza entra por arriba"])),
        ("fuente-sin-captura", "12", "contenido", "01/10 PowerPoint de superficie de vinilo: «mencionaste 3 fuentes sin adjuntar la foto de evidencia»",
         ["--tipo", "leer"],
         lambda: pptx(j("vinilo_antes.pptx"), [{"textos": [("El Insert lleva 0,18 m² de vinilo", 24), ("Según la planilla de Mesa de Corte", 12)], "foto": False}], carpeta=d),
         "fuente-sin-captura",
         lambda: pptx(j("vinilo_despues.pptx"), [{"textos": [("El Insert lleva 0,18 m² de vinilo", 24), ("Según la planilla de Mesa de Corte", 12)], "foto": True}], carpeta=d)),
        ("mail-testamento", "13", "contenido", "20/08-10/09 mails: «le mandaste un testamento»; «simplificá la descripción del mail»",
         ["--tipo", "mail"],
         lambda: texto(j("mail_antes.txt"), TESTAMENTO), "voz",
         lambda: texto(j("mail_despues.txt"), "Carlos, te paso el plan de control del apoyacabezas. Cualquier cosa me avisás.")),
        ("listado-firma", "14", "contenido", "08/10 listado de hojas de proceso: «Claude» en CREADO POR y la pestaña oculta _CONTEXTO_CLAUDE",
         ["--tipo", "excel", "--oficial"],
         lambda: xlsx(j("listado_antes.xlsx"), [{"nombre": "INDICE", "celdas": {"B2": "Listado de hojas de proceso", "B4": "HO", "C4": "CREADO POR", "B5": 996, "C5": "Claude"}},
                                               {"nombre": "_CONTEXTO_CLAUDE", "celdas": {"A1": "contexto para Claude"}, "oculta": True}]), "firma-ia",
         lambda: xlsx(j("listado_despues.xlsx"), [{"nombre": "INDICE", "celdas": {"B2": "Listado de hojas de proceso", "B4": "HO", "C4": "CREADO POR", "B5": 996, "C5": "F.Santoro"}}])),
    ]


def pptx_tabla(ruta, pt):
    """Una lamina 16:9 con una tabla de numeros a `pt` puntos."""
    from pptx import Presentation
    from pptx.util import Cm, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(33.867), Cm(19.05)
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    tb = sl.shapes.add_textbox(Cm(1.5), Cm(1.5), Cm(30), Cm(2))
    run = tb.text_frame.paragraphs[0].add_run()
    run.text, run.font.size = "La OP 30 pasa de 2 a 1 operario", Pt(30)
    tabla = sl.shapes.add_table(2, 3, Cm(1.5), Cm(5), Cm(20), Cm(3)).table
    for i in range(2):
        for k in range(3):
            c = tabla.cell(i, k)
            c.text = str(10 * i + k)
            c.text_frame.paragraphs[0].runs[0].font.size = Pt(pt)
    cp = prs.core_properties
    cp.author, cp.last_modified_by, cp.comments, cp.title = AUTOR, AUTOR, "", ""
    prs.save(ruta)
    return ruta


# Contraejemplos de la auditoria del 18/10 (docs/auditorias/TANDA16_FORMATOS_auditor_2026-10-18.md): cada uno es un
# agujero que el auditor encontro y que quedo cerrado. (id, argumentos, armar, regla que tiene que dar ROJO | None = VERDE)
def contraejemplos(d):
    j = lambda n: os.path.join(d, n)  # noqa: E731
    fila = {f"{c}5": {"v": f"dato {c}", "fondo": "FFFF00", "nota": "cambiado"} for c in "ABCDEF"}
    return [
        ("fila-de-datos-amarilla-entera", ["--oficial"],
         lambda: xlsx(j("c1.xlsx"), [{"nombre": "PC", "celdas": {"A1": "PLAN DE CONTROL I-AC-032", **{f"{c}4": f"Col {c}" for c in "ABCDEF"}, **fila}}]),
         "amarillo-y-antes"),
        ("banda-del-formulario-sigue-pasando", ["--oficial"],
         lambda: xlsx(j("c2.xlsx"), [{"nombre": "PC", "celdas": {**{f"{c}2": {"v": "PLAN DE CONTROL I-AC-032" if c == "A" else None, "fondo": "FFFF00"} for c in "ABCDEF"},
                                                                 **{f"{c}4": f"Col {c}" for c in "ABCDEF"}}}]), None),
        ("columna-fuente-debajo-de-un-encabezado", ["--oficial"],
         lambda: xlsx(j("c3.xlsx"), [{"nombre": "PC", "celdas": {"A1": "PLAN DE CONTROL I-AC-032", "D1": "FECHA: 15/10/2026",
                                                                 "A4": "Caracteristica", "B4": "Metodo", "C4": "Fuente", "A5": "Costura", "B5": "Visual", "C5": "plano"}}]),
         "columnas-inventadas"),
        ("mail-que-nombra-a-la-ia", ["--tipo", "mail"],
         lambda: texto(j("c4.txt"), "Carlos, te paso el plan. Lo revisé con Claude antes de mandarlo."), "firma-ia"),
        ("tabla-de-numeros-chicos-para-proyectar", ["--tipo", "proyectar"], lambda: pptx_tabla(j("c5.pptx"), 6), "letra-minima"),
        ("hora-con-el-dia-en-la-celda-de-al-lado", ["--tipo", "html"],
         lambda: texto(j("c6.html"), "<html><body><table><tr><td>dom 18/10</td></tr><tr><td>a las 13:29</td></tr></table></body></html>"), None),
        ("formulario-oficial-sin-banderas", [],
         lambda: xlsx(j("c7.xlsx"), [{"nombre": "PC", "celdas": {"A1": "PLAN DE CONTROL I-AC-032", "A3": "Caracteristica", "B3": "Clase", "A4": "Costura", "B4": "D"}}]), None),
        # revision por API del 18/10
        ("estado-verificado-en-una-fila-de-datos", ["--tipo", "excel"],
         lambda: xlsx(j("c9.xlsx"), [{"nombre": "Datos", "celdas": {"B2": "Estado de las piezas", "B4": "Pieza", "C4": "Estado",
                                                                    "B5": "Bracket", "C5": "Verificado"}}]), None),
        ("codigo-del-formulario-con-punto-final", ["--documento", "plan de control"],
         lambda: xlsx(j("c10.xlsx"), [{"nombre": "PC", "celdas": {"A1": "Formulario I-AC-032.", "A3": "Caracteristica", "B3": "Clase", "A4": "Costura", "B4": "D"}}]), None),
        ("el-codigo-del-flujograma-no-es-el-de-la-hoja", ["--documento", "flujograma", "--tipo", "leer"],
         lambda: pptx(j("c8.pptx"), [{"textos": [("Form: I-IN-002.4-R01 hoja de operaciones", None)], "foto": True}], carpeta=d), "formulario-oficial"),
    ]


def selftest(revisar, opciones):
    ok = True
    contenido_ok = 0
    print("Casos de respuesta conocida de revisar.py (los 13 entregables devueltos; informe G, tabla 1.3)")
    with tempfile.TemporaryDirectory(prefix="revisar_casos_", ignore_cleanup_errors=True) as d:
        for cid, n, clase, cita, args, antes, regla, despues in casos(d):
            linea = f"  #{n:<4} {clase:<9} {cid:<22}"
            caso_ok = True
            if antes:
                rep = revisar(antes(), opciones(["x", "--sin-render", *args]))
                frena = [x["regla"] for x in rep["reglas"] if x["estado"] == "ROJO"]
                bien = regla in frena
                caso_ok &= bien
                linea += f" antes: {'ROJO' if frena else 'VERDE'} ({', '.join(frena) or '-'}){'' if bien else f'  <-- esperaba ROJO en {regla}'}"
            rep = revisar(despues(), opciones(["x", "--sin-render", *args]))
            frena = [f"{x['regla']}: {x['donde']}" for x in rep["reglas"] if x["estado"] == "ROJO"]
            caso_ok &= not frena
            linea += f" | despues: {'VERDE' if not frena else 'ROJO <-- esperaba VERDE: ' + '; '.join(frena)}"
            print(linea)
            ok &= caso_ok
            if clase == "contenido" and caso_ok:
                contenido_ok += 1
        print("Contraejemplos de la auditoria del 18/10")
        for cid, args, armar, regla in contraejemplos(d):
            rep = revisar(armar(), opciones(["x", "--sin-render", *args]))
            frena = [x["regla"] for x in rep["reglas"] if x["estado"] == "ROJO"]
            bien = (regla in frena) if regla else not frena
            ok &= bien
            esperado = f"ROJO en {regla}" if regla else "VERDE"
            print(f"  {cid:<42} {'ROJO (' + ', '.join(frena) + ')' if frena else 'VERDE'}{'' if bien else f'  <-- esperaba {esperado}'}")
    print(f"  por contenido: {contenido_ok} de 10 dan ROJO antes y VERDE despues; visuales: la version que quedo da VERDE")
    print("SELFTEST OK" if ok else "SELFTEST FALLA")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.path.insert(0, os.path.dirname(AQUI))
    import revisar as _r
    sys.exit(selftest(_r.revisar, _r.opciones))
