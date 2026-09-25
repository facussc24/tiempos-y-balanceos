"""
_hoNumeros.py — control de numeros de HO: carpetas y archivos de HOJAS DE OPERACIONES contra el
listado maestro. Solo lee.

    python scripts/_hoNumeros.py            # choques, numeros sin fila, proximo libre
    python scripts/_hoNumeros.py --q3       # ademas lee el cajetin (Q3) de cada Excel (lento, red)

Criterio de Fak (25/09/2026, hoja oculta _CONTEXTO_CLAUDE del listado):
- El numero es de la HOJA, no del codigo: varias piezas que se hacen igual van en la misma HO,
  una pagina por codigo (025, 952, 984, 913).
- Si dos hojas tienen el mismo numero, se lo queda la fila del listado. Un numero que se deja de
  usar no se reusa.
- No se pasa de 999: despues del 999 van los huecos que nunca tuvieron hoja, y despues la serie 800.

Que cuenta como CHOQUE: el mismo numero como numero PRINCIPAL (el de la carpeta "HO NNN - ...",
o el del nombre del archivo si no hay carpeta, o el Q3 de la primera pestaña con --q3) en dos
carpetas de HO distintas. Una pestaña copiada de otra HO adentro de un libro (926 con pestañas de
la 027) no es choque. Sale con codigo 1 si hay choques.
"""
import os
import re
import sys
import warnings

H = r'Y:\BARACK\CALIDAD\DOCUMENTACION SGC\HOJAS DE OPERACIONES'
LISTADO = os.path.join(H, '3- LISTADO', 'Listado hojas de proceso.xlsx')
ZONAS = ['1- CLIENTES', '2- SECTORES', '4- RETRABAJOS']
HISTORICO = '9- HISTORICO (NO USAR)'
OBSOLETO = re.compile(r'^(obsoletos?|osboleto|obolseto|old)$', re.I)
CARPETA_HO = re.compile(r'^HO (\d{3}) - ')
NUM_NOMBRE = re.compile(r'\bH\.?O\.?\s*[-_]?\s*(\d{2,3})(?![\d-])', re.I)
EXT = ('.xlsx', '.xlsm', '.xls', '.pptx', '.pdf')
# numeros que Fak acepto compartidos entre carpetas (misma hoja, o su traduccion)
COMPARTIDOS = {
    25: 'embalaje de corte: la misma hoja para las piezas Amarok y Taos',
    112: 'traduccion al ingles de la HO 112',
    214: 'traduccion al ingles de la HO 214',
    935: 'un solo archivo para el IP Amarok corto y largo',
    945: 'un solo archivo para el IP Taos corto y largo',
}
HUECOS = [903, 921, 922, 923, 924, 925, 928, 930, 936, 947]


def numeros_del_listado():
    """{numero: [(fila, descripcion, tipo)]} de la hoja INDICE (una celda puede traer '915 / 958')."""
    import openpyxl
    wb = openpyxl.load_workbook(LISTADO, read_only=True, data_only=True)
    ws = wb['INDICE HOJAS DE PROCESO']
    out = {}
    for i, fila in enumerate(ws.iter_rows(min_row=7, values_only=True), start=7):
        d, g, t = fila[3], fila[6], fila[7]
        if d is None:
            continue
        for n in re.findall(r'\d{2,3}', str(d).split('.')[0] if isinstance(d, float) else str(d)):
            out.setdefault(int(n), []).append((i, str(g or '').strip(), str(t or '').strip()))
    wb.close()
    return out


def q3_primera(ruta):
    import openpyxl
    try:
        wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
        ws = next(w for w in wb.worksheets if w.sheet_state == 'visible')
        m = re.search(r'(\d{3})', str(ws['Q3'].value or ''))
        wb.close()
        return int(m.group(1)) if m else None
    except Exception:
        return None


def recorrer(con_q3):
    """[(numero, carpeta de la HO, archivo)] de las zonas vivas, sin OBSOLETO."""
    usos = []
    for z in ZONAS:
        for d, subs, files in os.walk(os.path.join(H, z)):
            subs[:] = [s for s in subs if not OBSOLETO.match(s)]
            partes = os.path.relpath(d, H).split(os.sep)
            ho_dir = None
            for k in range(len(partes), 0, -1):
                m = CARPETA_HO.match(partes[k - 1])
                if m:
                    ho_dir = (int(m.group(1)), os.sep.join(partes[:k]))
                    break
            for f in files:
                if f.startswith('~$') or not f.lower().endswith(EXT):
                    continue
                if ho_dir:
                    usos.append((ho_dir[0], ho_dir[1], f))
                    continue
                m = NUM_NOMBRE.search(f)
                n = int(m.group(1)) if m else None
                if n is None and con_q3 and f.lower().endswith(('.xlsx', '.xlsm')):
                    n = q3_primera(os.path.join(d, f))
                if n is not None:
                    usos.append((n, os.path.relpath(d, H), f))
    return usos


def numeros_del_historico():
    usados = set()
    for d, subs, files in os.walk(os.path.join(H, HISTORICO)):
        for f in files + subs:
            for m in NUM_NOMBRE.finditer(f):
                usados.add(int(m.group(1)))
    return usados


def main():
    warnings.filterwarnings('ignore')
    listado = numeros_del_listado()
    usos = recorrer('--q3' in sys.argv)
    por_num = {}
    for n, carpeta, f in usos:
        por_num.setdefault(n, {}).setdefault(carpeta, []).append(f)

    choques = {n: c for n, c in por_num.items() if len(c) > 1 and n not in COMPARTIDOS}
    print(f'HO en carpetas vivas: {len(por_num)} numeros · filas del listado: {len(listado)} numeros')
    print(f'\nCHOQUES (mismo numero en dos carpetas de HO): {len(choques)}')
    for n in sorted(choques):
        print(f'  HO {n:03d}')
        for carpeta, fs in choques[n].items():
            print(f'     {carpeta}  ({", ".join(fs[:2])}{"..." if len(fs) > 2 else ""})')

    def activa(g, t):
        return t.upper() != 'OBSOLETO' and 'SIN HOJA VIGENTE' not in g.upper()

    sin_fila = sorted(n for n in por_num if n not in listado)
    print(f'\nEN CARPETAS PERO SIN FILA EN EL LISTADO: {len(sin_fila)}')
    for n in sin_fila:
        print(f'  HO {n:03d}: {", ".join(list(por_num[n])[:2])}')
    dobles = sorted(n for n, filas in listado.items() if n not in COMPARTIDOS
                    and len({g.split("]")[-1].strip().upper() for _, g, t in filas if activa(g, t)}) > 1)
    print(f'\nLISTADO: numeros activos con descripciones distintas: {len(dobles)}')
    for n in dobles:
        print(f'  HO {n:03d}: ' + ' | '.join(f'fila {i} {g}' for i, g, t in listado[n] if activa(g, t)))

    usados = set(listado) | set(por_num) | numeros_del_historico()
    libres = [n for n in range(992, 1000) if n not in usados] + [n for n in HUECOS if n not in usados]
    if not libres:
        libres = [n for n in range(801, 900) if n not in usados][:1]
    print(f'\nPROXIMO LIBRE (hojas generales): {libres[0] if libres else "ninguno"}'
          f'  · quedan {len(libres)} antes de pasar a la 800')
    return 1 if choques else 0


if __name__ == '__main__':
    sys.exit(main())
