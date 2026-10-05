# -*- coding: utf-8 -*-
"""que_falta_filmar_ho971.py — la lista de lo que falta filmar o preguntar en planta para terminar la
HO-971 (APB de puerta Patagonia) en el formato A3. Una lamina A4 apaisada por tabla, para llevar a planta.

    py -3 scripts/ho971/que_falta_filmar_ho971.py      # pptx + pdf en exports/HO971_UNIFICADA

Fuente: el inventario de material del 05/10/2026 (fotos y videos de la biblioteca y de la HO vieja) y
la lectura a ciegas de las laminas de la mesa de corte.
"""
import os
import sys

from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(AQUI, "..", "hotmelt"))
import generar_hojas_hotmelt_a3 as base      # noqa: E402

SALIDA = os.path.abspath(os.path.join(AQUI, "..", "..", "exports", "HO971_UNIFICADA"))
AZUL = RGBColor(0x1F, 0x38, 0x64)

FILMAR = [
    ("50", "Inyección de piezas plásticas", "La inyectora con el molde del sustrato del APB: carga del molde, ciclo y retiro de la pieza."),
    ("60", "Troquelado de cinta (nueva)", "El troquelado del Tesa en Conversión de Cinta: carga del rollo, troquelado y las tiras que salen."),
    ("70", "Inyección PU", "La inyectora de PU con esta pieza: desmoldante, colocación del plástico, colada y desmolde."),
    ("71", "Armado de cinta y espuma (nueva)", "El armado de la cinta y la espuma sobre el plástico mientras cura la pieza siguiente."),
    ("80", "Adhesivado hot melt", "La pieza del APB pasando por la línea hot melt (rodillo): carga, pasada y retiro."),
    ("82", "Reproceso: falta de adhesivo", "Cómo se repone el adhesivo en la zona sin pegar."),
    ("91", "Tapizado automático", "La prensa BMA103 con el APB de puerta cargado y en ciclo (las fotos que hay son del APB de Taos)."),
    ("92", "Refilado manual", "El refilado a mano del APB de puerta después del tapizado (las fotos que hay son de Taos)."),
    ("101", "Reproceso: reactivación por calor", "La pistola de calor sobre una zona despegada del APB de puerta."),
    ("110", "Embalaje", "La caja con los APB terminados (falta que Carlos diga cuál ficha queda: cartón o plástica)."),
]
PREGUNTAR = [
    ("20", "Mesa de corte", "Cuál de los dos Start de la pantalla de la máquina de capas se aprieta; y fotos de la planilla, del cajón de scrap y del botón de bajada del rollo."),
    ("21", "Mesa de corte", "Qué archivo de Cutting Control se abre para cada código; cómo termina el corte (retirar piezas, apagar la succión, el nylon); fotos de cerca del botón Plato y del láser."),
    ("40 / 41", "Costura", "Fotos de los conos de hilo con la aguja cargada y del corte de hilos al terminar la pieza."),
    ("22 / 51 / 81 / 100", "Controles", "Son láminas de Calidad: la 51 delantero y trasero ya está; pedirle a Calidad la 22, la 81 y la 100."),
]


def tabla(prs, titulo, cab, filas):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Cm(1.0), Cm(0.6), Cm(27.7), Cm(1.4)).text_frame
    tb.text = titulo
    tb.paragraphs[0].runs[0].font.size = Pt(22)
    tb.paragraphs[0].runs[0].font.bold = True
    tb.paragraphs[0].runs[0].font.color.rgb = AZUL
    t = s.shapes.add_table(len(filas) + 1, 3, Cm(1.0), Cm(2.3), Cm(27.7), Cm(1.0 * (len(filas) + 1))).table
    for j, w in enumerate((2.8, 6.4, 18.5)):
        t.columns[j].width = Cm(w)
    for i, fila in enumerate([cab] + filas):
        for j, v in enumerate(fila):
            c = t.cell(i, j)
            c.text = v
            r = c.text_frame.paragraphs[0].runs[0]
            r.font.size = Pt(12 if i else 12.5)
            r.font.bold = (i == 0 or j == 0)
            if i == 0:
                c.fill.solid()
                c.fill.fore_color.rgb = AZUL
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(29.7), Cm(21.0)
    tabla(prs, "HO-971 APB de puerta — QUÉ FALTA FILMAR", ("OP", "Operación", "Qué filmar o fotografiar (de esta pieza)"), FILMAR)
    tabla(prs, "HO-971 APB de puerta — QUÉ FALTA PREGUNTAR O COMPLETAR", ("OP", "Sector", "Qué falta"), PREGUNTAR)
    os.makedirs(SALIDA, exist_ok=True)
    ruta = os.path.join(SALIDA, "HO-971 - QUE FALTA FILMAR Y PREGUNTAR.pptx")
    prs.save(ruta)
    print(ruta)
    base.exportar_pdf_com(ruta)


if __name__ == "__main__":
    main()
