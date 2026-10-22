"""
_pendientesPcApb.py — hoja simple que acompaña la propuesta del plan de control del APB trasero central
Patagonia (19/10/2026): lo que falta definir para que el plan quede definitivo, separado por quien lo tiene
que contestar, y aparte la propuesta de siglas del AMFE 150 (para Fak, que es quien las asigna).

Palabras de planta, cada punto con su papel, sin contar como se llego a cada diferencia
(core-prohibiciones §1; memoria mails a Calidad sin flancos).

Uso:  python scripts/apb/_pendientesPcApb.py <carpeta_salida> <amfe150.json>   -> .xlsx + .pdf (y la propuesta de siglas .xlsx)
"""
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

LOGO = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
                    'tools', 'flowchart', 'assets', 'barack_logo.png')
NOMBRE = 'Pendientes plan de control APB trasero central Patagonia 22-10-2026'
NOMBRE_SIGLAS = 'Propuesta de siglas AMFE 150 APB trasero central 22-10-2026'

# (para quien, operacion, que falta, de donde sale lo que hay hoy)
PENDIENTES = [
    ('Calidad', '10 Recepción', 'Plan de recepción del vinilo Sansuy 197.025.0425-6: tolerancia del ancho del rollo y código del medidor de espesor, '
     'la balanza y la cinta métrica de recepción',
     'El plan toma del plan Rev. 1 el espesor (1,8 a 2,5 mm) y el gramaje (800 a 1000 g/m²). En el cronograma 2026 Rev. 7 los medidores de '
     'espesor MC006 y MC369 están en Laboratorio, la cinta métrica MC406 en Mesa de corte y la de recepción es la MC403'),
    ('Calidad', '10 Recepción', 'Planes de recepción de los fieltros 427TEL003COR01 y RES30020TIT, del Aplix APLIX_225LLB, del cierre 427VAR001COS01 y de la cinta sarga 5825',
     'No hay planes de estos códigos. El plan usa los valores del AMFE (480 g/m², 270 a 320 g/m², cierre 360 ± 1 mm). La cinta: el arb tiene '
     'la 5825 de 24 mm; el AMFE y el plan Rev. 1 nombran la 5826 de 25 mm con espesor 1,2 ± 0,1 mm'),
    ('Calidad', '10 Recepción', 'Qué se controla al recibir la bolsa PP 427VAR002TAP01 y las etiquetas ET-SATO',
     'No tienen plan de recepción ni control en el AMFE o en el plan Rev. 1'),
    ('Calidad', '10 Recepción', 'Vencimiento del isocianato 1394 y del poliol FF-629',
     'AMFE y plan Rev. 1: más de 6 meses desde la recepción. Plan de recepción 819: 6 meses desde producido. Plan 822: según certificado. '
     'Especificación de Mas-Tin: 6 meses de almacenamiento'),
    ('Calidad', '10 Recepción', 'Unidades en los planes de recepción: el 819 escribe la densidad del isocianato en "gr/m³" y el 822 no tiene unidades; '
     'el 1043 escribe la flamabilidad del hilo "100 mm/min" sin el "<"',
     'El plan va con g/cm³, cps, kg/m³ y s (especificación y certificados de Mas-Tin) y < 100 mm/min (TL 1010)'),
    ('Calidad', 'Instrumentos', 'Cámara de flamabilidad MC184 vencida desde el 11/09/2026 y calibre MC413 vencido desde el 08/07/2026',
     'Cronograma de calibración 2026 Rev. 7'),
    ('Calidad', 'Instrumentos', 'Códigos MC de la regla de costura, la plantilla Mylar del APB trasero central, el calibre de tapizado, la regla de largo de capa, '
     'las balanzas (recepción y pieza espumada) y el termómetro',
     'No figuran en el cronograma 2026 Rev. 7 asignados a esta pieza. Para la pieza espumada están la balanza MC423 (Inyección PUR, vencida el '
     '12/10/2026) y la MC481 (Inyección VWA). El calibre de forma MC505 es de esta pieza: el cronograma lo ubica en Inyección VWA y la hoja de operaciones en Tapizado'),
    ('Calidad', '30, 80 y 82', 'Frecuencia de los controles de armado de kits y de los dos retrabajos', 'El plan las deja en TBD'),
    ('Calidad', '10 a 82', 'Muestras límite (o fotos de aceptación y rechazo) de: limpieza de los materiales, aspecto de los kits, costura (unión, '
     'cierre, funda), vinilo sin roturas, pieza espumada (adherencias, Aplix, aspecto, orificio), tapizado (funda, felt, desgarros, cierre, rebaba), '
     'portavasos y marco (encastre y aspecto), control final y retrabajos',
     'No hay muestra patrón de esta pieza (la nombran el AMFE 150, la hoja de operaciones y el plan Rev. 1, pero no existe). El plan escribe '
     '"muestra límite (TBD: la define Calidad)" en cada una (devolución de A. Aguayo del 22/10/2026)'),
    ('Calidad', '10 Recepción', 'Si existen los patrones de aspecto y color de los consignados, del vinilo y del cierre', 'El plan Rev. 1 los nombra; el plan '
     'escribe "patrón de aspecto (TBD: confirmar con Calidad)". El hilo FX284-E0PTO tiene patrón en su plan de recepción 1043'),
    ('Calidad', '40 y 70', 'Tolerancia de la cinta de manipulación: 25 mm cosidos y 75 mm afuera de la pieza',
     'Especificación de costura y cinta del 05/10/2026 (LAH.000.881.N v6 y corte A-A del plano). El plano de VW no acota el lazo: el bloque de '
     'tolerancias generales del plano no se le aplica. Se mide con regla en la 40 y en la 70'),
    ('Calidad', '61 y 70', 'Fuerza y recorrido de la prueba de apertura y cierre del portavasos', 'El plan pide abrirlo y cerrarlo a mano hasta el tope '
     'sin que se trabe. La norma de laboratorio del portavasos no es el criterio del puesto'),
    ('Calidad', 'Encabezado', 'Número de plan, equipo y contacto. La próxima revisión la abre Calidad', 'El plan sigue en la Rev. 1 del 31/07/2026'),
    ('Calidad', 'Lay out', 'Cotas del plano y cuáles ensayos del plan de validación entran en el lay out. La flamabilidad horizontal va con el valor de la TL 1010 (< 100 mm/min); el plan Rev. 1 decía ≤ 100 mm',
     'Plan de validación DV&PV preliminar del 07/04/2025 (TL 1010, VW 50180, VW 50185, PV 2005, VW 50190, resistencia, rigidez, durabilidad, portavasos)'),
    ('Calidad', '80 y 82 Retrabajos', 'Si VW aprobó los retrabajos de la pieza', 'Los retrabajos están en el flujograma 151, el AMFE 150 y la hoja de operaciones 986'),
    ('Leo Lattanzi', '10 / 50 Inyección PU', 'Qué sistema de poliuretano lleva la pieza',
     'El arb y el plan Rev. 1 dicen ISO-PLUS 1394 + POLI-PLUS FF-629 (el plan va con esos); la BOM005 Rev. A dice ISO-PLUS 1355 + POLI-PLUS FF-636'),
    ('Leo Lattanzi', '50 Inyección PU', 'Tolerancias del tiempo de colada, el caudal, las presiones y los coeficientes de corrección; temperatura de poliol, isocianato y molde',
     'Validación de espumado de marzo 2026: programa N° 9, colada 7,00 s, caudal 4,00 kg/min, relación 200/40, presión 0,8 / 1,3 bar, corrección 100 % / 65 %. '
     'La especificación de Mas-Tin da 27 a 35 °C para los componentes y 27 a 40 °C para el molde (condiciones de laboratorio)'),
    ('Leo Lattanzi', '50 Inyección PU', 'Peso de la pieza espumada con su tolerancia',
     'Plan Rev. 1: 530 ± 30 g. Validación de marzo 2026: 444 g de espuma y 1.608 g la pieza'),
    ('Planta', '20 Corte', 'Valores de velocidad de la máquina y tensión del vinilo, y cantidad de capas',
     'La hoja de operaciones dice cargar las capas de la planilla de corte; el plan Rev. 1 dice 20 capas'),
    ('Planta', '40 / 50', 'Número de la máquina de coser y de la inyectora', 'Para completar la columna de máquina (la mesa de corte es la BM 149)'),
    ('Planta', '60 Tapizado', 'Temperatura y tiempo del horno', 'El set up del horno lo nombran la hoja de operaciones 986, el AMFE 150 y el plan Rev. 1, sin valores'),
    ('Planta', '80 Retrabajo', 'Largo máximo del hilo sobrante', 'El AMFE pide definirlo en la instrucción de retrabajo'),
    ('Ingeniería', 'Hoja de operaciones 986', 'Embalaje: la hoja dice 126 piezas por medio 111970; la ficha vigente GE-102 Rev. C dice 30 piezas por cajón CKV3755',
     'El plan va con la GE-102 Rev. C (07/10/2026)'),
    ('Ingeniería', 'Hoja de operaciones 986', 'Hojas 30, 80 y 82 sin ciclo de control; hoja 40 sin el control de puntadas y margen; peso de la hoja 50 en TBD',
     'El plan ya tiene los controles de costura y peso'),
    ('Ingeniería', 'AMFE 150', 'Códigos que no coinciden con la BOM del arb: hilo FX284TK-E0PTO (arb: FX284-E0PTO), cinta 5826 de 25 mm (arb: 5825 de 24 mm), '
     'fieltros PA3840 y 427VIN026COR01 (arb: 427TEL003COR01 y RES30020TIT)', 'BOM arb último nivel del 22/09/2026'),
    ('Ingeniería', 'AMFE 150', 'Op 10 nombra "tela tapizado"; la BOM no tiene tela (el tapizado es el vinilo)', 'BOM arb último nivel del 22/09/2026'),
    ('Ingeniería', 'AMFE 150', 'Sigla de la flamabilidad del vinilo, los fieltros, el Aplix, la cinta y la del lay out', 'El plan Rev. 1 las marca D; el AMFE 150 no las tiene asignadas'),
    ('Ingeniería', 'AMFE 150', 'Op 40 tiene modos de falla de costura vista y de dobladillo; la pieza no lleva costura vista. '
     'Op 40 tiene "sellado del orificio inferior" en la máquina de coser (es de la op 50)', 'Especificación de costura y cinta del 05/10/2026'),
    ('Ingeniería', 'AMFE 150', 'Op 90 dice "apoyacabezas" en la prevención del embalaje; op 50 nombra un "inserto EPP"; op 20 tiene una causa de '
     '"instrucción del reproceso"', 'La pieza no lleva inserto EPP (BOM arb del 22/09/2026)'),
]


def borde():
    s = Side(style='thin', color='999999')
    return Border(left=s, right=s, top=s, bottom=s)


def hoja(ws, titulo, subtitulo, encabezados, filas, anchos):
    azul = PatternFill('solid', fgColor='1F3864')
    for i, w in enumerate(anchos):
        ws.column_dimensions[chr(65 + i)].width = w
    if os.path.exists(LOGO):
        img = Image(LOGO)
        img.height, img.width = 38, 38 * img.width / img.height
        ws.add_image(img, 'A1')
    ws.row_dimensions[1].height = 32
    ws['B1'] = titulo
    ws['B1'].font = Font(name='Arial', size=13, bold=True)
    ws['B2'] = subtitulo
    ws['B2'].font = Font(name='Arial', size=10)
    fila = 4
    for i, t in enumerate(encabezados):
        c = ws.cell(fila, i + 1, t)
        c.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        c.fill = azul
        c.alignment = Alignment(vertical='center', wrap_text=True)
        c.border = borde()
    for f in filas:
        fila += 1
        for i, t in enumerate(f):
            c = ws.cell(fila, i + 1, t)
            c.font = Font(name='Arial', size=9, bold=(i == 0))
            c.alignment = Alignment(vertical='top', wrap_text=True)
            c.border = borde()
        largo = max(len(str(t)) / max(anchos[i] * 1.1, 1) for i, t in enumerate(f))
        ws.row_dimensions[fila].height = max(15, 12.5 * (int(largo) + 1))
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True


def siglas(amfe_json):
    """Lo que el criterio de I-AC-005 rev.B marca y el AMFE 150 no tiene asignado (S 9-10 = critica; S 5-8 y O >= 4 = SC)."""
    d = json.load(open(amfe_json, encoding='utf-8'))['data']
    grupos = {}
    for op in d['operations']:
        for we in op.get('workElements', []):
            for fn in we.get('functions', []):
                for fm in fn.get('failures', []):
                    for c in fm.get('causes', []):
                        s = int(c.get('severity') or fm.get('severity') or 0)   # la S de la causa, no la del modo de falla
                        o = int(c.get('occurrence') or 0)
                        if (c.get('specialChar') or '').strip():
                            continue
                        crit = 'Crítica (D)' if s >= 9 else ('SC' if 5 <= s <= 8 and o >= 4 else '')
                        if not crit:
                            continue
                        k = (str(op['opNumber']), fm['description'], crit)
                        grupos.setdefault(k, set()).add((s, o))
    filas = []
    for (op, fm, crit), so in grupos.items():
        filas.append((op, fm, ', '.join(f'S{s} O{o}' for s, o in sorted(so)), crit))
    filas.sort(key=lambda f: (0 if f[3].startswith('Cr') else 1, int(f[0])))
    return filas


def exportar_pdf(xlsx):
    import win32com.client as win32
    pdf = os.path.splitext(xlsx)[0] + '.pdf'
    if os.path.exists(pdf):
        os.remove(pdf)
    xl = win32.DispatchEx('Excel.Application')
    xl.Visible = False
    xl.DisplayAlerts = False
    try:
        wb = xl.Workbooks.Open(xlsx, 0, True)
        ps = wb.Worksheets(1).PageSetup
        ps.Zoom = False
        ps.FitToPagesWide = 1
        ps.FitToPagesTall = False
        wb.Worksheets(1).ExportAsFixedFormat(0, pdf)
        wb.Close(False)
    finally:
        xl.Quit()
    if not os.path.exists(pdf):
        raise RuntimeError('no se genero el PDF')


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(2)
    salida, amfe_json = os.path.abspath(sys.argv[1]), sys.argv[2]
    os.makedirs(salida, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Pendientes'
    hoja(ws, 'Plan de control APB trasero central Patagonia — lo que falta definir',
         '2HC.885.081 RL1 · propuesta de Ingeniería del 19/10/2026', ('Quién', 'Operación', 'Qué falta', 'Lo que hay hoy'),
         [x for x in PENDIENTES if x[0] != 'Ingeniería'], (14, 18, 62, 62))   # lo de Ingenieria no va a Calidad (auditor, 19/10)
    wb.properties.creator = 'Barack Mercosul'
    wb.properties.lastModifiedBy = 'Barack Mercosul'
    destino = os.path.join(salida, NOMBRE + '.xlsx')
    wb.save(destino)
    exportar_pdf(destino)
    print(destino)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Propuesta de siglas'
    filas = siglas(amfe_json)
    hoja(ws, 'AMFE 150 APB trasero central — siglas que marca el criterio y no están asignadas',
         'Criterio I-AC-005 rev.B: S 9 o 10 = crítica (D); S 5 a 8 con O 4 o más = SC. Las asigna Ingeniería.',
         ('Op', 'Modo de falla', 'S y O de sus causas', 'El criterio marca'), filas, (6, 80, 22, 16))
    # aparte, para Ingenieria: las correcciones de nuestros documentos que salieron del cruce (no van en el mail a Calidad)
    ws2 = wb.create_sheet('Correcciones de Ingeniería')
    hoja(ws2, 'Correcciones de Ingeniería que salieron del plan de control', '2HC.885.081 RL1 · 19/10/2026',
         ('Quién', 'Documento', 'Qué corregir', 'De dónde sale'), [x for x in PENDIENTES if x[0] == 'Ingeniería'], (14, 22, 62, 50))
    wb.properties.creator = 'Barack Mercosul'
    wb.properties.lastModifiedBy = 'Barack Mercosul'
    destino = os.path.join(salida, NOMBRE_SIGLAS + '.xlsx')
    wb.save(destino)
    print(destino, len(filas), 'modos de falla')


if __name__ == '__main__':
    main()
