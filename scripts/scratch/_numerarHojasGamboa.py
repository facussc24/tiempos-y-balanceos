# -*- coding: utf-8 -*-
"""Copia de trabajo de las hojas A3 de Pablo Gamboa (mail del 01/10/2026 09:11) con el N° de
operacion del flujograma en el cajetin, y su PDF. No toca los originales: lee del rar extraido
y escribe en la carpeta de salida.

    python scripts/scratch/_numerarHojasGamboa.py <carpeta HOJAS DE GRAMPAS> <salida>

Criterio (Fak, 01/10/2026): APB de puerta = prensa y despues grampas -> el engrampado es la
operacion 83. Top Roll: el tapizado y engrampado es como se hace hoy el plegado de bordes (50).
"""
import os
import shutil
import sys

from pptx import Presentation

SRC, OUT = (os.path.abspath(a) for a in sys.argv[1:3])

# (origen relativo, nombre de salida sin extension, N° de operacion o None si no lleva cajetin)
HOJAS = [
    (r"APOYABRAZOS\DELANTERO\HO-971_Engrampado_APB_Delantero_Patagonia_A3.pptx",
     "HO-971 - OP 83 ENGRAMPADO - APB DE PUERTA DELANTERO", "83"),
    (r"APOYABRAZOS\TRASERO\HO-971_Engrampado_APB_Trasero_Patagonia_A3.pptx",
     "HO-971 - OP 83 ENGRAMPADO - APB DE PUERTA TRASERO", "83"),
    (r"TOP ROLL\TRASERO\HO-992_TAPIZADO_Y_ENGRAMPADO_DE_PANEL_A3.pptx",
     "HO-992 - OP 50 TAPIZADO Y ENGRAMPADO DE PANEL - TOP ROLL TRASERO", "50"),
    (r"APOYABRAZOS\DELANTERO\Camino_Inspeccion_Grampas_APB_Delantero_27.pptx",
     "CAMINO DE INSPECCION DE GRAMPAS - APB DE PUERTA DELANTERO - 27", None),
    (r"APOYABRAZOS\TRASERO\Camino_Inspeccion_Grampas_APB_Trasero_21.pptx",
     "CAMINO DE INSPECCION DE GRAMPAS - APB DE PUERTA TRASERO - 21", None),
    (r"INSERTO\DELANTERO\Camino_Inspeccion_Grampas_Inserto_Del_Der_36.pptx",
     "CAMINO DE INSPECCION DE GRAMPAS - INSERT DELANTERO - 36", None),
    (r"INSERTO\TRASERO\Camino_Inspeccion_Grampas_Inserto_Tras_23.pptx",
     "CAMINO DE INSPECCION DE GRAMPAS - INSERT TRASERO - 23", None),
    (r"TOP ROLL\DELANTERO\Camino_Inspeccion_Grampas_Top_Roll_Del_27.pptx",
     "CAMINO DE INSPECCION DE GRAMPAS - TOP ROLL DELANTERO - 27", None),
    (r"TOP ROLL\TRASERO\Camino_Inspeccion_Grampas_Top_Roll_Tras_34.pptx",
     "CAMINO DE INSPECCION DE GRAMPAS - TOP ROLL TRASERO - 34", None),
]
# La soldadura de soporte llego solo en PDF: se copia tal cual (no se puede numerar sin el pptx).
SOLO_PDF = [
    (r"TOP ROLL\TRASERO\HO-992_SOLDADURA_POR_ULTRASONIDO_DE_SOPORTE_A3.pdf",
     "HO-992 - SOLDADURA POR ULTRASONIDO DE SOPORTE - TOP ROLL TRASERO (sin N de operacion)"),
]


def numerar(origen, destino, n_op):
    prs = Presentation(origen)
    tocadas = 0
    for sh in prs.slides[0].shapes:
        if sh.name == "HO|f|nOp" and sh.has_text_frame:
            antes = sh.text_frame.text.strip()
            runs = [r for p in sh.text_frame.paragraphs for r in p.runs]
            if len(runs) != 1 or antes != "-":
                sys.exit(f"ABORTADO: el cajetin de {os.path.basename(origen)} no es el esperado "
                         f"(texto {antes!r}, {len(runs)} tramos)")
            runs[0].text = n_op
            tocadas += 1
    if tocadas != 1:
        sys.exit(f"ABORTADO: esperaba 1 casillero de N° de operacion en {origen}, hay {tocadas}")
    prs.save(destino)


def main():
    os.makedirs(OUT, exist_ok=True)
    hechos = []
    for rel, nombre, n_op in HOJAS:
        origen = os.path.join(SRC, rel)
        destino = os.path.join(OUT, nombre + ".pptx")
        if n_op:
            numerar(origen, destino, n_op)
        else:
            shutil.copy2(origen, destino)
        hechos.append(destino)
        print("pptx", "op " + n_op if n_op else "copia", "->", os.path.basename(destino))
    for rel, nombre in SOLO_PDF:
        shutil.copy2(os.path.join(SRC, rel), os.path.join(OUT, nombre + ".pdf"))
        print("pdf  copia ->", nombre + ".pdf")

    import win32com.client
    app = win32com.client.Dispatch("PowerPoint.Application")
    habia = app.Presentations.Count
    for p in hechos:
        pdf = os.path.splitext(p)[0] + ".pdf"
        if os.path.exists(pdf):
            os.remove(pdf)
        pres = app.Presentations.Open(p, ReadOnly=True, WithWindow=False)
        try:
            pres.SaveAs(pdf, 32)
        finally:
            pres.Close()
        print("pdf  ok    ->", os.path.basename(pdf), os.path.getsize(pdf), "bytes")
    if habia == 0 and app.Presentations.Count == 0:
        app.Quit()


if __name__ == "__main__":
    main()
