# -*- coding: utf-8 -*-
"""
matriz_hilos.py - Matriz de hilos por pieza (pedido de Carlos Baptista, 05/10/2026): que hilo aplica a que
pieza, separado por tabla (Patagonia, Amarok, Taos, SMRC, PWA, otros), el resumen de cada hilo (para ver
cuales no tienen ninguna pieza) y la lista pieza - hilo - uso para validar con Produccion.

Uso:
    python scripts/hilos/matriz_hilos.py <carpeta con las copias> <salida.xlsx>

La carpeta lleva COPIAS locales (nunca se abre el original del servidor):
    PE.xlsx, PROGRAMA.xlsx   ver matriz_datos.py
    PATAGONIA.xlsx           opcional: "Codigos y colores Patagonia_AAAAMMDD.xlsx" (hoja Ayuda Visual)
    COMPRAS.csv              opcional: ultima compra de cada hilo -> codigo;fecha;oc;proveedor;ocs
    fuentes.json             fechas de cada fuente, para la hoja Fuentes (ver FUENTES_DEFECTO)
No escribe en el arb ni en el servidor: solo el xlsx de salida.

El arb es UNA de las dos columnas del cruce: cada hilo de cada pieza se compara con lo que dicen los
documentos de esa pieza (hilos_documentos.py). El color de la celda dice si coinciden.
"""
import collections
import csv
import io
import json
import math
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)
import hilos_documentos as HD  # noqa: E402
import matriz_datos as MD  # noqa: E402
from criterio_hilo import es_hilo  # noqa: E402

AR = MD.AR
AZUL = '1F3A5F'
GRIS = 'F2F2F2'
BORDE = Border(*(Side(style='thin', color='999999'),) * 4)

OK, SOLO_ARB, FALTA, SIN_DOC, CONFLICTO = 'ok', 'solo_arb', 'falta', 'sin_doc', 'conflicto'
COLOR = {OK: 'E2F0D9', SOLO_ARB: 'FFF2CC', FALTA: 'F8CBAD', SIN_DOC: 'DDEBF7', CONFLICTO: 'FF8B8B'}
LEYENDA = [
    (OK, 'Coincide', 'Está en el arb y lo pide el documento de la pieza'),
    (CONFLICTO, 'Revisar', 'Está en el arb, pero un plano, una especificación, una BOM o un mail dice otro hilo (ver Observaciones)'),
    (FALTA, 'Falta en el arb', 'Lo pide el documento y no está cargado en el arb'),
    (SOLO_ARB, 'Solo en el arb', 'Está en el arb y el documento de la pieza no lo nombra'),
    (SIN_DOC, 'Aplica', 'Está en el arb; esa pieza no se cruzó con ningún documento'),
]
SITUACION = [
    ('uso', 'En uso: tiene pieza activa', 'E2F0D9'),
    ('sin_demanda', 'Solo en piezas sin programa, pedido ni producción en los últimos 3 meses', 'FFF2CC'),
    ('fuera', 'Solo en productos que no están en el plan de entrega', 'FCE4D6'),
    ('sin_pieza', 'Sin ninguna pieza en el arb', 'F8CBAD'),
]
FUENTES_DEFECTO = {'patagonia': '', 'programa': '', 'pe': '', 'maestro': '', 'compras': ''}


# ---------------------------------------------------------------------------- cruce

def cruzar(carpeta):
    d = MD.armar(carpeta)
    fuentes = dict(FUENTES_DEFECTO)
    pf = os.path.join(carpeta, 'fuentes.json')
    if os.path.exists(pf):
        with io.open(pf, encoding='utf-8') as fh:
            fuentes.update(json.load(fh))
    pat = os.path.join(carpeta, 'PATAGONIA.xlsx')
    docs, sin_listado = HD.armar(pat if os.path.exists(pat) else None,
                                 'Códigos y colores Patagonia' + (' (%s)' % fuentes['patagonia'] if fuentes['patagonia'] else ''))
    arb = d['arb']
    en_tablas = {p['art'] for p in d['piezas']}
    huerfanos = sorted(set(docs) - en_tablas)
    for p in d['piezas']:
        doc = docs.get(p['art'], {})
        celdas = collections.OrderedDict()
        for cod, h in p['hilos'].items():
            if cod in doc:
                celdas[cod] = {'estado': OK, 'uso': HD.texto_uso(doc[cod]['usos']), 'fuentes': doc[cod]['fuentes']}
            elif doc:
                celdas[cod] = {'estado': SOLO_ARB, 'uso': '', 'fuentes': sorted({f for x in doc.values() for f in x['fuentes']})}
            else:
                celdas[cod] = {'estado': SIN_DOC, 'uso': '', 'fuentes': []}
        for cod, x in doc.items():
            if cod not in p['hilos']:
                celdas[cod] = {'estado': FALTA, 'uso': HD.texto_uso(x['usos']), 'fuentes': x['fuentes']}
        p['notas'] = [n for n in HD.NOTAS if p['art'] in n['piezas']]
        for n in p['notas']:
            if n.get('conflicto') and n.get('codigo') in celdas and celdas[n['codigo']]['estado'] != FALTA:
                celdas[n['codigo']]['estado'] = CONFLICTO
        p['celdas'] = celdas
    d.update(docs=docs, sin_listado=sin_listado, huerfanos=huerfanos, fuentes=fuentes,
             compras=leer_compras(os.path.join(carpeta, 'COMPRAS.csv')))
    d['hilos'] = resumen_hilos(d)
    movimientos(d, carpeta)
    return d


def _dma(f):
    return f.strftime('%d/%m/%Y') if f else ''


def movimientos(d, carpeta):
    """Agrega a cada hilo su ultima compra y sus ultimos movimientos de stock (bases de PCP, si estan)."""
    idx = {}
    for x in d['hilos']:
        idx[AR.clave(x['cod'])] = x
        idx.setdefault(AR.compacto(x['cod']), x)

    def hilo(cod):
        return idx.get(AR.clave(cod)) or idx.get(AR.compacto(cod))
    d['rango_kardex'] = d['rango_ocs'] = None
    pk = os.path.join(carpeta, 'db_mov_depositos.xlsx')
    if os.path.exists(pk):
        mov, d['rango_kardex'] = MD.leer_kardex(pk, lambda c: hilo(c) is not None)
        for k, m in mov.items():
            x = hilo(k)
            for campo, f in m.items():
                if f and (x['mov'].get(campo) is None or f > x['mov'][campo]):
                    x['mov'][campo] = f
    po = os.path.join(carpeta, 'db_ordenes_compra.xlsx')
    if os.path.exists(po):
        ocs, d['rango_ocs'] = MD.leer_ocs(po, lambda c: hilo(c) is not None)
        for k, o in ocs.items():
            x = hilo(k)
            if not x['compra']:      # COMPRAS.csv (todas las OC) le gana a la base, que trae solo el año
                x['compra'] = {'fecha': _dma(o['fecha']), 'oc': o['oc'], 'proveedor': o['proveedor'], 'ocs': str(len(o['ocs']))}


def leer_compras(ruta):
    out = {}
    if not os.path.exists(ruta):
        return out
    with io.open(ruta, encoding='utf-8-sig') as fh:
        for r in csv.DictReader(fh, delimiter=';'):
            out[AR.clave(r['codigo'])] = r
    return out


def desc_hilo(arb, cod):
    return ' '.join((arb.rel.descripcion(cod) or '').split())


def resumen_hilos(d):
    """Una fila por codigo de hilo: en cuantas piezas esta y que situacion tiene."""
    arb = d['arb']
    hilos = collections.OrderedDict()

    def h(cod):
        return hilos.setdefault(AR.clave(cod), {
            'cod': cod, 'desc': desc_hilo(arb, cod), 'un': arb.rel.unidad(cod), 'por_tabla': collections.Counter(),
            'con': [], 'sin': [], 'fuera': [], 'falta': [], 'mov': {}})
    for p in d['piezas']:
        for cod, c in p['celdas'].items():
            x = h(cod)
            if c['estado'] == FALTA:
                x['falta'].append(p['art'])
                continue
            x['por_tabla'][p['tabla']] += 1
            (x['con'] if p['demanda'] else x['sin']).append(p['art'])
    # productos del arb que no son articulos del plan de entrega (solo los que nadie mas usa: no semielaborados)
    pe_compacto = {AR.compacto(p['art']): p['art'] for p in d['piezas']}
    for raiz in arb.productos():
        if AR.clave(raiz) in d['en_pe'] or arb.rel.tipo(raiz) != 'PRODUCTO':
            continue
        for cod in arb.hilos(raiz):
            gemelo = pe_compacto.get(AR.compacto(raiz))
            h(cod)['fuera'].append(raiz + (' (otra grafía de %s)' % gemelo if gemelo else ''))
    # hilos del maestro de insumos que no estan en ninguna BOM
    for k, (cod, desc, rubro) in arb.rel.insumos.items():
        if es_hilo(cod, desc) and k not in arb.rel._padres and k not in hilos:
            x = h(cod)
            x['desc'] = ' '.join(desc.split())
    for x in hilos.values():
        x['situacion'] = 'uso' if x['con'] else 'sin_demanda' if x['sin'] else 'fuera' if x['fuera'] else 'sin_pieza'
        x['compra'] = d['compras'].get(AR.clave(x['cod']))
    orden = [s[0] for s in SITUACION]
    return sorted(hilos.values(), key=lambda x: (orden.index(x['situacion']), -len(x['con']) - len(x['sin']), -len(x['fuera']), x['cod']))


# ---------------------------------------------------------------------------- planilla

def _c(ws, fila, col, valor, bold=False, fondo=None, centro=False, color='000000', size=10, italic=False, borde=True):
    c = ws.cell(row=fila, column=col, value=valor)
    c.font = Font(name='Calibri', size=size, bold=bold, color=color, italic=italic)
    if fondo:
        c.fill = PatternFill('solid', fgColor=fondo)
    c.alignment = Alignment(horizontal='center' if centro else 'left', vertical='center', wrap_text=True)
    if borde:
        c.border = BORDE
    return c


def _titulos(ws, fila, titulos, alto=32):
    for j, t in enumerate(titulos):
        _c(ws, fila, 2 + j, t, bold=True, fondo=AZUL, centro=True, color='FFFFFF')
    ws.row_dimensions[fila].height = alto


def _anchos(ws, anchos):
    ws.column_dimensions['A'].width = 2
    for j, a in enumerate(anchos):
        ws.column_dimensions[get_column_letter(2 + j)].width = a


def _alto(ws, fila, celdas, minimo=15.0):
    """Alto de la fila segun el texto mas largo: celdas = [(texto, ancho de columna)]."""
    lineas = 1
    for texto, ancho in celdas:
        if texto in (None, ''):
            continue
        n = 0
        for parte in str(texto).split('\n'):
            n += max(1, int(math.ceil(len(parte) / max(1.0, ancho * 1.15))))
        lineas = max(lineas, n)
    ws.row_dimensions[fila].height = max(minimo, 13.5 * lineas + 2)


def _pagina(ws, fila_titulos, col_fija=2, papel=None):
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = papel or ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = '%d:%d' % (fila_titulos, fila_titulos)
    ws.freeze_panes = ws.cell(row=fila_titulos + 1, column=col_fija)
    ws.sheet_view.showGridLines = False
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5


def _cabecera(ws, titulo, bajada):
    ws['B2'] = titulo
    ws['B2'].font = Font(name='Calibri', size=14, bold=True, color=AZUL)
    ws['B3'] = bajada
    ws['B3'].font = Font(name='Calibri', size=10, color='444444')


def _n(v):
    return '-' if v is None else int(round(v))


def texto_celda(c):
    if c['estado'] == OK:
        return c['uso']
    if c['estado'] == FALTA:
        return c['uso'] + '\n(falta en el arb)'
    if c['estado'] == SOLO_ARB:
        return 'Solo en el arb'
    if c['estado'] == CONFLICTO:
        return (c['uso'] or 'En el arb') + '\n(revisar)'
    return 'Aplica'


def _min(s):
    return s[0].lower() + s[1:] if s else s


def texto_nota(n):
    return '%s: %s. En el arb: %s.' % (n['documento'], n['dice'], n['arb'])


def observacion(p):
    partes = []
    for cod, c in p['celdas'].items():
        if c['estado'] == FALTA:
            partes.append('%s pide %s (%s) y no está en el arb.' % (' / '.join(c['fuentes']), cod, _min(c['uso'])))
        elif c['estado'] == SOLO_ARB and not any(n.get('codigo') == cod for n in p['notas']):
            partes.append('%s está en el arb; %s no lo nombra.' % (cod, ' / '.join(c['fuentes'])))
    partes += [texto_nota(n) for n in p['notas']]
    if not p['demanda']:
        partes.append('Sin programa, pedido ni producción en los últimos 3 meses.')
    if p['como'] != 'exacto' and p['arb']:
        partes.append('En el arb figura como %s.' % p['arb'])
    return ' '.join(partes)


def _leyenda(ws, fila):
    """Leyenda vertical: muestra de la celda (con su texto) y que quiere decir. Devuelve la ultima fila usada."""
    for est, muestra, txt in LEYENDA:
        _c(ws, fila, 2, muestra, fondo=COLOR[est], centro=True, bold=True, size=9)
        ws.merge_cells(start_row=fila, start_column=3, end_row=fila, end_column=5)
        _c(ws, fila, 3, txt, size=9, borde=False)
        ws.row_dimensions[fila].height = 24
        fila += 1
    return fila


def hoja_matriz(wb, d, tabla):
    piezas = [p for p in d['piezas'] if p['tabla'] == tabla]
    con = [p for p in piezas if p['celdas']]
    sin = [p for p in piezas if not p['celdas']]
    cuenta = collections.Counter(cod for p in con for cod in p['celdas'])
    cods = sorted(cuenta, key=lambda c: (-cuenta[c], c))
    ws = wb.create_sheet(tabla)
    _cabecera(ws, 'Matriz de hilos por pieza - %s' % tabla,
              'Qué hilo aplica a cada pieza: el hilo cargado en el arb, cruzado con los documentos de la pieza. Celda vacía = no aplica. '
              'Pieza en gris = sin programa, pedido ni producción en los últimos 3 meses.')
    F0 = _leyenda(ws, 5) + 1
    fijos = ['Sector', 'Pieza', 'Descripción', 'Programa\nsem. %s' % d['semana'].replace('W', ''), 'Pedido\n(firme +\nforecast)']
    anchos = [15, 19, 30, 9, 9]
    if d['rango_prod']:
        fijos.append('Última\nproducción\ndeclarada')
        anchos.append(11)
    C0 = 2 + len(fijos)          # primera columna de hilo
    _titulos(ws, F0, fijos + ['%s\n%s' % (c, desc_hilo(d['arb'], c)[:44]) for c in cods] + ['Observaciones'], alto=66)
    anchos = anchos + [17] * len(cods) + [70]
    _anchos(ws, anchos)
    f = F0 + 1
    ult = None
    for p in con:
        if p['sector'] != ult:
            ult = p['sector']
        obs = observacion(p)
        gris = None if p['demanda'] else GRIS
        vals = [p['sector'] or '-', p['art'], p['desc'], _n(p['prog_total']), _n(p['pedido'])]
        if d['rango_prod']:
            vals.append(_dma(p['ult_prod']) or '-')
        for j, v in enumerate(vals):
            _c(ws, f, 2 + j, v, centro=j >= 3, fondo=gris, bold=(j == 1))
        for j, cod in enumerate(cods):
            c = p['celdas'].get(cod)
            if c:
                _c(ws, f, C0 + j, texto_celda(c), centro=True, fondo=COLOR[c['estado']], bold=True, size=9)
            else:
                _c(ws, f, C0 + j, '', centro=True)
        revisar = any(n['conflicto'] and not n.get('codigo') for n in p['notas'])
        _c(ws, f, C0 + len(cods), obs, size=9, fondo=COLOR[CONFLICTO] if revisar else None)
        _alto(ws, f, [(obs, 70), (p['desc'], 30)] + [(texto_celda(c), 17) for c in p['celdas'].values()], minimo=18)
        f += 1
    if sin:
        f += 1
        _c(ws, f, 2, 'Piezas de estos sectores que no tienen hilo en el arb', bold=True, borde=False)
        f += 1
        por_sector = collections.OrderedDict()
        for p in sin:
            por_sector.setdefault(p['sector'] or '-', []).append(p)
        ancho_txt = sum(anchos[1:])
        for sector, ps in por_sector.items():
            txt = ' · '.join('%s%s%s%s' % (p['art'], ' (%s)' % p['desc'] if p['desc'] not in ('', '-') else '',
                                           '' if p['arb'] else ' [sin lista de materiales en el arb]',
                                           ' [ver hoja A revisar]' if p['notas'] else '') for p in ps)
            _c(ws, f, 2, sector, fondo=GRIS)
            ws.merge_cells(start_row=f, start_column=3, end_row=f, end_column=C0 + len(cods))
            _c(ws, f, 3, txt, size=9)
            _alto(ws, f, [(txt, ancho_txt)])
            f += 1
    _pagina(ws, F0, col_fija=4, papel=ws.PAPERSIZE_A3)
    return len(con), len(sin), cods


def hoja_resumen(wb, d):
    ws = wb.create_sheet('Resumen de hilos', 0)
    tablas = list(MD.TABLAS)
    _cabecera(ws, 'Hilos: en qué piezas aplica cada uno',
              'Todos los códigos de hilo del arb. El número de cada tabla es la cantidad de piezas del plan de entrega que lo llevan. '
              'Pieza activa = con programa en la semana, pedido o producción en los últimos 3 meses.')
    hay_compras = any(x['compra'] for x in d['hilos'])
    hay_mov = bool(d['rango_kardex'])
    tit = ['Hilo', 'Descripción en el arb', 'Unidad'] + tablas + ['Piezas\nactivas', 'Piezas\nno activas', 'Otros\nproductos\ndel arb']
    anchos = [21, 44, 7] + [10] * len(tablas) + [10, 10, 10]
    if hay_compras:
        tit += ['Última\ncompra']
        anchos += [11]
    if hay_mov:
        tit += ['Último\ningreso de\nproveedor', 'Último pase\nentre\ndepósitos']
        anchos += [11, 11]
    tit += ['Situación', 'Dónde figura', 'Observaciones']
    anchos += [34, 60, 50]
    F0 = 5
    _titulos(ws, F0, tit, alto=48)
    _anchos(ws, anchos)
    nombre = {s[0]: s[1] for s in SITUACION}
    color = {s[0]: s[2] for s in SITUACION}
    f = F0 + 1
    for x in d['hilos']:
        if x['situacion'] == 'uso':
            donde = ''
        elif x['situacion'] == 'sin_demanda':
            donde = ' · '.join(x['sin'][:8]) + (' y %d más' % (len(x['sin']) - 8) if len(x['sin']) > 8 else '')
        elif x['situacion'] == 'fuera':
            donde = ' · '.join(x['fuera'][:6]) + (' y %d más' % (len(x['fuera']) - 6) if len(x['fuera']) > 6 else '')
        else:
            donde = 'Está en el maestro de insumos y en ninguna lista de materiales.'
        vals = [x['cod'], x['desc'], x['un']] + [x['por_tabla'].get(t) or '' for t in tablas] + [len(x['con']) or '', len(x['sin']) or '', len(x['fuera']) or '']
        if hay_compras:
            vals += [(x['compra'] or {}).get('fecha', '') or '-']
        if hay_mov:
            vals += [_dma(x['mov'].get('ingreso')) or '-', _dma(x['mov'].get('pase')) or '-']
        nota = HD.NOTAS_HILO.get(x['cod'], '')
        vals += [nombre[x['situacion']], donde, nota]
        n = len(vals)
        for j, v in enumerate(vals):
            _c(ws, f, 2 + j, v, centro=(2 <= j < n - 3), bold=(j == 0),
               fondo=color[x['situacion']] if j == n - 3 else None, size=9 if j >= n - 2 else 10)
        _alto(ws, f, [(donde, 60), (x['desc'], 44), (nota, 50)])
        f += 1
    _pagina(ws, F0, col_fija=3, papel=ws.PAPERSIZE_A3)


def grupos(d):
    """Junta las piezas del mismo sector que llevan exactamente los mismos hilos con el mismo uso."""
    out = collections.OrderedDict()
    for p in d['piezas']:
        if not p['celdas']:
            continue
        firma = tuple((cod, c['estado'], c['uso']) for cod, c in sorted(p['celdas'].items()))
        g = out.setdefault((p['tabla'], p['sector'], p['demanda'], firma), {'piezas': [], 'ref': p})
        g['piezas'].append(p)
    return out


def hoja_lista(wb, d):
    ws = wb.create_sheet('Lista para Producción')
    _cabecera(ws, 'Hilo y uso por pieza - para validar con Producción',
              'Una fila por hilo. Las piezas que llevan los mismos hilos con el mismo uso van juntas. Completar las tres columnas de la derecha.')
    tit = ['Tabla', 'Sector', 'Piezas', 'Pieza\nactiva', 'Hilo', 'Descripción del hilo', 'Uso', 'Según', 'En el\narb',
           '¿Es así en\nplanta?\n(Sí / No)', 'Si no: qué hilo\nse usa', 'Observaciones de Producción']
    F0 = 5
    _titulos(ws, F0, tit, alto=48)
    anchos = [12, 15, 40, 9, 19, 36, 26, 30, 6, 11, 20, 36]
    _anchos(ws, anchos)
    f = F0 + 1
    n = 0
    for (tabla, sector, demanda, firma), g in grupos(d).items():
        piezas = ' · '.join(p['art'] for p in g['piezas'])
        for cod, c in g['ref']['celdas'].items():
            conocido = c['estado'] in (OK, FALTA) or (c['estado'] == CONFLICTO and c['uso'])
            uso = (c['uso'] + (' (revisar: ver hoja A revisar)' if c['estado'] == CONFLICTO else '')) if conocido else 'A confirmar'
            segun = ' / '.join(c['fuentes']) if conocido else ''
            vals = [tabla, sector or '-', piezas, 'Sí' if demanda else 'No', cod, desc_hilo(d['arb'], cod), uso, segun,
                    'No' if c['estado'] == FALTA else 'Sí', '', '', '']
            for j, v in enumerate(vals):
                _c(ws, f, 2 + j, v, centro=j in (3, 8), bold=(j == 4), size=9 if j in (2, 5, 7) else 10,
                   fondo=COLOR[c['estado']] if j == 8 and c['estado'] in (FALTA, SOLO_ARB, CONFLICTO) else (None if demanda else GRIS) if j < 4 else None)
            _alto(ws, f, [(piezas, 40), (vals[5], 36), (uso, 26), (segun, 30)], minimo=20)
            f += 1
            n += 1
    _pagina(ws, F0, col_fija=5, papel=ws.PAPERSIZE_A3)
    return n


def hoja_diferencias(wb, d):
    ws = wb.create_sheet('A revisar')
    _cabecera(ws, 'Diferencias entre el arb y los documentos',
              'Lo que el documento de la pieza dice distinto de lo que está cargado. No se cambió nada en el arb.')
    tit = ['Tabla', 'Piezas', 'Pieza\nactiva', 'Hilo', 'Qué dice el documento', 'Qué tiene el arb', 'Documento']
    F0 = 5
    _titulos(ws, F0, tit, alto=36)
    anchos = [12, 44, 9, 19, 52, 44, 44]
    _anchos(ws, anchos)
    f = F0 + 1
    n = 0
    filas = []
    for (tabla, sector, demanda, firma), g in grupos(d).items():
        piezas = ' · '.join(p['art'] for p in g['piezas'])
        for cod, c in g['ref']['celdas'].items():
            if c['estado'] == FALTA:
                filas.append((tabla, piezas, demanda, cod, 'Pide %s (%s): %s' % (cod, desc_hilo(d['arb'], cod) or 'no está en el maestro', c['uso'].lower()),
                              'No lo tiene cargado en esas piezas.', ' / '.join(c['fuentes'])))
            elif c['estado'] == SOLO_ARB:
                filas.append((tabla, piezas, demanda, cod, 'No lo nombra para esas piezas.',
                              'Tiene %s (%s).' % (cod, desc_hilo(d['arb'], cod)), ' / '.join(c['fuentes'])))
    notas_cod = {(n.get('codigo'), p) for n in HD.NOTAS for p in n['piezas']}
    filas = [r for r in filas if not any((r[3], art) in notas_cod for art in r[1].split(' · '))]
    for nota in HD.NOTAS:
        ps = [p for p in d['piezas'] if p['art'] in nota['piezas']]
        if not ps:
            continue
        filas.append((ps[0]['tabla'], ' · '.join(p['art'] for p in ps), any(p['demanda'] for p in ps), nota.get('codigo') or '',
                      nota['dice'], nota['arb'], nota['documento']))
    orden = list(MD.TABLAS)
    filas.sort(key=lambda r: (orden.index(r[0]), not r[2]))
    for r in filas:
        vals = [r[0], r[1], 'Sí' if r[2] else 'No', r[3], r[4], r[5], r[6]]
        for j, v in enumerate(vals):
            _c(ws, f, 2 + j, v, centro=(j == 2), bold=(j == 3), size=9 if j in (1, 4, 5, 6) else 10, fondo=None if r[2] else (GRIS if j < 3 else None))
        _alto(ws, f, [(vals[1], 44), (vals[4], 52), (vals[5], 44), (vals[6], 44)], minimo=20)
        f += 1
        n += 1
    _pagina(ws, F0, col_fija=3)
    return n


def hoja_evidencia(wb, carpeta):
    """Captura y cita de cada diferencia de la hoja A revisar (evidencia/evidencia.json en la carpeta de entrada:
    [{titulo, texto, imagen?, ancho?}], las imagenes al lado). Las capturas no van al repo: son papeles del cliente."""
    ruta = os.path.join(carpeta, 'evidencia', 'evidencia.json')
    if not os.path.exists(ruta):
        return 0
    from openpyxl.drawing.image import Image
    from openpyxl.worksheet.pagebreak import Break
    with io.open(ruta, encoding='utf-8') as fh:
        items = json.load(fh)
    ws = wb.create_sheet('Evidencia')
    _cabecera(ws, 'Evidencia de cada diferencia', 'La cita textual y la captura del documento, en el orden de la hoja A revisar.')
    _anchos(ws, [150])
    f = 5
    for k, it in enumerate(items):
        if k:
            ws.row_breaks.append(Break(id=f - 1))     # una diferencia por hoja impresa
        _c(ws, f, 2, it['titulo'], bold=True, fondo=GRIS, borde=False)
        ws.row_dimensions[f].height = 20
        f += 1
        _c(ws, f, 2, it['texto'], size=10, borde=False)
        _alto(ws, f, [(it['texto'], 150)], minimo=18)
        f += 1
        if it.get('imagen'):
            img = Image(os.path.join(carpeta, 'evidencia', it['imagen']))
            ancho = it.get('ancho', 800)
            img.height = int(img.height * ancho / img.width)
            img.width = ancho
            ws.add_image(img, 'B%d' % f)
            f += int(math.ceil(img.height / 20.0)) + 1
        f += 1
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return len(items)


def hoja_fuentes(wb, d, conteos):
    ws = wb.create_sheet('Fuentes')
    _cabecera(ws, 'De dónde sale cada dato', '')
    fu = d['fuentes']
    hilos = d['hilos']
    por_sit = collections.Counter(x['situacion'] for x in hilos)
    lineas = [
        ('Qué piezas entran', 'Los %d artículos vigentes del plan de entrega de PCP (%s%s). Van a las tablas los de sectores que cosen '
         '(%d artículos); los demás sectores no llevan hilo (%s).'
         % (d['n_pe'], 'semana ' + d['semana'].replace('W', ''), ', ' + fu['pe'] if fu['pe'] else '', len(d['piezas']),
            ', '.join('%s %d' % kv for kv in sorted(d['fuera'].items())))),
        ('Pieza activa', 'Tiene programa en la semana (Programa de Producción de PCP%s, en cualquier proceso), pedido (firme de la '
         'semana más todo el forecast del plan de entrega) o producción declarada en los últimos 3 meses%s. Una pieza en gris '
         'no tiene ninguno de los tres.'
         % (' (%s)' % fu['programa'] if fu['programa'] else '',
            ' (base de producción de PCP, %s a %s)' % (_dma(d['rango_prod'][0]), _dma(d['rango_prod'][1])) if d['rango_prod'] else '')),
        ('Hilo de cada pieza', 'Lista de materiales del arb, abierta hasta el último nivel. %s' % fu.get('arb', '')),
        ('Hilos sin pieza', 'Maestro de insumos del arb%s.' % (' (%s)' % fu['maestro'] if fu['maestro'] else '')),
        ('Documentos de cada pieza', 'Hojas de operaciones vigentes del SGC, BOM oficiales de Ingeniería, planos y especificaciones del '
         'cliente (FAKOM, CMF, LSC), el listado "Códigos y colores Patagonia"%s y, para traducir un color de VW al código de Linhanyl, '
         'la tabla de Linhanyl del 30/01/2026. Donde ningún documento nombra el hilo, la pieza queda sin cruzar (celeste) y el uso '
         'va "a confirmar".' % (' del %s' % fu['patagonia'] if fu['patagonia'] else '')),
    ]
    if any(x['compra'] for x in hilos):
        lineas.append(('Última compra', 'Órdenes de compra%s.' % (' (%s)' % fu['compras'] if fu['compras'] else '')))
    if d.get('rango_kardex'):
        lineas.append(('Movimientos de stock', 'Movimientos de depósito del arb (base de PCP, %s a %s): ingreso con remito del '
                       'proveedor y pase entre depósitos. No cuenta lo que descuenta el sistema por la lista de materiales.'
                       % (_dma(d['rango_kardex'][0]), _dma(d['rango_kardex'][1]))))
    lineas += [
        ('Un mismo hilo, dos escrituras', 'Las hojas de operaciones de Patagonia escriben FX284TK-E0PTO; en el arb, en las órdenes de compra y '
         'en la tabla de Linhanyl es FX284-E0PTO. Se tomó como el mismo hilo.'),
        ('Cuántos hilos', '%d códigos de hilo: %s.' % (len(hilos), '; '.join('%d %s' % (por_sit[s[0]], s[1][0].lower() + s[1][1:]) for s in SITUACION))),
        ('Cuántas piezas', '; '.join('%s: %d con hilo y %d sin hilo' % (t, c[0], c[1]) for t, c in conteos.items()) + '.'),
        ('Qué no es', 'No es una lista de hilos dados de baja: un hilo sin pieza o sin pedido es un candidato a revisar con Producción y Compras. '
         'En el arb no se cargó ni se cambió nada.'),
    ]
    _anchos(ws, [26, 150])
    f = 4
    for k, v in lineas:
        _c(ws, f, 2, k, bold=True, fondo=GRIS)
        _c(ws, f, 3, v)
        _alto(ws, f, [(v, 150)], minimo=18)
        f += 1
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def generar(carpeta, salida):
    d = cruzar(carpeta)
    wb = Workbook()
    wb.remove(wb.active)
    conteos = collections.OrderedDict()
    for tabla in MD.TABLAS:
        con, sin, cods = hoja_matriz(wb, d, tabla)
        conteos[tabla] = (con, sin, len(cods))
    hoja_resumen(wb, d)
    n_lista = hoja_lista(wb, d)
    n_dif = hoja_diferencias(wb, d)
    hoja_evidencia(wb, carpeta)
    hoja_fuentes(wb, d, conteos)
    wb.save(salida)
    return d, conteos, n_lista, n_dif


def main():
    sys.stdout.reconfigure(encoding='utf-8')
    if len(sys.argv) < 3:
        raise SystemExit(__doc__)
    carpeta, salida = os.path.abspath(sys.argv[1]), os.path.abspath(sys.argv[2])
    d, conteos, n_lista, n_dif = generar(carpeta, salida)
    for l in d['arb'].sello.lineas():
        print(l)
    print('plan de entrega %s: %d articulos, %d en las tablas' % (d['semana'], d['n_pe'], len(d['piezas'])))
    for t, (con, sin, n) in conteos.items():
        print('  %-20s %3d piezas con hilo, %3d sin hilo, %2d hilos' % (t, con, sin, n))
    sit = collections.Counter(x['situacion'] for x in d['hilos'])
    print('hilos: %d  ->  %s' % (len(d['hilos']), dict(sit)))
    est = collections.Counter(c['estado'] for p in d['piezas'] for c in p['celdas'].values())
    print('celdas: %s' % dict(est))
    print('lista para Produccion: %d filas   diferencias: %d filas' % (n_lista, n_dif))
    if d['huerfanos']:
        print('ATENCION - piezas con documento que no estan en el plan de entrega: %s' % d['huerfanos'])
    dobles = d['arb'].grafias_dobles()
    if dobles:
        print('productos con dos grafias en el arb (se uso el codigo exacto del plan de entrega): %d' % len(dobles))
    print('guardado: %s' % salida)


if __name__ == '__main__':
    main()
