"""
_pendientesPcIpPad.py — hoja simple que acompaña la propuesta del plan de control de costura y tapizado del IP PAD
Patagonia (22/10/2026): lo que falta definir para que el plan quede definitivo, separado por quien lo tiene que contestar,
y aparte la propuesta de siglas del AMFE 149 y las correcciones de Ingenieria (para Fak, que es quien las asigna y las hace).

Palabras de planta, cada punto con su papel, sin contar como se llego a cada diferencia (core-prohibiciones §1).
Mismo formato que scripts/apb/_pendientesPcApb.py (reusa sus funciones de hoja y de PDF).

Uso:  python scripts/ippad/_pendientesPcIpPad.py <carpeta_salida> <amfe149.json>   -> .xlsx + .pdf (y la propuesta de siglas .xlsx)
"""
import json
import os
import sys

from openpyxl import Workbook

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'apb'))
from _pendientesPcApb import exportar_pdf, hoja  # noqa: E402

NOMBRE = 'Pendientes plan de control IP PAD Patagonia costura y tapizado 22-10-2026'
NOMBRE_SIGLAS = 'Propuesta de siglas AMFE 149 IP PAD y correcciones de Ingenieria 22-10-2026'

# (para quien, operacion, que falta, lo que hay hoy)
PENDIENTES = [
    ('Calidad', '10 Recepción', 'Planes de recepción de los vinilos Sanleather IV (1246030227, 1246030229, 1246030230) y del Sanflex G KL S202 (124.505.0372-7)',
     'El plan toma espesor, gramaje, ancho y flamabilidad de los certificados de Sansuy CA-363026, CA-389326, CA-388926 (EPS-2206 Rev. 2) '
     'y CA-1568326 (EPS-1890 Rev. 8). La carpeta VINILO\\Patagonia de los planes de recepción está vacía'),
    ('Calidad', '10 Recepción', 'Planes de recepción del cordón de piping 427VAR002COS01, la grampa Dorking 84/06 (DK/1840600), la espuma de polietileno '
     'EMBEPE2X1X20 y las etiquetas ET-SATO', 'No hay planes vigentes de estos códigos (de la grampa 84/06 hay solo uno obsoleto). '
     'El plan deja en TBD la tolerancia del diámetro del cordón, el dimensional de la grampa y el espesor de la espuma de polietileno'),
    ('Calidad', '10 Recepción', 'Muestra de los consignados de VW (logo airbag, clip, tornillo y carcasa del difusor)',
     'El plan va al 100 % por entrega, como el control de detección del AMFE 149'),
    ('Calidad', '10 Recepción', 'Planes de recepción con datos a revisar: el 1043 escribe la flamabilidad del hilo "100 mm/min" sin el "<"; '
     'los números 1063 y 1064 están repetidos (hilos y estructuras de apoyacabezas) y el 1062 también; el 729 dice adhesivo "x 15 kg" y el arb "x 18 L"',
     'El plan nombra cada hilo con su código al lado del número de plan'),
    ('Calidad', 'Instrumentos', 'Cámara de flamabilidad MC184 (vencida el 11/09/2026) y garrafa MC185 (18/09/2026), que el plan usa en la recepción del '
     'vinilo y de la espuma', 'Cronograma de calibración 2026 Rev. 7. La cinta de recepción MC403 tiene la próxima fecha en 2027 con la última '
     'calibración del 12/06/2025 a 12 meses'),
    ('Calidad', 'Instrumentos', 'Códigos MC de: regla de costura, regla de largo de capa, plantilla Mylar del Low, calibre de control de la pieza tapizada '
     '(control final y costura alineada), calibre de espesor de la espuma troquelada, balanza y cronómetro del adhesivado',
     'En el cronograma están la plantilla MC517 (solo High), el calibre de forma MC509 (sustrato, sector Inyección VWA), la atornilladora MC400 de '
     'Ensamble IP y la regla flexible MC162 de mesa de corte. El calibre de costura MC218 dice "Amarok PA2 & IP" y no se asignó'),
    ('Calidad', '30, 31, 50, 60, 70, 80, 81 y 110', 'Frecuencia de los controles de preparación y corte del vinilo, troquelado, ensamble sustrato + espuma, '
     'mezcla del adhesivo, control y retrabajo de adhesivado, y del ensayo de resistencia de la soldadura',
     'El plan las deja en TBD. La hoja de operaciones dice frecuencia "1" en las hojas 30, 32 y 34, sin unidad'),
    ('Calidad', '32 Control con Mylar', 'Frecuencia del control del contorno',
     'El plan pone las dos: 1 pieza en arranque, mantenimiento, corte de energía y problema de calidad (hoja de operaciones) y 100 % por lote (AMFE 149)'),
    ('Calidad', '10 y 110 Logo airbag (D)', 'Registro de los controles del logo airbag, que llevan D, y si el P-09/I alcanza como plan de reacción',
     'El plan va con autocontrol al 100 %, como el AMFE 149'),
    ('Calidad', 'Encabezado', 'Número de plan y equipo. La próxima revisión la abre Calidad', 'El plan va con la Rev. 0 del plan preliminar (05/04/2025) '
     'y el contacto y el equipo del plan de inyección del IP PAD (G. Cal; C. Rodríguez / A. Aguayo)'),
    ('Calidad', 'Lay out', 'Qué ensayos de la hoja ADV del cliente entran en el lay out anual y con qué cantidad',
     'El plan pone flamabilidad ≤ 70 mm/min (CVTC 52034, plano C00751801), adhesión ≥ 525 N/m (CVTC 35041), rayado, solidez a la luz, fogging '
     'y olor, con la cantidad de la hoja ADV, en laboratorio (la hoja ADV pide laboratorio acreditado por la terminal). El plan preliminar decía '
     'flamabilidad con 1 muestra por auditoría de producto. Los materiales se reciben con < 100 mm/min (planes de recepción y certificados)'),
    ('Calidad', '10 a 130', 'Muestras límite (o fotos de aceptación y rechazo) de: limpieza de los materiales y estado de la carcasa en recepción; '
     'aspecto del cordón y del logo; costura completa (40, 41 y 42); vinilo sin roturas; vivo y unión de paños; espuma troquelada; Tesa en el ensamble; '
     'superficie y cantidad de adhesivo (70); costura encastrada (90); tapizado (91); virolado y zonas críticas (100); zona soldada (110); '
     'refilado y puntas selladas (120); despegues, aspecto, costura vista y logo en el control final (130)',
     'Del IP PAD no hay muestra patrón, muestra límite, biblia de defectos ni fotos OK/NOK (buscado en el SGC, el legajo y los mails el 22/10/2026). '
     'El plan escribe "Visual / muestra límite (TBD: la define Calidad)" en cada una. En la op 80 usa las piezas NOK colgadas en el puesto '
     '(falta de adhesivo en vinilo y en sustrato, exceso): confirmar que son del IP PAD y que siguen ahí'),
    ('Calidad', '10 Recepción', 'Patrón de color y de grabado de los vinilos Sanleather IV, patrón de color del Sanflex G KL, del cordón de piping '
     'y de la cinta Tesa (color y liner)', 'El plan de recepción 1048 controla el color de la Tesa a la vista. Los hilos tienen patrón según sus planes '
     'de recepción (1043, 1063 y 1064): confirmar que existe y usarlo también en la costura vista (op 42)'),
    ('Calidad', '40, 41 y 90', 'Tolerancias de la toma de costura de unión (6 mm), de la distancia de la 2ª costura al vivo (6 mm) y del ancho de la zona '
     'pegada debajo de la costura (5 cm)', 'Los valores de costura de Patagonia dicen 6 mm justos y la hoja de operaciones 5 cm, sin tolerancia'),
    ('Calidad', '110 Logo airbag (D)', 'Criterio de la prueba de fijación del logo (fuerza o método de aceptación)',
     'El plan controla presencia y posición contra el plano C00751801 (lista de materiales y los 2 alojamientos con nervio antierror) y la '
     'fijación con una prueba manual en la 110 y en el control final, con la fuerza TBD'),
    ('Calidad', '81 y 140 Retrabajos', 'Si VW aprobó los retrabajos de la pieza', 'Están en el flujograma 157, el AMFE 149 y la hoja de operaciones 985'),
    ('Leo Lattanzi', '10 Recepción', 'Densidad, espesor y código de la espuma de soporte',
     'Ficha del proveedor (Mentvil, calidad 45) y ficha técnica Barack Rev. A (19/11/2025): 45 kg/m³ y 3 mm; BOM006 y BOM V3: 427ESP002TRO01 '
     'de 3,15 mm 45 kg/m³; arb: 427ESP001TRO01 de 35 kg/m³'),
    ('Leo Lattanzi', '70 Adhesivado', 'Relación de mezcla del adhesivo FA con el reticulante GV, tiempo de agitación y vida útil de la mezcla',
     'La auditoría VDA 6.3 de VW (2025) marcó que la mezcla y la agitación no estaban en el plan de control. El plan pone la vida útil de la '
     'ficha técnica vigente del reticulante (HT.REGV Rev. 4: 4 a 8 h); la relación de mezcla y el tiempo de agitación no figuran en ningún papel de la pieza'),
    ('Leo Lattanzi', '90, 91 y 100', 'Temperatura (y tiempo) de la pistola de calor del pre-fixing y del fijado superior; temperatura y presión de la máquina de virolado',
     'Ningún papel de la pieza tiene valores. El plan preliminar de 2025 hablaba de un horno a 60-70 °C'),
    ('Leo Lattanzi', '110 Ultrasonido y ensamble', 'Parámetros del ultrasonido del logo airbag, método y valor del ensayo de resistencia de la soldadura, '
     'y torque de los tornillos 2HC.075.701.HE', 'La lista de equipos del 14/11/2025 nombra el equipo sin valores; la atornilladora MC400 es de 0,5 a 2 Nm'),
    ('Planta', '30 a 110', 'Número de cada máquina: mesa de corte, máquinas de coser de unión y de vista, troqueladora y troquel, dispositivo de ensamble, '
     'pistola de adhesivado, máquina de virolado y ultrasonido', 'Para completar la columna de máquina. La mesa de corte figura BMA090 en el AMFE 149 '
     'y BM 149 en el plan preliminar'),
    ('Planta', '31 Corte', 'Velocidad de la máquina y tensión del vinilo', 'La hoja de operaciones pide verificarlas con un instructivo de parámetros que no está'),
    ('Planta', '150 Embalaje', 'Medio de embalaje y cantidad por medio', 'No hay ficha de embalaje de serie: la GE-279 es borrador y la GE-288 es de repuestos'),
]

# para Ingenieria (Fak): no van a Calidad
INGENIERIA = [
    ('Ingeniería', 'Hoja de operaciones 985', 'Hojas 50, 60, 70, 81, 90, 91, 100, 110, 120, 130, 140 y 150 sin ciclo de control; 60, 70, 130, 140 y 150 '
     '"por completar"; la 120 del Excel en TBD (la hoja A3 del 01/10 está completa)', 'El plan ya tiene los controles de esas operaciones'),
    ('Ingeniería', 'Hoja de operaciones 985', 'Numeración vieja: pestañas 20 a 26 y 30 a 38', 'Flujograma 157 Rev. B: 20 y 21; 30, 31 y 32'),
    ('Ingeniería', 'AMFE 149', 'Op 120: el control de detección de "aplicar calor excesivamente" dice "undefined"', 'Supabase, updated_at 15/10/2026'),
    ('Ingeniería', 'AMFE 149', 'Op 70: el elemento "Parámetros presión y temperatura" tiene la función "Asegurar alineación de costura"', 'Supabase, updated_at 15/10/2026'),
    ('Ingeniería', 'AMFE 149', 'Op 10: "Identificación o documentación de lote faltante" repetido 3 veces; "Piping tubing negro (código TBD)" es el 427VAR002COS01; '
     'la espuma dice 3,15 mm 45 kg/m³', 'BOM arb del 06/10/2026'),
    ('Ingeniería', 'AMFE 149', 'Op 110 no tiene el ensamble de la carcasa, los clips y los tornillos', 'El flujograma dice "soldadura con ultrasonido y ensamble"'),
    ('Ingeniería', 'AMFE 149', 'La flamabilidad nombra la TL 1010; el plano pide CVTC 52034 ≤ 70 mm/min', 'Plano C00751801 y hoja ADV PKK1'),
    ('Ingeniería', 'Siglas', 'El AMFE 149 marca D/TLD la flamabilidad del vinilo y de la espuma y el logo airbag; el plano C00751801 no designa D/TLD '
     '(la tabla KCDS marca solo AQC en la perilla del aireador, que es consignada)', 'El plan copia la D del AMFE'),
    ('Ingeniería', 'Costura', 'Toma de la costura de unión: 6 mm exactos (desvío de Barack) contra 8 ± 1 mm de la LAH.000.881.N, sin nota en el plano',
     'La LAH pide que los desvíos acordados con VW queden en el plano'),
    ('Ingeniería', 'Vivo L2 / L3', 'El informe del 23/03/2026 propone pegar el vivo con hot-melt; la hoja de operaciones y el flujograma lo cosen', 'El plan va cosido'),
    ('Ingeniería', 'Embalaje', 'GE-279 en borrador con datos de relleno ("CAJA AAAA", 5 por caja) y los números viejos C00694582 / C00694583', 'Carlos Baptista'),
    ('Ingeniería', 'BOM', 'Sanflex G KL: el certificado dice "Laranja S-202" y la BOM V3 "Argan Oil"', 'Confirmar que es el mismo vinilo de piping'),
    ('Ingeniería', 'Siglas', 'El certificado del Sanflex G KL (CA-1568326) marca la flamabilidad con "YC", una sigla de diseño del proveedor',
     'El AMFE 149 no tiene sigla en el piping'),
    ('Ingeniería', 'Hoja de operaciones 985', 'Hilo de unión escrito FX284TK-E0PTO; en el arb es FX284-E0PTO', 'El plan va con el código del arb'),
    ('Ingeniería', 'Alcance', 'El AEKO RZ00345 agrega una L0 (2HC.858.417.D) que no está en el plan', 'El plan cubre L1, L2 y L3'),
]


def siglas(amfe_json):
    """Lo que el criterio de I-AC-005 rev.B marca y el AMFE 149 no tiene asignado (S 9-10 = critica; S 5-8 y O >= 4 = SC)."""
    j = json.load(open(amfe_json, encoding='utf-8'))
    d = j.get('doc') or j.get('data')
    grupos = {}
    for op in d['operations']:
        for we in op.get('workElements', []):
            for fn in we.get('functions', []):
                for fm in fn.get('failures', []):
                    for c in fm.get('causes', []):
                        s = int(c.get('severity') or fm.get('severity') or 0)   # la S de la causa, no la del modo de falla
                        o = int(c.get('occurrence') or 0)
                        if (c.get('specialChar') or '').strip():
                            continue
                        crit = 'Crítica (D)' if s >= 9 else ('SC' if 5 <= s <= 8 and o >= 4 else '')
                        if not crit:
                            continue
                        k = (str(op['opNumber']), fm['description'], crit)
                        grupos.setdefault(k, set()).add((s, o))
    filas = []
    for (op, fm, crit), so in grupos.items():
        filas.append((op, fm, ', '.join(f'S{s} O{o}' for s, o in sorted(so)), crit))

    def num(op):
        return int(str(op).split('-')[0])
    filas.sort(key=lambda f: (0 if f[3].startswith('Cr') else 1, num(f[0])))
    return filas


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(2)
    salida, amfe_json = os.path.abspath(sys.argv[1]), sys.argv[2]
    os.makedirs(salida, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Pendientes'
    hoja(ws, 'Plan de control IP PAD Patagonia costura y tapizado — lo que falta definir',
         '2HC.858.417.B FAM / 2HC.858.417.C GKK / 2HC.858.417.C GKN · propuesta de Ingeniería del 22/10/2026',
         ('Quién', 'Operación', 'Qué falta', 'Lo que hay hoy'), PENDIENTES, (14, 18, 62, 62))
    wb.properties.creator = 'Barack Mercosul'
    wb.properties.lastModifiedBy = 'Barack Mercosul'
    destino = os.path.join(salida, NOMBRE + '.xlsx')
    wb.save(destino)
    exportar_pdf(destino)
    print(destino)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Propuesta de siglas'
    filas = siglas(amfe_json)
    hoja(ws, 'AMFE 149 IP PAD — siglas que marca el criterio y no están asignadas',
         'Criterio I-AC-005 rev.B: S 9 o 10 = crítica (D); S 5 a 8 con O 4 o más = SC. Las asigna Ingeniería. '
         'Las S 10 de la op 70 (ventilación) y la op 120 (cortadura, quemadura) son de seguridad del operario.',
         ('Op', 'Modo de falla', 'S y O de sus causas', 'El criterio marca'), filas, (8, 80, 22, 16))
    ws2 = wb.create_sheet('Correcciones de Ingeniería')
    hoja(ws2, 'Correcciones de Ingeniería que salieron del plan de control del IP PAD', '22/10/2026',
         ('Quién', 'Documento', 'Qué corregir', 'De dónde sale'), INGENIERIA, (14, 22, 62, 50))
    wb.properties.creator = 'Barack Mercosul'
    wb.properties.lastModifiedBy = 'Barack Mercosul'
    destino = os.path.join(salida, NOMBRE_SIGLAS + '.xlsx')
    wb.save(destino)
    print(destino, len(filas), 'modos de falla')


if __name__ == '__main__':
    main()
